"""v3/control1_modos.py — AUDITORÍA DE ASIGNACIONES (CONTROL 1): modo PRELIMINAR
(mes abierto) y modo CIERRE DEFINITIVO. FASE 12E.6 (2026-09-18).

POR QUÉ EXISTE (sin tocar V2 — control_asignaciones.py sigue congelado):

`control_asignaciones.ejecutar_control()` (V2) está pensado para UN cierre
único: (a) empareja las decisiones previas del auditor por FILA_GLOBAL y solo
si el SHA256 del GLOBAL coincide, así que con un mes abierto — donde GLOBAL se
regenera al entrar más cierres y cambian tanto el SHA como las filas — TODAS
las decisiones se descartarían en silencio; y (b) en cuanto todas las alertas
están validadas, cierra solo: corrige el GLOBAL, actualiza el histórico y deja
el periodo como YA_PROCESADO. Un auditor que trabaja a mitad de mes necesita
lo contrario: revisar, decidir y volver a correr las veces que haga falta.

Este módulo agrega DOS modos como capa V3 que reutiliza las funciones de V2
tal cual (lectura del GLOBAL, exclusiones, detección de alertas, Excel de
revisión, resoluciones, histórico) y NO las modifica:

  PRELIMINAR (por defecto). Usa el GLOBAL oficial actual y el histórico
    maestro de la raíz; recupera la revisión previa y CONSERVA las decisiones
    del auditor emparejando por IDENTIDAD de la alerta (asignación + cuenta +
    fecha valor + importe + glosa + n-ésima ocurrencia) en vez de por
    fila/SHA, así sobreviven a un GLOBAL regenerado; agrega solo las alertas
    nuevas; una alerta que ya no está en el GLOBAL NO se borra: pasa a la hoja
    "ALERTAS_NO_VIGENTES" del mismo Excel (y vuelve a la hoja REVISION con su
    decisión si reaparece). NUNCA escribe el histórico maestro, NUNCA modifica
    el GLOBAL, NUNCA marca el periodo como procesado.

  CERRAR. Acción explícita (modo_control1="cerrar" + confirmación). Primero
    reancla las decisiones al GLOBAL actual (mismo emparejamiento por
    identidad) y luego delega en `ctrl1.ejecutar_control()` (V2, sin cambios),
    que ya exige todas las alertas resueltas (INCORRECTA sin ASIGNACION_
    CORRECTA sigue PENDIENTE), corrige la columna R del GLOBAL, actualiza el
    histórico y cierra. Un periodo ya cerrado (su GLOBAL ya figura en el
    histórico) responde YA_CERRADO sin escribir nada.

El periodo está CERRADO si y solo si el histórico maestro contiene filas de su
`archivo_global` (única fuente de verdad, la misma que usa V2): no hay un
segundo marcador que pueda desincronizarse.
"""

import datetime
import json
import os
import re
import sys
from decimal import Decimal, InvalidOperation

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import control_asignaciones as ctrl1  # noqa: E402  (reutilizado tal cual — V2, sin cambios)

PRELIMINAR = "preliminar"
CERRAR = "cerrar"

HOJA_NO_VIGENTES = "ALERTAS_NO_VIGENTES"
_COLUMNAS_NO_VIGENTES = list(ctrl1._COLUMNAS_REVISION_XLSX) + ["FECHA_RETIRO", "MOTIVO_RETIRO", "SHA256_GLOBAL_RETIRO"]
_CAMPOS_DECISION = ("VALIDACION_AUDITOR", "ASIGNACION_CORRECTA", "OBSERVACION_AUDITOR", "FECHA_VALIDACION")


# ---------------------------------------------------------------------------
# Modo
# ---------------------------------------------------------------------------

def validar_modo(modo, confirmacion_cierre=False):
    """PRELIMINAR por defecto. CERRAR exige una señal explícita doble:
    modo_control1="cerrar" Y confirmacion_cierre=True (que el frontend solo
    envía tras la confirmación humana). Nunca se infiere el cierre."""
    modo = PRELIMINAR if modo in (None, "") else modo
    if modo not in (PRELIMINAR, CERRAR):
        raise ValueError(f"MODO_CONTROL1_INVALIDO: {modo!r} (use 'preliminar' o 'cerrar')")
    if modo == CERRAR and confirmacion_cierre is not True:
        raise ValueError("CIERRE_SIN_CONFIRMACION: el cierre definitivo requiere confirmacion_cierre=true explícito")
    return modo


def periodo_cerrado(historico, nombre_archivo_global):
    """Cerrado ⇔ el histórico maestro ya tiene filas de este GLOBAL/periodo."""
    return any(f.get("archivo_global") == nombre_archivo_global for f in historico)


# ---------------------------------------------------------------------------
# Emparejamiento por identidad de la alerta (sobrevive a un GLOBAL regenerado)
# ---------------------------------------------------------------------------

_RE_FECHA_ISO = re.compile(r"^(\d{4}-\d{2}-\d{2})")


def _norm_fecha(valor):
    t = ctrl1._texto_o_vacio(valor).strip()
    m = _RE_FECHA_ISO.match(t)
    return m.group(1) if m else t


def _norm_importe(valor):
    t = ctrl1._texto_o_vacio(valor).strip()
    try:
        return str(Decimal(t).quantize(Decimal("0.01")))
    except (InvalidOperation, ValueError):
        return t


def _clave(fila):
    t = ctrl1._texto_o_vacio
    return (
        t(fila.get("ASIGNACION_ORIGINAL")).strip(), t(fila.get("CUENTA_MAYOR")).strip(),
        _norm_fecha(fila.get("FECHA_VALOR")), _norm_importe(fila.get("IMPORTE")), t(fila.get("GLOSA")).strip(),
    )


def _tiene_decision(fila):
    return any(ctrl1._texto_o_vacio(fila.get(c)).strip() for c in _CAMPOS_DECISION)


def fusionar_por_identidad(filas_nuevas, filas_vigentes_previas, filas_no_vigentes_previas, sha_actual, ahora):
    """Devuelve (filas_fusionadas, no_vigentes_actualizadas, traza).

    Cada fila NUEVA toma las decisiones de la fila previa con su misma
    identidad (vigentes primero; si no, una no vigente que reaparece). Las
    filas vigentes previas sin equivalente se RETIRAN (nunca se borran): pasan
    a no_vigentes con fecha y motivo. Las no vigentes que siguen sin
    reaparecer se conservan tal cual (el historial de retiros acumula)."""
    def agrupar(filas):
        grupos = {}
        for f in sorted(filas, key=lambda x: int(x.get("FILA_GLOBAL") or 0)):
            grupos.setdefault(_clave(f), []).append(f)
        return grupos

    vigentes = agrupar(filas_vigentes_previas)
    no_vigentes = agrupar(filas_no_vigentes_previas)

    fusionadas = []
    conservadas = nuevas = restauradas = 0
    for nueva in sorted(filas_nuevas, key=lambda x: int(x["FILA_GLOBAL"])):
        fila = dict(nueva)
        k = _clave(nueva)
        previa = None
        if vigentes.get(k):
            previa = vigentes[k].pop(0)
        elif no_vigentes.get(k):
            previa = no_vigentes[k].pop(0)
            restauradas += 1
        if previa is not None:
            for campo in _CAMPOS_DECISION:
                fila[campo] = previa.get(campo, "")
            if _tiene_decision(previa):
                conservadas += 1
        else:
            nuevas += 1
        fila["SHA256_GLOBAL"] = sha_actual
        fusionadas.append(fila)

    no_vigentes_actualizadas = [f for lista in no_vigentes.values() for f in lista]
    retiradas = 0
    for lista in vigentes.values():
        for f in lista:
            retirada = dict(f)
            retirada["FECHA_RETIRO"] = ahora
            retirada["MOTIVO_RETIRO"] = "YA_NO_PRESENTE_EN_GLOBAL_ACTUAL"
            retirada["SHA256_GLOBAL_RETIRO"] = sha_actual
            no_vigentes_actualizadas.append(retirada)
            retiradas += 1
    no_vigentes_actualizadas.sort(key=lambda f: (str(f.get("FECHA_RETIRO", "")), int(f.get("FILA_GLOBAL") or 0)))
    traza = {"decisiones_conservadas": conservadas, "alertas_nuevas": nuevas,
             "alertas_restauradas": restauradas, "alertas_retiradas": retiradas}
    return fusionadas, no_vigentes_actualizadas, traza


# ---------------------------------------------------------------------------
# Excel: hoja REVISION (V2) + hoja ALERTAS_NO_VIGENTES (V3)
# ---------------------------------------------------------------------------

def leer_no_vigentes(ruta_xlsx):
    if not ruta_xlsx or not os.path.isfile(ruta_xlsx):
        return []
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    try:
        if HOJA_NO_VIGENTES not in wb.sheetnames:
            return []
        ws = wb[HOJA_NO_VIGENTES]
        cab = [c.value for c in ws[1]]
        filas = []
        for celdas in ws.iter_rows(min_row=2):
            f = {}
            for i, c in enumerate(celdas):
                if i < len(cab) and cab[i]:
                    f[cab[i]] = ctrl1._texto_o_vacio(c.value)
            if not any(f.values()):
                continue
            try:
                f["FILA_GLOBAL"] = int(float(f.get("FILA_GLOBAL", "")))
            except (TypeError, ValueError):
                continue
            filas.append(f)
        return filas
    finally:
        wb.close()


def escribir_revision(ruta_xlsx, filas, no_vigentes):
    """Excel de revisión con el formato de V2 (`ctrl1.guardar_revision_xlsx`)
    + la hoja de trazabilidad ALERTAS_NO_VIGENTES cuando hay retiros."""
    ctrl1.guardar_revision_xlsx(ruta_xlsx, filas)
    agregar_hoja_no_vigentes(ruta_xlsx, no_vigentes)


def agregar_hoja_no_vigentes(ruta_xlsx, no_vigentes):
    if not no_vigentes or not os.path.isfile(ruta_xlsx):
        return
    wb = openpyxl.load_workbook(ruta_xlsx)
    try:
        if HOJA_NO_VIGENTES in wb.sheetnames:
            del wb[HOJA_NO_VIGENTES]
        ws = wb.create_sheet(HOJA_NO_VIGENTES)
        ws.append(_COLUMNAS_NO_VIGENTES)
        for celda in ws[1]:
            celda.font = openpyxl.styles.Font(bold=True)
        for f in no_vigentes:
            ws.append([f.get(c, "") for c in _COLUMNAS_NO_VIGENTES])
        ws.freeze_panes = "A2"
        tmp = f"{ruta_xlsx}.tmp"
        wb.save(tmp)
        os.replace(tmp, ruta_xlsx)
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Núcleo común: leer GLOBAL, alertas actuales y fusión con la revisión previa
# ---------------------------------------------------------------------------

def _preparar(ruta_global, ruta_historico, directorio_revision):
    """Devuelve (contexto | None, error | None). No escribe nada."""
    if not ruta_global or not os.path.isfile(ruta_global):
        return None, {"estado": "ERROR_TECNICO", "estado_control1": "ERROR", "problemas": ["GLOBAL_NO_ENCONTRADO"],
                      "ruta_global": ruta_global}
    nombre = os.path.basename(ruta_global)
    sha = ctrl1._hash_archivo(ruta_global)
    periodo = ctrl1._derivar_periodo(nombre)
    if periodo is None:
        return None, {"estado": "ERROR_TECNICO", "estado_control1": "ERROR", "problemas": ["GLOBAL_NOMBRE_NO_CANONICO"],
                      "archivo_global": nombre, "sha256_archivo": sha}
    historico = ctrl1.cargar_historico(ruta_historico)
    return {"nombre": nombre, "sha": sha, "periodo": periodo, "historico": historico,
            "ruta_xlsx": ctrl1.ruta_revision_xlsx(directorio_revision, periodo)}, None


def _alertas_actuales(ctx, ruta_global):
    try:
        partidas = ctrl1.leer_partidas_global(ruta_global)
    except ctrl1.HojaNoEncontradaError:
        return None, {"estado": "ERROR_TECNICO", "estado_control1": "ERROR", "problemas": ["GLOBAL_HOJA_1_NO_ENCONTRADA"],
                      "archivo_global": ctx["nombre"], "sha256_archivo": ctx["sha"]}
    except Exception as exc:  # noqa: BLE001 — GLOBAL ilegible: se reporta y se detiene
        return None, {"estado": "ERROR_TECNICO", "estado_control1": "ERROR", "problemas": [f"GLOBAL_ILEGIBLE:{exc}"],
                      "archivo_global": ctx["nombre"], "sha256_archivo": ctx["sha"]}
    candidatas, excluidas, sin_asig = [], 0, 0
    for p in partidas:
        if ctrl1.es_excluida(p["asignacion"], p["cuenta_mayor"]):
            excluidas += 1
        elif p["asignacion"] is None:
            sin_asig += 1
        else:
            candidatas.append(p)
    filas_nuevas = ctrl1.construir_filas_revision(candidatas, ctx["historico"], ctx["periodo"], ctx["sha"])
    return {"partidas": partidas, "candidatas": candidatas, "excluidas": excluidas, "sin_asignacion": sin_asig,
            "filas_nuevas": filas_nuevas}, None


def _contar(filas):
    res = ctrl1._resoluciones_por_fila(filas)
    pendientes = sum(1 for r in res.values() if not r["resuelta"])
    correctas = sum(1 for r in res.values() if r["validacion"] == "CORRECTA")
    incorrectas = sum(1 for r in res.values() if r["validacion"] == "INCORRECTA" and r["resuelta"])
    return pendientes, correctas, incorrectas


def _escribir_detalle(ruta, resumen, filas):
    if not ruta:
        return None
    os.makedirs(os.path.dirname(os.path.abspath(ruta)) or ".", exist_ok=True)
    alertas = [{"fila_global": f["FILA_GLOBAL"], "asignacion": f["ASIGNACION_ORIGINAL"], "tipo_alerta": f["TIPO_ALERTA"],
                "validacion_auditor": f.get("VALIDACION_AUDITOR", ""), "asignacion_correcta": f.get("ASIGNACION_CORRECTA", "")}
               for f in filas]
    with open(ruta, "w", encoding="utf-8") as fh:
        json.dump({**resumen, "alertas": alertas}, fh, ensure_ascii=False, indent=2, default=str)
    return ruta


# ---------------------------------------------------------------------------
# PRELIMINAR
# ---------------------------------------------------------------------------

def ejecutar_control1_preliminar(ruta_global, ruta_historico, directorio_revision,
                                 ruta_detalle_json=None, dry_run=False):
    """Revisión de mes abierto. Escribe SOLO la revisión (y el detalle) del
    periodo, en `directorio_revision`. No toca el histórico ni el GLOBAL."""
    ctx, error = _preparar(ruta_global, ruta_historico, directorio_revision)
    if error:
        return {**error, "modo_control1": PRELIMINAR}
    ahora = datetime.datetime.now().isoformat(timespec="seconds")

    if periodo_cerrado(ctx["historico"], ctx["nombre"]):
        return {"estado": "PERIODO_YA_CERRADO", "estado_control1": "YA_CERRADO", "modo_control1": PRELIMINAR,
                "periodo_cerrado": True, "archivo_global": ctx["nombre"], "periodo": ctx["periodo"],
                "sha256_archivo": ctx["sha"], "dry_run": dry_run, "historico_actualizado": False,
                "global_modificado": False, "revision_actualizada": False,
                "mensaje": "El periodo ya fue cerrado definitivamente: la revisión preliminar no se ejecuta ni modifica nada."}

    act, error = _alertas_actuales(ctx, ruta_global)
    if error:
        return {**error, "modo_control1": PRELIMINAR}

    previas = ctrl1.cargar_revision_xlsx(ctx["ruta_xlsx"])
    no_vigentes_previas = leer_no_vigentes(ctx["ruta_xlsx"])
    filas, no_vigentes, traza = fusionar_por_identidad(act["filas_nuevas"], previas, no_vigentes_previas, ctx["sha"], ahora)

    pendientes, correctas, incorrectas = _contar(filas)
    if not filas:
        estado, estado_control1, estado_validacion = "OK_SIN_DUPLICADOS", "PRELIMINAR_SIN_ALERTAS", None
    elif pendientes:
        estado, estado_control1, estado_validacion = "REVISAR_DUPLICADOS_ENCONTRADOS", "PRELIMINAR_PENDIENTE", "PENDIENTE_VALIDACION_AUDITOR"
    else:
        estado, estado_control1, estado_validacion = "REVISAR_DUPLICADOS_ENCONTRADOS", "PRELIMINAR_LISTO_PARA_CERRAR", "TODAS_VALIDADAS_PENDIENTE_CIERRE"

    revision_actualizada = False
    if not dry_run and (filas or no_vigentes or previas):
        escribir_revision(ctx["ruta_xlsx"], filas, no_vigentes)
        revision_actualizada = True

    resumen = {
        "estado": estado, "estado_control1": estado_control1, "estado_validacion": estado_validacion,
        "modo_control1": PRELIMINAR, "periodo_cerrado": False,
        "archivo_global": ctx["nombre"], "sha256_archivo": ctx["sha"],
        "sha256_global_original": ctx["sha"], "sha256_global_final": ctx["sha"],
        "global_modificado": False, "correcciones_aplicadas": 0,
        "periodo": ctx["periodo"], "fecha_ejecucion": ahora,
        "filas_leidas_global": len(act["partidas"]), "asignaciones_evaluadas": len(act["candidatas"]),
        "asignaciones_excluidas": act["excluidas"], "partidas_sin_asignacion": act["sin_asignacion"],
        "alertas_mismo_mes": sum(1 for f in filas if f["TIPO_ALERTA"] in ("DUPLICADA_MISMO_MES", "AMBAS")),
        "alertas_contra_historico": sum(1 for f in filas if f["TIPO_ALERTA"] in ("DUPLICADA_CON_HISTORICO", "AMBAS")),
        "filas_revisadas": len(filas), "filas_correctas": correctas, "filas_incorrectas": incorrectas,
        "filas_pendientes": pendientes, "alertas_pendientes_validacion": pendientes,
        **traza, "alertas_no_vigentes_total": len(no_vigentes),
        "dry_run": dry_run, "historico_actualizado": False, "filas_incorporadas_historico": 0,
        "ruta_revision": ctx["ruta_xlsx"] if revision_actualizada else None,
        "fuente_revision": "xlsx", "problemas_revision_json": [],
        "revision_actualizada": revision_actualizada, "nuevas_alertas_generadas": traza["alertas_nuevas"],
        "detalle_json": None,
    }
    resumen["detalle_json"] = _escribir_detalle(ruta_detalle_json, resumen, filas)
    return resumen


# ---------------------------------------------------------------------------
# CIERRE DEFINITIVO
# ---------------------------------------------------------------------------

def ejecutar_control1_cierre(ruta_global, ruta_historico, directorio_revision, ruta_detalle_json=None,
                             dry_run=False, ruta_revision_json=None):
    """Cierre definitivo del periodo. Reancla las decisiones previas al GLOBAL
    actual y delega el cierre real en `ctrl1.ejecutar_control()` (V2)."""
    ctx, error = _preparar(ruta_global, ruta_historico, directorio_revision)
    if error:
        return {**error, "modo_control1": CERRAR}

    if periodo_cerrado(ctx["historico"], ctx["nombre"]):
        misma_version = ctrl1.ya_procesado(ctx["historico"], ctx["sha"])
        return {
            "estado": "YA_PROCESADO_SIN_CAMBIOS" if misma_version else "GLOBAL_MODIFICADO_REQUIERE_REVISION",
            "estado_control1": "YA_CERRADO" if misma_version else "CIERRE_BLOQUEADO_GLOBAL_DISTINTO_DEL_CERRADO",
            "modo_control1": CERRAR, "periodo_cerrado": True, "archivo_global": ctx["nombre"], "periodo": ctx["periodo"],
            "sha256_archivo": ctx["sha"], "dry_run": dry_run, "historico_actualizado": False,
            "global_modificado": False, "revision_actualizada": False,
            "mensaje": ("El periodo ya está cerrado: no se hizo nada." if misma_version else
                        "El periodo ya está cerrado y el GLOBAL actual difiere del cerrado: requiere decisión humana; no se modificó nada."),
        }

    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    no_vigentes = []
    if not ruta_revision_json:
        act, error = _alertas_actuales(ctx, ruta_global)
        if error:
            return {**error, "modo_control1": CERRAR}
        previas = ctrl1.cargar_revision_xlsx(ctx["ruta_xlsx"])
        no_vigentes_previas = leer_no_vigentes(ctx["ruta_xlsx"])
        filas, no_vigentes, _traza = fusionar_por_identidad(act["filas_nuevas"], previas, no_vigentes_previas, ctx["sha"], ahora)
        if not dry_run and (filas or previas):
            # Reancla las decisiones al GLOBAL actual ANTES de que V2 decida.
            ctrl1.guardar_revision_xlsx(ctx["ruta_xlsx"], filas)

    r = ctrl1.ejecutar_control(ruta_global, ruta_historico, directorio_revision=directorio_revision,
                               dry_run=dry_run, ruta_detalle_json=ruta_detalle_json,
                               ruta_revision_json=ruta_revision_json)

    if not dry_run and no_vigentes and os.path.isfile(ctx["ruta_xlsx"]):
        agregar_hoja_no_vigentes(ctx["ruta_xlsx"], no_vigentes)  # V2 reescribe el xlsx: se repone la trazabilidad

    validacion = r.get("estado_validacion")
    if r.get("estado") in ("ERROR_TECNICO", "GLOBAL_MODIFICADO_REQUIERE_REVISION"):
        estado_control1 = "ERROR" if r.get("estado") == "ERROR_TECNICO" else "CIERRE_BLOQUEADO_GLOBAL_DISTINTO_DEL_CERRADO"
    elif r.get("estado") == "YA_PROCESADO_SIN_CAMBIOS":
        estado_control1 = "YA_CERRADO"
    elif dry_run:
        estado_control1 = "CIERRE_SIMULACRO"
    elif r.get("historico_actualizado") is True:
        estado_control1 = "CERRADO"
    elif validacion == "PENDIENTE_VALIDACION_AUDITOR":
        estado_control1 = "CIERRE_BLOQUEADO_PENDIENTES"
    else:
        estado_control1 = "CIERRE_BLOQUEADO"
    r["estado_control1"] = estado_control1
    r["modo_control1"] = CERRAR
    r["periodo_cerrado"] = estado_control1 in ("CERRADO", "YA_CERRADO")
    if estado_control1 == "CIERRE_BLOQUEADO_PENDIENTES":
        r["mensaje"] = ("No se cerró: hay alertas sin resolver (CORRECTA, o INCORRECTA con ASIGNACION_CORRECTA distinta "
                        f"de la original). Pendientes: {r.get('filas_pendientes')}. El histórico y el GLOBAL no se tocaron.")
    return r
