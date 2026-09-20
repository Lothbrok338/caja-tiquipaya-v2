"""v3/control3_modos.py — AUDITORÍA CxC / CxP (CONTROL 3): modo PRELIMINAR
(mes abierto) y modo CIERRE DEFINITIVO. FASE 12E.7 (2026-09-19).

QUÉ NO CAMBIA: las reglas contables. Este módulo solo agrega persistencia e
idempotencia por encima de `control_cxc_cxp.ejecutar_control()` (V2, sin
cambios): universo de 6 cuentas, llave CUENTA+ASIGNACION, saldo acumulado,
ABIERTO/CERRADO/REVISAR, cierre manual por observación del auditor, libro de
periodos PENDIENTE→APLICADO. CONTROL 3 sigue sin modificar jamás el GLOBAL.

POR QUÉ EXISTE: V2 acumula el periodo en el histórico y lo sella en el libro en
cuanto corre ("un solo cierre por periodo+SHA"). Con un mes abierto (GLOBAL que
se regenera al entrar más SAP) eso convertiría cada corrida en un cierre y
cualquier nuevo GLOBAL daría GLOBAL_MODIFICADO_REQUIERE_REVISION. Aquí:

  PRELIMINAR (por defecto). Ejecuta V2 sobre una COPIA de trabajo del histórico
    y del libro maestros (que solo contienen periodos ya cerrados) y descarta
    la copia. Solo escribe el reporte del periodo (xlsx + json). NUNCA toca
    HISTORICO_CXC_CXP.csv ni HISTORICO_CXC_CXP_PERIODOS.json maestros; se puede
    repetir cuantas veces haga falta. Las observaciones del auditor escritas en
    el reporte previo del periodo se conservan (se re-aplican al GLOBAL actual);
    las de una llave que ya no existe pasan a la hoja OBSERVACIONES_NO_VIGENTES.

  CERRAR. Acción explícita (modo_control3="cerrar" + confirmacion_cierre=True).
    Verifica en seco (V2 dry_run) que el cálculo es coherente y que las
    observaciones del auditor son válidas; solo entonces delega en V2, que
    actualiza el histórico maestro y sella el periodo APLICADO en el libro.
    Un periodo ya sellado responde YA_CERRADO sin escribir nada.

El periodo está CERRADO si y solo si el libro maestro tiene su entrada
APLICADO: única fuente de verdad (la misma de V2), sin un segundo marcador.

REQUISITOS PROPIOS DEL CIERRE: CONTROL 3 no tiene un flujo de "revisión
pendiente" como CONTROL 1: REVISAR/asignación faltante son hallazgos que el
auditor resuelve con OBSERVACION_AUDITOR (p. ej. CERRADO MANUALMENTE) y se
informan como `advertencias`, no bloquean. Bloquean el cierre: GLOBAL ilegible o
de nombre no canónico, observaciones del auditor inválidas (V2 las aplica
todo-o-nada; cerrar descartándolas perdería la memoria del auditor) y un
periodo ya sellado con otro GLOBAL.
"""

import datetime
import json
import os
import shutil
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import control_cxc_cxp as ctrl3  # noqa: E402  (reutilizado tal cual — V2, sin cambios)

PRELIMINAR = "preliminar"
CERRAR = "cerrar"

NOMBRE_HISTORICO = "HISTORICO_CXC_CXP.csv"
NOMBRE_PERIODOS = "HISTORICO_CXC_CXP_PERIODOS.json"
HOJA_NO_VIGENTES = "OBSERVACIONES_NO_VIGENTES"
_COLUMNAS_NO_VIGENTES = ["CUENTA", "ASIGNACION", "OBSERVACION_AUDITOR", "FECHA_RETIRO", "MOTIVO_RETIRO"]
_SCRATCH = "_preliminar_tmp"


def validar_modo(modo, confirmacion_cierre=False):
    """PRELIMINAR por defecto. CERRAR exige modo_control3="cerrar" Y
    confirmacion_cierre=True (la interfaz solo lo envía tras la confirmación
    humana). Nunca se infiere el cierre."""
    modo = PRELIMINAR if modo in (None, "") else modo
    if modo not in (PRELIMINAR, CERRAR):
        raise ValueError(f"MODO_CONTROL3_INVALIDO: {modo!r} (use 'preliminar' o 'cerrar')")
    if modo == CERRAR and confirmacion_cierre is not True:
        raise ValueError("ERROR_CONFIRMACION_CIERRE_REQUERIDA: el cierre definitivo requiere confirmacion_cierre=true explícito")
    return modo


def nombre_reporte(periodo, extension):
    return f"CONTROL_CXC_CXP_{periodo}.{extension}"


# ---------------------------------------------------------------------------
# Reporte previo del periodo (memoria del auditor) -> puente de observaciones
# ---------------------------------------------------------------------------

def leer_observaciones_xlsx(ruta_xlsx):
    """[(cuenta, asignacion, observacion_auditor)] de la hoja CONTROL del reporte
    previo (las filas 'ASIGNACION FALTANTE' no son llaves y se omiten)."""
    if not ruta_xlsx or not os.path.isfile(ruta_xlsx):
        return []
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    try:
        if ctrl3._HOJA_CONTROL not in wb.sheetnames:
            return []
        ws = wb[ctrl3._HOJA_CONTROL]
        cab = [c.value for c in ws[1]]
        if not all(c in cab for c in ("CUENTA", "ASIGNACION", "OBSERVACION_AUDITOR")):
            return []
        i_c, i_a, i_o = cab.index("CUENTA"), cab.index("ASIGNACION"), cab.index("OBSERVACION_AUDITOR")
        salida = []
        for fila in ws.iter_rows(min_row=2, values_only=True):
            cuenta = ctrl3._normalizar_cuenta(fila[i_c])
            asignacion = ctrl3._texto_celda(fila[i_a])
            if not cuenta or not asignacion or asignacion == ctrl3._ASIGNACION_FALTANTE_MARCADOR:
                continue
            obs = fila[i_o]
            salida.append((cuenta, asignacion, "" if obs is None else str(obs)))
        return salida
    finally:
        wb.close()


def leer_no_vigentes(ruta_xlsx):
    if not ruta_xlsx or not os.path.isfile(ruta_xlsx):
        return []
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    try:
        if HOJA_NO_VIGENTES not in wb.sheetnames:
            return []
        ws = wb[HOJA_NO_VIGENTES]
        cab = [c.value for c in ws[1]]
        salida = []
        for fila in ws.iter_rows(min_row=2, values_only=True):
            f = {cab[i]: ("" if v is None else str(v)) for i, v in enumerate(fila) if i < len(cab) and cab[i]}
            if f.get("CUENTA") and f.get("ASIGNACION"):
                salida.append(f)
        return salida
    finally:
        wb.close()


def agregar_hoja_no_vigentes(ruta_xlsx, no_vigentes):
    if not no_vigentes or not os.path.isfile(ruta_xlsx):
        return
    wb = openpyxl.load_workbook(ruta_xlsx)
    try:
        if HOJA_NO_VIGENTES in wb.sheetnames:
            del wb[HOJA_NO_VIGENTES]
        ws = wb.create_sheet(HOJA_NO_VIGENTES)
        ws.append(_COLUMNAS_NO_VIGENTES)
        for celda in ws[1]:
            celda.font = openpyxl.styles.Font(bold=True)
        for f in no_vigentes:
            ws.append([f.get(c, "") for c in _COLUMNAS_NO_VIGENTES])
        ws.freeze_panes = "A2"
        tmp = f"{ruta_xlsx}.tmp"
        wb.save(tmp)
        os.replace(tmp, ruta_xlsx)
    finally:
        wb.close()


def _claves_vigentes(ruta_global, historico):
    """Llaves que existirán tras acumular este GLOBAL: las del histórico maestro
    + las de sus partidas. None si el GLOBAL no se pudo leer (V2 lo reportará)."""
    try:
        candidatas, _faltantes = ctrl3.clasificar_partidas(ctrl3.leer_partidas_global(ruta_global))
        return set(historico.keys()) | set(ctrl3.agrupar_movimientos_mes(candidatas).keys())
    except Exception:  # noqa: BLE001 — se deja que ejecutar_control() informe el error técnico
        return None


def construir_puente(periodo, sha, obs_xlsx, no_vigentes_previas, claves, ahora):
    """(datos_puente | None, no_vigentes). Re-ancla las observaciones del reporte
    previo al SHA del GLOBAL actual. Las de llaves que ya no existen NO se
    borran: pasan a no_vigentes; las no vigentes que reaparecen (y cuya celda
    actual está vacía) se restauran."""
    if claves is None:
        return None, list(no_vigentes_previas)
    en_reporte = {(c, a): o for c, a, o in obs_xlsx}
    observaciones, no_vigentes, ya_retiradas = [], [], set()
    for previa in no_vigentes_previas:
        clave = (previa["CUENTA"], previa["ASIGNACION"])
        if clave in claves and not en_reporte.get(clave, "").strip():
            en_reporte[clave] = previa.get("OBSERVACION_AUDITOR", "")  # restaurada
        else:
            no_vigentes.append(previa)
            ya_retiradas.add((clave, previa.get("OBSERVACION_AUDITOR", "")))
    for (cuenta, asignacion), texto in en_reporte.items():
        if (cuenta, asignacion) in claves:
            observaciones.append({"cuenta": cuenta, "asignacion": asignacion, "observacion_auditor": texto})
        elif texto.strip() and (((cuenta, asignacion), texto) not in ya_retiradas):
            no_vigentes.append({"CUENTA": cuenta, "ASIGNACION": asignacion, "OBSERVACION_AUDITOR": texto,
                                "FECHA_RETIRO": ahora, "MOTIVO_RETIRO": "LLAVE_YA_NO_PRESENTE_EN_GLOBAL_ACTUAL"})
    datos = {"periodo": periodo, "sha256_global": sha, "observaciones": observaciones} if observaciones else None
    return datos, no_vigentes


# ---------------------------------------------------------------------------
# Contexto común
# ---------------------------------------------------------------------------

def _contexto(ruta_global, entrada, caja=None):
    if not ruta_global or not os.path.isfile(ruta_global):
        return None, {"estado": "ERROR_TECNICO", "estado_control3": "ERROR", "problemas": ["GLOBAL_NO_ENCONTRADO"],
                      "ruta_global": ruta_global}
    nombre = os.path.basename(ruta_global)
    sha = ctrl3._hash_archivo(ruta_global)
    periodo = ctrl3._derivar_periodo(nombre, caja)
    if periodo is None:
        return None, {"estado": "ERROR_TECNICO", "estado_control3": "ERROR", "problemas": ["GLOBAL_NOMBRE_NO_CANONICO"],
                      "archivo_global": nombre, "sha256_global": sha}
    ruta_hist = os.path.join(entrada, NOMBRE_HISTORICO)
    ruta_libro = os.path.join(entrada, NOMBRE_PERIODOS)
    libro = ctrl3.cargar_json(ruta_libro)
    return {
        "nombre": nombre, "sha": sha, "periodo": periodo, "ruta_hist": ruta_hist, "ruta_libro": ruta_libro,
        "libro": libro if isinstance(libro, dict) else {},
        "ruta_xlsx": os.path.join(entrada, nombre_reporte(periodo, "xlsx")),
        "ruta_json": os.path.join(entrada, nombre_reporte(periodo, "json")),
    }, None


def _advertencias(r):
    adv = []
    if r.get("revisar"):
        adv.append(f"{r['revisar']} llave(s) en REVISAR")
    if r.get("asignaciones_faltantes"):
        adv.append(f"{r['asignaciones_faltantes']} partida(s) sin asignación (no incorporadas al histórico)")
    return adv


def _reescribir_json(ruta, resumen):
    if ruta and os.path.isfile(ruta):
        with open(ruta, "w", encoding="utf-8") as f:
            json.dump(resumen, f, ensure_ascii=False, indent=2, default=str)


# ---------------------------------------------------------------------------
# PRELIMINAR
# ---------------------------------------------------------------------------

def ejecutar_control3_preliminar(ruta_global, directorio_entrada, dry_run=False, ruta_observaciones_json=None,
                                 caja=None):
    """Mes abierto. Escribe SOLO el reporte del periodo (xlsx + json) en
    `directorio_entrada`; el histórico y el libro maestros no se tocan.

    `caja` (config_cajas.CajaConfig o su `codigo`, por defecto TIQUIPAYA)
    decide el prefijo del nombre canónico exigido del GLOBAL."""
    ctx, error = _contexto(ruta_global, directorio_entrada, caja)
    if error:
        return {**error, "modo_control3": PRELIMINAR}
    base = {"modo_control3": PRELIMINAR, "archivo_global": ctx["nombre"], "periodo": ctx["periodo"],
            "sha256_global": ctx["sha"], "dry_run": dry_run,
            "historico_actualizado": False, "periodos_actualizado": False}

    registro = ctx["libro"].get(ctx["periodo"])
    if registro is not None:
        cerrado = registro.get("estado") == ctrl3._PERIODO_APLICADO
        return {**base, "estado": "PERIODO_YA_CERRADO" if cerrado else "CIERRE_PENDIENTE_DE_RECUPERAR",
                "estado_control3": "YA_CERRADO" if cerrado else "CIERRE_PENDIENTE_DE_RECUPERAR",
                "periodo_cerrado": cerrado, "archivo_control_xlsx": None, "archivo_control_json": None,
                "mensaje": ("El periodo ya fue cerrado definitivamente: la revisión preliminar no se ejecuta ni modifica nada."
                            if cerrado else
                            "Un cierre anterior quedó a medias: ejecute CERRAR AUDITORÍA para completarlo; el preliminar no se ejecuta.")}

    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    historico = ctrl3.cargar_historico(ctx["ruta_hist"])
    puente, no_vigentes = construir_puente(
        ctx["periodo"], ctx["sha"], leer_observaciones_xlsx(ctx["ruta_xlsx"]),
        leer_no_vigentes(ctx["ruta_xlsx"]), _claves_vigentes(ruta_global, historico), ahora)

    scratch = os.path.join(directorio_entrada, _SCRATCH)
    shutil.rmtree(scratch, ignore_errors=True)
    os.makedirs(scratch)
    try:
        for nombre in (NOMBRE_HISTORICO, NOMBRE_PERIODOS):
            origen = os.path.join(directorio_entrada, nombre)
            if os.path.isfile(origen):
                shutil.copyfile(origen, os.path.join(scratch, nombre))
        ruta_obs = ruta_observaciones_json
        if ruta_obs is None and puente is not None:
            ruta_obs = os.path.join(scratch, "observaciones_puente.json")
            with open(ruta_obs, "w", encoding="utf-8") as f:
                json.dump(puente, f, ensure_ascii=False)
        r = ctrl3.ejecutar_control(
            ruta_global, os.path.join(scratch, NOMBRE_HISTORICO),
            ruta_salida_xlsx=ctx["ruta_xlsx"], ruta_salida_json=ctx["ruta_json"],
            ruta_observaciones_json=ruta_obs, dry_run=dry_run, caja=caja)
    finally:
        shutil.rmtree(scratch, ignore_errors=True)

    r.update(base)
    r["periodo_cerrado"] = False
    r["historico_actualizado"] = False   # V2 lo reporta True sobre la copia de trabajo descartada
    r["advertencias"] = _advertencias(r)
    if r.get("estado") == "OK":
        r["estado_control3"] = "PRELIMINAR_SIMULACRO" if dry_run else "PRELIMINAR_OK"
        r["observaciones_no_vigentes_total"] = len(no_vigentes)
        if not dry_run:
            agregar_hoja_no_vigentes(ctx["ruta_xlsx"], no_vigentes)
    else:
        r["estado_control3"] = "ERROR" if r.get("estado") == "ERROR_TECNICO" else "PRELIMINAR_BLOQUEADO"
    _reescribir_json(r.get("archivo_control_json"), r)
    return r


# ---------------------------------------------------------------------------
# CIERRE DEFINITIVO
# ---------------------------------------------------------------------------

def ejecutar_control3_cierre(ruta_global, directorio_entrada, dry_run=False, ruta_observaciones_json=None,
                             caja=None):
    """Cierre definitivo del periodo sobre el histórico y el libro MAESTROS
    materializados en `directorio_entrada` (que el backend publica después).

    `caja` (config_cajas.CajaConfig o su `codigo`, por defecto TIQUIPAYA)
    decide el prefijo del nombre canónico exigido del GLOBAL."""
    ctx, error = _contexto(ruta_global, directorio_entrada, caja)
    if error:
        return {**error, "modo_control3": CERRAR}
    base = {"modo_control3": CERRAR, "archivo_global": ctx["nombre"], "periodo": ctx["periodo"],
            "sha256_global": ctx["sha"], "dry_run": dry_run,
            "historico_actualizado": False, "periodos_actualizado": False,
            "ruta_historico": ctx["ruta_hist"], "ruta_periodos": ctx["ruta_libro"]}

    registro = ctx["libro"].get(ctx["periodo"])
    if registro is not None and registro.get("estado") == ctrl3._PERIODO_APLICADO:
        misma = registro.get("sha256_global") == ctx["sha"]
        return {**base, "periodo_cerrado": True, "archivo_control_xlsx": None, "archivo_control_json": None,
                "estado": "YA_PROCESADO_SIN_CAMBIOS" if misma else "GLOBAL_MODIFICADO_REQUIERE_REVISION",
                "estado_control3": "YA_CERRADO" if misma else "CIERRE_BLOQUEADO_GLOBAL_DISTINTO_DEL_CERRADO",
                "mensaje": ("El periodo ya está cerrado: no se hizo nada." if misma else
                            "El periodo ya está cerrado y el GLOBAL actual difiere del cerrado: requiere decisión humana; no se modificó nada.")}

    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    historico = ctrl3.cargar_historico(ctx["ruta_hist"])

    if registro is None and ctrl3._historico_ya_refleja_periodo(historico, ctx["periodo"], ctx["sha"]):
        # Publicación interrumpida: el histórico maestro ya trae este periodo+SHA pero el libro no llegó a Drive.
        # Se sella el libro SIN reacumular (reacumular duplicaría los importes).
        if not dry_run:
            ctx["libro"][ctx["periodo"]] = {"sha256_global": ctx["sha"], "estado": ctrl3._PERIODO_APLICADO,
                                            "fecha_ejecucion": ahora, "recuperado": True}
            ctrl3._guardar_libro_periodos(ctx["ruta_hist"], ctx["libro"])
        return {**base, "estado": "OK", "estado_control3": "CIERRE_SIMULACRO" if dry_run else "CERRADO",
                "periodo_cerrado": not dry_run, "recuperado": True, "periodos_actualizado": not dry_run,
                "archivo_control_xlsx": None, "archivo_control_json": None,
                "mensaje": "El histórico ya reflejaba este periodo (publicación interrumpida): solo se selló el libro de periodos."}

    puente, no_vigentes = construir_puente(
        ctx["periodo"], ctx["sha"], leer_observaciones_xlsx(ctx["ruta_xlsx"]),
        leer_no_vigentes(ctx["ruta_xlsx"]), _claves_vigentes(ruta_global, historico), ahora)
    ruta_obs = ruta_observaciones_json
    tmp_obs = None
    if ruta_obs is None and puente is not None:
        tmp_obs = ctx["ruta_hist"] + ".observaciones_puente.json"
        with open(tmp_obs, "w", encoding="utf-8") as f:
            json.dump(puente, f, ensure_ascii=False)
        ruta_obs = tmp_obs
    try:
        # 1) Verificación en seco: nada se escribe.
        r0 = ctrl3.ejecutar_control(ruta_global, ctx["ruta_hist"], ruta_observaciones_json=ruta_obs, dry_run=True,
                                    caja=caja)
        if r0.get("estado") != "OK":
            estado_c3 = "ERROR" if r0.get("estado") == "ERROR_TECNICO" else "CIERRE_BLOQUEADO"
            return {**r0, **base, "estado_control3": estado_c3, "periodo_cerrado": False}
        if r0.get("problemas_observaciones_json"):
            return {**r0, **base, "estado_control3": "CIERRE_BLOQUEADO_OBSERVACIONES", "periodo_cerrado": False,
                    "mensaje": "No se cerró: las observaciones del auditor no son válidas (" +
                               ", ".join(r0["problemas_observaciones_json"]) + "). Histórico y libro no se tocaron."}
        if dry_run:
            r0.update(base)
            r0.update({"estado_control3": "CIERRE_SIMULACRO", "periodo_cerrado": False,
                       "advertencias": _advertencias(r0)})
            return r0
        # 2) Cierre real: V2 escribe PENDIENTE -> histórico -> APLICADO en el directorio materializado.
        r = ctrl3.ejecutar_control(ruta_global, ctx["ruta_hist"], ruta_salida_xlsx=ctx["ruta_xlsx"],
                                   ruta_salida_json=ctx["ruta_json"], ruta_observaciones_json=ruta_obs, caja=caja)
    finally:
        if tmp_obs and os.path.isfile(tmp_obs):
            os.remove(tmp_obs)

    cerrado = r.get("estado") == "OK" and r.get("historico_actualizado") is True
    r.update(base)
    r["historico_actualizado"] = cerrado
    r["periodos_actualizado"] = cerrado
    r["periodo_cerrado"] = cerrado
    r["estado_control3"] = "CERRADO" if cerrado else ("ERROR" if r.get("estado") == "ERROR_TECNICO" else "CIERRE_BLOQUEADO")
    r["advertencias"] = _advertencias(r)
    if cerrado:
        agregar_hoja_no_vigentes(ctx["ruta_xlsx"], no_vigentes)
        _reescribir_json(r.get("archivo_control_json"), r)
    return r
