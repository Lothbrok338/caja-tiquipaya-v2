"""tests_v3/test_global_institucional.py — GLOBAL INSTITUCIONAL
(v3/global_institucional.py).

Usa una plantilla SAP SINTÉTICA (tests/xlsx_fixtures.crear_plantilla_sap) y
GLOBAL TIQ/AME sintéticos construidos localmente en este archivo (nunca
datos contables reales, mismo patrón que tests/test_consolidador_mensual.py
para los SAP diarios). Verifica exclusivamente v3/global_institucional.py:
no repite ninguna regla de consolidador_mensual.py/control_asignaciones.py.

Uso: python -m unittest tests_v3.test_global_institucional -v
"""

import json
import os
import shutil
import sys
import tempfile
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import config_cajas as cfg
import consolidador_mensual as cm
import control_asignaciones as ctrl1
from tests.xlsx_fixtures import crear_plantilla_sap
from v3 import global_institucional as gi


# ---------------------------------------------------------------------------
# Fixture: SAP GLOBAL sintético mínimo válido (hoja "1", cabecera fija
# B10=BO01/C10=DB/H10=BOB/L10=<nombre_sap de la caja>, partidas cuadradas
# desde fila 16) — mismo layout que un SAP diario, referencia de cabecera
# distinta (la del GLOBAL, no la del diario).
# ---------------------------------------------------------------------------

def _crear_sap_global(ruta, partidas, caja):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1"

    ws["B10"] = "BO01"
    ws["C10"] = "DB"
    ws["H10"] = "BOB"
    ws["L10"] = caja.nombre_sap

    fila = 16
    for p in partidas:
        ws[f"B{fila}"] = p.get("sociedad", "BO01")
        ws[f"C{fila}"] = p.get("cuenta_mayor", "110101001")
        ws[f"D{fila}"] = p.get("texto_posicion")
        cargo = p.get("cargo", "0.00")
        haber = p.get("haber", "0.00")
        ws[f"E{fila}"] = Decimal(str(cargo)) if cargo is not None else None
        ws[f"F{fila}"] = Decimal(str(haber)) if haber is not None else None
        ws[f"L{fila}"] = p.get("centro_beneficio", "10010101")
        ws[f"R{fila}"] = p.get("asignacion")
        ws[f"U{fila}"] = p.get("xref1")
        ws[f"V{fila}"] = p.get("xref2")
        ws[f"W{fila}"] = p.get("xref3")
        fila += 1

    wb.save(ruta)


def _partida_par(cargo, haber="0.00", **extra):
    base = {"cargo": cargo, "haber": haber}
    base.update(extra)
    return base


def _dos_partidas_cuadradas(importe="100.00", asignacion="ASIG-1", **kwargs):
    return [
        _partida_par(importe, "0.00", cuenta_mayor="110101001",
                     texto_posicion="RECAUDACION", asignacion=asignacion,
                     xref1="X1", **kwargs),
        _partida_par("0.00", importe, cuenta_mayor="210201005",
                     texto_posicion="CONTRAPARTIDA", asignacion=asignacion),
    ]


class _GlobalInstitucionalTestBase(unittest.TestCase):
    ANIO = 2026
    MES = 9

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="global_institucional_test_")
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)

        self.ruta_plantilla = os.path.join(self.tmpdir, "Plantilla_SAP_maestra.xlsx")
        crear_plantilla_sap(self.ruta_plantilla)

        self.ruta_tiq = os.path.join(
            self.tmpdir, cm.nombre_sap_global(self.ANIO, self.MES, cfg.TIQUIPAYA)
        )
        self.ruta_ame = os.path.join(
            self.tmpdir, cm.nombre_sap_global(self.ANIO, self.MES, cfg.AMERICA)
        )
        self.ruta_salida = os.path.join(
            self.tmpdir, cm.nombre_sap_global(self.ANIO, self.MES, gi.CAJA_INSTITUCIONAL)
        )

        self.partidas_tiq = _dos_partidas_cuadradas("100.00", asignacion="TIQ-1")
        self.partidas_ame = (
            _dos_partidas_cuadradas("50.00", asignacion="AME-1")
            + _dos_partidas_cuadradas("75.00", asignacion="AME-2")
        )
        _crear_sap_global(self.ruta_tiq, self.partidas_tiq, cfg.TIQUIPAYA)
        _crear_sap_global(self.ruta_ame, self.partidas_ame, cfg.AMERICA)

    def _fusionar(self, **overrides):
        kwargs = dict(
            ruta_global_tiq=overrides.get("ruta_global_tiq", self.ruta_tiq),
            ruta_global_ame=overrides.get("ruta_global_ame", self.ruta_ame),
            anio=overrides.get("anio", self.ANIO),
            mes=overrides.get("mes", self.MES),
            ruta_plantilla=overrides.get("ruta_plantilla", self.ruta_plantilla),
            ruta_salida=overrides.get("ruta_salida", self.ruta_salida),
            ruta_mapa_origen=overrides.get("ruta_mapa_origen"),
            force=overrides.get("force", False),
        )
        return gi.fusionar_global_institucional(**kwargs)


class TestFusionBasica(_GlobalInstitucionalTestBase):

    def test_tiq_y_ame_se_fusionan(self):
        resultado = self._fusionar()
        self.assertTrue(os.path.isfile(resultado["ruta_global_institucional"]))
        self.assertEqual(resultado["cantidad_partidas_tiq"], 2)
        self.assertEqual(resultado["cantidad_partidas_ame"], 4)
        self.assertEqual(resultado["cantidad_partidas_total"], 6)

    def test_orden_deterministico_tiq_primero_ame_despues(self):
        resultado = self._fusionar()
        wb = openpyxl.load_workbook(resultado["ruta_global_institucional"], data_only=True)
        ws = wb["1"]
        asignaciones = []
        fila = cm._FILA_PRIMERA_PARTIDA
        while True:
            valor = ws[f"R{fila}"].value
            if valor in (None, "") and ws[f"C{fila}"].value in (None, ""):
                break
            asignaciones.append(valor)
            fila += 1
        wb.close()
        self.assertEqual(
            asignaciones,
            ["TIQ-1", "TIQ-1", "AME-1", "AME-1", "AME-2", "AME-2"],
        )

    def test_cantidad_de_filas_correcta(self):
        resultado = self._fusionar()
        wb = openpyxl.load_workbook(resultado["ruta_global_institucional"], data_only=True)
        ws = wb["1"]
        cuentas = []
        fila = cm._FILA_PRIMERA_PARTIDA
        while ws[f"C{fila}"].value not in (None, ""):
            cuentas.append(ws[f"C{fila}"].value)
            fila += 1
        wb.close()
        self.assertEqual(len(cuentas), 6)

    def test_columnas_trazabilidad_caja_archivo_fila(self):
        resultado = self._fusionar()
        wb = openpyxl.load_workbook(resultado["ruta_global_institucional"], data_only=True)
        ws = wb["1"]

        filas = []
        for offset in range(6):
            fila = cm._FILA_PRIMERA_PARTIDA + offset
            filas.append((
                ws[f"{gi._COL_CAJA}{fila}"].value,
                ws[f"{gi._COL_ARCHIVO_ORIGEN}{fila}"].value,
                ws[f"{gi._COL_FILA_ORIGEN}{fila}"].value,
            ))
        wb.close()

        nombre_tiq = os.path.basename(self.ruta_tiq)
        nombre_ame = os.path.basename(self.ruta_ame)
        self.assertEqual(filas, [
            ("tiquipaya", nombre_tiq, 16),
            ("tiquipaya", nombre_tiq, 17),
            ("america", nombre_ame, 16),
            ("america", nombre_ame, 17),
            ("america", nombre_ame, 18),
            ("america", nombre_ame, 19),
        ])

    def test_posiciones_control1_control3_no_alteradas(self):
        """B..W (lo que ya escribe escribir_sap_global) queda exactamente
        igual que en un GLOBAL normal de consolidador_mensual.py; la
        trazabilidad va SOLO en AA/AB/AC."""
        resultado = self._fusionar()
        wb = openpyxl.load_workbook(resultado["ruta_global_institucional"], data_only=True)
        ws = wb["1"]
        primera = {
            "cuenta_mayor": ws["C16"].value,
            "texto_posicion": ws["D16"].value,
            "cargo": ws["E16"].value,
            "haber": ws["F16"].value,
            "asignacion": ws["R16"].value,
            "xref1": ws["U16"].value,
        }
        wb.close()
        self.assertEqual(primera["cuenta_mayor"], "110101001")
        self.assertEqual(primera["texto_posicion"], "RECAUDACION")
        self.assertEqual(primera["asignacion"], "TIQ-1")
        self.assertEqual(primera["xref1"], "X1")


class TestSidecarMapaOrigen(_GlobalInstitucionalTestBase):

    def test_sidecar_correcto(self):
        resultado = self._fusionar()
        with open(resultado["ruta_mapa_origen"], encoding="utf-8") as f:
            mapa = json.load(f)

        nombre_tiq = os.path.basename(self.ruta_tiq)
        nombre_ame = os.path.basename(self.ruta_ame)

        self.assertEqual(mapa["16"], {"caja": "tiquipaya", "archivo_origen": nombre_tiq, "fila_origen": 16})
        self.assertEqual(mapa["17"], {"caja": "tiquipaya", "archivo_origen": nombre_tiq, "fila_origen": 17})
        self.assertEqual(mapa["18"], {"caja": "america", "archivo_origen": nombre_ame, "fila_origen": 16})
        self.assertEqual(mapa["21"], {"caja": "america", "archivo_origen": nombre_ame, "fila_origen": 19})
        self.assertEqual(len(mapa), 6)

    def test_nombre_sidecar_canonico(self):
        resultado = self._fusionar()
        self.assertEqual(
            os.path.basename(resultado["ruta_mapa_origen"]),
            f"MAPA_ORIGEN_INSTITUCIONAL_SEPTIEMBRE_{self.ANIO}.json",
        )


class TestDeterminismo(_GlobalInstitucionalTestBase):

    def test_mismo_input_mismo_sha(self):
        resultado1 = self._fusionar()
        sha1 = resultado1["sha256_global_institucional"]

        os.remove(self.ruta_salida)
        os.remove(resultado1["ruta_mapa_origen"])

        resultado2 = self._fusionar()
        sha2 = resultado2["sha256_global_institucional"]

        self.assertEqual(sha1, sha2)
        self.assertEqual(cm._sha256_archivo(self.ruta_salida), sha1)

    def test_mismo_input_mismo_mapa_origen(self):
        resultado1 = self._fusionar()
        with open(resultado1["ruta_mapa_origen"], encoding="utf-8") as f:
            contenido1 = f.read()

        os.remove(self.ruta_salida)
        os.remove(resultado1["ruta_mapa_origen"])

        resultado2 = self._fusionar()
        with open(resultado2["ruta_mapa_origen"], encoding="utf-8") as f:
            contenido2 = f.read()

        self.assertEqual(contenido1, contenido2)

    def test_force_permite_regenerar_con_mismo_resultado(self):
        resultado1 = self._fusionar()
        resultado2 = self._fusionar(force=True)
        self.assertEqual(
            resultado1["sha256_global_institucional"],
            resultado2["sha256_global_institucional"],
        )


class TestOrigenesSoloLectura(_GlobalInstitucionalTestBase):

    def test_no_modifica_los_dos_global_origen(self):
        hash_tiq_antes = cm._sha256_archivo(self.ruta_tiq)
        hash_ame_antes = cm._sha256_archivo(self.ruta_ame)

        self._fusionar()

        self.assertEqual(cm._sha256_archivo(self.ruta_tiq), hash_tiq_antes)
        self.assertEqual(cm._sha256_archivo(self.ruta_ame), hash_ame_antes)

    def test_no_modifica_la_plantilla(self):
        hash_plantilla_antes = cm._sha256_archivo(self.ruta_plantilla)
        self._fusionar()
        self.assertEqual(cm._sha256_archivo(self.ruta_plantilla), hash_plantilla_antes)


class TestLectoresControl1Control3(_GlobalInstitucionalTestBase):

    def test_control1_deriva_periodo_con_pseudo_caja_institucional(self):
        resultado = self._fusionar()
        nombre = os.path.basename(resultado["ruta_global_institucional"])
        periodo = ctrl1._derivar_periodo(nombre, gi.CAJA_INSTITUCIONAL)
        self.assertEqual(periodo, f"SEPTIEMBRE_{self.ANIO}")

    def test_control1_no_deriva_periodo_sin_pseudo_caja(self):
        resultado = self._fusionar()
        nombre = os.path.basename(resultado["ruta_global_institucional"])
        # Sin pasar la caja institucional, el nombre no matchea el patrón
        # de TIQUIPAYA (default) ni el de AMERICA: sigue fallando cerrado,
        # tal como antes de que este módulo existiera.
        self.assertIsNone(ctrl1._derivar_periodo(nombre, None))
        self.assertIsNone(ctrl1._derivar_periodo(nombre, cfg.AMERICA))

    def test_control1_lee_partidas_institucional(self):
        resultado = self._fusionar()
        partidas = ctrl1.leer_partidas_global(resultado["ruta_global_institucional"])
        self.assertEqual(len(partidas), 6)
        self.assertEqual(
            [p["asignacion"] for p in partidas],
            ["TIQ-1", "TIQ-1", "AME-1", "AME-1", "AME-2", "AME-2"],
        )

    def test_pseudo_caja_institucional_no_registrada_globalmente(self):
        self.assertNotIn("institucional", cfg.CAJAS)
        with self.assertRaises(ValueError):
            cfg.resolver_caja("institucional")


class TestValidacionesFailClosed(_GlobalInstitucionalTestBase):

    def test_falta_tiq(self):
        os.remove(self.ruta_tiq)
        with self.assertRaisesRegex(RuntimeError, "GLOBAL_TIQ_NO_ENCONTRADO"):
            self._fusionar()

    def test_falta_ame(self):
        os.remove(self.ruta_ame)
        with self.assertRaisesRegex(RuntimeError, "GLOBAL_AME_NO_ENCONTRADO"):
            self._fusionar()

    def test_nombre_tiq_incorrecto(self):
        ruta_mala = os.path.join(self.tmpdir, "SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        shutil.copyfile(self.ruta_tiq, ruta_mala)
        with self.assertRaisesRegex(RuntimeError, "NOMBRE_GLOBAL_TIQ_NO_CANONICO"):
            self._fusionar(ruta_global_tiq=ruta_mala)

    def test_nombre_ame_incorrecto(self):
        ruta_mala = os.path.join(self.tmpdir, "SAP_GLOBAL_AME_AGOSTO_2026.xlsx")
        shutil.copyfile(self.ruta_ame, ruta_mala)
        with self.assertRaisesRegex(RuntimeError, "NOMBRE_GLOBAL_AME_NO_CANONICO"):
            self._fusionar(ruta_global_ame=ruta_mala)

    def test_nombre_salida_incorrecto(self):
        ruta_mala = os.path.join(self.tmpdir, "SAP_GLOBAL_INSTITUCIONAL.xlsx")
        with self.assertRaisesRegex(RuntimeError, "NOMBRE_SALIDA_INSTITUCIONAL_NO_CANONICO"):
            self._fusionar(ruta_salida=ruta_mala)

    def test_global_tiq_corrupto(self):
        _crear_sap_global(self.ruta_tiq, self.partidas_tiq, cfg.AMERICA)  # L10 equivocado
        with self.assertRaisesRegex(RuntimeError, "GLOBAL_TIQ_CORRUPTO"):
            self._fusionar()

    def test_global_ame_corrupto_partidas_descuadradas(self):
        partidas_descuadradas = [_partida_par("100.00", "0.00", cuenta_mayor="110101001")]
        _crear_sap_global(self.ruta_ame, partidas_descuadradas, cfg.AMERICA)
        with self.assertRaisesRegex(RuntimeError, "GLOBAL_AME_CORRUPTO"):
            self._fusionar()

    def test_origen_ambiguo_mismo_archivo(self):
        with self.assertRaisesRegex(RuntimeError, "ORIGEN_AMBIGUO"):
            self._fusionar(ruta_global_ame=self.ruta_tiq)

    def test_plantilla_no_encontrada(self):
        with self.assertRaisesRegex(RuntimeError, "PLANTILLA_NO_ENCONTRADA"):
            self._fusionar(ruta_plantilla=os.path.join(self.tmpdir, "no_existe.xlsx"))

    def test_salida_ya_existe_sin_force(self):
        self._fusionar()
        with self.assertRaisesRegex(RuntimeError, "SALIDA_YA_EXISTE_SIN_FORCE"):
            self._fusionar()

    def test_falla_cerrado_no_escribe_nada_parcial(self):
        partidas_descuadradas = [_partida_par("100.00", "0.00", cuenta_mayor="110101001")]
        _crear_sap_global(self.ruta_ame, partidas_descuadradas, cfg.AMERICA)
        with self.assertRaises(RuntimeError):
            self._fusionar()
        self.assertFalse(os.path.exists(self.ruta_salida))


if __name__ == "__main__":
    unittest.main()
