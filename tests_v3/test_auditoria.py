"""tests_v3/test_auditoria.py — pruebas propias del módulo 07 · AUDITORIA
de V3 (v3/auditoria.py). Este módulo es un ADAPTADOR de solo-lectura +
consolidación: NUNCA decide, NUNCA cambia diferencia/bloqueadores/
clasificación/correcciones/SAP/publicación — solo observa lo que los
módulos 01-06 ya decidieron y lo deja trazable. CONTROL 1 y CONTROL 3 se
delegan tal cual a control_asignaciones.py/control_cxc_cxp.py — NO se
reimplementan.

Uso: python -m pytest tests_v3/test_auditoria.py -q
"""

import glob
import hashlib
import json
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import control_asignaciones as ctrl1  # noqa: E402
import control_cxc_cxp as ctrl3  # noqa: E402
from v3.auditoria import (  # noqa: E402
    consolidar_auditoria_cierre, consolidar_auditoria_lote,
    ejecutar_control1_correccion, evaluar_control3_idempotencia,
)


def _sha256(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 16), b""):
            h.update(bloque)
    return h.hexdigest()


def _item_completo(**overrides):
    base = {
        "fecha": "2026-09-01", "archivo_esperado": "CIERRE 01-09-2026.xlsm",
        "estado_ingesta": "ENCONTRADO", "estado_materializacion": "MATERIALIZADO",
        "estado_motor": "PROCESADO", "estado_final": "LISTO_PARA_PUBLICAR",
        "requiere_revision": False, "correccion_aplicada": None, "resultado_reproceso": None,
        "estado_publicacion": "PUBLICADO", "publicado": True,
        "ruta_resultado": "/dev/resultado.json", "ruta_sap": "/dev/sap.xlsx",
        "ruta_marker": "/dev/PROCESADO_abc.json", "sha256": "a" * 64,
        "mensaje": "Cierre publicado en DEV.",
    }
    base.update(overrides)
    return base


def _crear_global_minimo(ruta, asignacion_r16="VIEJA_ASIG"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1"
    ws["C16"] = "110201002"
    ws["D16"] = "GLOSA DEMO"
    ws["E16"] = 100
    ws["F16"] = None
    ws["O16"] = "2026-09-01"
    ws["R16"] = asignacion_r16
    wb.save(ruta)


# 1) genera registro de auditoria por cierre
def test_genera_registro_de_auditoria_por_cierre(tmp_path):
    item = _item_completo()
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert os.path.isfile(r["ruta_auditoria"])
    assert os.path.basename(r["ruta_auditoria"]) == f"AUDITORIA_2026-09-01_{item['sha256']}.json"
    with open(r["ruta_auditoria"], "r", encoding="utf-8") as f:
        contenido = json.load(f)
    assert contenido["fecha"] == "2026-09-01"


# 2) conserva SHA256
def test_conserva_sha256(tmp_path):
    item = _item_completo(sha256="f" * 64)
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["sha256"] == "f" * 64


# 3) conserva estado de ingesta
def test_conserva_estado_ingesta(tmp_path):
    item = _item_completo(estado_ingesta="SIN_ARCHIVO")
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["estado_ingesta"] == "SIN_ARCHIVO"


# 4) conserva estado de motor
def test_conserva_estado_motor(tmp_path):
    item = _item_completo(estado_motor="ERROR_MOTOR")
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["estado_motor"] == "ERROR_MOTOR"


# 5) conserva clasificacion (estado_final/requiere_revision)
def test_conserva_clasificacion(tmp_path):
    item = _item_completo(estado_final="ERROR_REVISAR", requiere_revision=True)
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["estado_final"] == "ERROR_REVISAR"
    assert r["requiere_revision"] is True


# 6) conserva correcciones
def test_conserva_correcciones(tmp_path):
    item = _item_completo(correccion_aplicada=True, resultado_reproceso="LISTO_PARA_PUBLICAR")
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["correccion_aplicada"] is True
    assert r["resultado_reproceso"] == "LISTO_PARA_PUBLICAR"


# 7) conserva publicacion
def test_conserva_publicacion(tmp_path):
    item = _item_completo(estado_publicacion="YA_PUBLICADO", publicado=False)
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["estado_publicacion"] == "YA_PUBLICADO"
    assert r["publicado"] is False


# 8) registra usuario auditor
def test_registra_usuario_auditor(tmp_path):
    item = _item_completo()
    item.pop("usuario_auditor", None)
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"), usuario_auditor="auditor.dev")
    assert r["usuario_auditor"] == "auditor.dev"

    # el usuario dentro del item (si viene, p. ej. de Modulo 06) tiene prioridad
    item2 = _item_completo(usuario_auditor="ana.torrico")
    r2 = consolidar_auditoria_cierre(item2, str(tmp_path / "dev"), usuario_auditor="otro")
    assert r2["usuario_auditor"] == "ana.torrico"


# 9) registra rutas DEV
def test_registra_rutas_dev(tmp_path):
    item = _item_completo(ruta_resultado="/dev/x/resultado.json", ruta_sap="/dev/x/sap.xlsx",
                           ruta_marker="/dev/x/PROCESADO_x.json")
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["ruta_resultado"] == "/dev/x/resultado.json"
    assert r["ruta_sap"] == "/dev/x/sap.xlsx"
    assert r["ruta_marker"] == "/dev/x/PROCESADO_x.json"


# 10) registra errores sin alterar el resultado
def test_registra_errores_sin_alterar_resultado(tmp_path):
    item = _item_completo(estado_final="ERROR_TECNICO", mensaje="MOTOR_EXCEPCION: fallo tecnico X")
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["estado_final"] == "ERROR_TECNICO"  # NUNCA se reinterpreta ni se "arregla"
    assert "MOTOR_EXCEPCION" in r["mensajes"][0]


# 11) CONTROL 1 mantiene comportamiento V2 (todo-o-nada, solo columna R)
def test_control1_mantiene_comportamiento_v2(tmp_path):
    ruta_global = str(tmp_path / "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
    _crear_global_minimo(ruta_global, asignacion_r16="VIEJA_ASIG")
    hash_antes = _sha256(ruta_global)

    with pytest.raises(ctrl1.CorreccionInvalidaError):
        ejecutar_control1_correccion(ruta_global, [(16, "ASIGNACION_DESACTUALIZADA", "NUEVA_ASIG")])
    assert _sha256(ruta_global) == hash_antes  # todo-o-nada: nada se toco

    ejecutar_control1_correccion(ruta_global, [(16, "VIEJA_ASIG", "NUEVA_ASIG")])
    partidas = ctrl1.leer_partidas_global(ruta_global)
    assert partidas[0]["asignacion"] == "NUEVA_ASIG"
    assert partidas[0]["cuenta_mayor"] == "110201002"  # ninguna otra celda cambio


# 12) CONTROL 3 mantiene comportamiento V2 (idempotencia periodo+SHA)
def test_control3_mantiene_comportamiento_v2():
    historico_dict = {
        "110201002|3P66536982": {
            "periodo_ultimo_movimiento": "SEPTIEMBRE_2026", "sha256_global_ultimo": "sha_septiembre",
        }
    }
    libro_periodos = {"AGOSTO_2026": {"sha256_global": "sha_agosto", "estado": "APLICADO"}}

    assert evaluar_control3_idempotencia(historico_dict, libro_periodos, "AGOSTO_2026", "sha_agosto") \
        == "YA_PROCESADO_SIN_CAMBIOS"
    assert evaluar_control3_idempotencia(historico_dict, libro_periodos, "AGOSTO_2026", "sha_diferente") \
        == "GLOBAL_MODIFICADO_REQUIERE_REVISION"


# 13) auditoria no modifica SAP
def test_auditoria_no_modifica_sap(tmp_path):
    ruta_sap = tmp_path / "sap.xlsx"
    openpyxl.Workbook().save(str(ruta_sap))
    hash_antes = _sha256(str(ruta_sap))
    item = _item_completo(ruta_sap=str(ruta_sap))
    consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert _sha256(str(ruta_sap)) == hash_antes


# 14) auditoria no modifica resultado
def test_auditoria_no_modifica_resultado(tmp_path):
    ruta_resultado = tmp_path / "resultado.json"
    ruta_resultado.write_text(json.dumps({"diferencia": "0.00", "blockers": 0}), encoding="utf-8")
    hash_antes = _sha256(str(ruta_resultado))
    item = _item_completo(ruta_resultado=str(ruta_resultado))
    consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert _sha256(str(ruta_resultado)) == hash_antes
    assert json.loads(ruta_resultado.read_text(encoding="utf-8")) == {"diferencia": "0.00", "blockers": 0}


# 15) auditoria no modifica el original (.xlsm)
def test_auditoria_no_modifica_original(tmp_path):
    ruta_cierre = tmp_path / "CIERRE 01-09-2026.xlsm"
    openpyxl.Workbook().save(str(ruta_cierre))
    hash_antes = _sha256(str(ruta_cierre))
    item = _item_completo(ruta_cierre_local=str(ruta_cierre))
    item.pop("sha256")  # fuerza a que auditoria calcule el hash LEYENDO el original
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    assert r["sha256"] == hash_antes  # se leyo correctamente
    assert _sha256(str(ruta_cierre)) == hash_antes  # y nunca se toco


# 16) auditoria no publica nada
def test_auditoria_no_publica_nada(tmp_path):
    base_dir_dev = str(tmp_path / "dev")
    item = _item_completo(estado_final="LISTO_PARA_PUBLICAR", estado_publicacion=None, publicado=None)
    consolidar_auditoria_cierre(item, base_dir_dev)
    assert glob.glob(os.path.join(base_dir_dev, "**", "PROCESADO_*.json"), recursive=True) == []
    assert glob.glob(os.path.join(base_dir_dev, "publicacion", "**"), recursive=True) == []


# 17) fallo de auditoria de un cierre no altera otros cierres
def test_fallo_de_un_cierre_no_altera_otros(tmp_path):
    item_ok = _item_completo(fecha="2026-09-01")
    # Un byte nulo en el nombre de archivo hace que open() falle con
    # ValueError SOLO al escribir el JSON de ESTE cierre (el directorio
    # compartido auditoria/ ya se crea sin problema antes de llegar ahi).
    item_roto = _item_completo(fecha="2026-09-02\x00roto")

    resultado = consolidar_auditoria_lote([item_ok, item_roto], str(tmp_path / "dev"))
    assert resultado["total_cierres"] == 2  # el lote se completo igual, con ambos aislados
    ok = next(r for r in resultado["cierres"] if r.get("fecha") == "2026-09-01")
    roto = next(r for r in resultado["cierres"] if "error_auditoria" in r)
    assert "ruta_auditoria" in ok and "error_auditoria" not in ok
    assert "ValueError" in roto["error_auditoria"]
    # y el cierre OK si quedo escrito en disco, sin verse afectado por el otro
    assert os.path.isfile(ok["ruta_auditoria"])


# 18) consolidado de lote contiene todos los cierres
def test_consolidado_de_lote_contiene_todos_los_cierres(tmp_path):
    cierres = [_item_completo(fecha=f"2026-09-0{n}", sha256=f"{n}" * 64) for n in range(1, 4)]
    resultado = consolidar_auditoria_lote(cierres, str(tmp_path / "dev"))
    assert os.path.isfile(resultado["ruta_lote"])
    with open(resultado["ruta_lote"], "r", encoding="utf-8") as f:
        consolidado = json.load(f)
    assert consolidado["total_cierres"] == 3
    assert {c["fecha"] for c in consolidado["cierres"]} == {"2026-09-01", "2026-09-02", "2026-09-03"}


def test_controles_ejecutados_se_registran_en_el_lote_sin_reejecutarlos(tmp_path):
    item = _item_completo()
    resultado_c1 = {"estado": "APLICADO", "filas_actualizadas": 1}
    resultado_c3 = "YA_PROCESADO_SIN_CAMBIOS"
    resultado = consolidar_auditoria_lote(
        [item], str(tmp_path / "dev"), control1_resultado=resultado_c1, control3_resultado=resultado_c3,
    )
    with open(resultado["ruta_lote"], "r", encoding="utf-8") as f:
        consolidado = json.load(f)
    assert consolidado["controles_ejecutados"] == [
        {"control": "CONTROL_1", "resultado": resultado_c1},
        {"control": "CONTROL_3", "resultado": resultado_c3},
    ]


def test_todas_las_filas_tienen_al_menos_las_claves_pedidas(tmp_path):
    item = _item_completo()
    r = consolidar_auditoria_cierre(item, str(tmp_path / "dev"))
    esperado_minimo = {
        "fecha", "archivo_esperado", "sha256", "estado_ingesta", "estado_materializacion",
        "estado_motor", "estado_final", "requiere_revision", "correccion_aplicada",
        "resultado_reproceso", "estado_publicacion", "publicado", "ruta_resultado",
        "ruta_sap", "ruta_marker", "usuario_auditor", "mensajes", "timestamp_auditoria",
    }
    assert esperado_minimo.issubset(set(r.keys()))
