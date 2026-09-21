"""tests_v3/test_control3_institucional.py — CONTROL 3 institucional
(v3/control3_institucional.py): combina SAP_GLOBAL_TIQ_... y
SAP_GLOBAL_AME_... SOLO EN MEMORIA por la llave CUENTA+ASIGNACION ya
existente en control_cxc_cxp.py. Nunca modifica ningún GLOBAL.

Uso: python -m pytest tests_v3/test_control3_institucional.py -q
"""
import os
import shutil
import sys
import tempfile
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import config_cajas as cfg
import consolidador_mensual
import control_cxc_cxp as c3
from v3 import control3_institucional as c3i

ANIO, MES = 2026, 9

# Cuentas controladas por control_cxc_cxp: CxC UNIDADES / CxP UNIDADES.
CTA_CXC = "110201002"
CTA_CXP = "210103002"


def _crear_global(ruta, partidas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1"
    fila = 16
    for p in partidas:
        ws[f"C{fila}"] = p.get("cuenta")
        ws[f"D{fila}"] = p.get("glosa", "GLOSA")
        if p.get("debe") is not None:
            ws[f"E{fila}"] = Decimal(str(p["debe"]))
        if p.get("haber") is not None:
            ws[f"F{fila}"] = Decimal(str(p["haber"]))
        ws[f"R{fila}"] = p.get("asignacion")
        fila += 1
    wb.save(ruta)


def _partida(cuenta, asignacion, debe=None, haber=None):
    return {"cuenta": cuenta, "asignacion": asignacion, "debe": debe, "haber": haber}


class BaseInstitucional(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.ruta_tiq = os.path.join(self.tmp, consolidador_mensual.nombre_sap_global(ANIO, MES, cfg.TIQUIPAYA))
        self.ruta_ame = os.path.join(self.tmp, consolidador_mensual.nombre_sap_global(ANIO, MES, cfg.AMERICA))
        self.ruta_historico = os.path.join(self.tmp, c3i.nombre_historico_institucional())

    def _ejecutar(self, tiq_partidas, ame_partidas):
        _crear_global(self.ruta_tiq, tiq_partidas)
        _crear_global(self.ruta_ame, ame_partidas)
        return c3i.ejecutar_control3_institucional(self.ruta_tiq, self.ruta_ame, self.ruta_historico)


class TestSoloUnaCaja(BaseInstitucional):
    def test_solo_tiq(self):
        r = self._ejecutar([_partida(CTA_CXC, "A1", debe="100.00")], [])
        self.assertEqual(r["estado"], "OK")
        self.assertEqual(r["llaves_evaluadas"], 1)
        self.assertEqual(r["abiertas"], 1)

    def test_solo_ame(self):
        r = self._ejecutar([], [_partida(CTA_CXP, "B1", haber="50.00")])
        self.assertEqual(r["estado"], "OK")
        self.assertEqual(r["llaves_evaluadas"], 1)
        self.assertEqual(r["abiertas"], 1)


class TestCompensacionInstitucional(BaseInstitucional):
    def test_tiq_mas_ame_misma_llave_se_combinan(self):
        r = self._ejecutar(
            [_partida(CTA_CXP, "MISMA", debe="0.00", haber=None)],
            [_partida(CTA_CXP, "MISMA", haber="0.00")],
        )
        self.assertEqual(r["estado"], "OK")
        self.assertEqual(r["llaves_evaluadas"], 1, "una sola llave institucional CUENTA+ASIGNACION")

    def test_obligacion_ame_compensada_por_tiq(self):
        # AME HABER 100 de una CxP; TIQ DEBE 100 de la misma CUENTA+ASIGNACION
        # -> saldo institucional final = 0 (CERRADO).
        r = self._ejecutar(
            [_partida(CTA_CXP, "COMPENSA", debe="100.00")],
            [_partida(CTA_CXP, "COMPENSA", haber="100.00")],
        )
        self.assertEqual(r["estado"], "OK")
        self.assertEqual(r["cerradas"], 1)
        self.assertEqual(r["abiertas"], 0)
        historico = c3.cargar_historico(self.ruta_historico)
        fila = historico[(CTA_CXP, "COMPENSA")]
        self.assertEqual(Decimal(fila["saldo"]), Decimal("0.00"))
        self.assertEqual(fila["estado"], c3._ESTADO_CERRADO)

    def test_trazabilidad_origen_incluye_ambas_cajas(self):
        r = self._ejecutar(
            [_partida(CTA_CXP, "COMPENSA", debe="100.00")],
            [_partida(CTA_CXP, "COMPENSA", haber="100.00")],
        )
        origenes = r["trazabilidad_origen"][f"{CTA_CXP}|COMPENSA"]
        cajas = {o["caja"] for o in origenes}
        self.assertEqual(cajas, {"tiquipaya", "america"})


class TestHistoricoUnico(BaseInstitucional):
    def test_historico_y_periodos_son_archivos_unicos(self):
        self._ejecutar([_partida(CTA_CXC, "A1", debe="10.00")], [_partida(CTA_CXP, "B1", haber="20.00")])
        self.assertTrue(os.path.isfile(self.ruta_historico))
        ruta_periodos = c3._ruta_libro_periodos(self.ruta_historico)
        self.assertTrue(os.path.isfile(ruta_periodos))
        historico = c3.cargar_historico(self.ruta_historico)
        self.assertEqual(len(historico), 2)


class TestIdempotenciaParSha(BaseInstitucional):
    def test_mismo_par_sha_es_idempotente(self):
        r1 = self._ejecutar([_partida(CTA_CXC, "A1", debe="10.00")], [_partida(CTA_CXP, "B1", haber="20.00")])
        self.assertTrue(r1["historico_actualizado"])
        r2 = c3i.ejecutar_control3_institucional(self.ruta_tiq, self.ruta_ame, self.ruta_historico)
        self.assertEqual(r2["estado"], "YA_PROCESADO_SIN_CAMBIOS")
        self.assertFalse(r2["historico_actualizado"])

    def test_cambia_ame_requiere_revision(self):
        self._ejecutar([_partida(CTA_CXC, "A1", debe="10.00")], [_partida(CTA_CXP, "B1", haber="20.00")])
        _crear_global(self.ruta_ame, [_partida(CTA_CXP, "B1", haber="99.00")])
        r2 = c3i.ejecutar_control3_institucional(self.ruta_tiq, self.ruta_ame, self.ruta_historico)
        self.assertEqual(r2["estado"], "GLOBAL_MODIFICADO_REQUIERE_REVISION")

    def test_cambia_tiq_requiere_revision(self):
        self._ejecutar([_partida(CTA_CXC, "A1", debe="10.00")], [_partida(CTA_CXP, "B1", haber="20.00")])
        _crear_global(self.ruta_tiq, [_partida(CTA_CXC, "A1", debe="99.00")])
        r2 = c3i.ejecutar_control3_institucional(self.ruta_tiq, self.ruta_ame, self.ruta_historico)
        self.assertEqual(r2["estado"], "GLOBAL_MODIFICADO_REQUIERE_REVISION")

    def test_sha_compuesto_orden_fijo_tiq_ame(self):
        _crear_global(self.ruta_tiq, [_partida(CTA_CXC, "A1", debe="10.00")])
        _crear_global(self.ruta_ame, [_partida(CTA_CXP, "B1", haber="20.00")])
        sha_tiq = c3._hash_archivo(self.ruta_tiq)
        sha_ame = c3._hash_archivo(self.ruta_ame)
        r = c3i.ejecutar_control3_institucional(self.ruta_tiq, self.ruta_ame, self.ruta_historico)
        self.assertEqual(r["sha_par"], f"{sha_tiq}|{sha_ame}")


class TestGlobalNuncaModificado(BaseInstitucional):
    def test_global_tiq_y_ame_no_cambian(self):
        _crear_global(self.ruta_tiq, [_partida(CTA_CXC, "A1", debe="10.00")])
        _crear_global(self.ruta_ame, [_partida(CTA_CXP, "B1", haber="20.00")])
        sha_tiq_antes = c3._hash_archivo(self.ruta_tiq)
        sha_ame_antes = c3._hash_archivo(self.ruta_ame)
        c3i.ejecutar_control3_institucional(self.ruta_tiq, self.ruta_ame, self.ruta_historico)
        self.assertEqual(c3._hash_archivo(self.ruta_tiq), sha_tiq_antes)
        self.assertEqual(c3._hash_archivo(self.ruta_ame), sha_ame_antes)


class TestPeriodos(BaseInstitucional):
    def test_periodos_distintos_tiq_ame_falla_cerrado(self):
        _crear_global(self.ruta_tiq, [_partida(CTA_CXC, "A1", debe="10.00")])
        ruta_ame_otro_mes = os.path.join(
            self.tmp, consolidador_mensual.nombre_sap_global(ANIO, 10, cfg.AMERICA))
        _crear_global(ruta_ame_otro_mes, [_partida(CTA_CXP, "B1", haber="20.00")])
        r = c3i.ejecutar_control3_institucional(self.ruta_tiq, ruta_ame_otro_mes, self.ruta_historico)
        self.assertEqual(r["estado"], "ERROR_TECNICO")
        self.assertIn("PERIODOS_DISTINTOS_TIQ_AME", r["problemas"])


if __name__ == "__main__":
    unittest.main()
