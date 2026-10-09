"""tests_v3/test_global_caja_america.py — GENERAR GLOBAL mensual por caja
(TIQUIPAYA y AMERICA) y contrato JSON de /webhook/tiq-v3-dev/global.

Bug real (Railway, 2026-10-09): /global con caja=america, 2026-09 respondió
HTTP 200 sin cuerpo JSON ("Respuesta no-JSON del backend (HTTP 200)") y no
se creó SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx en Drive. Estos tests cubren:

  - GLOBAL TIQ sigue igual (nombres, cabecera L10, SAP legacy);
  - GLOBAL AME acepta SAP_AME_DD-MM-YYYY.xlsx, exige L10 "CAJA AMERICA" y
    genera SAP_GLOBAL_AME_<MES>_<AÑO>.xlsx / RESULTADO_GLOBAL_AME_<MES>_<AÑO>.json;
  - un fallo del consolidador (o del propio proceso Python) siempre deja
    un JSON válido para el endpoint, nunca texto plano;
  - caja desconocida falla cerrado;
  - 07D (publicar archivo oficial) llega a la rama "crear" cuando el
    archivo todavía no existe en Drive (causa raíz del bug).

Sin Drive, sin n8n real: solo archivos sintéticos en tmp_path.
"""
import datetime
import json
import os
import re
import stat
import subprocess
import sys
from decimal import Decimal

import openpyxl
import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.join(RAIZ, "tests"))

import config_cajas as cfg  # noqa: E402
import consolidador_mensual as cm  # noqa: E402
from xlsx_fixtures import crear_plantilla_sap  # noqa: E402
from v3 import dev_api  # noqa: E402
from v3 import consolidador_mensual_v3 as cm3  # noqa: E402


def _crear_sap_diario(ruta, referencia, dia, cuenta_haber="110101001", c10="SA"):
    """SAP diario sintético con el layout real (hoja "1", cabecera fila 10,
    partidas desde la 16) y C10="SA" como el SAP diario real."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1"
    ws["B10"] = "BO01"
    ws["C10"] = c10
    ws["H10"] = "BOB"
    ws["L10"] = referencia
    fecha_valor = datetime.date(2026, 9, dia)
    for fila, (cuenta, cargo, haber) in enumerate((
        ("210201005", Decimal("150.00"), None),
        (cuenta_haber, None, Decimal("150.00")),
    ), start=16):
        ws[f"B{fila}"] = "BO01"
        ws[f"C{fila}"] = cuenta
        ws[f"D{fila}"] = "PARTIDA DE PRUEBA"
        ws[f"E{fila}"] = cargo
        ws[f"F{fila}"] = haber
        ws[f"L{fila}"] = "10010101"
        ws[f"O{fila}"] = fecha_valor
        ws[f"R{fila}"] = f"ASIG-{dia}"
    wb.save(ruta)


@pytest.fixture()
def base(tmp_path):
    plantilla = tmp_path / "Plantilla SAP maestra.xlsx"
    crear_plantilla_sap(str(plantilla))
    return {"dir": str(tmp_path / "dev_workdir"), "plantilla": str(plantilla)}


def _materializar(base, caja, archivos):
    """Simula lo que n8n materializa desde la carpeta SAP oficial de Drive."""
    entrada = dev_api.preparar_global_entrada(2026, 9, base["dir"], caja)["dir_entrada"]
    for nombre, referencia, dia in archivos:
        cuenta = cfg.resolver_caja(caja).cuenta_haber
        _crear_sap_diario(os.path.join(entrada, nombre), referencia, dia, cuenta_haber=cuenta)
    return entrada


def _cli_generar_global(tmp_path, base, caja):
    entrada = tmp_path / "in.json"
    salida = tmp_path / "out.json"
    entrada.write_text(json.dumps({
        "anio": 2026, "mes": 9, "base_dir_dev": base["dir"],
        "ruta_plantilla_origen": base["plantilla"], "caja": caja,
    }), encoding="utf-8")
    assert dev_api.main(["--accion", "generar_global", "--input", str(entrada), "--output", str(salida)]) == 0
    return json.loads(salida.read_text(encoding="utf-8"))


def _l10(ruta_xlsx):
    wb = openpyxl.load_workbook(ruta_xlsx, read_only=True)
    try:
        return wb["1"]["L10"].value
    finally:
        wb.close()


def _assert_sobre_error(r):
    assert r["resultado"] == "ERROR"
    assert r["ok"] is False
    assert r["estado"] == "ERROR_REVISAR"
    assert r["codigo"] == "GLOBAL_ERROR_TECNICO"
    assert isinstance(r["mensaje"], str) and r["mensaje"]
    assert isinstance(r["detalle"], str) and r["detalle"]
    assert "Traceback" not in json.dumps(r)


# ---------------------------------------------------------------------------
# Configuración explícita por caja (sin duplicar el consolidador)
# ---------------------------------------------------------------------------

def test_configuracion_explicita_por_caja():
    assert (cfg.TIQUIPAYA.codigo, cfg.TIQUIPAYA.prefijo_archivo, cfg.TIQUIPAYA.nombre_sap) == \
        ("tiquipaya", "TIQ", "CAJA TIQUIPAYA")
    assert (cfg.AMERICA.codigo, cfg.AMERICA.prefijo_archivo, cfg.AMERICA.nombre_sap) == \
        ("america", "AME", "CAJA AMERICA")
    assert cm.nombre_sap_global(2026, 9, "america") == "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx"
    assert cm.nombre_resultado_json(2026, 9, "america") == "RESULTADO_GLOBAL_AME_SEPTIEMBRE_2026.json"
    assert cm.nombre_sap_global(2026, 9) == "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx"
    assert cm.nombre_resultado_json(2026, 9) == "RESULTADO_GLOBAL_TIQ_SEPTIEMBRE_2026.json"
    assert cm.construir_metadata_cabecera_global(2026, 9, "america")["referencia"] == "CAJA AMERICA"
    assert cm.construir_metadata_cabecera_global(2026, 9)["referencia"] == "CAJA TIQUIPAYA"


# ---------------------------------------------------------------------------
# TIQUIPAYA — sin cambios
# ---------------------------------------------------------------------------

def test_global_tiq_sigue_funcionando_con_nombres_cabecera_y_legacy(tmp_path, base):
    _materializar(base, "tiquipaya", [
        ("SAP_01-09-2026.xlsx", "CAJA TIQUIPAYA", 1),       # legacy, solo TIQ
        ("SAP_TIQ_02-09-2026.xlsx", "CAJA TIQUIPAYA", 2),
        ("SAP_AME_03-09-2026.xlsx", "CAJA AMERICA", 3),     # otra caja: ignorado
    ])
    r = _cli_generar_global(tmp_path, base, "tiquipaya")
    assert r["resultado"] == "OK"
    assert r["estado"] == "VALIDADO_PENDIENTE_PUBLICACION"
    assert r["sap_incluidos"] == ["SAP_01-09-2026.xlsx", "SAP_TIQ_02-09-2026.xlsx"]
    assert r["ruta_global_generado"] == os.path.abspath(
        os.path.join(base["dir"], "global", "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx"))
    assert os.path.basename(r["ruta_resultado_json"]) == "RESULTADO_GLOBAL_TIQ_SEPTIEMBRE_2026.json"
    assert _l10(r["ruta_global_generado"]) == "CAJA TIQUIPAYA"


def test_global_tiq_sin_caja_es_identico_a_caja_tiquipaya(tmp_path, base):
    _materializar(base, "tiquipaya", [("SAP_TIQ_02-09-2026.xlsx", "CAJA TIQUIPAYA", 2)])
    r = dev_api.generar_global(2026, 9, base["dir"], base["plantilla"])
    assert os.path.basename(r["ruta_global_generado"]) == "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx"
    assert r["estado"] == "VALIDADO_PENDIENTE_PUBLICACION"


# ---------------------------------------------------------------------------
# AMERICA
# ---------------------------------------------------------------------------

def test_global_ame_acepta_sap_ame_y_genera_nombres_ame(tmp_path, base):
    _materializar(base, "america", [
        ("SAP_AME_01-09-2026.xlsx", "CAJA AMERICA", 1),
        ("SAP_AME_02-09-2026.xlsx", "CAJA AMERICA", 2),
    ])
    r = _cli_generar_global(tmp_path, base, "america")
    assert r["resultado"] == "OK"
    assert r["estado"] == "VALIDADO_PENDIENTE_PUBLICACION", r["blockers"]
    assert r["sap_incluidos"] == ["SAP_AME_01-09-2026.xlsx", "SAP_AME_02-09-2026.xlsx"]
    assert r["ruta_global_generado"] == os.path.abspath(
        os.path.join(base["dir"], "global", "america", "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx"))
    assert os.path.isfile(r["ruta_global_generado"])
    assert os.path.basename(r["ruta_resultado_json"]) == "RESULTADO_GLOBAL_AME_SEPTIEMBRE_2026.json"
    assert os.path.isfile(r["ruta_resultado_json"])
    assert _l10(r["ruta_global_generado"]) == "CAJA AMERICA"
    assert r["cargo_global"] == r["haber_global"] == "300.00"


def test_global_ame_valida_cabecera_caja_america(tmp_path, base):
    _materializar(base, "america", [("SAP_AME_01-09-2026.xlsx", "CAJA TIQUIPAYA", 1)])
    r = _cli_generar_global(tmp_path, base, "america")
    assert r["estado"] == "ERROR_REVISAR"
    assert r["ruta_global_generado"] is None
    assert any("CABECERA_L10_ESPERADO_'CAJA AMERICA'" in b for b in r["blockers"]), r["blockers"]
    assert not os.path.exists(os.path.join(base["dir"], "global", "america", "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx"))


def test_global_ame_nunca_toma_sap_tiq_ni_legacy(tmp_path, base):
    _materializar(base, "america", [
        ("SAP_TIQ_01-09-2026.xlsx", "CAJA TIQUIPAYA", 1),
        ("SAP_02-09-2026.xlsx", "CAJA TIQUIPAYA", 2),
        ("SAP_AME_03-09-2026.xlsx", "CAJA AMERICA", 3),
    ])
    r = _cli_generar_global(tmp_path, base, "america")
    assert r["sap_incluidos"] == ["SAP_AME_03-09-2026.xlsx"]
    assert r["estado"] == "VALIDADO_PENDIENTE_PUBLICACION"


def test_global_tiq_y_ame_no_se_pisan(tmp_path, base):
    _materializar(base, "tiquipaya", [("SAP_TIQ_02-09-2026.xlsx", "CAJA TIQUIPAYA", 2)])
    _materializar(base, "america", [("SAP_AME_02-09-2026.xlsx", "CAJA AMERICA", 2)])
    r_tiq = _cli_generar_global(tmp_path, base, "tiquipaya")
    r_ame = _cli_generar_global(tmp_path, base, "america")
    assert r_tiq["ruta_global_generado"] != r_ame["ruta_global_generado"]
    assert _l10(r_tiq["ruta_global_generado"]) == "CAJA TIQUIPAYA"
    assert _l10(r_ame["ruta_global_generado"]) == "CAJA AMERICA"


def test_guardarriel_salida_usa_patron_de_la_caja(tmp_path, base):
    sap = tmp_path / "SAP_AME_01-09-2026.xlsx"
    _crear_sap_diario(str(sap), "CAJA AMERICA", 1, cuenta_haber=cfg.AMERICA.cuenta_haber)
    with pytest.raises(RuntimeError, match="RUTA_SALIDA_NOMBRE_SAP_DIARIO"):
        cm3.ejecutar_consolidacion_v3(
            2026, 9, base["plantilla"], str(tmp_path / "SAP_AME_30-09-2026.xlsx"),
            [str(sap)], True, caja="america",
        )


# ---------------------------------------------------------------------------
# Contrato JSON del endpoint (Python)
# ---------------------------------------------------------------------------

def test_fallo_del_consolidador_devuelve_json_valido(tmp_path, base, monkeypatch):
    _materializar(base, "america", [("SAP_AME_01-09-2026.xlsx", "CAJA AMERICA", 1)])

    def _explota(*_a, **_k):
        raise RuntimeError("fallo simulado del consolidador")

    monkeypatch.setattr(dev_api, "generar_global_mensual", _explota)
    r = _cli_generar_global(tmp_path, base, "america")
    _assert_sobre_error(r)
    assert "fallo simulado del consolidador" in r["mensaje"]
    assert r["detalle"].startswith("RuntimeError")


def test_caja_desconocida_falla_cerrado_con_json(tmp_path, base):
    r = _cli_generar_global(tmp_path, base, "bolivia")
    _assert_sobre_error(r)
    assert "CAJA_DESCONOCIDA" in r["mensaje"]
    assert not os.path.exists(os.path.join(base["dir"], "global"))


def test_entrada_no_materializada_devuelve_json(tmp_path, base):
    r = _cli_generar_global(tmp_path, base, "america")
    _assert_sobre_error(r)
    assert "GLOBAL_ENTRADA_NO_MATERIALIZADA" in r["mensaje"]


def test_periodo_cerrado_conserva_mensaje_que_lee_el_frontend(tmp_path, base, monkeypatch):
    _materializar(base, "america", [("SAP_AME_01-09-2026.xlsx", "CAJA AMERICA", 1)])

    def _cerrado(*_a, **_k):
        raise RuntimeError("PERIODO_CERRADO_CONTROL1: cerrado")

    monkeypatch.setattr(dev_api, "_verificar_periodo_no_cerrado", _cerrado)
    r = _cli_generar_global(tmp_path, base, "america")
    _assert_sobre_error(r)
    assert "PERIODO_CERRADO_CONTROL1" in r["mensaje"]


def test_otras_acciones_conservan_su_contrato_de_error(tmp_path):
    entrada = tmp_path / "in.json"
    salida = tmp_path / "out.json"
    entrada.write_text(json.dumps({"anio": 2026, "mes": 9, "base_dir_dev": str(tmp_path), "caja": "bolivia"}))
    dev_api.main(["--accion", "preparar_global_entrada", "--input", str(entrada), "--output", str(salida)])
    r = json.loads(salida.read_text())
    assert r == {"resultado": "ERROR", "codigo": "ValueError", "mensaje": r["mensaje"]}


# ---------------------------------------------------------------------------
# Workflow n8n (snapshot desplegado en Railway + snapshot v3-final)
# ---------------------------------------------------------------------------

SNAPSHOTS = ("railway-shadow", "v3-final")


def _nodo(snapshot, archivo, nombre):
    wf = json.load(open(os.path.join(RAIZ, "snapshots", snapshot, archivo), encoding="utf-8"))
    return next(n for n in wf["nodes"] if n["name"] == nombre)


@pytest.mark.parametrize("snapshot", SNAPSHOTS)
def test_07d_buscar_emite_item_aunque_el_archivo_no_exista(snapshot):
    """Causa raíz: con 0 resultados de búsqueda, n8n no ejecuta los nodos
    siguientes salvo que el NODO tenga alwaysOutputData (dentro de
    `parameters` n8n lo ignora). Sin esto, la rama "crear" de 07D nunca
    corre: el primer SAP_GLOBAL_AME nunca se crea y /global termina sin
    RESPONDER (HTTP 200 vacío)."""
    nodo = _nodo(snapshot, "HhuQCVP2oCubavzY_07d_publicar_oficial.json", "BUSCAR - Archivo por nombre en carpeta")
    assert nodo.get("alwaysOutputData") is True
    assert "alwaysOutputData" not in nodo["parameters"]


def _renderizar_comando(comando, repo_root, python_bin, ruta_input, ruta_output):
    assert comando.startswith("=")
    c = comando[1:]
    c = c.replace("{{ $env.TIQ_REPO_ROOT || '/workspaces/caja-tiquipaya-v2' }}", repo_root)
    c = c.replace("cd /workspaces/caja-tiquipaya-v2 && /workspaces/.venv-caja/bin/python3",
                  f"cd {repo_root} && {python_bin}")
    c = c.replace("{{ $env.TIQ_PYTHON_BIN || '/workspaces/.venv-caja/bin/python3' }}", python_bin)
    c = c.replace("{{ $('CONSTRUIR payload global').item.json.ruta_input_tmp }}", ruta_input)
    c = c.replace("{{ $('CONSTRUIR payload global').item.json.ruta_output_tmp }}", ruta_output)
    assert "{{" not in c, c
    return c


@pytest.mark.parametrize("snapshot", SNAPSHOTS)
def test_ejecutar_global_si_python_muere_sin_salida_deja_json_valido(snapshot, tmp_path):
    comando = _nodo(snapshot, "aLs1f3GMqswbaENA_backend_dev.json", "EJECUTAR global (Python)")["parameters"]["command"]
    falso = tmp_path / "python_que_muere"
    falso.write_text("#!/bin/sh\necho 'Traceback (most recent call last): boom' >&2\nexit 1\n")
    falso.chmod(falso.stat().st_mode | stat.S_IEXEC)
    salida = tmp_path / "out.json"
    c = _renderizar_comando(comando, RAIZ, str(falso), str(tmp_path / "in.json"), str(salida))
    p = subprocess.run(["bash", "-c", c], capture_output=True, text=True)
    assert p.returncode == 0
    _assert_sobre_error(json.loads(salida.read_text(encoding="utf-8")))


@pytest.mark.parametrize("snapshot", SNAPSHOTS)
def test_ejecutar_global_no_pisa_la_salida_real_de_python(snapshot, tmp_path, base):
    _materializar(base, "america", [("SAP_AME_01-09-2026.xlsx", "CAJA AMERICA", 1)])
    entrada = tmp_path / "in.json"
    entrada.write_text(json.dumps({
        "anio": 2026, "mes": 9, "base_dir_dev": base["dir"],
        "ruta_plantilla_origen": base["plantilla"], "caja": "america",
    }))
    comando = _nodo(snapshot, "aLs1f3GMqswbaENA_backend_dev.json", "EJECUTAR global (Python)")["parameters"]["command"]
    salida = tmp_path / "out.json"
    c = _renderizar_comando(comando, RAIZ, sys.executable, str(entrada), str(salida))
    p = subprocess.run(["bash", "-c", c], capture_output=True, text=True)
    assert p.returncode == 0, p.stderr
    r = json.loads(salida.read_text(encoding="utf-8"))
    assert r["resultado"] == "OK"
    assert os.path.basename(r["ruta_global_generado"]) == "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx"


def test_destino_ame_no_hardcodea_folder_id():
    js = _nodo("railway-shadow", "aLs1f3GMqswbaENA_backend_dev.json", "CONSTRUIR - Lista publicaciones GLOBAL")["parameters"]["jsCode"]
    assert "$env.DRIVE_GLOBAL_AME" in js
    assert not re.search(r"['\"][A-Za-z0-9_-]{25,}['\"]", js), "no debe haber folder IDs literales"
