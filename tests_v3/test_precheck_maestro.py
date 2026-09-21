"""tests_v3/test_precheck_maestro.py — pruebas del precheck de cobertura
del maestro (v3/precheck_maestro.py, FASE 10C/10E), la precondición V3 que
corre entre el Módulo 02 MATERIALIZACION y el Módulo 03 MOTOR PYTHON.

Uso: python -m pytest tests_v3/test_precheck_maestro.py -q
"""

import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tests"))

import openpyxl

import run_batch  # noqa: E402
import xlsx_fixtures as fx  # noqa: E402
from v3.precheck_maestro import (  # noqa: E402
    evaluar_cobertura_maestro, aplicar_precheck_maestro, filtrar_aptos_para_motor,
    MAESTRO_APTO, BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA,
)
from v3.materializacion import MATERIALIZADO, SIN_ARCHIVO
from v3.motor import NO_PROCESADO


def _maestro(tmp_path, macros_filas, atc_filas, nombre="MAESTRO.xlsm"):
    ruta = tmp_path / nombre
    fx.crear_maestro_unico(str(ruta), macros_filas=macros_filas, atc_filas=atc_filas)
    return str(ruta)


def _fila_atc(fecha, monto="0.00"):
    return (fecha, "BANCO (NETO)", "110103012", f"ATC {fecha}", monto, "ASIG")


def _fila_macros(fecha, credito="0.01"):
    return (fecha, "COD-TEST", credito)


def _cierre(tmp_path, fecha, cobros_atc="0.00", nombre=None):
    """Cierre sintetico con cobros_atc controlado en SFC101 (SFC102 en
    0.00) -- mismo campo que motor_tiquipaya.cruzar_atc_preconciliado()
    suma para decidir si el dia tuvo movimiento ATC (ver
    v3.precheck_maestro._cierre_tiene_movimiento_atc)."""
    nombre = nombre or run_batch.nombre_cierre_esperado(fecha)
    ruta = tmp_path / nombre
    sfc_con_atc = {"total_movimiento": "0.00", "cobros_atc": cobros_atc, "dolares": "0.00", "depositos": []}
    sfc_vacio = {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []}
    fx.crear_cierre(str(ruta), sfc_con_atc, sfc_vacio)
    return str(ruta)


# 1) cierre <= fecha maxima de MACROS, CON movimiento ATC y CON fila ATC -> MAESTRO_APTO
def test_cierre_dentro_de_cobertura_da_maestro_apto(tmp_path):
    macros_filas = [_fila_macros(f) for f in ("2026-09-01", "2026-09-05", "2026-09-10")]
    atc_filas = [_fila_atc(f) for f in ("2026-09-01", "2026-09-05", "2026-09-10")]
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre(tmp_path, "2026-09-10", cobros_atc="500.00")

    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-10", ruta_cierre)
    assert r["estado"] == MAESTRO_APTO
    assert r["fecha_maxima_macros"] == "2026-09-10"
    assert r["fecha_maxima_atc"] == "2026-09-10"

    # una fecha ANTERIOR a la maxima, con fila ATC propia, tambien es apta
    ruta_cierre_01 = _cierre(tmp_path, "2026-09-01", cobros_atc="300.00")
    r2 = evaluar_cobertura_maestro(ruta_maestro, "2026-09-01", ruta_cierre_01)
    assert r2["estado"] == MAESTRO_APTO


# 2) cierre > fecha maxima de MACROS -> bloqueado (independiente de ATC)
def test_cierre_posterior_a_macros_bloquea(tmp_path):
    macros_filas = [_fila_macros("2026-09-05")]
    atc_filas = [_fila_atc(f) for f in ("2026-09-01", "2026-09-10")]  # ATC SI llega hasta 09-10
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre(tmp_path, "2026-09-10", cobros_atc="0.00")  # sin ATC: no cambia el resultado

    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-10", ruta_cierre)
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert r["fecha_maxima_macros"] == "2026-09-05"
    assert "MACROS" in r["mensaje"]


# 3) FASE 10E — CORREGIDO: cierre CON movimiento ATC pero SIN fila ATC
#    real para esa fecha exacta -> bloqueado, con el mensaje simple pedido.
#    (antes se comparaba contra una fecha maxima global; ahora se exige la
#    fila EXACTA solo cuando el cierre realmente tuvo movimiento ATC).
def test_cierre_con_atc_y_fila_atc_ausente_bloquea(tmp_path):
    macros_filas = [_fila_macros(f) for f in ("2026-09-01", "2026-09-12")]  # MACROS SI llega hasta 09-12
    atc_filas = [_fila_atc("2026-09-01")]  # SIN fila para 2026-09-12
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre(tmp_path, "2026-09-12", cobros_atc="1234.56")  # el cierre SI tuvo movimiento ATC

    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-12", ruta_cierre)
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert r["mensaje"] == (
        "No se encontró información ATC del maestro para la fecha del cierre. "
        "Actualice/verifique el maestro y vuelva a procesar."
    )


# caso real 12/09/2026: cierre SIN movimiento ATC, SIN fila ATC para esa
# fecha -> APTO (la ausencia de fila es legitima, no se exige nada de ATC).
def test_cierre_sin_atc_y_fila_atc_ausente_continua(tmp_path):
    macros_filas = [_fila_macros(f) for f in ("2026-09-01", "2026-09-12")]
    atc_filas = [_fila_atc("2026-09-01")]  # SIN fila para 2026-09-12, igual que el maestro real
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre(tmp_path, "2026-09-12", cobros_atc="0.00")  # el cierre NO tuvo movimiento ATC

    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-12", ruta_cierre)
    assert r["estado"] == MAESTRO_APTO


# cierre CON movimiento ATC y CON fila ATC presente -> continua (caso base)
def test_cierre_con_atc_y_fila_atc_presente_continua(tmp_path):
    macros_filas = [_fila_macros(f) for f in ("2026-09-01", "2026-09-12")]
    atc_filas = [_fila_atc("2026-09-12")]
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre(tmp_path, "2026-09-12", cobros_atc="999.00")

    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-12", ruta_cierre)
    assert r["estado"] == MAESTRO_APTO


# 4) MACROS insuficiente sigue bloqueando aunque el cierre no tenga ATC
def test_macros_insuficiente_bloquea_incluso_sin_atc(tmp_path):
    macros_filas = [_fila_macros("2026-09-03")]
    atc_filas = []
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre(tmp_path, "2026-09-10", cobros_atc="0.00")

    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-10", ruta_cierre)
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert r["fecha_maxima_macros"] == "2026-09-03"


# 5) maestro ilegible (archivo inexistente) -> bloqueo de precondicion
def test_maestro_inexistente_bloquea(tmp_path):
    ruta_cierre = _cierre(tmp_path, "2026-09-10", cobros_atc="0.00")
    r = evaluar_cobertura_maestro("/no/existe/MAESTRO.xlsm", "2026-09-10", ruta_cierre)
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert "MAESTRO_ILEGIBLE" in r["mensaje"]
    assert r["fecha_maxima_macros"] is None


def test_maestro_sin_ninguna_fecha_en_macros_bloquea(tmp_path):
    # Maestro sintacticamente valido (hojas correctas) pero MACROS sin
    # ninguna fila de datos -- "no adivinar": sin evidencia, no hay MAESTRO_APTO.
    ruta_maestro = _maestro(tmp_path, macros_filas=[], atc_filas=[])
    ruta_cierre = _cierre(tmp_path, "2026-09-10", cobros_atc="0.00")
    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-10", ruta_cierre)
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert "MAESTRO_SIN_FECHAS_REGISTRADAS" in r["mensaje"]


def test_cierre_ilegible_para_determinar_atc_bloquea(tmp_path):
    ruta_maestro = _maestro(tmp_path, [_fila_macros("2026-09-10")], [_fila_atc("2026-09-10")])
    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-10", "/no/existe/CIERRE.xlsm")
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert "movimiento ATC" in r["mensaje"]


# 6) un hueco de ATC para una fecha SIN movimiento ATC nunca bloquea,
#    incluso si otras fechas SI tienen fila -- reproduce el hallazgo real
#    del 12/09/2026 (MACROS hasta 09-14, ATC hasta 09-11, cierre del 09-12
#    confirmado SIN movimiento ATC por el usuario).
def test_hueco_de_atc_en_dia_sin_movimiento_no_bloquea(tmp_path):
    fechas_macros = ["2026-09-01", "2026-09-10", "2026-09-11", "2026-09-12", "2026-09-13", "2026-09-14"]
    fechas_atc = ["2026-09-01", "2026-09-10", "2026-09-11"]  # SIN fila para 09-12/13/14
    macros_filas = [_fila_macros(f) for f in fechas_macros]
    atc_filas = [_fila_atc(f) for f in fechas_atc]
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)

    # 09-11: tiene fila ATC propia -> apto
    ruta_11 = _cierre(tmp_path, "2026-09-11", cobros_atc="100.00", nombre="CIERRE 11-09-2026 variante.xlsm")
    assert evaluar_cobertura_maestro(ruta_maestro, "2026-09-11", ruta_11)["estado"] == MAESTRO_APTO

    # 09-12: SIN fila ATC, pero el cierre real NO tuvo movimiento ATC -> apto
    ruta_12 = _cierre(tmp_path, "2026-09-12", cobros_atc="0.00", nombre="CIERRE 12-09-2026 variante A.xlsm")
    r12 = evaluar_cobertura_maestro(ruta_maestro, "2026-09-12", ruta_12)
    assert r12["estado"] == MAESTRO_APTO
    assert r12["fecha_maxima_atc"] == "2026-09-11"  # informativo, ya no es motivo de bloqueo

    # pero si ESE MISMO dia (09-12) SI hubiera tenido movimiento ATC, bloquearia
    ruta_12_con_atc = _cierre(tmp_path, "2026-09-12", cobros_atc="777.00", nombre="CIERRE 12-09-2026 variante B.xlsm")
    r12b = evaluar_cobertura_maestro(ruta_maestro, "2026-09-12", ruta_12_con_atc)
    assert r12b["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA


# ---------------------------------------------------------------------------
# aplicar_precheck_maestro() / filtrar_aptos_para_motor() — orquestacion
# por lote, tal como se inserta entre v3.materializacion y v3.motor.
# ---------------------------------------------------------------------------

def test_aplicar_precheck_solo_anota_items_materializados(tmp_path):
    ruta_maestro = _maestro(tmp_path, [_fila_macros("2026-09-10")], [_fila_atc("2026-09-10")], "APTO.xlsm")
    ruta_cierre = _cierre(tmp_path, "2026-09-10", cobros_atc="0.00")
    items = [
        {"fecha": "2026-09-01", "estado_materializacion": SIN_ARCHIVO},
        {"fecha": "2026-09-10", "estado_materializacion": MATERIALIZADO,
         "ruta_maestro_local": ruta_maestro, "ruta_cierre_local": ruta_cierre},
    ]
    anotados = aplicar_precheck_maestro(items)
    assert "estado_precheck_maestro" not in anotados[0]  # nunca se evalua lo que no hace falta
    assert anotados[1]["estado_precheck_maestro"] == MAESTRO_APTO


# 7) un cierre BLOQUEADO por el precheck nunca llega al Modulo 03 (motor):
#    filtrar_aptos_para_motor() lo excluye explicitamente de la lista que
#    v3.motor.ejecutar_motor() recibe -- no es "no procesado dentro del
#    motor", es "el motor jamas lo ve".
def test_cierre_bloqueado_nunca_pasa_a_filtrar_aptos_para_motor(tmp_path):
    ruta_vieja = _maestro(tmp_path, [_fila_macros("2026-09-05")], [], "VIEJO.xlsm")
    ruta_apta = _maestro(tmp_path, [_fila_macros("2026-09-11")], [], "APTO.xlsm")
    cierre_11 = _cierre(tmp_path, "2026-09-11", cobros_atc="0.00", nombre="CIERRE 11-09-2026 variante.xlsm")
    cierre_05 = _cierre(tmp_path, "2026-09-05", cobros_atc="0.00", nombre="CIERRE 05-09-2026 variante.xlsm")
    items = [
        {"fecha": "2026-09-11", "estado_materializacion": MATERIALIZADO,
         "ruta_maestro_local": ruta_vieja, "ruta_cierre_local": cierre_11},
        {"fecha": "2026-09-05", "estado_materializacion": MATERIALIZADO,
         "ruta_maestro_local": ruta_apta, "ruta_cierre_local": cierre_05},
    ]
    anotados = aplicar_precheck_maestro(items)
    bloqueado, apto = anotados
    assert bloqueado["estado_precheck_maestro"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert bloqueado["estado_motor"] == NO_PROCESADO
    assert bloqueado["resultado"] is None
    assert apto["estado_precheck_maestro"] == MAESTRO_APTO

    aptos_para_motor = filtrar_aptos_para_motor(anotados)
    fechas_aptas = {c["fecha"] for c in aptos_para_motor}
    assert fechas_aptas == {"2026-09-05"}  # el bloqueado (09-11) nunca aparece aqui
    assert len(aptos_para_motor) == 1


def test_cierre_bloqueado_en_procesar_lote_no_invoca_motor_real(tmp_path, monkeypatch):
    # Integracion con v3.dev_api.procesar_lote(): el cierre bloqueado
    # termina con estado_final BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA y
    # pipeline_tiquipaya.procesar_cierre_completo() NUNCA se llama para esa
    # fecha (se intercepta la funcion real para probarlo, no se confia
    # solo en el resultado final).
    import pipeline_tiquipaya as pipeline
    from v3 import dev_api

    llamadas = []
    original = pipeline.procesar_cierre_completo

    def _espia(*args, **kwargs):
        llamadas.append(kwargs.get("ruta_cierre") or (args[0] if args else None))
        return original(*args, **kwargs)

    monkeypatch.setattr(pipeline, "procesar_cierre_completo", _espia)

    origen_dir = tmp_path / "origen"
    origen_dir.mkdir()
    nombre = run_batch.nombre_cierre_esperado("2026-09-11")
    fx.crear_cierre(
        str(origen_dir / nombre),
        {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []},
        {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []},
    )
    # maestro con cobertura MACROS SOLO hasta 09-10 (como el maestro real de FASE 10B)
    ruta_maestro = _maestro(tmp_path, [_fila_macros("2026-09-10")], [_fila_atc("2026-09-10")])
    ruta_plantilla = tmp_path / "Plantilla.xlsx"
    fx.crear_plantilla_sap(str(ruta_plantilla))

    base_dir_dev = str(tmp_path / "dev")
    r = dev_api.crear_lote_pendiente("2026-09-11", "2026-09-11", "auditor.dev", base_dir_dev)
    lote = dev_api.procesar_lote(
        r["lote_id"], base_dir_dev, str(origen_dir), ruta_maestro, str(ruta_plantilla),
    )
    assert lote["cierres"][0]["estado_final"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert lote["cierres"][0]["fecha_maxima_macros"] == "2026-09-10"
    assert llamadas == []  # el motor real JAMAS se invoco para este cierre
    # ni corregible ni publicable: no es un problema del cierre que el
    # auditor deba resolver desde el formulario de correccion.
    assert lote["cierres"][0]["requiere_revision"] is False
    assert lote["cierres"][0]["publicable"] is False


def test_cierre_apto_en_procesar_lote_si_llega_al_motor(tmp_path):
    from v3 import dev_api

    origen_dir = tmp_path / "origen"
    origen_dir.mkdir()
    nombre = run_batch.nombre_cierre_esperado("2026-09-10")
    fx.crear_cierre(
        str(origen_dir / nombre),
        {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []},
        {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []},
    )
    ruta_maestro = _maestro(tmp_path, [_fila_macros("2026-09-10")], [_fila_atc("2026-09-10")])
    ruta_plantilla = tmp_path / "Plantilla.xlsx"
    fx.crear_plantilla_sap(str(ruta_plantilla))

    base_dir_dev = str(tmp_path / "dev")
    r = dev_api.crear_lote_pendiente("2026-09-10", "2026-09-10", "auditor.dev", base_dir_dev)
    lote = dev_api.procesar_lote(
        r["lote_id"], base_dir_dev, str(origen_dir), ruta_maestro, str(ruta_plantilla),
    )
    assert lote["cierres"][0]["estado_precheck_maestro"] == MAESTRO_APTO
    assert lote["cierres"][0]["estado_final"] != BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA


# caso real: cierre 12/09/2026 SIN movimiento ATC (confirmado por el
# usuario) sobre el maestro real (MACROS hasta 09-14, ATC hasta 09-11) ->
# debe ser APTO, no bloqueado. Reproduce el hallazgo de FASE 10E con datos
# equivalentes a los reales (mismo patron de cobertura), sin depender del
# archivo real de Drive para que el test sea autocontenido/reproducible.
def test_caso_real_12_09_sin_movimiento_atc_es_apto(tmp_path):
    fechas_macros = ["2026-09-%02d" % d for d in range(1, 15) if d not in (6,)]  # ~hasta 09-14
    fechas_atc = ["2026-09-%02d" % d for d in range(1, 12) if d not in (6,)]  # hasta 09-11
    ruta_maestro = _maestro(
        tmp_path,
        [_fila_macros(f) for f in fechas_macros],
        [_fila_atc(f) for f in fechas_atc],
    )
    ruta_cierre_12 = _cierre(tmp_path, "2026-09-12", cobros_atc="0.00")
    r = evaluar_cobertura_maestro(ruta_maestro, "2026-09-12", ruta_cierre_12)
    assert r["estado"] == MAESTRO_APTO


# ---------------------------------------------------------------------------
# BUG real (bloque caja-america): el precheck ignoraba `caja` y siempre
# interpretaba SFC101/SFC102 (TIQUIPAYA) además de exigir la fila ATC con
# CAJA=TIQUIPAYA, sin importar la caja real del lote. Los fixtures de
# AMERICA viven aquí (y no en tests/xlsx_fixtures.py), igual que en
# tests/test_caja_america.py, para no tocar los fixtures compartidos por
# las pruebas históricas de TIQUIPAYA.
# ---------------------------------------------------------------------------

def _hoja_resumen_ame(ws, cobros_atc):
    ws.append(["TOTAL MOVIMIENTO DEL DIA", "0.00"])
    ws.append(["COBROS ATC", cobros_atc])
    ws.append(["TOTAL COMUNICACIONES INTERNAS", "0.00"])
    ws.append(["DOLARES", "0.00"])
    ws.append(["POSGRADO RESERVA", "0.00"])  # AMERICA exige este campo (reserva_posgrado=True)
    ws.append([None, None, None, None, None])
    ws.append(["COMPOSICION DE DEPOSITOS", "IMPORTE Bs", "FECHA DE DEPOSITO",
               "ASIGNACION", "BANCO"])


def _cierre_america(tmp_path, fecha, cobros_atc_107="0.00", cobros_atc_108="0.00", nombre=None):
    """CIERRE de América: hojas SFC107/SFC108 (+ sus Comunicaciones Internas),
    igual patrón que tests/test_caja_america.py:crear_cierre_america()."""
    nombre = nombre or run_batch.nombre_cierre_esperado(fecha)
    ruta = tmp_path / nombre
    wb = openpyxl.Workbook()
    ws107 = wb.active
    ws107.title = "SFC107"
    _hoja_resumen_ame(ws107, cobros_atc_107)
    ws108 = wb.create_sheet("SFC108")
    _hoja_resumen_ame(ws108, cobros_atc_108)
    for sfc in ("SFC107", "SFC108"):
        ws_ci = wb.create_sheet(f"COMUNICACIONES INTERNAS {sfc}")
        ws_ci.append(["N°", "N° DE FACTURA", "TOTAL C.I.",
                      "CUENTA CONTABLE BANCO", "ASIGNACION", "BANCO"])
    wb.save(str(ruta))
    return str(ruta)


def _maestro_con_caja(tmp_path, macros_filas, atc_filas, nombre="MAESTRO_CON_CAJA.xlsm"):
    """Maestro mensual en formato NUEVO (columna CAJA en "ATC TIQUIPAYA"),
    igual patrón que tests/test_caja_america.py:crear_maestro_con_caja()."""
    ruta = tmp_path / nombre
    wb = openpyxl.Workbook()
    ws_macros = wb.active
    ws_macros.title = "Tablas Dinamicas Profesional"
    ws_macros.append(["Fecha", "Código de Asignación", "Créditos"])
    for fila in macros_filas:
        ws_macros.append(list(fila))
    ws_atc = wb.create_sheet("ATC TIQUIPAYA")
    ws_atc.append(["FECHA", "TIPO", "CUENTA CONTABLE", "DETALLE", "MONTO",
                   "ASIGNACION", "CAJA"])
    for fila in atc_filas:
        ws_atc.append(list(fila))
    wb.save(str(ruta))
    return str(ruta)


def _fila_atc_caja(fecha, caja, monto="0.00"):
    return (fecha, "BANCO (NETO)", "110103012", f"ATC {fecha} {caja}", monto, "ASIG", caja)


def test_default_sin_caja_sigue_siendo_tiquipaya_sfc101_102(tmp_path):
    """Compatibilidad obligatoria: una llamada histórica sin `caja` sigue
    interpretando SFC101/SFC102 exactamente como siempre (caja=None)."""
    macros_filas = [_fila_macros("2026-09-10")]
    atc_filas = [_fila_atc("2026-09-10")]
    ruta_maestro = _maestro(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre(tmp_path, "2026-09-10", cobros_atc="500.00")

    sin_caja = evaluar_cobertura_maestro(ruta_maestro, "2026-09-10", ruta_cierre)
    explicito_tiq = evaluar_cobertura_maestro(ruta_maestro, "2026-09-10", ruta_cierre, caja="tiquipaya")
    assert sin_caja == explicito_tiq
    assert sin_caja["estado"] == MAESTRO_APTO


def test_america_lee_sfc107_108_no_sfc101_102(tmp_path):
    """El bug real: el precheck sumaba cobros_atc de SFC101/SFC102
    (siempre 0.00 en un cierre de América) en vez de SFC107/SFC108, así
    que nunca detectaba movimiento ATC real de América."""
    fecha = "2026-09-15"
    macros_filas = [_fila_macros(fecha)]
    atc_filas = [_fila_atc_caja(fecha, "AMERICA", "1900.00")]
    ruta_maestro = _maestro_con_caja(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre_america(tmp_path, fecha, cobros_atc_107="1900.00", cobros_atc_108="0.00")

    r = evaluar_cobertura_maestro(ruta_maestro, fecha, ruta_cierre, caja="america")
    assert r["estado"] == MAESTRO_APTO


def test_atc_america_se_filtra_por_caja_america(tmp_path):
    """ATC AME se filtra por CAJA=AMERICA: una fila TIQUIPAYA para esa
    misma fecha no basta para dar cobertura a AMERICA."""
    fecha = "2026-09-16"
    macros_filas = [_fila_macros(fecha)]
    atc_filas = [_fila_atc_caja(fecha, "TIQUIPAYA", "900.00")]  # solo TIQUIPAYA, sin fila AMERICA
    ruta_maestro = _maestro_con_caja(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre_america(tmp_path, fecha, cobros_atc_107="1900.00", cobros_atc_108="0.00")

    r = evaluar_cobertura_maestro(ruta_maestro, fecha, ruta_cierre, caja="america")
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA


def test_america_con_atc_valido_pasa_precheck(tmp_path):
    fecha = "2026-09-17"
    macros_filas = [_fila_macros(fecha)]
    atc_filas = [_fila_atc_caja(fecha, "AMERICA", "2000.00")]
    ruta_maestro = _maestro_con_caja(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre_america(tmp_path, fecha, cobros_atc_107="2000.00", cobros_atc_108="0.00")

    r = evaluar_cobertura_maestro(ruta_maestro, fecha, ruta_cierre, caja="america")
    assert r["estado"] == MAESTRO_APTO


def test_america_sin_fila_atc_cuando_tuvo_atc_bloquea(tmp_path):
    fecha = "2026-09-18"
    macros_filas = [_fila_macros(fecha)]
    atc_filas = []  # sin fila AMERICA para esa fecha
    ruta_maestro = _maestro_con_caja(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre_america(tmp_path, fecha, cobros_atc_107="1500.00", cobros_atc_108="0.00")

    r = evaluar_cobertura_maestro(ruta_maestro, fecha, ruta_cierre, caja="america")
    assert r["estado"] == BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA


def test_america_sin_movimiento_atc_no_exige_fila(tmp_path):
    fecha = "2026-09-19"
    macros_filas = [_fila_macros(fecha)]
    atc_filas = []  # sin fila AMERICA -- legitimo, el cierre no tuvo movimiento ATC
    ruta_maestro = _maestro_con_caja(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre_america(tmp_path, fecha, cobros_atc_107="0.00", cobros_atc_108="0.00")

    r = evaluar_cobertura_maestro(ruta_maestro, fecha, ruta_cierre, caja="america")
    assert r["estado"] == MAESTRO_APTO


def test_no_contamina_tiquipaya_con_filas_america_ni_viceversa(tmp_path):
    """Mismo maestro, mismo día, ambas cajas con movimiento ATC real y su
    propia fila: cada caja ve solo la suya (misma garantía que
    tests/test_caja_america.py, pero pasando por el precheck completo)."""
    fecha = "2026-09-20"
    macros_filas = [_fila_macros(fecha)]
    atc_filas = [
        _fila_atc_caja(fecha, "TIQUIPAYA", "900.00"),
        _fila_atc_caja(fecha, "AMERICA", "1950.00"),
    ]
    ruta_maestro = _maestro_con_caja(tmp_path, macros_filas, atc_filas)

    ruta_cierre_tiq = _cierre(tmp_path, fecha, cobros_atc="900.00",
                               nombre="CIERRE TIQ " + os.path.basename(run_batch.nombre_cierre_esperado(fecha)))
    r_tiq = evaluar_cobertura_maestro(ruta_maestro, fecha, ruta_cierre_tiq, caja="tiquipaya")
    assert r_tiq["estado"] == MAESTRO_APTO

    ruta_cierre_ame = _cierre_america(tmp_path, fecha, cobros_atc_107="1950.00", cobros_atc_108="0.00",
                                       nombre="CIERRE AME " + os.path.basename(run_batch.nombre_cierre_esperado(fecha)))
    r_ame = evaluar_cobertura_maestro(ruta_maestro, fecha, ruta_cierre_ame, caja="america")
    assert r_ame["estado"] == MAESTRO_APTO


def test_aplicar_precheck_maestro_propaga_caja_america(tmp_path):
    """aplicar_precheck_maestro(..., caja="america") evalúa el lote entero
    como América (mismo camino que v3/dev_api.py:procesar_lote pasando
    caja_lote.codigo)."""
    fecha = "2026-09-21"
    macros_filas = [_fila_macros(fecha)]
    atc_filas = [_fila_atc_caja(fecha, "AMERICA", "1900.00")]
    ruta_maestro = _maestro_con_caja(tmp_path, macros_filas, atc_filas)
    ruta_cierre = _cierre_america(tmp_path, fecha, cobros_atc_107="1900.00", cobros_atc_108="0.00")

    items = [{
        "fecha": fecha, "estado_materializacion": MATERIALIZADO,
        "ruta_maestro_local": ruta_maestro, "ruta_cierre_local": ruta_cierre,
    }]
    anotados = aplicar_precheck_maestro(items, caja="america")
    assert anotados[0]["estado_precheck_maestro"] == MAESTRO_APTO
