"""tests_v3/test_motor.py — pruebas propias del módulo 03 · MOTOR PYTHON de
V3 (v3/motor.py). Este módulo es un ADAPTADOR sobre el motor real de V2
(motor_tiquipaya/sap_writer/pipeline_tiquipaya, sin modificar): estas
pruebas verifican que el adaptador invoca ese código real correctamente y
no introduce ninguna regresión — NO reimplementan ni reinterpretan
ninguna regla contable.

Uso: python -m pytest tests_v3/test_motor.py -q
"""

import hashlib
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tests"))

import xlsx_fixtures as fx
from v3.motor import ejecutar_motor, ejecutar_motor_cierre, PROCESADO, NO_PROCESADO, ERROR_MOTOR
from v3.ingesta import SIN_ARCHIVO, AMBIGUO, ERROR_INGESTA
from v3.materializacion import MATERIALIZADO, ERROR_MATERIALIZACION


def _sha256(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 16), b""):
            h.update(bloque)
    return h.hexdigest()


def _item_materializado(tmp_path, sfc101, sfc102, nombre="CIERRE 01-09-2026.xlsm",
                         macros_filas=(), atc_filas=(), fecha="2026-09-01"):
    ruta_cierre = tmp_path / nombre
    fx.crear_cierre(str(ruta_cierre), sfc101, sfc102)
    ruta_maestro = tmp_path / "MAESTRO_SEPTIEMBRE.xlsm"
    fx.crear_maestro_unico(str(ruta_maestro), macros_filas=list(macros_filas), atc_filas=list(atc_filas))
    ruta_plantilla = tmp_path / "Plantilla_SAP_maestra.xlsx"
    fx.crear_plantilla_sap(str(ruta_plantilla))
    return {
        "fecha": fecha, "archivo_esperado": nombre, "estado_ingesta": "ENCONTRADO",
        "estado_materializacion": MATERIALIZADO,
        "ruta_cierre_local": str(ruta_cierre), "ruta_maestro_local": str(ruta_maestro),
        "ruta_template_sap_local": str(ruta_plantilla), "ruta_markers_local": None,
    }


_SFC_VACIO = {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []}


# 1) MATERIALIZADO -> ejecuta motor
def test_materializado_ejecuta_motor(tmp_path):
    item = _item_materializado(tmp_path, _SFC_VACIO, _SFC_VACIO)
    base_dir_dev = str(tmp_path / "dev")
    r = ejecutar_motor_cierre(item, base_dir_dev)
    assert r["estado_motor"] == PROCESADO
    assert r["resultado"] == "VALIDADO_PENDIENTE_PUBLICACION"
    assert r["diferencia"] == "0.00"
    assert r["bloqueadores"] == 0
    assert r["cargo"] == r["haber"]
    assert os.path.isfile(r["ruta_resultado"])
    assert os.path.isfile(r["ruta_sap"])


# 2/3/4) SIN_ARCHIVO / AMBIGUO / ERROR_MATERIALIZACION -> no ejecuta
@pytest.mark.parametrize("estado_mat", [SIN_ARCHIVO, AMBIGUO, ERROR_MATERIALIZACION, ERROR_INGESTA])
def test_estados_no_materializados_no_ejecutan_motor(tmp_path, estado_mat):
    item = {
        "fecha": "2026-09-02", "archivo_esperado": "CIERRE 02-09-2026.xlsm",
        "estado_ingesta": "ENCONTRADO", "estado_materializacion": estado_mat,
        "ruta_cierre_local": None, "ruta_maestro_local": None,
        "ruta_template_sap_local": None, "ruta_markers_local": None,
    }
    r = ejecutar_motor_cierre(item, str(tmp_path / "dev"))
    assert r["estado_motor"] == NO_PROCESADO
    assert r["resultado"] == estado_mat
    assert r["ruta_resultado"] is None
    assert r["ruta_sap"] is None


# 5) Cargo = Haber para caso valido
def test_cargo_igual_haber_caso_valido(tmp_path):
    item = _item_materializado(tmp_path, _SFC_VACIO, _SFC_VACIO)
    r = ejecutar_motor_cierre(item, str(tmp_path / "dev"))
    assert r["estado_motor"] == PROCESADO
    assert r["cargo"] == r["haber"]
    assert r["diferencia"] == "0.00"


# 6) ALQUILERES conserva comportamiento V2 (exclusion del cuadre)
def test_alquileres_conserva_comportamiento_v2(tmp_path):
    sfc101 = dict(_SFC_VACIO)
    sfc101 = {
        "total_movimiento": "50.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": [],
        "ci": [{"n": 1, "factura": "F-1", "total": "50.00", "cuenta": None,
                "asignacion": None, "banco": "ALQUILERES"}],
    }
    item = _item_materializado(tmp_path, sfc101, _SFC_VACIO)
    r = ejecutar_motor_cierre(item, str(tmp_path / "dev"))
    assert r["estado_motor"] == PROCESADO
    assert r["resultado"] == "VALIDADO_PENDIENTE_PUBLICACION"
    assert r["diferencia"] == "0.00"  # universo (50) - alquileres (50) = 0 = recaudacion (0)

    import json
    with open(r["ruta_resultado"], "r", encoding="utf-8") as f:
        resultado_json = json.load(f)
    assert resultado_json["alquileres"] == "50.00"
    assert resultado_json["universo_ajustado"] == "0.00"


# 7) comision ATC conserva cuenta 110201008 (default cuando el maestro no la especifica)
def test_comision_atc_conserva_cuenta_110201008(tmp_path):
    sfc101 = {"total_movimiento": "1000.00", "cobros_atc": "1000.00", "dolares": "0.00", "depositos": []}
    atc_filas = [
        ("2026-09-01", "BANCO (NETO)", "110103012", "NETO ATC", "950.00", "3P00000001"),
        ("2026-09-01", "COMISION ATC", None, "COMISION ATC", "50.00", None),  # sin cuenta -> default 110201008
    ]
    item = _item_materializado(tmp_path, sfc101, _SFC_VACIO, atc_filas=atc_filas)
    r = ejecutar_motor_cierre(item, str(tmp_path / "dev"))
    assert r["estado_motor"] == PROCESADO
    assert r["resultado"] == "VALIDADO_PENDIENTE_PUBLICACION"

    wb = openpyxl.load_workbook(r["ruta_sap"], data_only=True)
    ws = wb["1"]
    cuentas = [ws[f"C{fila}"].value for fila in range(16, 16 + 10) if ws[f"C{fila}"].value]
    assert "110201008" in cuentas  # la comision ATC SI llego a esa cuenta


# 8) 110201003 no se bloquea globalmente (CI legitima con esa cuenta)
def test_cuenta_110201003_no_bloqueada_globalmente_en_ci(tmp_path):
    sfc101 = {
        "total_movimiento": "75.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": [],
        "ci": [{"n": 1, "factura": "F-1", "total": "75.00", "cuenta": "110201003",
                "asignacion": "REF-CI", "banco": "BNB"}],
    }
    item = _item_materializado(tmp_path, sfc101, _SFC_VACIO)
    r = ejecutar_motor_cierre(item, str(tmp_path / "dev"))
    assert r["estado_motor"] == PROCESADO
    assert r["resultado"] == "VALIDADO_PENDIENTE_PUBLICACION"  # nunca bloqueado por usar esa cuenta
    assert r["diferencia"] == "0.00"


# 9) errores del motor quedan aislados por cierre
def test_error_de_un_cierre_no_detiene_el_lote(tmp_path):
    (tmp_path / "a").mkdir()
    item_ok = _item_materializado(tmp_path / "a", _SFC_VACIO, _SFC_VACIO, fecha="2026-09-01")

    item_roto = {
        "fecha": "2026-09-02", "archivo_esperado": "CIERRE 02-09-2026.xlsm",
        "estado_ingesta": "ENCONTRADO", "estado_materializacion": MATERIALIZADO,
        "ruta_cierre_local": str(tmp_path / "no_existe.xlsm"),  # archivo inexistente
        "ruta_maestro_local": item_ok["ruta_maestro_local"],
        "ruta_template_sap_local": item_ok["ruta_template_sap_local"],
        "ruta_markers_local": None,
    }
    resultados = ejecutar_motor([item_ok, item_roto], str(tmp_path / "dev"))
    por_fecha = {r["fecha"]: r["estado_motor"] for r in resultados}
    assert por_fecha["2026-09-01"] == PROCESADO
    assert por_fecha["2026-09-02"] == ERROR_MOTOR


def test_materializacion_incompleta_da_error_motor_no_crashea(tmp_path):
    item = {
        "fecha": "2026-09-03", "archivo_esperado": "CIERRE 03-09-2026.xlsm",
        "estado_ingesta": "ENCONTRADO", "estado_materializacion": MATERIALIZADO,
        "ruta_cierre_local": None,  # inconsistencia: MATERIALIZADO pero sin ruta
        "ruta_maestro_local": None, "ruta_template_sap_local": None, "ruta_markers_local": None,
    }
    r = ejecutar_motor_cierre(item, str(tmp_path / "dev"))
    assert r["estado_motor"] == ERROR_MOTOR
    assert "MATERIALIZACION_INCOMPLETA" in r["mensaje"]


# 10) original/copia de entrada no se modifica
def test_copia_de_entrada_no_se_modifica(tmp_path):
    item = _item_materializado(tmp_path, _SFC_VACIO, _SFC_VACIO)
    hash_cierre_antes = _sha256(item["ruta_cierre_local"])
    hash_maestro_antes = _sha256(item["ruta_maestro_local"])
    hash_plantilla_antes = _sha256(item["ruta_template_sap_local"])

    ejecutar_motor_cierre(item, str(tmp_path / "dev"))

    assert _sha256(item["ruta_cierre_local"]) == hash_cierre_antes
    assert _sha256(item["ruta_maestro_local"]) == hash_maestro_antes
    assert _sha256(item["ruta_template_sap_local"]) == hash_plantilla_antes


# 11) outputs solo quedan dentro del directorio DEV
def test_outputs_quedan_dentro_del_directorio_dev(tmp_path):
    item = _item_materializado(tmp_path, _SFC_VACIO, _SFC_VACIO)
    base_dir_dev = str(tmp_path / "dev_workdir")
    r = ejecutar_motor_cierre(item, base_dir_dev)
    base_abs = os.path.abspath(base_dir_dev)
    assert os.path.commonpath([os.path.abspath(r["ruta_resultado"]), base_abs]) == base_abs
    assert os.path.commonpath([os.path.abspath(r["ruta_sap"]), base_abs]) == base_abs


def test_todas_las_filas_tienen_al_menos_las_claves_pedidas(tmp_path):
    # FASE 8: ejecutar_motor_cierre() ahora preserva TODO lo que el item ya
    # traía del Módulo 02 (carry-forward: fecha, archivo_esperado,
    # estado_ingesta, estado_materializacion, ruta_cierre_local, etc.)
    # además de sus propias claves — necesario para que el Módulo 04
    # (CLASIFICACION) reciba estado_ingesta/estado_materializacion y el
    # Módulo 05 (REVISION) reciba las rutas locales (ver FASE 8, hallazgo
    # de incompatibilidad 03->04 corregido en v3/motor.py).
    item = _item_materializado(tmp_path, _SFC_VACIO, _SFC_VACIO)
    r = ejecutar_motor_cierre(item, str(tmp_path / "dev"))
    esperado_minimo = {
        "fecha", "estado_motor", "resultado", "diferencia", "bloqueadores",
        "ruta_resultado", "ruta_sap", "cargo", "haber", "mensaje", "mensajes",
    }
    assert esperado_minimo.issubset(set(r.keys()))
    # y las claves de entrada (Modulo 02) sobreviven intactas
    assert r["archivo_esperado"] == item["archivo_esperado"]
    assert r["estado_ingesta"] == item["estado_ingesta"]
    assert r["ruta_maestro_local"] == item["ruta_maestro_local"]
    assert r["ruta_template_sap_local"] == item["ruta_template_sap_local"]
