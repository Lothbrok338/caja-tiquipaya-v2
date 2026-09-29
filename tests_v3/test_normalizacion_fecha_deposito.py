"""tests_v3/test_normalizacion_fecha_deposito.py — normalización operativa
de FECHA DE DEPOSITO antes del precheck (hallazgo real CAJA AMÉRICA,
2026-09-28: '28/09/26' con año de 2 dígitos rompía
v3/precheck_maestro.py, reportado engañosamente como "Maestro sin
cobertura"). Ver v3/normalizacion_fecha_deposito.py: reutiliza LITERAL el
regex y la regla de expansión del año de 2 dígitos del skill
`claude/auditor-cierres-skill` (control5_fecha_deposito.py /
control5_regla.md, paso 3 "TEXTO válido"), y el mecanismo de escritura
XML seguro vendorizado en v3/_xlsm_xml.py (nunca openpyxl.save() sobre un
cierre real).

Uso: python -m pytest tests_v3/test_normalizacion_fecha_deposito.py -q
"""
import datetime
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tests"))

import openpyxl  # noqa: E402

import xlsx_fixtures as fx  # noqa: E402
from v3 import normalizacion_fecha_deposito as nfd  # noqa: E402
from v3 import precheck_maestro as pm  # noqa: E402
import excel_io  # noqa: E402


# ---------------------------------------------------------------------------
# interpretar_fecha_texto: la regla pura, extraída literal del auditor
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("texto,anio_cierre,esperado", [
    ("28/09/26", 2026, datetime.date(2026, 9, 28)),   # caso real CIERRE 26-09-2026 SFC108!F3
    ("25/09/26", 2026, datetime.date(2026, 9, 25)),   # caso real CIERRE 25-09-2026 SFC108!F3
    ("5/9/26", 2026, datetime.date(2026, 9, 5)),       # D/M/YY (un solo dígito)
    ("05/9/2026", 2026, datetime.date(2026, 9, 5)),    # DD/M/YYYY
    ("28/09/2026", 2020, datetime.date(2026, 9, 28)),  # año de 4 dígitos: nunca depende del año del cierre
    ("28/09/2026", None, datetime.date(2026, 9, 28)),  # idem, aunque no se conozca el año del cierre
])
def test_formatos_admitidos_se_interpretan_correctamente(texto, anio_cierre, esperado):
    fecha, motivo = nfd.interpretar_fecha_texto(texto, anio_cierre)
    assert fecha == esperado
    assert motivo is None


@pytest.mark.parametrize("texto,anio_cierre,motivo_esperado", [
    ("28/09/26", 2025, "TEXTO_ANIO_2_DIGITOS_NO_COINCIDE_CON_ANIO_DEL_CIERRE"),  # NO se autocorrige: año no coincide
    ("28/09/26", None, "TEXTO_ANIO_2_DIGITOS_NO_COINCIDE_CON_ANIO_DEL_CIERRE"),  # año del cierre desconocido: tampoco
    ("31/02/26", 2026, "TEXTO_FECHA_INEXISTENTE_EN_CALENDARIO"),                # 31 de febrero no existe
    ("2026-09-28", 2026, "TEXTO_FORMATO_NO_ADMITIDO"),                          # ISO: no es el formato que esta regla admite
    (" 28/09/26", 2026, "TEXTO_FORMATO_NO_ADMITIDO"),                           # espacio: ambiguo, no se admite
    ("28/09/26 ", 2026, "TEXTO_FORMATO_NO_ADMITIDO"),
    ("28-09-2026", 2026, "TEXTO_FORMATO_NO_ADMITIDO"),                          # separador '-': no admitido aquí
    ("28/9", 2026, "TEXTO_FORMATO_NO_ADMITIDO"),
    ("hola", 2026, "TEXTO_FORMATO_NO_ADMITIDO"),
])
def test_formatos_ambiguos_o_no_admitidos_nunca_se_autocorrigen(texto, anio_cierre, motivo_esperado):
    fecha, motivo = nfd.interpretar_fecha_texto(texto, anio_cierre)
    assert fecha is None
    assert motivo == motivo_esperado


# ---------------------------------------------------------------------------
# clasificar_fechas_deposito / aplicar / normalizar_cierre_para_precheck
# ---------------------------------------------------------------------------

def _cierre_composicion_depositos(ruta, hojas):
    """Fixture MÍNIMO (solo el bloque COMPOSICION DE DEPOSITOS que este
    módulo lee) -- deliberadamente distinto de xlsx_fixtures.crear_cierre
    (que arma un cierre COMPLETO para excel_io.leer_cierre): alcanza para
    probar clasificar_fechas_deposito()/aplicar() en aislamiento.
    `hojas`: {"SFC107": [(etiqueta, importe, fecha_valor, asignacion), ...]}."""
    wb = openpyxl.Workbook()
    primero = True
    for nombre_hoja, filas in hojas.items():
        ws = wb.active if primero else wb.create_sheet(nombre_hoja)
        if primero:
            ws.title = nombre_hoja
            primero = False
        ws.append(["COMPOSICION DE DEPOSITOS", "IMPORTE Bs", "FECHA DE DEPOSITO", "ASIGNACION", "BANCO"])
        for etiqueta, importe, fecha_valor, asignacion in filas:
            ws.append([etiqueta, importe, fecha_valor, asignacion, "BNB 123"])
    wb.save(ruta)


def test_clasifica_solo_las_celdas_de_texto_ambiguo_pero_resoluble(tmp_path):
    ruta = str(tmp_path / "CIERRE 26-09-2026.xlsm")
    _cierre_composicion_depositos(ruta, {
        "SFC107": [("DEPOSITO", "100.00", "28/09/26", "AB1")],   # america, año coincide -> normaliza
        "SFC108": [("DEPOSITO", "50.00", "28/09/25", "AB2")],    # america, año NO coincide -> NO se toca
    })
    resultado = nfd.clasificar_fechas_deposito(ruta, "america", "2026-09-26")
    assert len(resultado["plan"]) == 1
    assert resultado["plan"][0]["hoja"] == "SFC107"
    assert resultado["plan"][0]["fecha_nueva"] == datetime.date(2026, 9, 28)
    assert len(resultado["trazabilidad"]) == 1
    traza = resultado["trazabilidad"][0]
    assert traza["valor_original"] == "28/09/26"
    assert traza["valor_normalizado"] == "2026-09-28"
    assert traza["motivo"] == nfd.MOTIVO_ANIO_2_DIGITOS
    assert traza["hoja"] == "SFC107" and traza["celda"]


def test_tiquipaya_usa_sfc101_102_no_sfc107_108(tmp_path):
    ruta = str(tmp_path / "CIERRE 26-09-2026.xlsm")
    _cierre_composicion_depositos(ruta, {
        "SFC101": [("DEPOSITO", "100.00", "26/09/26", "AB1")],
        "SFC107": [("DEPOSITO", "100.00", "26/09/26", "AB1")],  # misma fecha, pero NO es hoja de tiquipaya
    })
    resultado = nfd.clasificar_fechas_deposito(ruta, "tiquipaya", "2026-09-26")
    assert len(resultado["plan"]) == 1
    assert resultado["plan"][0]["hoja"] == "SFC101"


def test_normalizar_cierre_para_precheck_escribe_copia_nueva_y_no_toca_el_original(tmp_path):
    ruta = str(tmp_path / "CIERRE 26-09-2026.xlsm")
    _cierre_composicion_depositos(ruta, {"SFC108": [("DEPOSITO", "100.00", "28/09/26", "AB1")]})
    bytes_originales = open(ruta, "rb").read()

    resultado = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-26", "america")

    assert resultado["normalizado"] is True
    assert resultado["ruta_cierre_local"] != ruta
    assert resultado["ruta_cierre_local"].endswith(".normalizado.xlsm")
    assert os.path.isfile(resultado["ruta_cierre_local"])
    # el original queda BYTE A BYTE intacto -- nunca se modifica ni se sobrescribe
    assert open(ruta, "rb").read() == bytes_originales

    # la copia nueva SI tiene la fecha real (no texto) al releerla --
    # openpyxl puede devolver datetime.datetime para una celda de fecha
    # corta (igual que ya asume excel_io._fecha_iso), nunca un str.
    wb = openpyxl.load_workbook(resultado["ruta_cierre_local"], data_only=True)
    valor = wb["SFC108"]["C2"].value
    assert isinstance(valor, (datetime.date, datetime.datetime)), f"deberia ser una fecha real, no texto: {valor!r}"
    fecha_leida = valor.date() if isinstance(valor, datetime.datetime) else valor
    assert fecha_leida == datetime.date(2026, 9, 28)


def test_sin_nada_que_normalizar_no_crea_ningun_archivo_nuevo(tmp_path):
    """Caso normal (incluida toda Caja Tiquipaya hoy): cero fechas de
    texto ambiguas -> ruta_cierre_local queda EXACTAMENTE igual, sin
    ningun archivo nuevo -- cero costo, cero cambio de comportamiento."""
    ruta = str(tmp_path / "CIERRE 26-09-2026.xlsm")
    _cierre_composicion_depositos(ruta, {"SFC101": [("DEPOSITO", "100.00", datetime.date(2026, 9, 26), "AB1")]})
    resultado = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-26", "tiquipaya")
    assert resultado["normalizado"] is False
    assert resultado["ruta_cierre_local"] == ruta
    assert resultado["cambios"] == []
    assert not os.path.isfile(ruta.replace(".xlsm", ".normalizado.xlsm"))


def test_anio_que_no_coincide_no_se_autocorrige_ni_crea_copia(tmp_path):
    ruta = str(tmp_path / "CIERRE 26-09-2026.xlsm")
    _cierre_composicion_depositos(ruta, {"SFC108": [("DEPOSITO", "100.00", "28/09/25", "AB1")]})  # año 25, cierre 2026
    resultado = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-26", "america")
    assert resultado["normalizado"] is False
    assert resultado["ruta_cierre_local"] == ruta


def test_normalizar_cierres_materializados_solo_toca_items_materializados():
    items = [
        {"fecha": "2026-09-26", "estado_materializacion": "MATERIALIZADO", "ruta_cierre_local": "/no/existe/CIERRE.xlsm"},
        {"fecha": "2026-09-27", "estado_materializacion": "SIN_ARCHIVO", "ruta_cierre_local": None},
    ]
    # ruta_cierre_local inexistente -> clasificar_fechas_deposito lanza,
    # normalizar_cierre_para_precheck lo aisla (no normaliza, no lanza).
    resultado = nfd.normalizar_cierres_materializados(items, caja="america")
    assert resultado[0]["ruta_cierre_local"] == "/no/existe/CIERRE.xlsm"
    assert resultado[0]["normalizacion_fecha_deposito"] == []
    assert resultado[1]["estado_materializacion"] == "SIN_ARCHIVO"
    assert "normalizacion_fecha_deposito" not in resultado[1]


# ---------------------------------------------------------------------------
# End-to-end REAL: precheck_maestro sobre el cierre SIN normalizar (falla
# con CIERRE_FECHA_INVALIDA) vs. sobre la copia normalizada (decide con
# datos reales) -- el mismo escenario que rompía en producción (CAJA
# AMÉRICA, 2026-09-28), reproducido aquí con TIQUIPAYA (misma lógica,
# mismos módulos, sin datos reales de ninguna caja).
# ---------------------------------------------------------------------------

def test_end_to_end_precheck_antes_y_despues_de_normalizar(tmp_path):
    macros_filas = [("2026-09-26", "COD-TEST", "0.01"), ("2026-09-27", "COD-TEST", "0.01")]
    atc_filas = [("2026-09-26", "BANCO (NETO)", "110103012", "ATC 2026-09-26", "0.00", "ASIG")]
    ruta_maestro = str(tmp_path / "MAESTRO.xlsm")
    fx.crear_maestro_unico(str(ruta_maestro), macros_filas=macros_filas, atc_filas=atc_filas)

    ruta_cierre = str(tmp_path / "CIERRE 26-09-2026.xlsm")
    sfc_con_deposito_ambiguo = {
        "total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00",
        "depositos": [{"importe": "100.00", "fecha": "26/09/26", "asignacion": "AB1"}],
    }
    sfc_vacio = {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []}
    fx.crear_cierre(ruta_cierre, sfc_con_deposito_ambiguo, sfc_vacio)

    # ANTES de normalizar: el cierre no se puede leer para el chequeo ATC
    # (año de 2 dígitos que excel_io._fecha_iso rechaza) -> CIERRE_FECHA_INVALIDA,
    # NUNCA "Maestro sin cobertura".
    r_antes = pm.evaluar_cobertura_maestro(ruta_maestro, "2026-09-26", ruta_cierre, caja="tiquipaya")
    assert r_antes["estado"] == pm.CIERRE_FECHA_INVALIDA
    assert r_antes["estado"] != pm.BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert "Fecha de depósito inválida en el cierre" in r_antes["mensaje"]

    # DESPUES de normalizar (mismo año que el cierre: 26/09/26 -> 2026-09-26,
    # inequívoco): decide con datos reales -> MAESTRO_APTO.
    norm = nfd.normalizar_cierre_para_precheck(ruta_cierre, "2026-09-26", "tiquipaya")
    assert norm["normalizado"] is True
    r_despues = pm.evaluar_cobertura_maestro(ruta_maestro, "2026-09-26", norm["ruta_cierre_local"], caja="tiquipaya")
    assert r_despues["estado"] == pm.MAESTRO_APTO


def test_ano_no_coincide_sigue_bloqueado_como_cierre_fecha_invalida_tras_intentar_normalizar(tmp_path):
    """Un año de 2 dígitos que NO coincide con el año del cierre nunca se
    autocorrige (regla explícita): sigue siendo CIERRE_FECHA_INVALIDA
    incluso después de pasar por la normalización (que, correctamente, no
    tocó nada)."""
    macros_filas = [("2026-09-26", "COD-TEST", "0.01")]
    atc_filas = [("2026-09-26", "BANCO (NETO)", "110103012", "ATC", "0.00", "ASIG")]
    ruta_maestro = str(tmp_path / "MAESTRO.xlsm")
    fx.crear_maestro_unico(str(ruta_maestro), macros_filas=macros_filas, atc_filas=atc_filas)

    ruta_cierre = str(tmp_path / "CIERRE 26-09-2026.xlsm")
    sfc = {
        "total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00",
        "depositos": [{"importe": "100.00", "fecha": "26/09/25", "asignacion": "AB1"}],  # año 25, cierre 2026
    }
    sfc_vacio = {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []}
    fx.crear_cierre(ruta_cierre, sfc, sfc_vacio)

    norm = nfd.normalizar_cierre_para_precheck(ruta_cierre, "2026-09-26", "tiquipaya")
    assert norm["normalizado"] is False

    r = pm.evaluar_cobertura_maestro(ruta_maestro, "2026-09-26", norm["ruta_cierre_local"], caja="tiquipaya")
    assert r["estado"] == pm.CIERRE_FECHA_INVALIDA


def test_excel_io_no_se_toca_para_nada(tmp_path):
    """excel_io.py es código compartido con V2/Tiquipaya con pruebas
    propias que exigen diff CERO contra HEAD (ver
    tests_v3/test_auditoria_mensual.py::test_P_Q_v2_y_control3_permanecen_sin_cambios) --
    _fecha_iso() sigue lanzando exactamente el mismo ValueError genérico
    de siempre, con el mismo mensaje; v3.precheck_maestro distingue el
    caso por ese mensaje, nunca por un tipo de excepción nuevo."""
    with pytest.raises(ValueError) as exc_info:
        excel_io._fecha_iso("28/09/26")  # 2 dígitos: excel_io por sí solo nunca expande el año
    assert type(exc_info.value) is ValueError  # exactamente ValueError, ninguna subclase nueva
    assert "Fecha en formato no reconocido" in str(exc_info.value)


# ===========================================================================
# CONTROL 5 — FECHA EXCEL REAL CON DD/MM INVERTIDOS (hallazgo real CAJA
# AMÉRICA, cierres 09, 10 y 11/09/2026: la celda FECHA DE DEPOSITO ya era
# una fecha Excel REAL -- no texto -- guardada como 2026-10-09, 2026-11-09
# y 2026-12-09 en vez de 2026-09-10, 2026-09-11 y 2026-09-12).
# Regla LITERAL del auditor (control5_regla.md, NORMALIZAR_INVERSION_DDMM):
# solo con voucher UNICO en MACROS (asignacion normalizada + importe
# exacto) cuya fecha == fecha invertida. Sin eso, NO se corrige.
# ===========================================================================

D = datetime.date


def _indice(*entradas):
    """Índice de vouchers de MACROS en memoria: (codigo, importe_2dec) -> [fechas].
    Cada `entrada` es (codigo, importe, fecha_o_None)."""
    idx = {}
    for codigo, importe, fecha in entradas:
        idx.setdefault((codigo, importe), []).append(fecha)
    return idx


def _cierre_dep(tmp_path, hojas, nombre="CIERRE 09-09-2026.xlsm"):
    ruta = str(tmp_path / nombre)
    _cierre_composicion_depositos(ruta, hojas)
    return ruta


def _fecha_celda_en_copia(ruta, hoja, celda):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    valor = wb[hoja][celda].value
    return valor.date() if isinstance(valor, datetime.datetime) else valor


def test_inversion_confirmada_por_voucher_unico_se_normaliza_en_la_copia_y_no_toca_el_original(tmp_path):
    # Caso real 09/09: SFC107!F3 = 2026-10-09 (Excel real), voucher unico 2026-09-10.
    ruta = _cierre_dep(tmp_path, {
        "SFC107": [("DEPOSITO 1", 51230.1, D(2026, 10, 9), "3P90102072")],
        "SFC108": [("DEPOSITO 1", 22071.0, D(2026, 10, 9), "3P90094768")],
    })
    original = open(ruta, "rb").read()
    indice = _indice(("3P90102072", "51230.10", D(2026, 9, 10)), ("3P90094768", "22071.00", D(2026, 9, 10)))

    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=indice)

    assert r["normalizado"] is True
    assert r["ruta_cierre_local"] == ruta.replace(".xlsm", ".normalizado.xlsm")
    assert open(ruta, "rb").read() == original                                   # original byte a byte intacto
    assert _fecha_celda_en_copia(r["ruta_cierre_local"], "SFC107", "C2") == D(2026, 9, 10)
    assert _fecha_celda_en_copia(r["ruta_cierre_local"], "SFC108", "C2") == D(2026, 9, 10)
    assert len(r["cambios"]) == 2
    t = r["cambios"][0]
    assert t["clase"] == nfd.CLASE_INVERSION_DDMM == "NORMALIZAR_INVERSION_DDMM"
    assert t["motivo"] == nfd.MOTIVO_INVERSION_DDMM
    assert (t["valor_original"], t["valor_normalizado"], t["fecha_voucher"]) == ("2026-10-09", "2026-09-10", "2026-09-10")
    assert (t["asignacion"], t["importe"], t["hoja"]) == ("3P90102072", "51230.10", "SFC107")
    assert r["revision"] == []


def test_inversion_solo_normaliza_la_celda_autorizada_y_nada_mas(tmp_path):
    """Una celda invertida con voucher que confirma + otra celda NO confirmada en el mismo cierre:
    solo cambia la primera; el resto del paquete queda idéntico (verificar_integridad)."""
    ruta = _cierre_dep(tmp_path, {
        "SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1"), ("DEPOSITO 2", 200.0, D(2026, 11, 9), "AB2")],
    })
    indice = _indice(("AB1", "100.00", D(2026, 9, 10)))          # AB2 sin voucher -> NO se corrige
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=indice)
    assert r["normalizado"] is True and len(r["cambios"]) == 1
    assert _fecha_celda_en_copia(r["ruta_cierre_local"], "SFC107", "C2") == D(2026, 9, 10)
    assert _fecha_celda_en_copia(r["ruta_cierre_local"], "SFC107", "C3") == D(2026, 11, 9)   # intacta
    assert [(x["celda"], x["motivo"]) for x in r["revision"]] == [("C3", "SIN_VOUCHER_EN_MACROS")]


@pytest.mark.parametrize("indice,motivo_esperado", [
    (_indice(), "SIN_VOUCHER_EN_MACROS"),                                                            # sin voucher
    (_indice(("AB1", "100.00", D(2026, 9, 10)), ("AB1", "100.00", D(2026, 9, 10))), "VOUCHER_NO_UNICO_2_COINCIDENCIAS"),
    (_indice(("AB1", "100.00", None)), "VOUCHER_SIN_FECHA_VALIDA"),
    (_indice(("AB1", "100.00", D(2026, 9, 20))), "FECHA_DISTINTA_A_VOUCHER_Y_SU_INVERSION_TAMPOCO_COINCIDE"),  # ni por cercanía
    (_indice(("AB1", "100.01", D(2026, 9, 10))), "SIN_VOUCHER_EN_MACROS"),                           # importe NO exacto
    (_indice(("XX9", "100.00", D(2026, 9, 10))), "SIN_VOUCHER_EN_MACROS"),                           # otra asignación
])
def test_sin_voucher_unico_que_confirme_la_inversion_no_se_corrige(tmp_path, indice, motivo_esperado):
    ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")]})
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=indice)
    assert r["normalizado"] is False and r["ruta_cierre_local"] == ruta and r["cambios"] == []
    assert not os.path.isfile(ruta.replace(".xlsm", ".normalizado.xlsm"))
    assert [x["motivo"] for x in r["revision"]] == [motivo_esperado]
    assert r["revision"][0]["fecha_excel"] == "2026-10-09" and r["revision"][0]["fecha_invertida_candidata"] == "2026-09-10"


def test_voucher_que_coincide_con_la_fecha_excel_o_dentro_de_1_dia_no_se_toca(tmp_path):
    # fecha Excel == voucher, y fecha Excel a 1 dia del voucher: tolerancia del auditor, valor intacto.
    for voucher in (D(2026, 10, 9), D(2026, 10, 8), D(2026, 10, 10)):
        ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")]}, nombre=f"C {voucher}.xlsm")
        r = nfd.normalizar_cierre_para_precheck(ruta, "2026-10-09", "america", indice_vouchers=_indice(("AB1", "100.00", voucher)))
        assert r["normalizado"] is False and r["revision"] == []


@pytest.mark.parametrize("fecha", [D(2026, 9, 25), D(2026, 9, 9), D(2026, 12, 13)])
def test_fecha_sin_inversion_posible_no_se_evalua_ni_lee_macros(tmp_path, fecha):
    """día > 12 (no existe la inversión) o día == mes (inversión idéntica): no hay nada que confirmar,
    y MACROS ni siquiera se lee (evidencia perezosa)."""
    ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, fecha, "AB1")]})
    llamadas = []

    def proveedor():
        llamadas.append(1)
        return _indice(("AB1", "100.00", D(2026, 9, 10)))

    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=proveedor)
    assert r["normalizado"] is False and r["revision"] == [] and llamadas == []


def test_el_indice_de_macros_se_pide_solo_por_celda_con_inversion_posible(tmp_path):
    ruta = _cierre_dep(tmp_path, {
        "SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")],
        "SFC108": [("DEPOSITO 1", 200.0, D(2026, 10, 9), "AB2")],
    })
    llamadas = []

    def proveedor():
        llamadas.append(1)
        return _indice(("AB1", "100.00", D(2026, 9, 10)), ("AB2", "200.00", D(2026, 9, 10)))

    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=proveedor)
    assert r["normalizado"] is True and len(r["cambios"]) == 2
    assert len(llamadas) == 2  # una petición por celda candidata (la memoización entre cierres se prueba con _proveedor_indice_vouchers)


def test_sin_indice_de_macros_la_inversion_nunca_se_corrige(tmp_path):
    ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")]})
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america")          # sin evidencia
    assert r["normalizado"] is False and r["ruta_cierre_local"] == ruta
    assert [x["motivo"] for x in r["revision"]] == ["MACROS_NO_DISPONIBLE_PARA_CONFIRMAR"]


def test_proveedor_de_macros_que_lanza_no_corrige_ni_rompe(tmp_path):
    ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")]})

    def proveedor():
        raise ValueError("maestro corrupto")

    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=proveedor)
    assert r["normalizado"] is False and r["ruta_cierre_local"] == ruta
    assert r["revision"][0]["motivo"] == "MACROS_ILEGIBLE:ValueError"


def test_asignacion_vacia_no_se_corrige(tmp_path):
    ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), None)]})
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=_indice(("AB1", "100.00", D(2026, 9, 10))))
    assert r["normalizado"] is False
    assert [x["motivo"] for x in r["revision"]] == ["ASIGNACION_VACIA_SIN_VOUCHER"]


def test_asignacion_se_compara_normalizada_sin_tildes_ni_espacios_ni_mayusculas(tmp_path):
    ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), " 3p 90102072 ")]})
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=_indice(("3P90102072", "100.00", D(2026, 9, 10))))
    assert r["normalizado"] is True


def test_inversion_no_aplica_a_texto_solo_a_fecha_excel_real(tmp_path):
    """control5_regla.md: 'La inversion solo se considera para fechas Excel reales, no para texto'.
    Un texto 09/10/2026 se sigue tratando como antes (formato inequivoco -> misma fecha), sin invertir."""
    ruta = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, "09/10/2026", "AB1")]})
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=_indice(("AB1", "100.00", D(2026, 9, 10))))
    assert r["normalizado"] is True
    assert _fecha_celda_en_copia(r["ruta_cierre_local"], "SFC107", "C2") == D(2026, 10, 9)   # NO 2026-09-10
    assert r["cambios"][0]["motivo"] == nfd.MOTIVO_FORMATO_NORMALIZADO


def test_normalizacion_dd_mm_yy_de_texto_se_conserva_junto_con_la_inversion(tmp_path):
    ruta = _cierre_dep(tmp_path, {
        "SFC107": [("DEPOSITO 1", 100.0, "28/09/26", "AB1")],                       # texto DD/MM/YY -> 2026-09-28
        "SFC108": [("DEPOSITO 1", 200.0, D(2026, 10, 9), "AB2")],                   # fecha Excel invertida
    })
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=_indice(("AB2", "200.00", D(2026, 9, 10))))
    assert r["normalizado"] is True
    assert sorted(x["motivo"] for x in r["cambios"]) == sorted([nfd.MOTIVO_ANIO_2_DIGITOS, nfd.MOTIVO_INVERSION_DDMM])
    assert _fecha_celda_en_copia(r["ruta_cierre_local"], "SFC107", "C2") == D(2026, 9, 28)
    assert _fecha_celda_en_copia(r["ruta_cierre_local"], "SFC108", "C2") == D(2026, 9, 10)


def test_tiquipaya_misma_regla_sobre_sfc101_102_y_america_no_toca_hojas_de_tiquipaya(tmp_path):
    ruta = _cierre_dep(tmp_path, {
        "SFC101": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")],
        "SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")],
    })
    indice = _indice(("AB1", "100.00", D(2026, 9, 10)))
    rt = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "tiquipaya", indice_vouchers=indice)
    assert [c["hoja"] for c in rt["cambios"]] == ["SFC101"]
    ra = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "america", indice_vouchers=indice)
    assert [c["hoja"] for c in ra["cambios"]] == ["SFC107"]


def test_cierre_tiquipaya_sin_inversion_no_cambia_nada_aunque_haya_evidencia(tmp_path):
    """Sin regresion: un cierre normal (fecha ya correcta) no crea copia, aunque MACROS este disponible."""
    ruta = _cierre_dep(tmp_path, {"SFC101": [("DEPOSITO 1", 100.0, D(2026, 9, 10), "AB1")]})
    r = nfd.normalizar_cierre_para_precheck(ruta, "2026-09-09", "tiquipaya", indice_vouchers=_indice(("AB1", "100.00", D(2026, 9, 10))))
    assert r["normalizado"] is False and r["ruta_cierre_local"] == ruta and r["cambios"] == [] and r["revision"] == []


def test_indice_vouchers_macros_usa_el_lector_del_precheck(tmp_path):
    ruta_maestro = str(tmp_path / "MAESTRO.xlsm")
    fx.crear_maestro_unico(ruta_maestro, macros_filas=[
        ("2026-09-10", "3P90102072", "51230.10"), ("2026-09-11", "AB 1", "10.00"), ("2026-09-12", "AB1", "10.00"),
    ], atc_filas=[])
    idx = nfd.indice_vouchers_macros(ruta_maestro)
    assert idx[("3P90102072", "51230.10")] == [D(2026, 9, 10)]
    assert idx[("AB1", "10.00")] == [D(2026, 9, 11), D(2026, 9, 12)]                 # "AB 1" y "AB1" -> misma clave -> NO unico
    with pytest.raises(Exception):
        nfd.indice_vouchers_macros(str(tmp_path / "no_existe.xlsm"))


def test_normalizar_cierres_materializados_lee_el_maestro_materializado_una_sola_vez(tmp_path, monkeypatch):
    ruta_maestro = str(tmp_path / "MAESTRO.xlsm")
    fx.crear_maestro_unico(ruta_maestro, macros_filas=[("2026-09-10", "AB1", "100.00"), ("2026-09-11", "AB2", "200.00")], atc_filas=[])
    c1 = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")]}, nombre="CIERRE 09-09-2026.xlsm")
    c2 = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 200.0, D(2026, 11, 9), "AB2")]}, nombre="CIERRE 10-09-2026.xlsm")
    lecturas = []
    real = nfd.indice_vouchers_macros
    monkeypatch.setattr(nfd, "indice_vouchers_macros", lambda ruta: (lecturas.append(ruta), real(ruta))[1])
    items = [
        {"fecha": "2026-09-09", "estado_materializacion": "MATERIALIZADO", "ruta_cierre_local": c1, "ruta_maestro_local": ruta_maestro},
        {"fecha": "2026-09-10", "estado_materializacion": "MATERIALIZADO", "ruta_cierre_local": c2, "ruta_maestro_local": ruta_maestro},
    ]
    salida = nfd.normalizar_cierres_materializados(items, caja="america")
    assert lecturas == [ruta_maestro]                                                 # una sola lectura para ambos cierres
    assert all(s["ruta_cierre_local"].endswith(".normalizado.xlsm") for s in salida)
    assert all(s["normalizacion_fecha_deposito"][0]["clase"] == "NORMALIZAR_INVERSION_DDMM" for s in salida)
    assert "normalizacion_fecha_deposito_revision" not in salida[0]


def test_normalizar_cierres_materializados_maestro_ausente_no_corrige_y_reporta_revision(tmp_path):
    c1 = _cierre_dep(tmp_path, {"SFC107": [("DEPOSITO 1", 100.0, D(2026, 10, 9), "AB1")]})
    items = [{"fecha": "2026-09-09", "estado_materializacion": "MATERIALIZADO", "ruta_cierre_local": c1,
              "ruta_maestro_local": str(tmp_path / "no_existe.xlsm")}]
    salida = nfd.normalizar_cierres_materializados(items, caja="america")
    assert salida[0]["ruta_cierre_local"] == c1 and salida[0]["normalizacion_fecha_deposito"] == []
    assert salida[0]["normalizacion_fecha_deposito_revision"][0]["motivo"] == "MACROS_NO_DISPONIBLE_PARA_CONFIRMAR"


def _cierre_america_dep(ruta, fecha_dep, asignacion, importe):
    """CIERRE de América completo (SFC107/SFC108 + CI) con UN depósito en SFC107, lo mínimo que
    excel_io.leer_cierre exige para que el precheck pueda leerlo."""
    wb = openpyxl.Workbook()
    for k, nombre in enumerate(("SFC107", "SFC108")):
        ws = wb.active if k == 0 else wb.create_sheet(nombre)
        ws.title = nombre
        ws.append(["TOTAL MOVIMIENTO DEL DIA", "0.00"])
        ws.append(["COBROS ATC", "0.00"])
        ws.append(["TOTAL COMUNICACIONES INTERNAS", "0.00"])
        ws.append(["DOLARES", "0.00"])
        ws.append(["POSGRADO RESERVA", "0.00"])
        ws.append([None, None, None, None, None])
        ws.append(["COMPOSICION DE DEPOSITOS", "IMPORTE Bs", "FECHA DE DEPOSITO", "ASIGNACION", "BANCO"])
        if nombre == "SFC107":
            ws.append(["DEPOSITO 1", importe, fecha_dep, asignacion, "BNB 123"])
    for nombre in ("SFC107", "SFC108"):
        wb.create_sheet(f"COMUNICACIONES INTERNAS {nombre}").append(
            ["N°", "N° DE FACTURA", "TOTAL C.I.", "CUENTA CONTABLE BANCO", "ASIGNACION", "BANCO"])
    wb.save(ruta)


def _maestro_america(ruta, macros_filas, atc_filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tablas Dinamicas Profesional"
    ws.append(["Fecha", "Código de Asignación", "Créditos"])
    for fila in macros_filas:
        ws.append(list(fila))
    ws_atc = wb.create_sheet("ATC TIQUIPAYA")
    ws_atc.append(["FECHA", "TIPO", "CUENTA CONTABLE", "DETALLE", "MONTO", "ASIGNACION", "CAJA"])
    for fila in atc_filas:
        ws_atc.append(list(fila))
    wb.save(ruta)


def test_precheck_usa_la_copia_corregida_y_sin_ella_bloquea_con_mensaje_sobre_el_cierre(tmp_path):
    """Secuencia completa: original -> copia operativa -> normalizacion con evidencia MACROS -> precheck."""
    ruta_maestro = str(tmp_path / "MAESTRO.xlsm")
    _maestro_america(ruta_maestro,
                     macros_filas=[("2026-09-10", "AB1", "100.00"), ("2026-09-27", "ZZ", "1.00")],
                     atc_filas=[("2026-09-09", "BANCO (NETO)", "110103012", "ATC", "0.00", "ASIG", "AMERICA")])
    ruta = str(tmp_path / "CIERRE 09-09-2026.xlsm")
    _cierre_america_dep(ruta, D(2026, 10, 9), "AB1", 100.0)
    original = open(ruta, "rb").read()

    antes = pm.evaluar_cobertura_maestro(ruta_maestro, "2026-09-09", ruta, caja="america")
    assert antes["estado"] == pm.BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert antes["codigo_bloqueo"] == pm.MACROS_NO_CUBRE_FECHA_DEPOSITO
    assert antes["posible_inversion_ddmm"] == [{"fecha_deposito": "2026-10-09", "fecha_invertida": "2026-09-10"}]

    items = [{"fecha": "2026-09-09", "estado_materializacion": "MATERIALIZADO", "ruta_cierre_local": ruta,
              "ruta_maestro_local": ruta_maestro}]
    norm = nfd.normalizar_cierres_materializados(items, caja="america")
    assert norm[0]["ruta_cierre_local"].endswith(".normalizado.xlsm")
    anotado = pm.aplicar_precheck_maestro(norm, caja="america")[0]
    assert anotado["estado_precheck_maestro"] == pm.MAESTRO_APTO                       # el precheck leyó la COPIA corregida
    assert anotado["fecha_requerida_deposito"] == "2026-09-10"
    assert anotado["posible_inversion_ddmm"] == []
    assert open(ruta, "rb").read() == original                                          # el original (copia materializada) intacto


def test_sin_voucher_el_precheck_sigue_bloqueando_y_el_mensaje_apunta_a_la_fecha_del_cierre(tmp_path):
    ruta_maestro = str(tmp_path / "MAESTRO.xlsm")
    _maestro_america(ruta_maestro,
                     macros_filas=[("2026-09-27", "ZZ", "1.00")],                       # NO hay voucher AB1/100.00
                     atc_filas=[("2026-09-09", "BANCO (NETO)", "110103012", "ATC", "0.00", "ASIG", "AMERICA")])
    ruta = str(tmp_path / "CIERRE 09-09-2026.xlsm")
    _cierre_america_dep(ruta, D(2026, 10, 9), "AB1", 100.0)
    items = [{"fecha": "2026-09-09", "estado_materializacion": "MATERIALIZADO", "ruta_cierre_local": ruta,
              "ruta_maestro_local": ruta_maestro}]
    norm = nfd.normalizar_cierres_materializados(items, caja="america")
    assert norm[0]["ruta_cierre_local"] == ruta                                          # NO se corrigió
    anotado = pm.aplicar_precheck_maestro(norm, caja="america")[0]
    assert anotado["estado_precheck_maestro"] == pm.BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert anotado["codigo_bloqueo_precheck"] == pm.MACROS_NO_CUBRE_FECHA_DEPOSITO
    msg = anotado["mensaje_precheck_maestro"]
    assert "día y mes invertidos" in msg and "2026-10-09 ↔ 2026-09-10" in msg
    assert "NO se corrigió automáticamente" in msg
    assert "actualice MACROS en Drive y vuelva a procesar" not in msg                    # no afirma que MACROS esté desactualizado
    assert norm[0]["normalizacion_fecha_deposito_revision"][0]["motivo"] == "SIN_VOUCHER_EN_MACROS"


# ---------------------------------------------------------------------------
# End-to-end por dev_api.procesar_lote (misma cadena que /procesar): original
# -> copia materializada -> normalizacion con evidencia MACROS -> precheck ->
# motor. TIQUIPAYA (misma logica comun; sin datos reales de ninguna caja).
# ---------------------------------------------------------------------------

def _sfc_dep(importe, fecha, asignacion):
    return {"total_movimiento": importe, "cobros_atc": "0.00", "dolares": "0.00",
            "depositos": [{"deposito": "DEPOSITO 1", "importe": importe, "fecha": fecha, "asignacion": asignacion, "banco": "BNB"}]}


_SFC_SIN_DEP = {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00", "depositos": []}


def _procesar_lote_12(tmp_path, macros_filas, fecha_dep):
    import run_batch
    from v3 import dev_api
    carpeta = tmp_path / "cierres_origen"
    carpeta.mkdir()
    ruta_cierre = str(carpeta / run_batch.nombre_cierre_esperado("2026-09-12"))
    fx.crear_cierre(ruta_cierre, _sfc_dep("11096.00", fecha_dep, "3P9E113705"), _SFC_SIN_DEP)
    ruta_maestro = str(tmp_path / "maestro.xlsm")
    fx.crear_maestro_unico(ruta_maestro, macros_filas=macros_filas, atc_filas=[])
    plantilla = tmp_path / "Plantilla.xlsx"
    fx.crear_plantilla_sap(str(plantilla))
    base = str(tmp_path / "dev")
    lote = dev_api.crear_lote_pendiente("2026-09-12", "2026-09-12", "auditor.dev", base)
    bytes_origen = open(ruta_cierre, "rb").read()
    resultado = dev_api.procesar_lote(lote["lote_id"], base, str(carpeta), ruta_maestro, str(plantilla))
    assert open(ruta_cierre, "rb").read() == bytes_origen                          # el origen nunca se toca
    return resultado["cierres"][0]


def test_dev_api_inversion_confirmada_el_precheck_y_el_motor_usan_la_copia_corregida(tmp_path):
    c = _procesar_lote_12(tmp_path, [("2026-09-12", "3P9E113705", "11096.00"), ("2026-09-17", "3P97000009", "7.00")],
                          datetime.date(2026, 12, 9))                              # 09/12 en vez de 12/09
    assert c["ruta_cierre_local"].endswith(".normalizado.xlsm")
    assert c["normalizacion_fecha_deposito"][0]["clase"] == "NORMALIZAR_INVERSION_DDMM"
    assert c["estado_precheck_maestro"] == pm.MAESTRO_APTO
    assert c["fecha_requerida_deposito"] == "2026-09-12"
    assert c["estado_final"] == "LISTO_PARA_PUBLICAR"


def test_dev_api_sin_voucher_no_corrige_y_bloquea_con_mensaje_sobre_el_cierre(tmp_path):
    c = _procesar_lote_12(tmp_path, [("2026-09-17", "3P97000009", "7.00")], datetime.date(2026, 12, 9))
    assert not c["ruta_cierre_local"].endswith(".normalizado.xlsm")
    assert c["estado_precheck_maestro"] == pm.BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA
    assert c["codigo_bloqueo_precheck"] == pm.MACROS_NO_CUBRE_FECHA_DEPOSITO
    assert c["posible_inversion_ddmm"] == [{"fecha_deposito": "2026-12-09", "fecha_invertida": "2026-09-12"}]
    assert "día y mes invertidos" in c["mensaje_precheck_maestro"]
    assert c["normalizacion_fecha_deposito_revision"][0]["motivo"] == "SIN_VOUCHER_EN_MACROS"
