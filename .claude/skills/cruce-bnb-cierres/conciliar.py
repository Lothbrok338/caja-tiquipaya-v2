"""Cruce diario de depósitos declarados en cierres de cajeros vs abonos BNB.

Uso:
    python conciliar.py --cierres DIR --macro ARCHIVO.xlsm \
        --desde 27-10-2026 --hasta 28-10-2026 --salida DIR_SALIDA

--cierres  carpeta con los cierres descargados, nombrados "CIERRE dd-mm-aaaa.xlsm"
--macro    libro de macros con la hoja "Tablas Dinamicas Profesional" (detalle BNB)
--desde / --hasta  ventana por fecha de caja del cierre (incluye ambos extremos)
--cajas    cajas a revisar (por defecto SFC101,SFC102,SFC107,SFC108)

Escribe resultado.json y reporte.html en --salida.
"""
import argparse
import datetime as dt
import glob
import html
import json
import os
import re

import openpyxl

HOJA_BNB = "Tablas Dinamicas Profesional"
CAJAS_DEFAULT = ["SFC101", "SFC102", "SFC107", "SFC108"]


def parse_fecha(valor):
    """Acepta datetime, date o texto dd/mm/aa(aa). Devuelve date o None."""
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, str):
        m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{2,4})", valor)
        if m:
            anio = int(m.group(3))
            anio = anio + 2000 if anio < 100 else anio
            try:
                return dt.date(anio, int(m.group(2)), int(m.group(1)))
            except ValueError:
                return None
    return None


def fecha_desde_nombre(nombre):
    m = re.search(r"CIERRE (\d{2})-(\d{2})-(\d{4})", os.path.basename(nombre))
    if not m:
        return None
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))


def leer_bnb(ruta):
    """Devuelve lista de abonos con asignación y monto, y la cobertura de fechas."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb[HOJA_BNB]
    abonos = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        asig, monto = r[2], r[3]
        if isinstance(asig, str) and asig.strip().startswith("3P") and isinstance(monto, (int, float)):
            abonos.append({
                "fecha": parse_fecha(r[0]),
                "asig": asig.strip(),
                "monto": float(monto),
                "caja": (r[5] or "").strip() if isinstance(r[5], str) else r[5],
            })
    fechas = [a["fecha"] for a in abonos if a["fecha"]]
    cobertura = (min(fechas), max(fechas)) if fechas else (None, None)
    return abonos, cobertura


def leer_depositos(ruta, cajas):
    """Depósitos con código de asignación en las hojas de caja del cierre."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    depositos = []
    for sn in cajas:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and re.fullmatch(r"DEPOSITO \d", c.value.strip()):
                    importe = ws.cell(row=c.row, column=c.column + 1).value
                    fecha = ws.cell(row=c.row, column=c.column + 2).value
                    asig = ws.cell(row=c.row, column=c.column + 3).value
                    if isinstance(asig, str) and asig.strip() and isinstance(importe, (int, float)) and importe > 0:
                        depositos.append({
                            "caja": sn,
                            "deposito": c.value.strip(),
                            "importe": float(importe),
                            "fecha_dep": parse_fecha(fecha),
                            "fecha_dep_texto": fecha if isinstance(fecha, str) else (fecha.strftime("%d-%m-%Y") if fecha else None),
                            "asig": asig.strip(),
                        })
    return depositos


def es_dia_mes_invertido(declarada, real):
    return declarada and real and declarada.day == real.month and declarada.month == real.day


def clasificar(dep, bnb_por_asig, cobertura_hasta):
    """Devuelve (estado, detalle) para un depósito declarado."""
    abono = bnb_por_asig.get(dep["asig"])
    if abono:
        if abs(abono["monto"] - dep["importe"]) >= 0.01:
            return "MONTO_DISTINTO", f"BNB abonó {abono['monto']:,.2f}"
        if dep["fecha_dep"] == abono["fecha"]:
            return "OK", ""
        if es_dia_mes_invertido(dep["fecha_dep"], abono["fecha"]):
            return "FECHA_INVERTIDA", f"Cierre: {dep['fecha_dep_texto']} · BNB: {abono['fecha']:%d-%m-%Y}"
        return "FECHA_DISTINTA", f"Cierre: {dep['fecha_dep_texto']} · BNB: {abono['fecha']:%d-%m-%Y}"
    if dep["fecha_dep"] and cobertura_hasta and dep["fecha_dep"] > cobertura_hasta:
        return "PENDIENTE", f"El detalle BNB llega hasta {cobertura_hasta:%d-%m-%Y}"
    if dep["fecha_dep"] is None:
        return "REVISAR", "Fecha de depósito ilegible en el cierre"
    return "SIN_ABONO", "No existe abono BNB con este código de asignación"


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--cierres", required=True)
    p.add_argument("--macro", required=True)
    p.add_argument("--desde", required=True)
    p.add_argument("--hasta", required=True)
    p.add_argument("--salida", required=True)
    p.add_argument("--cajas", default=",".join(CAJAS_DEFAULT))
    args = p.parse_args()

    desde = dt.datetime.strptime(args.desde, "%d-%m-%Y").date()
    hasta = dt.datetime.strptime(args.hasta, "%d-%m-%Y").date()
    cajas = [c.strip() for c in args.cajas.split(",") if c.strip()]
    os.makedirs(args.salida, exist_ok=True)

    abonos, (cob_desde, cob_hasta) = leer_bnb(args.macro)
    bnb_por_asig = {}
    for a in abonos:
        bnb_por_asig.setdefault(a["asig"], a)

    resultados = []
    cierres_usados = []
    for ruta in sorted(glob.glob(os.path.join(args.cierres, "*.xlsm"))):
        fecha_caja = fecha_desde_nombre(ruta)
        if fecha_caja is None or not (desde <= fecha_caja <= hasta):
            continue
        cierres_usados.append({"archivo": os.path.basename(ruta), "fecha_caja": fecha_caja.isoformat()})
        for dep in leer_depositos(ruta, cajas):
            estado, detalle = clasificar(dep, bnb_por_asig, cob_hasta)
            resultados.append({
                "cierre": os.path.basename(ruta),
                "fecha_caja": fecha_caja.isoformat(),
                **dep,
                "fecha_dep": dep["fecha_dep"].isoformat() if dep["fecha_dep"] else None,
                "estado": estado,
                "detalle": detalle,
            })

    conteo = {}
    for r in resultados:
        conteo[r["estado"]] = conteo.get(r["estado"], 0) + 1

    salida = {
        "ventana": {"desde": args.desde, "hasta": args.hasta},
        "cobertura_bnb": {
            "desde": cob_desde.isoformat() if cob_desde else None,
            "hasta": cob_hasta.isoformat() if cob_hasta else None,
        },
        "cierres": cierres_usados,
        "cajas": cajas,
        "conteo": conteo,
        "depositos": resultados,
    }
    with open(os.path.join(args.salida, "resultado.json"), "w", encoding="utf-8") as f:
        json.dump(salida, f, ensure_ascii=False, indent=2)

    with open(os.path.join(args.salida, "reporte.html"), "w", encoding="utf-8") as f:
        f.write(generar_html(salida))
    print(json.dumps(conteo, ensure_ascii=False))


ETIQUETA = {
    "OK": ("ok", "Cuadra"),
    "FECHA_INVERTIDA": ("bad", "Fecha invertida"),
    "FECHA_DISTINTA": ("bad", "Fecha distinta"),
    "MONTO_DISTINTO": ("bad", "Monto distinto"),
    "SIN_ABONO": ("bad", "Sin abono BNB"),
    "PENDIENTE": ("warn", "Pendiente de extracto"),
    "REVISAR": ("warn", "Revisar"),
}


def generar_html(s):
    filas = []
    for r in s["depositos"]:
        if r["estado"] == "OK":
            continue
        clase, texto = ETIQUETA.get(r["estado"], ("warn", r["estado"]))
        filas.append(
            "<tr>"
            f"<td><span class='tag {clase}'>{texto}</span></td>"
            f"<td>{html.escape(r['cierre'])}</td>"
            f"<td>{html.escape(r['caja'])}</td>"
            f"<td>{html.escape(r['asig'])}</td>"
            f"<td class='num'>{r['importe']:,.2f}</td>"
            f"<td>{html.escape(r['detalle'])}</td>"
            "</tr>"
        )
    total = len(s["depositos"])
    ok = s["conteo"].get("OK", 0)
    tabla = "".join(filas) or "<tr><td colspan='6'>Sin discrepancias ni pendientes en la ventana.</td></tr>"
    cob = s["cobertura_bnb"]
    return f"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<title>Cruce BNB {html.escape(s['ventana']['desde'])} al {html.escape(s['ventana']['hasta'])}</title>
<style>
body{{font-family:Arial,sans-serif;color:#1d2433;margin:28px;font-size:13px}}
h1{{font-size:20px;margin:0 0 4px}} .sub{{color:#6b7385;margin-bottom:16px}}
.cards{{display:flex;gap:10px;margin:12px 0}} .card{{flex:1;background:#f6f7fb;border-radius:8px;padding:10px 12px}}
.card .n{{font-size:22px;font-weight:700}} .card .t{{color:#6b7385;font-size:11px;text-transform:uppercase}}
table{{width:100%;border-collapse:collapse;margin-top:8px}} th,td{{padding:6px;border-bottom:1px solid #e3e6ee;text-align:left}}
th{{background:#f6f7fb;font-size:11px;text-transform:uppercase;color:#6b7385}} td.num{{text-align:right}}
.tag{{padding:2px 8px;border-radius:10px;font-size:11px;font-weight:600}}
.tag.ok{{background:#dcf3e7;color:#17693f}} .tag.bad{{background:#fde3e1;color:#a32a27}} .tag.warn{{background:#fff3cc;color:#7a5a00}}
</style></head><body>
<h1>Cruce de depósitos vs extracto BNB</h1>
<div class="sub">Cierres del {html.escape(s['ventana']['desde'])} al {html.escape(s['ventana']['hasta'])} · Cajas: {html.escape(', '.join(s['cajas']))} · Detalle BNB cubre {html.escape(str(cob['desde']))} a {html.escape(str(cob['hasta']))}</div>
<div class="cards">
<div class="card"><div class="n">{total}</div><div class="t">Depósitos revisados</div></div>
<div class="card"><div class="n">{ok}</div><div class="t">Cuadran</div></div>
<div class="card"><div class="n">{total - ok}</div><div class="t">Observados</div></div>
</div>
<table><thead><tr><th>Estado</th><th>Cierre</th><th>Caja</th><th>Asignación</th><th>Importe</th><th>Detalle</th></tr></thead>
<tbody>{tabla}</tbody></table>
</body></html>"""


if __name__ == "__main__":
    main()
