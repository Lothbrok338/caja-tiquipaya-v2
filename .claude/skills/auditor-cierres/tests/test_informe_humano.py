"""Pruebas del informe humano (PDF principal + DOCX editable). Solo presentacion: ninguna regla de auditoria cambia."""
import collections
import datetime
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings
import zipfile

warnings.simplefilter("ignore")
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import auditar as A  # noqa: E402
import controles_lectura as cl  # noqa: E402
import informe_humano as ih  # noqa: E402
from test_controles import BASE, ci, cierre, macros_xlsx  # noqa: E402

D = datetime.date


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="test_ih_")
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def ruta(self, *p):
        return os.path.join(self.tmp, *p)

    def corrida(self, casos, nombre="rep", extra_macros=None, **kw):
        """casos: [(caja, dia, kwargs de cierre())]. Devuelve (out, modelo_del_json)."""
        mac = self.ruta("MACROS SEPTIEMBRE.xlsm")
        macros_xlsx(mac, [(D(2026, 9, 10), "AB1", 100)] + (extra_macros or []))
        grupos = []
        for caja, dia, args in casos:
            d = self.ruta(caja)
            os.makedirs(d, exist_ok=True)
            r = os.path.join(d, "CIERRE %02d-09-2026.xlsm" % dia)
            cierre(r, **args)
            grupos.append((caja, r))
        rep = self.ruta(nombre)
        out = A.auditar(grupos, [mac], salida_dir=os.path.join(rep, "CIERRES_NORMALIZADOS"), reporte_dir=rep,
                        hoy=D(2026, 9, 26), **kw)
        with zipfile.ZipFile(out["archivos"]["zip"]) as z:
            modelo = json.loads(z.read("RESULTADOS_AUDITORIA.json").decode("utf-8"))["informe_humano"]
        return out, modelo


def _escenario():
    total_mal = dict(BASE)
    total_mal["TOTAL MOVIMIENTO DEL DIA"] = 1800.04
    return [
        ("tiquipaya", 9, dict(deposito=(100, D(2026, 9, 10), "AB1"))),                                   # sin observaciones
        ("tiquipaya", 10, dict(deposito=(100, D(2026, 9, 10), "AB1"), sfc={"SFC101": total_mal, "SFC102": dict(BASE)})),  # diferencia 0,04
        ("tiquipaya", 11, dict(deposito=(100, "23/09/26", "NO-ESTA"))),                                   # texto sin voucher -> normaliza
        ("tiquipaya", 12, dict(deposito=(100, D(2026, 10, 9), "AB1"))),                                   # DD/MM invertido confirmado
        ("tiquipaya", 14, dict(deposito=(100, D(2026, 9, 9), "AB1"))),                                    # tolerancia 1 dia (sin cambio)
        ("tiquipaya", 15, dict(deposito=(100, "09/09/26", "AB1"))),                                       # texto dentro de tolerancia
        ("tiquipaya", 16, dict(deposito=(100, D(2026, 9, 16), "ZZ9"))),                                   # voucher no encontrado
    ]


class TestContenido(Base):
    @classmethod
    def setUpClass(cls):
        cls._t = tempfile.mkdtemp(prefix="test_ih_cls_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls._t, True)

    def setUp(self):
        super().setUp()
        self.out, self.m = self.corrida(_escenario())

    def test_resumen_y_estados_humanos(self):
        self.assertEqual(self.m["resumen"], [["Cierres revisados", "7"], ["Sin observaciones", "1"], ["Requieren revisión", "2"],
                                             ["Informativos", "4"], ["Correcciones automáticas", "3"]])
        estados = {f[1]: f[2] for f in self.m["cierres"]}
        self.assertEqual(estados["09/09/2026"], "SIN OBSERVACIONES")
        self.assertEqual(estados["10/09/2026"], "REQUIERE REVISIÓN")
        self.assertEqual(estados["16/09/2026"], "REQUIERE REVISIÓN")
        for d in ("11/09/2026", "12/09/2026", "14/09/2026", "15/09/2026"):
            self.assertEqual(estados[d], "OBSERVACIÓN INFORMATIVA")
        self.assertTrue(set(estados.values()) <= {"SIN OBSERVACIONES", "REQUIERE REVISIÓN", "OBSERVACIÓN INFORMATIVA"})
        self.assertEqual(self.m["meta"][1], ["Periodo revisado", "09/09/2026 al 16/09/2026"])
        self.assertEqual(self.m["meta"][2], ["Cajas revisadas", "Caja Tiquipaya"])
        self.assertEqual(self.m["meta"][3], ["Cantidad total de cierres", "7"])

    def test_estado_del_informe_coincide_con_el_de_la_auditoria(self):
        """El informe humano no cambia ninguna decision: 'requiere revision' == cierres REVISAR de la auditoria."""
        self.assertEqual(self.out["cierres"]["requieren_revision"], 2)
        self.assertEqual(self.out["cierres"]["sin_observaciones"], 5)      # los informativos no cuentan como revision

    def test_casos_que_requieren_revision_en_lenguaje_natural(self):
        filas = {(f[1]): f for f in self.m["revisar"]}
        self.assertEqual(filas["10/09/2026"][2], "Existe una diferencia de Bs 0,04 entre el total calculado y el registrado en SFC101.")
        self.assertEqual(filas["10/09/2026"][3], "Revisar los importes de SFC101 antes de procesar el cierre.")
        self.assertIn("El voucher informado no fue encontrado en MACROS.", filas["16/09/2026"][2])
        self.assertEqual(len(self.m["revisar"]), 2)

    def test_correcciones_automaticas_con_valor_original_y_normalizado(self):
        c = {f[1]: f for f in self.m["correcciones"]}
        self.assertEqual(c["11/09/2026"], ["Tiquipaya", "11/09/2026", "Fecha de depósito\n(SFC101, celda F3)", "23/09/26", "23/09/2026",
                                           "Se normalizó el formato sin cambiar la fecha declarada. No se encontró el voucher en "
                                           "MACROS, por lo que la fecha no pudo contrastarse."])
        # inversion DD/MM confirmada por voucher, explicada en lenguaje natural
        self.assertEqual((c["12/09/2026"][3], c["12/09/2026"][4]), ("09/10/2026", "10/09/2026"))
        self.assertEqual(c["12/09/2026"][5], "Se invirtieron el día y el mes: la fecha registrada era 09/10/2026 y el voucher de MACROS "
                                             "confirma 10/09/2026.")
        # correccion dentro de la tolerancia de +-1 dia de MACROS
        self.assertEqual((c["15/09/2026"][3], c["15/09/2026"][4]), ("09/09/26", "09/09/2026"))
        self.assertEqual(c["15/09/2026"][5], "Se normalizó el formato sin cambiar la fecha declarada. El voucher en MACROS indica "
                                             "10/09/2026 (un día de diferencia, dentro de la tolerancia permitida).")
        self.assertEqual(len(self.m["correcciones"]), 3)

    def test_observaciones_informativas_no_requieren_accion(self):
        self.assertEqual(len(self.m["informativas"]), 1)
        caja, fecha, obs = self.m["informativas"][0]
        self.assertEqual((caja, fecha), ("Tiquipaya", "14/09/2026"))
        self.assertIn("dentro de la tolerancia permitida", obs)
        self.assertTrue(obs.endswith("Acción requerida: Ninguna."))

    def test_conclusion_a_partir_de_resultados_reales(self):
        c = self.m["conclusion"]
        self.assertTrue(c.startswith("Se revisaron 7 cierres de caja (Caja Tiquipaya)."))
        self.assertIn("La mayoría no presenta observaciones que requieran intervención.", c)
        self.assertIn("deben ser verificados antes de considerar concluido el control.", c)
        self.assertIn("Las normalizaciones automáticas realizadas (3) corrigieron únicamente la fecha de depósito y no alteraron "
                      "los importes registrados.", c)
        self.assertNotRegex(c.lower(), r"grave|alarm|urgente|cr[ií]tico|peligro")

    def test_sin_codigos_internos_ni_trazas_ni_rutas(self):
        textos = " \n".join(ih.textos_del_modelo(self.m)).replace("AUDITORIA_CIERRES_RESULTADOS.zip", "")
        for prohibido in (r"[A-Z]{2,}_[A-Z0-9_]{2,}", r"Error\b", r"Traceback", r"\.py\b", r"[A-Za-z]:\\", r"/Users/",
                          r"Exception", r"\bNone\b", r"\{|\}", r"celda_", r"NORMALIZAR_INVERSION", r"DENTRO_TOLERANCIA"):
            self.assertIsNone(re.search(prohibido, textos), "aparece %r en el informe" % prohibido)

    def test_todo_el_texto_es_representable_en_la_tipografia_del_pdf(self):
        for t in ih.textos_del_modelo(self.m):
            t.encode("cp1252")

    def test_el_excel_no_es_el_informe_principal(self):
        rep = os.path.dirname(self.out["archivos"]["zip"])
        self.assertEqual(sorted(os.listdir(rep)), ["AUDITORIA_CIERRES_RESULTADOS.zip", "PARA_CAJA_AMERICA.pdf",
                                                    "PARA_CAJA_TIQUIPAYA.pdf", "RESUMEN_GABO.pdf"])
        self.assertEqual(os.path.basename(self.out["reporte"]), "RESUMEN_GABO.pdf")
        with zipfile.ZipFile(self.out["archivos"]["zip"]) as z:
            self.assertIn("DETALLE_TECNICO_AUDITORIA.xlsx", z.namelist())
            self.assertIn("RECURSOS_USADOS.txt", z.namelist())
            self.assertIn("INFORME_AUDITORIA_CIERRES.pdf", z.namelist())
            self.assertIn("INFORME_AUDITORIA_CIERRES.docx", z.namelist())


class TestMismoContenido(Base):
    def setUp(self):
        super().setUp()
        self.out, self.m = self.corrida(_escenario(), extra_macros=[])
        self.pdf, self.docx = self.ruta("informe_tecnico.pdf"), self.ruta("informe_tecnico.docx")
        with zipfile.ZipFile(self.out["archivos"]["zip"]) as z:
            with open(self.pdf, "wb") as f:
                f.write(z.read("INFORME_AUDITORIA_CIERRES.pdf"))
            with open(self.docx, "wb") as f:
                f.write(z.read("INFORME_AUDITORIA_CIERRES.docx"))

    def test_docx_tiene_exactamente_el_contenido_del_modelo(self):
        from docx import Document
        doc = Document(self.docx)
        m = self.m
        esperado = [
            [[k, v] for k, v in m["meta"]],
            [[v for _k, v in m["resumen"]], [k for k, _v in m["resumen"]]],
            [ih.COLS_CIERRES] + m["cierres"],
            [ih.COLS_REVISAR] + m["revisar"],
            *([b["columnas"]] + b["filas"] for b in m["vouchers"] if b["filas"]),      # detalle de vouchers no encontrados
            [ih.COLS_CORRECCIONES] + m["correcciones"],
            [ih.COLS_INFO] + m["informativas"],
        ]
        if m["anexo"]["alquileres"]:
            esperado.append([ih.COLS_ALQUILERES] + m["anexo"]["alquileres"] + [["Total del periodo", "", m["anexo"]["total_alquileres"]]])
        real = [[[c.text for c in fila.cells] for fila in t.rows] for t in doc.tables]
        self.assertEqual(real, esperado)
        parrafos = "\n".join(p.text for p in doc.paragraphs)
        if m["vouchers"]:
            self.assertIn(ih.T_SEC2B, parrafos)
        for texto in (m["titulo"], ih.T_SEC1, ih.T_SEC2, ih.T_SEC3, ih.T_SEC4, ih.T_SEC5, ih.T_ANEXO, m["conclusion"], *m["leyenda"]):
            self.assertIn(texto, parrafos)

    def test_pdf_contiene_todo_el_texto_del_modelo(self):
        exe = shutil.which("pdftotext")
        if not exe:
            self.skipTest("pdftotext no disponible")
        txt = subprocess.run([exe, "-enc", "UTF-8", self.pdf, "-"], capture_output=True, text=True, encoding="utf-8").stdout
        modelo = collections.Counter(" ".join(ih.textos_del_modelo(self.m)).split())
        en_pdf = collections.Counter(txt.split())
        self.assertEqual(modelo - en_pdf, collections.Counter(), "palabras del modelo ausentes en el PDF")
        # y el PDF no agrega contenido propio salvo el pie y la numeracion de paginas
        sobrantes = en_pdf - modelo
        encabezados = {w for cols in (ih.COLS_CIERRES, ih.COLS_REVISAR, ih.COLS_VOUCHER, ih.COLS_CORRECCIONES, ih.COLS_INFO,
                               ih.COLS_ALQUILERES)
                       for c in cols for w in c.split()}                     # se repiten en cada pagina que continua la tabla
        permitidas = {"CAJAS", "GABO", "—", "Auditoría", "de", "cierres", "Página", "1", "2", "3"} | encabezados
        self.assertTrue(set(sobrantes) <= permitidas, sorted(set(sobrantes) - permitidas))
        self.assertIn("Página 1 de", txt)

    def test_pdf_es_a4_vertical(self):
        with open(self.pdf, "rb") as f:
            datos = f.read()
        self.assertTrue(datos.startswith(b"%PDF"))
        caja = re.search(rb"/MediaBox\s*\[\s*0\s+0\s+([\d.]+)\s+([\d.]+)\s*\]", datos)
        self.assertIsNotNone(caja)
        self.assertAlmostEqual(float(caja.group(1)), 595.27, delta=0.5)
        self.assertAlmostEqual(float(caja.group(2)), 841.89, delta=0.5)

    def test_docx_es_a4_con_tablas_que_repiten_encabezado_y_no_parten_filas_y_pie_numerado(self):
        from docx import Document
        doc = Document(self.docx)
        s = doc.sections[0]
        self.assertAlmostEqual(s.page_width.cm, 21.0, places=1)
        self.assertAlmostEqual(s.page_height.cm, 29.7, places=1)
        self.assertGreater(s.page_height, s.page_width)
        datos = doc.tables[2:]                       # cierres, revisar, correcciones, informativas, alquileres
        self.assertGreaterEqual(len(datos), 4)
        for t in datos:
            self.assertIsNotNone(t.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
            self.assertTrue(t.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
            for fila in t.rows:
                self.assertTrue(fila._tr.xpath("./w:trPr/w:cantSplit"))
        pie = s.footer._element.xml
        self.assertIn("CAJAS GABO — Auditoría de cierres", "".join(p.text for p in s.footer.paragraphs))
        self.assertIn("PAGE", pie)
        self.assertIn("NUMPAGES", pie)

    def test_pdf_y_docx_estan_en_el_zip_identicos_a_los_entregados(self):
        with zipfile.ZipFile(self.out["archivos"]["zip"]) as z:
            self.assertEqual(z.read("INFORME_AUDITORIA_CIERRES.pdf"), open(self.pdf, "rb").read())
            self.assertEqual(z.read("INFORME_AUDITORIA_CIERRES.docx"), open(self.docx, "rb").read())


class TestCasosLimite(Base):
    def test_sin_hallazgos_las_secciones_dicen_que_no_hay_nada(self):
        out, m = self.corrida([("tiquipaya", 9, dict(deposito=(100, D(2026, 9, 10), "AB1")))])
        self.assertEqual(m["resumen"], [["Cierres revisados", "1"], ["Sin observaciones", "1"], ["Requieren revisión", "0"],
                                        ["Informativos", "0"], ["Correcciones automáticas", "0"]])
        self.assertEqual(m["revisar"] + m["correcciones"] + m["informativas"], [])
        self.assertIn("Ningún cierre presenta observaciones que requieran intervención.", m["conclusion"])
        self.assertIn("No se realizaron correcciones automáticas.", m["conclusion"])
        self.assertEqual(m["vacios"]["revisar"], "No hay casos que requieran revisión.")
        from docx import Document
        with zipfile.ZipFile(out["archivos"]["zip"]) as z:
            doc = Document(io.BytesIO(z.read("INFORME_AUDITORIA_CIERRES.docx")))
        texto = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("No hay casos que requieran revisión.", texto)
        self.assertIn("No se realizaron correcciones automáticas.", texto)

    def test_muchas_filas_iguales_se_agrupan_en_una_sola_fila_del_informe(self):
        filas = [ci(i, 10, "BISA", "110103032", "ABC") for i in range(1, 31)]      # 30 asignaciones invalidas
        out, m = self.corrida([("tiquipaya", 9, dict(deposito=(100, D(2026, 9, 10), "AB1"), cis={"SFC102": filas}))])
        c4 = [f for f in m["revisar"] if "asignación no tiene el formato esperado" in f[2]]
        self.assertEqual(len(c4), 1)
        self.assertIn("y 25 más", c4[0][2])
        self.assertLess(len(c4[0][2]), 900)

    def test_sin_macros_se_explica_sin_terminologia_tecnica(self):
        r = self.ruta("CIERRE 23-09-2026.xlsm")
        cierre(r, deposito=(100, "23/09/26", "AB1"))
        out = A.auditar([("tiquipaya", r)], [], salida_dir=os.path.join(self.ruta("rep"), "CN"), reporte_dir=self.ruta("rep"),
                        hoy=D(2026, 9, 26))
        with zipfile.ZipFile(out["archivos"]["zip"]) as z:
            m = json.loads(z.read("RESULTADOS_AUDITORIA.json").decode("utf-8"))["informe_humano"]
        self.assertIn("Las fechas de depósito no se contrastaron con los vouchers porque no se dispuso del archivo MACROS.",
                      m["revisar"][0][2])
        self.assertEqual(m["cierres"][0][2], "REQUIERE REVISIÓN")
        self.assertEqual(len(m["correcciones"]), 1)          # igual normaliza el formato

    def test_si_falla_el_pdf_tecnico_se_avisa_y_el_detalle_permanece_en_el_zip(self):
        original = ih.escribir_pdf

        def sin_reportlab(m, ruta):
            raise ImportError("No module named 'reportlab'")

        ih.escribir_pdf = sin_reportlab
        try:
            out, _m = self.corrida([("tiquipaya", 9, dict(deposito=(100, "23/09/26", "AB1")))])
        finally:
            ih.escribir_pdf = original
        self.assertTrue(os.path.isfile(out["archivos"]["resumen_gabo"]))
        self.assertTrue(any("reportlab" in x for x in out["informe_humano_incidencias"]))
        self.assertTrue(os.path.isfile(out["archivos"]["zip"]))
        with zipfile.ZipFile(out["archivos"]["zip"]) as z:
            self.assertNotIn("INFORME_AUDITORIA_CIERRES.pdf", z.namelist())
            self.assertIn("INFORME_AUDITORIA_CIERRES.docx", z.namelist())
            self.assertIn("DETALLE_TECNICO_AUDITORIA.xlsx", z.namelist())
        self.assertEqual(sorted(os.listdir(os.path.dirname(out["archivos"]["zip"]))),
                         ["AUDITORIA_CIERRES_RESULTADOS.zip", "PARA_CAJA_AMERICA.pdf", "PARA_CAJA_TIQUIPAYA.pdf", "RESUMEN_GABO.pdf"])


class TestTraduccionDelMotor(unittest.TestCase):
    """Cada motivo de bloqueo del motor CAJAS GABO se explica con su causa real (nunca 'MACROS no cubre' por defecto)."""

    def h(self, **datos):
        base = {"tipo": "motor_observaciones", "estado_maestro": "BLOQUEADO_MAESTRO_COBERTURA_NO_CONFIRMADA", "codigo_bloqueo": None,
                "fecha_maxima_macros": None, "fecha_requerida_deposito": None, "fecha_cierre": "2026-09-09", "mensaje": "",
                "observaciones": []}
        base.update(datos)
        return {"fecha": D(2026, 9, 9), "caja": "tiquipaya", "sfc": None, "control": 7, "resultado": cl.REVISAR, "hallazgo": "x",
                "accion": "y", "fila": None, "datos": base}

    def texto(self, **datos):
        return ih.traducir([self.h(**datos)])[0]["ocurrio"]

    def test_motivos(self):
        self.assertIn("no pudo leer el archivo mensual de datos (MACROS / ATC)", self.texto(mensaje="MAESTRO_ILEGIBLE: ValueError: x"))
        self.assertIn("MACROS no tiene movimientos registrados", self.texto(mensaje="MAESTRO_SIN_FECHAS_REGISTRADAS: x"))
        self.assertEqual(self.texto(fecha_maxima_macros="2026-09-05", mensaje="El maestro no tiene registros de MACROS ..."),
                         "MACROS no cubre todavía la fecha del cierre (llega hasta el 05/09/2026).")
        self.assertEqual(self.texto(codigo_bloqueo="MACROS_NO_CUBRE_FECHA_DEPOSITO", fecha_maxima_macros="2026-09-20",
                                    fecha_requerida_deposito="2026-09-23"),
                         "MACROS no cubre todavía la fecha requerida para validar el cierre (los depósitos llegan hasta el 23/09/2026 "
                         "y MACROS hasta el 20/09/2026).")
        self.assertIn("información ATC de la fecha del cierre", self.texto(fecha_maxima_macros="2026-09-20",
                                                                          mensaje="No se encontró información ATC del maestro ..."))
        self.assertIn("no pudo determinar si el cierre tuvo movimiento ATC",
                      self.texto(fecha_maxima_macros="2026-09-20", mensaje="No se pudo leer el cierre para determinar si tuvo ..."))
        self.assertIn("no pudo confirmar la cobertura de los datos mensuales", self.texto(fecha_maxima_macros="2026-09-20", mensaje="?"))
        self.assertIn("año distinto al del cierre (15/09/2016)",
                      self.texto(estado_maestro="MAESTRO_APTO", observaciones=["FECHA_DEPOSITO_ANOMALA: 15/09/2016"]))


if __name__ == "__main__":
    unittest.main()
