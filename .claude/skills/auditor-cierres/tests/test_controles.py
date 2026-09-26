"""Pruebas de los Controles 1, 2, 3, 4 y 6 + reporte + salida compacta.

Ejecutar:  python -m unittest discover -s tests -v   (desde la carpeta auditor-cierres)
Fixtures sinteticos (openpyxl solo para CREARLOS en una carpeta temporal).
"""
import datetime
import hashlib
import io
import json
import os
import sys
import tempfile
import unittest
import warnings
import zipfile
from decimal import Decimal

warnings.simplefilter("ignore")
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))

import openpyxl  # noqa: E402

import auditar as A  # noqa: E402
import catalogo_bancos as cb  # noqa: E402
import control5_fecha_deposito as c5  # noqa: E402
import controles_lectura as cl  # noqa: E402
import macros_vouchers as mv  # noqa: E402
import xlsm_xml as X  # noqa: E402

D = datetime.date


def detalle(out):
    """Libro Excel interno (DETALLE_TECNICO_AUDITORIA.xlsx), que ahora viaja solo dentro del ZIP."""
    with zipfile.ZipFile(out["archivos"]["zip"]) as z:
        return openpyxl.load_workbook(io.BytesIO(z.read("DETALLE_TECNICO_AUDITORIA.xlsx")))


def sha(ruta):
    with open(ruta, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


BASE = {"TOTAL EFECTIVO DISPONIBLE": 1000, "DOLARES": 0, "CHEQUES": 0, "COBROS ATC": 500,
        "TOTAL COMUNICACIONES INTERNAS": 300, "FACTURAS ANULADAS": 0, "TOTAL MOVIMIENTO DEL DIA": 1800}
CI_HDR = ["N°", "FECHA2", "NUMERO DE FACTURA", "NIT/C.I.", "RAZON SOCIAL", "NOMBRE DEL ESTUDIANTE", "TIPO DE PAGO",
          "MONTO", "USUARIO", "ESTADO", "FACTURACION", "CLINICA2", "TOTAL C.I.3", "GLOSA ASIENTO COMUNICACIONES INTERNAS",
          "BANCO", "CUENTA CONTABLE BANCO", "ASIGNACION"]


def ci(n, importe, banco, cuenta, asig, fecha=D(2026, 9, 9)):
    """Fila de CI (17 columnas como en el cierre real)."""
    f = [None] * 17
    f[0], f[1], f[2] = n, fecha, "F%s" % n
    f[12], f[13], f[14], f[15], f[16] = importe, "F-%s" % n, banco, cuenta, asig
    return f


def cierre(ruta, sfc=None, cis=None, deposito=None, sfcs=("SFC101", "SFC102"), vba=True, banco=None):
    """sfc: {sfcN: dict etiqueta->valor} (None omite el campo); cis: {sfcN: [filas ci]}"""
    sfc = sfc or {}
    cis = cis or {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for s in sfcs:
        ws = wb.create_sheet(s)
        ws["D2"], ws["E2"], ws["F2"], ws["G2"], ws["H2"] = "COMPOSICIÓN DE DEPOSITOS", "IMPORTE Bs", "FECHA DE DEPOSITO", "ASIGNACION", "BANCO"
        if deposito and s == sfcs[0]:
            imp, f, cod = deposito
            ws["D3"], ws["E3"], ws["G3"] = "DEPOSITO 1", imp, cod
            ws["F3"] = f
            ws["F3"].number_format = "dd/mm/yyyy"
            if banco is not None:
                ws["H3"] = banco
        for i, (k, v) in enumerate((sfc.get(s, BASE)).items(), 3):
            ws.cell(i, 1, k)
            if v is not None:
                ws.cell(i, 2, v)
    for s in sfcs:
        ws = wb.create_sheet("COMUNICACIONES INTERNAS %s %s" % (s[:3], s[3:]))
        ws.append(CI_HDR)
        for fila in cis.get(s, []):
            ws.append(fila)
            ws.cell(ws.max_row, 2).number_format = "dd/mm/yyyy"
        ws.append([None] * 12 + [0, "F-  "])          # fila vacia como las del cierre real
        ws.append([None] * 12 + [999, None])           # fila de total
    tmp = ruta + ".x"
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as zi, zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            zo.writestr(i.filename, zi.read(i.filename))
        if vba:
            zo.writestr("xl/vbaProject.bin", b"VBA-FALSO" * 100)
    os.remove(tmp)


class Base(unittest.TestCase):
    def setUp(self):
        self._t = tempfile.TemporaryDirectory()
        self.d = self._t.name
        self.addCleanup(self._t.cleanup)

    def ruta(self, n):
        return os.path.join(self.d, n)

    def leer(self, cis=None, sfc=None, caja="tiquipaya", sfcs=("SFC101", "SFC102"), fecha=D(2026, 9, 9)):
        r = self.ruta("CIERRE 09-09-2026.xlsm")
        cierre(r, sfc=sfc, cis=cis, sfcs=sfcs)
        return cl.auditar_lectura(X.Paquete(r), sfcs, fecha, caja)

    @staticmethod
    def de(hall, control):
        return [h for h in hall if h["control"] == control]


# ---------------------------------------------------------------------------
class TestControl1(Base):
    def uno(self, valores):
        hall, _, _ = self.leer(sfc={"SFC101": valores, "SFC102": BASE})
        return [h for h in self.de(hall, 1) if h["sfc"] == "SFC101"][0]

    def test_cuadra(self):
        h = self.uno(BASE)
        self.assertEqual((h["resultado"], h["hallazgo"]), (cl.OK, "✅ Recaudación cuadra: Bs 1.800"))

    def test_no_cuadra_con_diferencia(self):
        h = self.uno({**BASE, "TOTAL MOVIMIENTO DEL DIA": 1790})
        self.assertEqual(h["resultado"], cl.NO_CUADRA)
        self.assertEqual(h["hallazgo"], "❌ Recaudación no cuadra — diferencia Bs -10")

    def test_facturas_anuladas_restan(self):
        self.assertEqual(self.uno({**BASE, "FACTURAS ANULADAS": 100, "TOTAL MOVIMIENTO DEL DIA": 1700})["resultado"], cl.OK)

    def test_dolares_y_cheques_suman(self):
        v = {**BASE, "DOLARES": 50, "CHEQUES": 25, "TOTAL MOVIMIENTO DEL DIA": 1875}
        self.assertEqual(self.uno(v)["resultado"], cl.OK)

    def test_falta_campo_no_asume_cero(self):
        v = dict(BASE)
        v["CHEQUES"] = None  # etiqueta presente, valor vacio
        h = self.uno(v)
        self.assertEqual((h["resultado"], h["hallazgo"]), (cl.REVISAR, "⚠️ No se pudo verificar el cuadre — falta CHEQUES."))

    def test_falta_etiqueta_completa(self):
        v = {k: x for k, x in BASE.items() if k != "COBROS ATC"}
        self.assertIn("falta COBROS ATC", self.uno(v)["hallazgo"])

    def test_varios_faltantes(self):
        v = {k: x for k, x in BASE.items() if k not in ("COBROS ATC", "DOLARES")}
        self.assertIn("falta DOLARES, COBROS ATC", self.uno(v)["hallazgo"])

    def test_campos_por_etiqueta_no_por_fila(self):
        v = dict(reversed(list(BASE.items())))  # orden distinto de filas
        self.assertEqual(self.uno(v)["resultado"], cl.OK)

    def test_ruido_de_coma_flotante_a_centavos(self):
        v = {**BASE, "TOTAL EFECTIVO DISPONIBLE": 19467, "COBROS ATC": 9181,
             "TOTAL COMUNICACIONES INTERNAS": 57470.52, "TOTAL MOVIMIENTO DEL DIA": 86118.51999999999}
        self.assertEqual(self.uno(v)["resultado"], cl.OK)

    def test_diferencia_de_un_centavo_se_detecta(self):
        v = {**BASE, "TOTAL MOVIMIENTO DEL DIA": 1800.01}
        h = self.uno(v)
        self.assertEqual((h["resultado"], h["hallazgo"]), (cl.NO_CUADRA, "❌ Recaudación no cuadra — diferencia Bs 0,01"))

    def test_cada_sfc_es_independiente(self):
        hall, _, _ = self.leer(sfc={"SFC101": BASE, "SFC102": {**BASE, "TOTAL MOVIMIENTO DEL DIA": 5}})
        r = {h["sfc"]: h["resultado"] for h in self.de(hall, 1)}
        self.assertEqual(r, {"SFC101": cl.OK, "SFC102": cl.NO_CUADRA})

    def test_bs_formato(self):
        self.assertEqual((cl.bs(Decimal("2350")), cl.bs(Decimal("4150.5")), cl.bs(Decimal("79024"))), ("2.350", "4.150,50", "79.024"))


# ---------------------------------------------------------------------------
class TestControl2(Base):
    def revisar(self, filas102):
        hall, alq, esp = self.leer(cis={"SFC102": filas102})
        return hall, alq, esp

    def test_falta_cuenta(self):
        hall, _, _ = self.revisar([ci(1, 100, "BNB MN", None, "3P123")])
        h = self.de(hall, 2)
        self.assertEqual([x["hallazgo"] for x in h], ["⚠️ SFC102 — fila 2: falta CUENTA CONTABLE."])

    def test_falta_asignacion_y_banco(self):
        hall, _, _ = self.revisar([ci(1, 100, "BNB MN", "110103012", None), ci(2, 100, None, "110103012", "3P1")])
        m = [x["hallazgo"] for x in self.de(hall, 2)]
        self.assertIn("⚠️ SFC102 — fila 2: falta ASIGNACION.", m)
        self.assertIn("⚠️ SFC102 — fila 3: falta BANCO.", m)

    def test_varios_faltantes_en_una_fila(self):
        hall, _, _ = self.revisar([ci(1, 100, None, None, None)])
        self.assertEqual([x["hallazgo"] for x in self.de(hall, 2)], ["⚠️ SFC102 — fila 2: falta BANCO, CUENTA CONTABLE y ASIGNACION."])

    def test_filas_vacias_y_total_se_ignoran(self):
        hall, _, _ = self.revisar([ci(1, 100, "BNB MN", "110103012", "3P1")])
        self.assertEqual(self.de(hall, 2), [])

    def test_alquileres_exento(self):
        hall, alq, _ = self.revisar([ci(1, 500, "ALQUILERES", None, None)])
        self.assertEqual((self.de(hall, 2), self.de(hall, 3), self.de(hall, 4)), ([], [], []))
        self.assertEqual(len(alq), 1)

    def test_especiales_no_generan_falsas_alertas_por_banco(self):
        for cat in ("OTROS INGRESOS", "GASTO.ADM", "POSGRADO.PLA"):
            hall, _, esp = self.revisar([ci(1, 100, cat, None, None)])
            self.assertEqual([h for h in hall if h["control"] in (3, 4)], [], cat)
            self.assertEqual([e["categoria"] for e in esp], [cat])
            self.assertEqual(esp[0]["falta"], ["CUENTA CONTABLE", "ASIGNACION"])

    def test_especiales_siempre_quedan_registrados_con_fecha_y_caja(self):
        _, _, esp = self.revisar([ci(1, 100, "otros ingresos", "999", "X1")])
        self.assertEqual((esp[0]["fecha"], esp[0]["caja"], esp[0]["sfc"], esp[0]["falta"]), (D(2026, 9, 9), "tiquipaya", "SFC102", []))

    def test_hallazgo_de_especiales_revisar_si_falta_e_info_si_completo(self):
        r = self.ruta("x.xlsm")
        cierre(r, cis={"SFC102": [ci(1, 100, "OTROS INGRESOS", None, None), ci(2, 50, "GASTO.ADM", "5", "A")]})
        _, _, esp = cl.auditar_lectura(X.Paquete(r), ("SFC101", "SFC102"), D(2026, 9, 9), "tiquipaya")
        h = A.hallazgos_especiales(esp, D(2026, 9, 9), "tiquipaya")
        self.assertEqual({(x["hallazgo"].split(":")[0]): x["resultado"] for x in h},
                         {"⚠️ OTROS INGRESOS": cl.REVISAR, "ℹ️ GASTO.ADM": cl.INFO})


# ---------------------------------------------------------------------------
class TestControl3(Base):
    def c3(self, banco, cuenta, asig="3P123"):
        hall, _, _ = self.leer(cis={"SFC102": [ci(1, 100, banco, cuenta, asig)]})
        return self.de(hall, 3)

    def test_toda_la_tabla_oficial_valida(self):
        for cuenta, _d, corto in cb.TABLA_OFICIAL:
            self.assertEqual(self.c3(corto, cuenta), [], corto)

    def test_incoherencia_formato_exacto(self):
        h = self.c3("BISA MN", "110103012")
        self.assertEqual(h[0]["hallazgo"], "⚠️ Revisar cuenta contable\nSFC102 — fila 2\nBanco: BISA MN\nCuenta registrada: 110103012\nCuenta esperada: 110103032")

    def test_sufijos_importan(self):
        self.assertEqual(self.c3("BNB ME", "110103012")[0]["hallazgo"].splitlines()[-1], "Cuenta esperada: 110104012")
        self.assertEqual(self.c3("BNB CLINICA", "110103012")[0]["hallazgo"].splitlines()[-1], "Cuenta esperada: 110103022")
        self.assertEqual(self.c3("BISA EURO", "110103032")[0]["hallazgo"].splitlines()[-1], "Cuenta esperada: 110105112")
        self.assertEqual(self.c3("BANECO AH", "110103062")[0]["hallazgo"].splitlines()[-1], "Cuenta esperada: 110103722")

    def test_sin_sufijo_se_asume_mn(self):
        for b, c in (("BNB", "110103012"), ("BUSA", "110103052"), ("BISA", "110103032"), ("BCP", "110103042"), ("BANECO", "110103062"), ("BMS", "110103072")):
            self.assertEqual(self.c3(b, c), [], b)
        self.assertEqual(self.c3("BNB", "110104012")[0]["hallazgo"].splitlines()[-1], "Cuenta esperada: 110103012")

    def test_alias(self):
        self.assertEqual(self.c3("BANCO ECONOMICO", "110103062"), [])
        self.assertEqual(self.c3("Banco Económico AH", "110103722"), [])
        self.assertEqual(self.c3("BANCO UNION", "110103052"), [])
        self.assertEqual(self.c3("BANCO UNION ME", "110104042"), [])
        self.assertEqual(self.c3("BMSC", "110103072"), [])
        self.assertEqual(self.c3("BANECO MN", "110103062"), [])

    def test_banco_desconocido_y_combinacion_fuera_de_tabla_se_reportan(self):
        self.assertIn("no figura en la tabla oficial", self.c3("CITIBANK", "110103012")[0]["hallazgo"])
        self.assertIn("no corresponde a una cuenta de la tabla oficial", self.c3("BANECO ME", "110103062")[0]["hallazgo"])

    def test_cuenta_faltante_no_duplica_alerta_en_c3(self):
        self.assertEqual(self.c3("BNB MN", None), [])

    def test_cuenta_numerica_en_celda(self):
        self.assertEqual(self.c3("BNB MN", 110103012), [])

    def test_no_corrige(self):
        r = self.ruta("CIERRE 09-09-2026.xlsm")
        cierre(r, cis={"SFC102": [ci(1, 100, "BISA MN", "110103012", "123")]})
        h0 = sha(r)
        cl.auditar_lectura(X.Paquete(r), ("SFC101", "SFC102"), D(2026, 9, 9), "tiquipaya")
        self.assertEqual(h0, sha(r))


# ---------------------------------------------------------------------------
class TestControl4(Base):
    def c4(self, banco, asig, cuenta="110103012"):
        hall, _, _ = self.leer(cis={"SFC102": [ci(1, 100, banco, cuenta, asig)]})
        return self.de(hall, 4)

    def test_formato_exacto_bisa_numerico(self):
        h = self.c4("BISA", "ABC123")
        self.assertEqual(h[0]["hallazgo"], "⚠️ REVISAR ASIGNACIÓN\nSFC102 — fila 2\nBanco: BISA\nAsignación: ABC123\nEsperado: código numérico")

    def test_familias_numericas(self):
        for b in ("BANECO", "BANECO AH", "BANCO ECONOMICO", "BUSA", "BUSA MN", "BUSA ME", "BANCO UNION", "BCP", "BCP MN", "BCP ME", "BISA", "BISA MN", "BISA ME", "BISA EURO"):
            self.assertEqual(self.c4(b, "74186779"), [], b)
            self.assertEqual(len(self.c4(b, "3P123")), 1, b)

    def test_familias_alfanumericas(self):
        for b in ("BNB", "BNB MN", "BNB ME", "BNB AH", "BNB CLINICA", "BMS", "BMSC"):
            self.assertEqual(self.c4(b, "3P06423289"), [], b)
            h = self.c4(b, "12345678")
            self.assertEqual(len(h), 1, b)
            self.assertTrue(h[0]["hallazgo"].endswith("Esperado: código alfanumérico (con letras)"), b)

    def test_asignacion_numerica_en_celda_es_numerica(self):
        self.assertEqual(self.c4("BUSA", 74186779), [])

    def test_especiales_y_alquileres_no_se_validan(self):
        for cat in ("OTROS INGRESOS", "GASTO.ADM", "ALQUILERES", "POSGRADO.PLA"):
            self.assertEqual(self.c4(cat, "ABC"), [], cat)

    def test_asignacion_faltante_no_duplica_en_c4(self):
        self.assertEqual(self.c4("BISA", None), [])

    def test_banco_desconocido_no_se_valida_formato(self):
        self.assertEqual(self.c4("CITIBANK", "ABC"), [])


# ---------------------------------------------------------------------------
class TestControl6(Base):
    def test_alquileres_con_fecha2_caja_sfc_importe(self):
        _, alq, _ = self.leer(cis={"SFC101": [ci(1, 2350, "ALQUILERES", None, None, D(2026, 9, 18))],
                                   "SFC102": [ci(1, 1800, "alquileres", None, None, D(2026, 9, 19)), ci(2, 10, "BNB MN", "110103012", "3P1")]})
        r = sorted((a["fecha"], a["sfc"], a["importe"], a["fecha_del_cierre"]) for a in alq)
        self.assertEqual(r, [(D(2026, 9, 18), "SFC101", Decimal("2350"), False), (D(2026, 9, 19), "SFC102", Decimal("1800"), False)])
        self.assertTrue(all(a["caja"] == "tiquipaya" for a in alq))
        self.assertTrue(all(a["marcado_en"] == "banco" for a in alq))

    def test_sin_fecha2_usa_fecha_del_cierre_marcada(self):
        _, alq, _ = self.leer(cis={"SFC102": [ci(1, 500, "ALQUILERES", None, None, None)]})
        self.assertEqual((alq[0]["fecha"], alq[0]["fecha_del_cierre"]), (D(2026, 9, 9), True))

    def test_alquiler_solo_se_busca_en_ci(self):
        # 'ALQUILERES' en otra columna (glosa) de CI no cuenta; en SFC tampoco
        fila = ci(1, 100, "BNB MN", "110103012", "3P1")
        fila[13] = "ALQUILERES"
        _, alq, _ = self.leer(cis={"SFC102": [fila]})
        self.assertEqual(alq, [])

    def test_alquiler_marcado_en_asignacion_con_banco_vacio(self):
        # caso real TIQ 23/09: BANCO y CUENTA vacios, ASIGNACION = ALQUILERES
        hall, alq, _ = self.leer(cis={"SFC102": [ci(1, 1560, None, None, "ALQUILERES", D(2026, 9, 23)),
                                                 ci(2, 700, None, None, " alquileres ", D(2026, 9, 23))]})
        self.assertEqual([a["importe"] for a in alq], [Decimal("1560"), Decimal("700")])
        self.assertEqual([a["marcado_en"] for a in alq], ["asignacion", "asignacion"])
        self.assertEqual((self.de(hall, 2), self.de(hall, 3), self.de(hall, 4)), ([], [], []))

    def test_asignacion_alquileres_con_banco_real_no_es_alquiler(self):
        hall, alq, _ = self.leer(cis={"SFC102": [ci(1, 100, "BISA MN", "110103032", "ALQUILERES")]})
        self.assertEqual(alq, [])
        self.assertEqual(len(self.de(hall, 4)), 1)  # sigue valorandose como asignacion de BISA (numerica)

    def test_banco_vacio_sin_alquileres_sigue_siendo_hallazgo(self):
        hall, alq, _ = self.leer(cis={"SFC102": [ci(1, 100, None, None, "3P1")]})
        self.assertEqual(alq, [])
        self.assertEqual(len(self.de(hall, 2)), 1)

    def test_alquiler_sin_importe_se_reporta(self):
        hall, alq, _ = self.leer(cis={"SFC102": [ci(1, None, "ALQUILERES", None, None)]})
        self.assertEqual(alq, [])
        self.assertEqual(self.de(hall, 6)[0]["resultado"], cl.REVISAR)


# ---------------------------------------------------------------------------
def macros_xlsx(ruta, vouchers):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = mv.HOJA_MACROS
    ws.append(["Fecha", "Código de Asignación", "Créditos"])
    for f, c, i in vouchers:
        ws.append([f, c, i])
    wb.save(ruta)


class TestOrquestadorYReporte(Base):
    def setUp(self):
        super().setUp()
        self.macros = self.ruta("MACROS.xlsx")
        macros_xlsx(self.macros, [(D(2026, 9, 10), "AB1", 100)])

    def cierre_dia(self, dia, cis=None, deposito=(100, D(2026, 10, 9), "AB1")):
        r = self.ruta("CIERRE %02d-09-2026.xlsm" % dia)
        cierre(r, cis=cis, deposito=deposito)
        return r

    def test_flujo_completo_no_toca_el_original_y_genera_reporte(self):
        r = self.cierre_dia(9, cis={"SFC101": [ci(1, 2350, "ALQUILERES", None, None, D(2026, 9, 18))],
                                    "SFC102": [ci(1, 100, "BISA MN", "110103012", "ABC"), ci(2, 40, "OTROS INGRESOS", None, None)]})
        h0 = sha(r)
        out = A.auditar([("tiquipaya", r)], [self.macros], salida_dir=self.ruta("salida"), reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(h0, sha(r))
        self.assertTrue(os.path.exists(os.path.join(self.ruta("salida"), "TIQUIPAYA", os.path.basename(r))))
        self.assertFalse(os.path.exists(os.path.join(self.ruta("salida"), os.path.basename(r))))
        self.assertEqual(os.path.basename(out["reporte"]), "RESUMEN_GABO.pdf")
        self.assertEqual(os.path.basename(out["archivos"]["para_caja_tiquipaya"]), "PARA_CAJA_TIQUIPAYA.pdf")
        self.assertEqual(os.path.basename(out["archivos"]["zip"]), "AUDITORIA_CIERRES_RESULTADOS.zip")
        self.assertEqual(out["fechas_normalizadas"], 1)
        self.assertEqual(out["cierres"], {"revisados": 1, "sin_observaciones": 0, "requieren_revision": 1})
        self.assertEqual(out["alquileres"]["total_periodo"], "Bs 2.350")
        textos = " ".join(out["hallazgos"])
        self.assertIn("Revisar cuenta contable", textos)
        self.assertIn("REVISAR ASIGNACIÓN", textos)
        self.assertIn("OTROS INGRESOS", " ".join(out["especiales_a_completar"]))
        wb = detalle(out)
        self.assertEqual(wb.sheetnames, ["Resumen", "Auditoria", "Alquileres", "Completar a mano"])
        self.assertEqual([c.value for c in wb["Auditoria"][1]], ["FECHA", "CAJA", "SFC", "RESULTADO", "HALLAZGO", "ACCION"])
        self.assertEqual([c.value for c in wb["Alquileres"][1]], ["FECHA", "CAJA", "SFC", "IMPORTE"])
        self.assertEqual(wb["Alquileres"]["D2"].value, 2350)
        self.assertEqual(wb["Completar a mano"]["D2"].value, "OTROS INGRESOS")

    def test_salidas_separadas_por_caja_sin_colisiones(self):
        # mismo nombre de archivo en las dos cajas: cada copia va a su propia carpeta
        t = os.path.join(self.d, "t")
        a = os.path.join(self.d, "a")
        os.makedirs(t)
        os.makedirs(a)
        rt = os.path.join(t, "CIERRE 09-09-2026.xlsm")
        ra = os.path.join(a, "CIERRE 09-09-2026.xlsm")
        cierre(rt, deposito=(100, D(2026, 10, 9), "AB1"))
        cierre(ra, sfcs=("SFC107", "SFC108"), deposito=(100, D(2026, 10, 9), "AB1"))
        out = A.auditar([("tiquipaya", rt), ("america", ra)], [self.macros], salida_dir=self.ruta("salida"), reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(out["fechas_normalizadas"], 2)
        pt = os.path.join(self.ruta("salida"), "TIQUIPAYA", "CIERRE 09-09-2026.xlsm")
        pa = os.path.join(self.ruta("salida"), "AMERICA", "CIERRE 09-09-2026.xlsm")
        self.assertTrue(os.path.exists(pt) and os.path.exists(pa))
        with zipfile.ZipFile(pt) as z:
            self.assertIn("SFC101", z.read("xl/workbook.xml").decode())
        with zipfile.ZipFile(pa) as z:
            self.assertIn("SFC107", z.read("xl/workbook.xml").decode())
        self.assertEqual(sorted(os.listdir(self.ruta("salida"))), ["AMERICA", "TIQUIPAYA"])

    def test_misma_caja_mismo_nombre_no_pisa_y_avisa(self):
        d1 = os.path.join(self.d, "d1")
        d2 = os.path.join(self.d, "d2")
        os.makedirs(d1)
        os.makedirs(d2)
        r1 = os.path.join(d1, "CIERRE 09-09-2026.xlsm")
        r2 = os.path.join(d2, "CIERRE 09-09-2026.xlsm")
        cierre(r1, deposito=(100, D(2026, 10, 9), "AB1"))
        cierre(r2, deposito=(100, D(2026, 10, 9), "AB1"))
        out = A.auditar([("tiquipaya", r1), ("tiquipaya", r2)], [self.macros], salida_dir=self.ruta("salida"), reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(out["fechas_normalizadas"], 1)
        self.assertIn("colisión de salida", " ".join(out["hallazgos"]))

    def test_dd_mm_yy_se_normaliza_en_el_flujo_y_no_se_toca_lo_manual(self):
        r = self.ruta("CIERRE 09-09-2026.xlsm")
        cierre(r, deposito=(100, "10/09/26", "AB1"))
        out = A.auditar([("tiquipaya", r)], [self.macros], salida_dir=self.ruta("salida"), reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(out["fechas_normalizadas"], 1)
        wb = openpyxl.load_workbook(os.path.join(self.ruta("salida"), "TIQUIPAYA", os.path.basename(r)))
        self.assertEqual(wb["SFC101"]["F3"].value, datetime.datetime(2026, 9, 10))
        r2 = self.ruta("CIERRE 10-09-2026.xlsm")
        cierre(r2, deposito=(100, "05/0/2026", "AB1"))
        out2 = A.auditar([("tiquipaya", r2)], [self.macros], salida_dir=self.ruta("salida2"), reporte_dir=self.ruta("rep2"), hoy=D(2026, 9, 26))
        self.assertEqual(out2["fechas_normalizadas"], 0)
        self.assertIn("no existe en el calendario", " ".join(out2["hallazgos"]))
        self.assertFalse(os.path.exists(self.ruta("salida2")))

    def test_sin_voucher_el_flujo_normaliza_y_deja_traza(self):
        r = self.ruta("CIERRE 09-09-2026.xlsm")
        cierre(r, deposito=(100, "23/09/26", "NO-ESTA-EN-MACROS"))
        out = A.auditar([("tiquipaya", r)], [self.macros], salida_dir=self.ruta("salida"), reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual((out["fechas_normalizadas"], out["normalizadas_sin_voucher_macros"]), (1, 1))
        self.assertEqual(out["hallazgos"], [])                          # es traza informativa, no alerta
        wb = openpyxl.load_workbook(os.path.join(self.ruta("salida"), "TIQUIPAYA", os.path.basename(r)))
        self.assertEqual(wb["SFC101"]["F3"].value, datetime.datetime(2026, 9, 23))
        aud = detalle(out)["Auditoria"]
        textos = [row[4].value for row in aud.iter_rows(min_row=2) if row[4].value and "SIN_VOUCHER" in row[4].value]
        self.assertEqual(len(textos), 1)
        self.assertIn("FECHA_NORMALIZADA_SIN_VOUCHER_MACROS", textos[0])
        self.assertIn("texto '23/09/26'", textos[0])
        self.assertIn("23/09/2026", textos[0])

    def test_sin_voucher_y_sin_salida_dir_queda_pendiente(self):
        r = self.ruta("CIERRE 09-09-2026.xlsm")
        cierre(r, deposito=(100, "23/09/26", "NO-ESTA"))
        out = A.auditar([("tiquipaya", r)], [self.macros], reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(out["fechas_normalizadas"], 0)
        self.assertIn("FECHA_NORMALIZADA_SIN_VOUCHER_MACROS", " ".join(out["hallazgos"]))

    def test_tolerancia_1_dia_deja_traza_y_no_cambia_el_valor(self):
        r = self.ruta("CIERRE 09-09-2026.xlsm")
        cierre(r, deposito=(100, "09/09/26", "AB1"))   # MACROS dice 10/09: 1 dia de diferencia
        out = A.auditar([("tiquipaya", r)], [self.macros], salida_dir=self.ruta("salida"), reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual((out["fechas_normalizadas"], out["dentro_tolerancia_1_dia"], out["hallazgos"]), (1, 1, []))
        wb = openpyxl.load_workbook(os.path.join(self.ruta("salida"), "TIQUIPAYA", os.path.basename(r)))
        self.assertEqual(wb["SFC101"]["F3"].value, datetime.datetime(2026, 9, 9))       # el valor NO cambia a 10/09
        aud = detalle(out)["Auditoria"]
        t = [row[4].value for row in aud.iter_rows(min_row=2) if row[4].value and "TOLERANCIA" in row[4].value]
        self.assertEqual(len(t), 1)
        self.assertIn("FECHA_DENTRO_TOLERANCIA_MACROS_1_DIA", t[0])
        self.assertIn("09/09/2026", t[0])
        self.assertIn("10/09/2026", t[0])

    def test_sin_salida_dir_no_escribe_cierres_y_avisa(self):
        r = self.cierre_dia(9)
        out = A.auditar([("tiquipaya", r)], [self.macros], reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(out["fechas_normalizadas"], 0)
        self.assertIn("por corregir", " ".join(out["hallazgos"]))
        self.assertFalse(os.path.exists(self.ruta("salida")))

    def test_sin_macros_no_corre_control5_pero_si_el_resto(self):
        r = self.cierre_dia(9)
        out = A.auditar([("tiquipaya", r)], [], reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertIn("no se indicó el archivo MACROS", " ".join(out["hallazgos"]))
        self.assertEqual(out["cuadres_ok"], 2)

    def test_maximo_15_hallazgos_en_stdout(self):
        filas = [ci(i, 10, "BISA", "110103032", "ABC") for i in range(1, 31)]  # 30 asignaciones invalidas
        r = self.cierre_dia(9, cis={"SFC102": filas})
        out = A.auditar([("tiquipaya", r)], [self.macros], salida_dir=self.ruta("s"), reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(len(out["hallazgos"]), 16)
        self.assertRegex(out["hallazgos"][-1], r"^y \d+ adicionales; ver reporte$")
        self.assertLess(len(json.dumps(out, ensure_ascii=False)), 6000)
        wb = detalle(out)
        self.assertGreaterEqual(wb["Auditoria"].max_row, 31)  # el reporte trae todo

    def test_varias_cajas_y_alquileres_por_dia_y_total(self):
        wb = self.ruta("a.xlsm")
        cierre(wb, sfcs=("SFC107", "SFC108"), cis={"SFC107": [ci(1, 1800, "ALQUILERES", None, None, D(2026, 9, 19))]})
        r1 = self.ruta("CIERRE 18-09-2026.xlsm")
        cierre(r1, cis={"SFC102": [ci(1, 2350, "ALQUILERES", None, None, D(2026, 9, 18))]})
        r2 = self.ruta("CIERRE 19-09-2026.xlsm")
        os.replace(wb, r2)
        out = A.auditar([("tiquipaya", r1), ("america", r2)], [], reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(out["alquileres"]["detalle"], ["18/09/2026 — Bs 2.350", "19/09/2026 — Bs 1.800"])
        self.assertEqual(out["alquileres"]["total_periodo"], "Bs 4.150")
        wbk = detalle(out)
        celdas = [r for r in wbk["Alquileres"].iter_rows(values_only=True) if r[0] is not None]
        self.assertEqual(celdas[-1][0], "TOTAL PERIODO")
        self.assertEqual(celdas[-1][3], 4150)

    def test_archivo_ilegible_no_detiene_la_corrida(self):
        malo = self.ruta("CIERRE 09-09-2026.xlsm")
        with open(malo, "wb") as f:
            f.write(b"no es un zip")
        bueno = self.cierre_dia(10)
        out = A.auditar([("tiquipaya", malo), ("tiquipaya", bueno)], [self.macros], reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(out["cierres"]["revisados"], 2)
        self.assertIn("No se pudo leer el cierre", " ".join(out["hallazgos"]))

    def test_salida_estandar_no_sobrescribe_la_de_otra_corrida(self):
        r = self.cierre_dia(9)
        a = A.auditar([("tiquipaya", r)], [], reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        contenido = open(a["archivos"]["pdf"], "rb").read()
        with self.assertRaisesRegex(FileExistsError, "SALIDA_YA_EXISTE"):
            A.auditar([("tiquipaya", r)], [], reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(open(a["archivos"]["pdf"], "rb").read(), contenido)   # la corrida previa queda intacta

    def test_cli_stdout_es_solo_json_utf8(self):
        import subprocess
        r = self.cierre_dia(9)
        p = subprocess.run([sys.executable, os.path.join(os.path.dirname(A.__file__), "auditar.py"), "--macros", self.macros,
                            "--caja", "tiquipaya", "--cierre", r, "--reporte-dir", self.ruta("rep")], capture_output=True)
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertEqual(json.loads(p.stdout.decode("utf-8"))["cierres"]["revisados"], 1)


# ---------------------------------------------------------------------------
REAL = r"C:\Users\PC\Desktop\prueba_auditor"


@unittest.skipUnless(os.path.exists(os.path.join(REAL, "tmp_orig", "CIERRE 09-09-2026.xlsm")), "copia real no disponible")
class TestCopiaReal(Base):
    def test_cierre_real_09_09_2026(self):
        r = os.path.join(REAL, "tmp_orig", "CIERRE 09-09-2026.xlsm")
        pq = X.Paquete(r)
        hall, alq, esp = cl.auditar_lectura(pq, ("SFC101", "SFC102"), D(2026, 9, 9), "tiquipaya")
        c1 = {h["sfc"]: h["hallazgo"] for h in hall if h["control"] == 1}
        self.assertEqual(c1, {"SFC101": "✅ Recaudación cuadra: Bs 79.024", "SFC102": "✅ Recaudación cuadra: Bs 86.118,52"})
        self.assertEqual([h for h in hall if h["control"] != 1], [])
        self.assertEqual((alq, esp), ([], []))
        h0 = sha(r)
        out = A.auditar([("tiquipaya", r)], [os.path.join(REAL, "MACROS SEPTIEMBRE.xlsm")], salida_dir=self.ruta("s"),
                        reporte_dir=self.ruta("rep"), hoy=D(2026, 9, 26))
        self.assertEqual(h0, sha(r))
        self.assertEqual((out["cuadres_ok"], out["fechas_normalizadas"], out["estado"]), (2, 1, "OK"))


if __name__ == "__main__":
    unittest.main()
