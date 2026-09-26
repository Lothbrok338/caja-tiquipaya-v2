"""Pruebas de la busqueda de CANDIDATOS en MACROS cuando un deposito no tiene voucher exacto (asignacion + importe).

Solo diagnostico: nunca se autocorrige nada. Casos cubiertos:
  * ejemplo real America 03/09/2026 SFC107 (candidato probable por importe + cuenta + caja + fecha esperada, sin autocorregir);
  * candidato unico fuerte / varios candidatos / ningun candidato / candidato fuera de +-1 dia pero importe unico en MACROS;
  * la clasificacion del Control 5 y su plan de cambios NO cambian; el informe PDF/DOCX muestra el detalle humano.
"""
import datetime
import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings
import zipfile

warnings.simplefilter("ignore")
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402

import auditar as A  # noqa: E402
import candidatos_voucher as cv  # noqa: E402
import control5_fecha_deposito as c5  # noqa: E402
import informe_humano as ih  # noqa: E402
import macros_vouchers as mv  # noqa: E402
from test_controles import cierre  # noqa: E402

D = datetime.date


def reg(fecha, codigo, importe, caja="SFC107", cuenta="110103012"):
    return {"fecha": fecha, "importe": importe, "codigo": codigo, "cuenta": cuenta, "caja": caja}


def macros_completo(ruta, filas):
    """MACROS con las columnas reales: filas (fecha, codigo, importe, caja, cuenta)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = mv.HOJA_MACROS
    ws.append(["Fecha", "Glosa Recortada", "Codigo de Asignacion", "Creditos", "Adicionales", "Numero de Caja", "Hora",
               "Codigo Asiento", "CUENTA CONTABLE"])
    for f, cod, imp, caja, cta in filas:
        ws.append([f, "ABONO %s" % caja, cod, imp, "Nombre:PRUEBA", caja, None, None, cta])
    wb.save(ruta)


# Registros de MACROS del ejemplo real (America, septiembre 2026)
MACROS_REAL = [
    reg(D(2026, 9, 1), "3P91155531", "25000.00"),        # la asignacion escrita en el cierre pertenece a este abono del 01/09
    reg(D(2026, 9, 3), "3P93097190", "19158.20"),
    reg(D(2026, 9, 4), "3P94114380", "9527.50"),         # el voucher real del deposito
    reg(D(2026, 9, 5), "3P95109062", "10641.60"),
]


class TestBuscar(unittest.TestCase):
    """Nucleo puro: candidatos_voucher.buscar."""

    def buscar(self, registros, importe="9527.50", cuenta="110103012", sfc="SFC107", fecha=D(2026, 4, 9), cierre_=D(2026, 9, 3),
               asignacion="3P91155531"):
        return cv.buscar(registros, importe, cuenta, sfc, fecha, cierre_, asignacion)

    def test_ejemplo_real_es_candidato_probable(self):
        r = self.buscar(MACROS_REAL)
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.PROBABLE, cv.FUERTE))
        self.assertEqual(len(r["candidatos"]), 1)
        c = r["candidatos"][0]
        self.assertEqual((c["fecha"], c["importe"], c["asignacion"], c["cuenta"], c["caja"]),
                         ("04/09/2026", "9527.50", "3P94114380", "110103012", "SFC107"))
        self.assertEqual(c["fuerza"], cv.FUERTE)
        # coincidencias: importe, cuenta, caja, fecha esperada, dia/mes invertidos, importe unico
        joint = " | ".join(c["coincidencias"])
        for esperado in ("Importe exacto", "Cuenta 110103012", "Caja SFC107", "ventana esperada (04/09/2026", "invertidos",
                         "Importe único en MACROS"):
            self.assertIn(esperado, joint)
        # diferencias: la asignacion y la fecha escritas
        self.assertIn("Asignación distinta (cierre: 3P91155531; MACROS: 3P94114380)", c["diferencias"])
        self.assertIn("Fecha distinta (cierre: 09/04/2026; MACROS: 04/09/2026)", c["diferencias"])
        self.assertTrue(c["fecha_invertida"] and c["asignacion_distinta"])
        self.assertEqual(r["declarado"], {"importe": "9527.50", "fecha": "09/04/2026", "asignacion": "3P91155531",
                                          "cuenta": "110103012", "sfc": "SFC107"})
        self.assertEqual((r["fecha_esperada"], r["ventana"]), ("04/09/2026", ["03/09/2026", "05/09/2026"]))
        # la asignacion escrita existe en MACROS como OTRO abono (pista de un error de copia)
        self.assertEqual(r["asignacion_en_macros"], [{"fecha": "01/09/2026", "importe": "25000.00", "caja": "SFC107"}])

    def test_varios_candidatos_fuertes(self):
        regs = MACROS_REAL + [reg(D(2026, 9, 5), "3P95000001", "9527.50")]
        r = self.buscar(regs)
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.VARIOS, None))
        self.assertEqual([c["asignacion"] for c in r["candidatos"]], ["3P94114380", "3P95000001"])   # el mas cercano a la fecha esperada primero
        self.assertTrue(all(c["fuerza"] == cv.FUERTE for c in r["candidatos"]))

    def test_varios_candidatos_sin_fuerte_en_otras_fechas(self):
        regs = [reg(D(2026, 9, 10), "AAA", "9527.50"), reg(D(2026, 9, 20), "BBB", "9527.50")]
        r = self.buscar(regs)
        self.assertEqual(r["veredicto"], cv.VARIOS)
        self.assertEqual({c["fuerza"] for c in r["candidatos"]}, {cv.FUERA_DE_VENTANA})

    def test_ningun_candidato(self):
        r = self.buscar([reg(D(2026, 9, 4), "X", "1000.00"), reg(D(2026, 9, 5), "Y", "9527.51")])
        self.assertEqual((r["veredicto"], r["candidatos"], r["importe_en_macros"]), (cv.NINGUNO, [], 0))
        self.assertEqual(self.buscar([])["veredicto"], cv.NINGUNO)

    def test_igual_importe_pero_otra_caja_y_otra_cuenta_no_es_razonable(self):
        r = self.buscar([reg(D(2026, 9, 4), "OTRO", "9527.50", caja="SFC101", cuenta="110103032")])
        self.assertEqual(r["veredicto"], cv.NINGUNO)
        self.assertEqual(r["importe_en_macros"], 1)          # se sabe que el importe existe, pero no cuenta como candidato

    def test_candidato_fuera_de_1_dia_pero_importe_unico_en_macros(self):
        regs = MACROS_REAL[:2] + [reg(D(2026, 9, 12), "3P9120001", "9527.50")]      # 8 dias despues de la fecha esperada
        r = self.buscar(regs)
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.POSIBLE, cv.FUERA_DE_VENTANA))      # fuera de +-1 dia: nunca "probable"
        c = r["candidatos"][0]
        self.assertEqual((c["fecha"], c["fuerza"]), ("12/09/2026", cv.FUERA_DE_VENTANA))
        self.assertEqual(c["difiere"], ["fecha"])
        self.assertIn("Importe único en MACROS", c["coincidencias"])
        self.assertIn("Cuenta 110103012", c["coincidencias"])
        self.assertTrue(any(d.startswith("Fecha fuera de la ventana esperada (03/09/2026 a 05/09/2026)") for d in c["diferencias"]))
        self.assertEqual(r["importe_en_macros"], 1)

    def test_fuera_de_ventana_se_omite_si_hay_un_fuerte(self):
        r = self.buscar(MACROS_REAL + [reg(D(2026, 9, 20), "LEJOS", "9527.50")])
        self.assertEqual((r["veredicto"], len(r["candidatos"])), (cv.PROBABLE, 1))
        self.assertEqual(r["candidatos"][0]["asignacion"], "3P94114380")
        self.assertNotIn("Importe único en MACROS", r["candidatos"][0]["coincidencias"])   # ahora el importe aparece dos veces

    def test_misma_fecha_pero_otra_caja_es_candidato_parcial(self):
        r = self.buscar([reg(D(2026, 9, 4), "3P94000000", "9527.50", caja="SFC108")])
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.POSIBLE, cv.PARCIAL))                # coincidencia parcial
        self.assertIn("Caja distinta (cierre: SFC107; MACROS: SFC108)", r["candidatos"][0]["diferencias"])
        self.assertEqual(r["candidatos"][0]["difiere"], ["caja"])

    def test_dato_no_verificable_no_descarta_pero_tampoco_es_fuerte(self):
        r = self.buscar([reg(D(2026, 9, 4), "3P94000000", "9527.50", caja=None, cuenta=None)])
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.POSIBLE, cv.PARCIAL))                # no se descarta, pero no es probable
        self.assertIn("Cuenta no verificable", r["candidatos"][0]["diferencias"])
        self.assertIn("Caja no verificable", r["candidatos"][0]["diferencias"])
        self.assertEqual(r["candidatos"][0]["no_verificable"], ["cuenta", "caja"])

    def test_sin_asignacion_en_el_cierre(self):
        r = self.buscar(MACROS_REAL, asignacion=None)
        self.assertIn("El cierre no declara asignación", r["candidatos"][0]["diferencias"])
        self.assertEqual(r["asignacion_en_macros"], [])

    def test_probable_solo_si_es_el_unico_candidato_y_es_fuerte(self):
        # un candidato fuerte + otro razonable en la misma ventana (otra caja): ya no es "unico" -> VARIOS
        r = self.buscar(MACROS_REAL + [reg(D(2026, 9, 4), "OTRA-CAJA", "9527.50", caja="SFC108")])
        self.assertEqual((r["veredicto"], r["fuerza"], len(r["candidatos"])), (cv.VARIOS, None, 2))
        self.assertEqual([c["fuerza"] for c in r["candidatos"]], [cv.FUERTE, cv.PARCIAL])
        # unico y fuerte: PROBABLE
        self.assertEqual(self.buscar(MACROS_REAL)["veredicto"], cv.PROBABLE)

    def test_cuenta_distinta_con_misma_caja_es_solo_posible(self):
        r = self.buscar([reg(D(2026, 9, 4), "3P94000000", "9527.50", cuenta="110103032")])
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.POSIBLE, cv.PARCIAL))
        self.assertEqual(r["candidatos"][0]["difiere"], ["cuenta"])

    def test_cuenta_del_cierre_desconocida_impide_probable(self):
        r = self.buscar(MACROS_REAL, cuenta=None)
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.POSIBLE, cv.PARCIAL))
        self.assertEqual(r["candidatos"][0]["no_verificable"], ["cuenta"])
        self.assertIn("Cuenta no verificable", r["candidatos"][0]["diferencias"])

    def test_fecha_del_cierre_desconocida_impide_probable(self):
        r = self.buscar(MACROS_REAL, cierre_=None)
        self.assertEqual((r["veredicto"], r["fuerza"]), (cv.POSIBLE, cv.FUERA_DE_VENTANA))
        self.assertEqual(r["candidatos"][0]["no_verificable"], ["fecha"])
        self.assertIsNone(r["ventana"])

    def test_voucher_de_macros_sin_fecha_valida_es_posible(self):
        r = self.buscar([reg(None, "3P94000000", "9527.50")])
        self.assertEqual((r["veredicto"], r["candidatos"][0]["fecha"]), (cv.POSIBLE, None))
        self.assertIn("El voucher de MACROS no tiene fecha válida", r["candidatos"][0]["diferencias"])

    def test_los_cuatro_veredictos_son_distintos(self):
        self.assertEqual(len({cv.PROBABLE, cv.POSIBLE, cv.VARIOS, cv.NINGUNO}), 4)
        self.assertEqual(cv.POSIBLE, "CANDIDATO POSIBLE — REVISIÓN MANUAL")

    def test_fecha_esperada_salta_domingos(self):
        self.assertEqual(cv.fecha_esperada(D(2026, 9, 3)), D(2026, 9, 4))     # jueves -> viernes
        self.assertEqual(cv.fecha_esperada(D(2026, 9, 5)), D(2026, 9, 7))     # sabado -> lunes (el domingo 6 no cuenta)
        self.assertIsNone(cv.fecha_esperada(None))

    def test_cuenta_de_banco(self):
        self.assertEqual(cv.cuenta_de_banco("110103012"), "110103012")
        self.assertEqual(cv.cuenta_de_banco(" BNB MN "), "110103012")
        self.assertIsNone(cv.cuenta_de_banco("BANCO INVENTADO"))
        self.assertIsNone(cv.cuenta_de_banco(None))

    def test_no_modifica_los_registros(self):
        regs = [dict(r) for r in MACROS_REAL]
        antes = json.dumps(regs, default=str)
        self.buscar(regs)
        self.assertEqual(json.dumps(regs, default=str), antes)


class Base(unittest.TestCase):
    def setUp(self):
        self._t = tempfile.TemporaryDirectory()
        self.d = self._t.name
        self.addCleanup(self._t.cleanup)

    def ruta(self, *p):
        return os.path.join(self.d, *p)


class TestLecturaMacros(Base):
    def test_registros_con_cuenta_y_caja(self):
        r = self.ruta("MACROS SEPTIEMBRE.xlsx")
        macros_completo(r, [(D(2026, 9, 4), "3P94114380", 9527.5, "SFC107", "110103012")])
        leido = mv.leer_indice_macros(r)
        self.assertEqual(leido["registros"], [{"fecha": D(2026, 9, 4), "importe": "9527.50", "codigo": "3P94114380",
                                                "cuenta": "110103012", "caja": "SFC107"}])
        self.assertEqual(leido["indice"], {("3P94114380", "9527.50"): [D(2026, 9, 4)]})       # el indice de siempre, intacto

    def test_macros_sin_columnas_opcionales(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = mv.HOJA_MACROS
        ws.append(["Fecha", "Código de Asignación", "Créditos"])
        ws.append([D(2026, 9, 4), "AB1", 100])
        r = self.ruta("M.xlsx")
        wb.save(r)
        reg_ = mv.leer_indice_macros(r)["registros"][0]
        self.assertEqual((reg_["cuenta"], reg_["caja"], reg_["importe"]), (None, None, "100.00"))


class TestControl5NoAutocorrige(Base):
    """El Control 5 sigue clasificando igual y NUNCA usa un candidato para corregir."""

    def analizar(self, deposito, registros, banco=110103012, fecha=D(2026, 9, 3)):
        r = self.ruta("CIERRE %02d-09-2026.xlsm" % fecha.day)
        cierre(r, deposito=deposito, sfcs=("SFC107", "SFC108"), banco=banco)
        return c5.analizar(r, {}, "america", fecha_cierre=fecha, registros_macros=registros)

    def test_ejemplo_real_queda_en_revision_con_candidato_probable(self):
        inf = self.analizar((9527.5, D(2026, 4, 9), "3P91155531"), MACROS_REAL)
        self.assertEqual(len(inf["filas"]), 1)
        f = inf["filas"][0]
        self.assertEqual((f["clase"], f["accion"], f["motivo"]), (c5.REQUIERE_REVISION, "NINGUNA", "SIN_VOUCHER_EN_MACROS"))
        self.assertIsNone(f["fecha_nueva"])
        self.assertEqual(inf["plan"], [])                       # nada que escribir: NO se autocorrige
        self.assertEqual(f["valor_actual"], "09/04/2026")
        c = f["candidatos"]
        self.assertEqual((c["veredicto"], c["fuerza"]), (cv.PROBABLE, cv.FUERTE))
        self.assertEqual(c["declarado"]["cuenta"], "110103012")
        self.assertEqual(c["candidatos"][0]["asignacion"], "3P94114380")

    def test_sin_registros_de_macros_no_hay_busqueda(self):
        inf = self.analizar((9527.5, D(2026, 4, 9), "3P91155531"), None)
        self.assertEqual((inf["filas"][0]["motivo"], inf["filas"][0]["candidatos"]), ("SIN_VOUCHER_EN_MACROS", None))

    def test_banco_escrito_como_nombre_se_traduce_a_cuenta(self):
        inf = self.analizar((9527.5, D(2026, 4, 9), "3P91155531"), MACROS_REAL, banco="BNB MN")
        self.assertEqual(inf["filas"][0]["candidatos"]["declarado"]["cuenta"], "110103012")
        self.assertEqual(inf["filas"][0]["candidatos"]["veredicto"], cv.PROBABLE)

    def test_asignacion_vacia_tambien_busca_candidatos(self):
        inf = self.analizar((9527.5, D(2026, 9, 4), None), MACROS_REAL)
        f = inf["filas"][0]
        self.assertEqual((f["motivo"], f["accion"]), ("ASIGNACION_VACIA_SIN_VOUCHER", "NINGUNA"))
        self.assertEqual(f["candidatos"]["veredicto"], cv.PROBABLE)

    def _sin_candidatos(self, inf):
        return [{k: v for k, v in f.items() if k != "candidatos"} for f in inf["filas"]]

    def test_texto_sin_voucher_se_normaliza_igual_y_ademas_busca_candidatos(self):
        registros = [reg(D(2026, 9, 4), "3P94114380", "9527.50")]
        con = self.analizar((9527.5, "04/09/26", "NO-ESTA"), registros)
        sin = self.analizar((9527.5, "04/09/26", "NO-ESTA"), None)
        f = con["filas"][0]
        self.assertEqual((f["clase"], f["accion"], f["motivo"], f["fecha_nueva"]),
                         (c5.CONVERTIR_TEXTO_SIN_VOUCHER, "CONVERTIR", c5.MOTIVO_SIN_VOUCHER, D(2026, 9, 4)))
        self.assertEqual((f["candidatos"]["veredicto"], f["candidatos"]["candidatos"][0]["asignacion"]), (cv.PROBABLE, "3P94114380"))
        # con o sin busqueda, la clasificacion y el plan son EXACTAMENTE los mismos
        self.assertIsNone(sin["filas"][0]["candidatos"])
        self.assertEqual(self._sin_candidatos(con), self._sin_candidatos(sin))
        self.assertEqual([(p["celda"], p["accion"], p["fecha_nueva"]) for p in con["plan"]],
                         [(p["celda"], p["accion"], p["fecha_nueva"]) for p in sin["plan"]])
        self.assertEqual(len(con["plan"]), 1)

    def test_texto_sin_voucher_sin_candidatos_tambien_lo_informa(self):
        f = self.analizar((9527.5, "04/09/26", "NO-ESTA"), [reg(D(2026, 9, 4), "X", "1.00")])["filas"][0]
        self.assertEqual((f["clase"], f["candidatos"]["veredicto"]), (c5.CONVERTIR_TEXTO_SIN_VOUCHER, cv.NINGUNO))

    def test_una_falla_en_la_busqueda_no_altera_la_normalizacion_ni_la_revision(self):
        original = cv.buscar

        def falla(*a, **k):
            raise RuntimeError("falla simulada")
        cv.buscar = falla
        try:
            texto = self.analizar((9527.5, "04/09/26", "NO-ESTA"), MACROS_REAL)
            revision = self.analizar((9527.5, D(2026, 4, 9), "3P91155531"), MACROS_REAL)
        finally:
            cv.buscar = original
        f = texto["filas"][0]
        self.assertEqual((f["clase"], f["accion"], f["fecha_nueva"], f["candidatos"]),
                         (c5.CONVERTIR_TEXTO_SIN_VOUCHER, "CONVERTIR", D(2026, 9, 4), None))
        self.assertEqual(len(texto["plan"]), 1)
        g = revision["filas"][0]
        self.assertEqual((g["clase"], g["motivo"], g["candidatos"]), (c5.REQUIERE_REVISION, "SIN_VOUCHER_EN_MACROS", None))

    def test_voucher_exacto_no_busca_candidatos(self):
        r = self.ruta("CIERRE 03-09-2026.xlsm")
        cierre(r, deposito=(9527.5, D(2026, 9, 4), "3P94114380"), sfcs=("SFC107", "SFC108"), banco=110103012)
        idx = {("3P94114380", "9527.50"): [D(2026, 9, 4)]}
        f = c5.analizar(r, idx, "america", fecha_cierre=D(2026, 9, 3), registros_macros=MACROS_REAL)["filas"][0]
        self.assertEqual((f["clase"], f["candidatos"]), (c5.CORRECTA, None))


def _sha(ruta):
    with open(ruta, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def _docx_tecnico(out):
    from docx import Document
    with zipfile.ZipFile(out["archivos"]["zip"]) as z:
        return Document(io.BytesIO(z.read("INFORME_AUDITORIA_CIERRES.docx")))


class TestInformeYCorrida(Base):
    """De punta a punta: auditar() -> stdout, JSON, PDF y DOCX con el detalle humano de cada veredicto."""

    @classmethod
    def setUpClass(cls):
        cls._t = tempfile.mkdtemp(prefix="test_cand_")
        d = cls._t
        mac = os.path.join(d, "MACROS SEPTIEMBRE.xlsm")
        macros_completo(mac, [
            (D(2026, 9, 1), "3P91155531", 25000, "SFC107", "110103012"),
            (D(2026, 9, 3), "3P93097190", 19158.2, "SFC107", "110103012"),
            (D(2026, 9, 4), "3P94114380", 9527.5, "SFC107", "110103012"),          # 03/09: candidato probable
            (D(2026, 9, 10), "3P9100001", 7000, "SFC107", "110103012"),            # 09/09: dos candidatos fuertes
            (D(2026, 9, 10), "3P9100002", 7000, "SFC107", "110103012"),
            (D(2026, 9, 28), "3P9280001", 4444.44, "SFC107", "110103012"),         # 17/09: importe unico, muy fuera de la ventana
        ])
        cls.grupos, cls.hashes = [], {}
        casos = [(3, (9527.5, D(2026, 4, 9), "3P91155531")),        # ejemplo real
                 (9, (7000, D(2026, 9, 10), "NO-ESTA")),            # varios candidatos
                 (16, (3333.33, D(2026, 9, 17), "NO-ESTA")),        # ningun candidato
                 (17, (4444.44, D(2026, 9, 18), "NO-ESTA"))]        # fuera de +-1 dia, importe unico
        os.makedirs(os.path.join(d, "america"))
        for dia, dep in casos:
            r = os.path.join(d, "america", "CIERRE %02d-09-2026.xlsm" % dia)
            cierre(r, deposito=dep, sfcs=("SFC107", "SFC108"), banco=110103012)
            cls.grupos.append(("america", r))
            cls.hashes[r] = _sha(r)
        cls.rep = os.path.join(d, "rep")
        cls.out = A.auditar(cls.grupos, [mac], salida_dir=os.path.join(cls.rep, "CIERRES_NORMALIZADOS"), reporte_dir=cls.rep,
                            hoy=D(2026, 9, 26))
        with zipfile.ZipFile(cls.out["archivos"]["zip"]) as z:
            cls.det = json.loads(z.read("RESULTADOS_AUDITORIA.json").decode("utf-8"))
            cls.nombres_zip = z.namelist()
            cls.pdf_tecnico = os.path.join(cls._t, "INFORME_AUDITORIA_CIERRES.pdf")
            with open(cls.pdf_tecnico, "wb") as f:
                f.write(z.read("INFORME_AUDITORIA_CIERRES.pdf"))
        cls.m = cls.det["informe_humano"]

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls._t, True)

    def test_no_se_autocorrige_nada_ni_se_tocan_los_originales(self):
        self.assertEqual(self.out["fechas_normalizadas"], 0)
        self.assertFalse([n for n in self.nombres_zip if n.startswith("CIERRES_NORMALIZADOS/")])
        for _caja, r in self.grupos:
            self.assertEqual(_sha(r), self.hashes[r])
        self.assertFalse(self.det["originales_drive_modificados"])
        self.assertEqual([c["estado"] for c in self.det["cierres"]], ["REVISAR"] * 4)

    def test_los_cuatro_veredictos_en_orden(self):
        self.assertEqual([b["veredicto"] for b in self.m["vouchers"]],
                         ["CANDIDATO PROBABLE", "VARIOS CANDIDATOS — REVISIÓN MANUAL", "SIN CANDIDATOS EN MACROS",
                          "CANDIDATO POSIBLE — REVISIÓN MANUAL"])
        self.assertEqual([b["titulo"] for b in self.m["vouchers"]],
                         ["América · 03/09/2026 · SFC107 · Depósito 1 (celda F3)", "América · 09/09/2026 · SFC107 · Depósito 1 (celda F3)",
                          "América · 16/09/2026 · SFC107 · Depósito 1 (celda F3)", "América · 17/09/2026 · SFC107 · Depósito 1 (celda F3)"])

    def test_ejemplo_real_detalle_humano(self):
        b = self.m["vouchers"][0]
        self.assertEqual(b["declarados"], "Datos declarados en el cierre: importe Bs 9.527,50; fecha de depósito 09/04/2026; "
                                          "asignación 3P91155531; cuenta/banco 110103012.")
        self.assertEqual(b["columnas"], ["Fecha", "Importe", "Asignación / Voucher", "Cuenta / Banco", "Caja / SFC", "Coincidencias",
                                         "Diferencias"])
        self.assertEqual(len(b["filas"]), 1)
        fecha, importe, asig, cuenta, caja, coin, dif = b["filas"][0]
        self.assertEqual((fecha, importe, asig, cuenta, caja), ("04/09/2026", "Bs 9.527,50", "3P94114380", "110103012", "SFC107"))
        for t in ("Importe exacto", "Cuenta 110103012", "Caja SFC107", "Fecha dentro de la ventana esperada", "invertidos"):
            self.assertIn(t, coin)
        self.assertIn("Asignación distinta (cierre: 3P91155531; MACROS: 3P94114380)", dif)
        self.assertIn("Fecha distinta (cierre: 09/04/2026; MACROS: 04/09/2026)", dif)
        self.assertEqual(b["interpretacion"],
                         "Interpretación probable: Es probable que el depósito corresponda al voucher de MACROS del 04/09/2026 por "
                         "Bs 9.527,50 (asignación 3P94114380): coinciden el importe, la cuenta y la caja, y la fecha está dentro de lo "
                         "esperado. El día y el mes de la fecha declarada (09/04/2026) parecen estar invertidos: MACROS indica "
                         "04/09/2026. La asignación escrita (3P91155531) no coincide con la del voucher (3P94114380). En MACROS, la "
                         "asignación escrita en el cierre corresponde a otro abono (01/09/2026, Bs 25.000).")
        self.assertTrue(b["accion"].startswith("Acción requerida al cajero: Confirmar con el comprobante bancario"))
        self.assertTrue(b["accion"].endswith("El auditor no modifica el cierre."))

    def test_varios_candidatos(self):
        b = self.m["vouchers"][1]
        self.assertEqual([f[2] for f in b["filas"]], ["3P9100001", "3P9100002"])
        self.assertIn("Hay 2 registros posibles en MACROS con el mismo importe", b["interpretacion"])
        self.assertIn("Revisar con el comprobante bancario cuál", b["accion"])

    def test_sin_candidatos(self):
        b = self.m["vouchers"][2]
        self.assertEqual(b["filas"], [])
        self.assertIn("No se encontró en MACROS ningún registro razonable con el mismo importe (Bs 3.333,33).", b["interpretacion"])
        self.assertIn("Verificar el importe, la fecha y la asignación", b["accion"])

    def test_fuera_de_ventana_con_importe_unico(self):
        b = self.m["vouchers"][3]
        self.assertEqual(len(b["filas"]), 1)
        fecha, _imp, asig, _cta, _caja, coin, dif = b["filas"][0]
        self.assertEqual((fecha, asig), ("28/09/2026", "3P9280001"))
        self.assertIn("Importe único en MACROS", coin)
        self.assertIn("Fecha fuera de la ventana esperada (17/09/2026 a 19/09/2026)", dif)
        self.assertIn("fuera de la ventana esperada (17/09/2026 a 19/09/2026)", b["interpretacion"])
        self.assertIn("Ese importe no aparece en ningún otro registro de MACROS.", b["interpretacion"])

    def test_la_tabla_de_revision_remite_al_detalle(self):
        filas = {f[1]: f[2] for f in self.m["revisar"] if "Búsqueda de candidatos en MACROS" in f[2]}
        self.assertEqual(sorted(filas), ["03/09/2026", "09/09/2026", "16/09/2026", "17/09/2026"])
        self.assertIn("El voucher informado no fue encontrado en MACROS.", filas["03/09/2026"])
        self.assertIn("Búsqueda de candidatos en MACROS: CANDIDATO PROBABLE (ver detalle en “2.1 Detalle de vouchers no encontrados "
                      "en MACROS”).", filas["03/09/2026"])

    def test_salida_compacta(self):
        self.assertEqual(self.out["candidatos_voucher"], [
            "03/09 AME SFC107 · CANDIDATO PROBABLE — MACROS 04/09/2026, Bs 9527.50, asignación 3P94114380",
            "09/09 AME SFC107 · VARIOS CANDIDATOS — REVISIÓN MANUAL",
            "16/09 AME SFC107 · SIN CANDIDATOS EN MACROS",
            "17/09 AME SFC107 · CANDIDATO POSIBLE — REVISIÓN MANUAL — MACROS 28/09/2026, Bs 4444.44, asignación 3P9280001"])

    def test_el_json_conserva_el_detalle_tecnico(self):
        hs = [h for h in self.det["hallazgos"] if (h.get("datos") or {}).get("candidatos")]
        self.assertEqual(len(hs), 4)
        self.assertEqual(hs[0]["datos"]["candidatos"]["candidatos"][0]["asignacion"], "3P94114380")
        self.assertEqual(hs[0]["accion"], "Revisar con el voucher; no se modifica")       # la accion de siempre

    def test_docx_muestra_el_detalle(self):
        doc = _docx_tecnico(self.out)
        parrafos = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn(ih.T_SEC2B, parrafos)
        self.assertIn(ih.INTRO_VOUCHERS, parrafos)
        for linea in ih.LEYENDA_VOUCHERS:                      # los cuatro veredictos se explican y se distinguen
            self.assertIn(linea, parrafos)
        for b in self.m["vouchers"]:
            for t in (b["titulo"], b["veredicto"], b["declarados"], b["interpretacion"], b["accion"]):
                self.assertIn(t, parrafos)
        tablas = [[[c.text for c in fila.cells] for fila in t.rows] for t in doc.tables]
        for b in self.m["vouchers"]:
            if b["filas"]:
                self.assertIn([b["columnas"]] + b["filas"], tablas)

    def test_pdf_muestra_el_detalle(self):
        exe = shutil.which("pdftotext")
        if not exe:
            self.skipTest("pdftotext no disponible")
        txt = subprocess.run([exe, "-enc", "UTF-8", self.pdf_tecnico, "-"], capture_output=True, text=True,
                             encoding="utf-8").stdout
        plano = " ".join(txt.split())
        self.assertIn(ih.T_SEC2B, plano)
        for palabra in ("CANDIDATO PROBABLE", "CANDIDATO POSIBLE", "VARIOS CANDIDATOS", "SIN CANDIDATOS EN MACROS", "3P94114380", "3P91155531", "9.527,50",
                        "Datos declarados en el cierre", "Interpretación probable", "Acción requerida al cajero"):
            self.assertIn(palabra, plano)

    def test_leyenda_de_los_cuatro_veredictos_en_el_modelo(self):
        self.assertEqual(self.m["vouchers_leyenda"], ih.LEYENDA_VOUCHERS)
        self.assertEqual([l.split(":")[0] for l in self.m["vouchers_leyenda"]],
                         ["CANDIDATO PROBABLE", "CANDIDATO POSIBLE — REVISIÓN MANUAL", "VARIOS CANDIDATOS — REVISIÓN MANUAL",
                          "SIN CANDIDATOS EN MACROS"])

    def test_candidato_posible_explica_por_que_no_es_probable(self):
        b = self.m["vouchers"][3]
        self.assertIn("Por eso no puede darse como probable", b["interpretacion"])
        self.assertIn("requiere revisión manual", b["interpretacion"])
        self.assertIn("(no todos los criterios coinciden o pudieron verificarse)", b["accion"])

    def test_sin_codigos_internos_en_el_informe(self):
        textos = " ".join(ih.textos_del_modelo(self.m))
        for prohibido in ("SIN_VOUCHER_EN_MACROS", "FUERA_DE_VENTANA", "PARCIAL", "FUERTE", "Traceback", "\\"):
            self.assertNotIn(prohibido, textos)


class TestTextoNormalizadoConCandidatos(Base):
    """Fecha de texto sin voucher que se normaliza sola: la normalizacion es la de siempre; los candidatos son evidencia adicional."""

    def corrida(self, nombre, filas_macros):
        mac = self.ruta(nombre, "MACROS SEPTIEMBRE.xlsm")
        os.makedirs(os.path.dirname(mac))
        macros_completo(mac, filas_macros)
        r = self.ruta(nombre, "america", "CIERRE 09-09-2026.xlsm")
        os.makedirs(os.path.dirname(r))
        cierre(r, deposito=(5000, "10/09/26", "NO-ESTA"), sfcs=("SFC107", "SFC108"), banco=110103012)
        rep = self.ruta(nombre, "rep")
        out = A.auditar([("america", r)], [mac], salida_dir=os.path.join(rep, "CIERRES_NORMALIZADOS"), reporte_dir=rep,
                        hoy=D(2026, 9, 26))
        with zipfile.ZipFile(out["archivos"]["zip"]) as z:
            det = json.loads(z.read("RESULTADOS_AUDITORIA.json").decode("utf-8"))
            copia = z.read("CIERRES_NORMALIZADOS/AMERICA/CIERRE 09-09-2026.xlsm")
        return out, det, copia

    def test_normaliza_igual_con_y_sin_candidatos(self):
        con, det_con, copia_con = self.corrida("corrida_a", [(D(2026, 9, 10), "3P9100001", 5000, "SFC107", "110103012")])
        sin, det_sin, copia_sin = self.corrida("corrida_b", [(D(2026, 9, 10), "3P9100001", 1234, "SFC107", "110103012")])
        for out in (con, sin):
            self.assertEqual((out["fechas_normalizadas"], out["normalizadas_sin_voucher_macros"]), (1, 1))
        # El contenido OOXML es idéntico; las marcas de tiempo internas del contenedor ZIP pueden variar entre corridas.
        with zipfile.ZipFile(io.BytesIO(copia_con)) as za, zipfile.ZipFile(io.BytesIO(copia_sin)) as zb:
            self.assertEqual({n: za.read(n) for n in za.namelist() if n != "docProps/core.xml"},
                             {n: zb.read(n) for n in zb.namelist() if n != "docProps/core.xml"})
        m_con, m_sin = det_con["informe_humano"], det_sin["informe_humano"]
        self.assertEqual(m_con["correcciones"], m_sin["correcciones"])           # y la correccion informada, identica
        self.assertEqual(m_con["correcciones"][0][3:5], ["10/09/26", "10/09/2026"])
        # solo cambia la evidencia adicional
        self.assertEqual([b["veredicto"] for b in m_con["vouchers"]], [cv.PROBABLE])
        self.assertEqual([b["veredicto"] for b in m_sin["vouchers"]], [cv.NINGUNO])

    def test_el_informe_muestra_el_detalle_del_caso_normalizado(self):
        out, det, _ = self.corrida("corrida_a", [(D(2026, 9, 10), "3P9100001", 5000, "SFC107", "110103012")])
        b = det["informe_humano"]["vouchers"][0]
        self.assertEqual(b["titulo"], "América · 09/09/2026 · SFC107 · Depósito 1 (celda F3)")
        self.assertEqual(b["filas"][0][:5], ["10/09/2026", "Bs 5.000", "3P9100001", "110103012", "SFC107"])
        self.assertIn("La normalización del formato de la fecha del cierre no depende de esta búsqueda", b["interpretacion"])
        self.assertEqual(out["candidatos_voucher"], ["09/09 AME SFC107 · CANDIDATO PROBABLE — MACROS 10/09/2026, Bs 5000.00, "
                                                     "asignación 3P9100001"])
        # el caso normalizado NO es un caso a revisar: no cambia el estado del cierre ni aparece como pendiente
        self.assertEqual(det["informe_humano"]["revisar"], [])
        self.assertEqual(out["cierres"], {"revisados": 1, "sin_observaciones": 1, "requieren_revision": 0})
        parrafos = "\n".join(p.text for p in _docx_tecnico(out).paragraphs)
        self.assertIn(ih.T_SEC2B, parrafos)


class TestSinCandidatosEnElInforme(Base):
    """Sin depositos sin voucher, el informe no agrega la seccion 2.1 (nada cambia respecto de antes)."""

    def test_no_hay_seccion_si_no_hay_casos(self):
        mac = self.ruta("MACROS SEPTIEMBRE.xlsm")
        macros_completo(mac, [(D(2026, 9, 4), "3P94114380", 9527.5, "SFC107", "110103012")])
        r = self.ruta("america", "CIERRE 03-09-2026.xlsm")
        os.makedirs(os.path.dirname(r))
        cierre(r, deposito=(9527.5, D(2026, 9, 4), "3P94114380"), sfcs=("SFC107", "SFC108"), banco=110103012)
        out = A.auditar([("america", r)], [mac], salida_dir=self.ruta("rep", "CIERRES_NORMALIZADOS"), reporte_dir=self.ruta("rep"),
                        hoy=D(2026, 9, 26))
        self.assertNotIn("candidatos_voucher", out)
        with zipfile.ZipFile(out["archivos"]["zip"]) as z:
            m = json.loads(z.read("RESULTADOS_AUDITORIA.json").decode("utf-8"))["informe_humano"]
        self.assertEqual(m["vouchers"], [])
        self.assertNotIn(ih.T_SEC2B, "\n".join(p.text for p in _docx_tecnico(out).paragraphs))


if __name__ == "__main__":
    unittest.main()
