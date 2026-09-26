"""informe_humano.py — Informe de auditoria para PERSONAS (PDF principal + DOCX editable).

SOLO PRESENTACION. No contiene ninguna regla de auditoria: traduce los hallazgos ya calculados (con sus `datos`
estructurados) a lenguaje natural y los maqueta. Un unico modelo de contenido (`construir_modelo`) alimenta los dos
escritores (`escribir_pdf`, `escribir_docx`), de modo que ambos archivos dicen exactamente lo mismo.

Requisitos: reportlab (PDF) y python-docx (DOCX). Sin codigos internos, sin trazas de Python, sin rutas.
"""
import datetime
import re
import unicodedata
from decimal import Decimal
from xml.sax.saxutils import escape

import candidatos_voucher as cv
import controles_lectura as cl

TITULO = "INFORME DE AUDITORÍA DE CIERRES DE CAJA"
PIE = "CAJAS GABO — Auditoría de cierres"
SIN, REQUIERE, INFORMATIVA = "SIN OBSERVACIONES", "REQUIERE REVISIÓN", "OBSERVACIÓN INFORMATIVA"
ACCION_NINGUNA = "Acción requerida: Ninguna."
CAJAS = {"tiquipaya": "Tiquipaya", "america": "América"}
MAX_ITEMS_EN_CELDA = 5

T_SEC1 = "1. Resumen ejecutivo"
T_SEC2 = "2. Casos que requieren revisión"
T_SEC2B = "2.1 Detalle de vouchers no encontrados en MACROS"
T_SEC3 = "3. Correcciones automáticas"
T_SEC4 = "4. Observaciones informativas"
T_SEC5 = "5. Conclusión"
T_ANEXO = "Anexo. Información complementaria"

COLS_CIERRES = ["Caja", "Fecha del cierre", "Estado", "Motivo principal"]
COLS_REVISAR = ["Caja", "Fecha", "Qué ocurrió", "Qué debe revisarse"]
COLS_VOUCHER = ["Fecha", "Importe", "Asignación / Voucher", "Cuenta / Banco", "Caja / SFC", "Coincidencias", "Diferencias"]
ANCHOS_VOUCHER = [0.13, 0.10, 0.14, 0.12, 0.09, 0.21, 0.21]
INTRO_VOUCHERS = ("Para cada depósito cuyo voucher no se encontró en MACROS (asignación + importe) se buscaron registros con el mismo "
                  "importe, la misma cuenta y caja, y una fecha cercana a la esperada (±1 día). Los candidatos son solo una ayuda para "
                  "la revisión: el auditor no corrige ningún dato del cierre.")
LEYENDA_VOUCHERS = [
    "%s: un único registro de MACROS con el importe exacto, la misma cuenta, la misma caja y la fecha esperada (±1 día)." % cv.PROBABLE,
    "%s: un único registro con el importe exacto, pero fuera de ±1 día, con coincidencia parcial o con algún dato que no pudo "
    "verificarse." % cv.POSIBLE,
    "%s: más de un registro razonable; no se puede elegir automáticamente." % cv.VARIOS,
    "%s: ningún registro razonable con ese importe." % cv.NINGUNO,
]
COLS_CORRECCIONES = ["Caja", "Fecha del cierre", "Campo", "Valor original", "Valor normalizado", "Motivo"]
COLS_INFO = ["Caja", "Fecha", "Observación"]
COLS_ALQUILERES = ["Fecha", "Caja", "Importe (Bs)"]


# ---------------------------------------------------------------------------
# utilidades de texto
# ---------------------------------------------------------------------------

def limpio(s):
    """Solo caracteres representables en WinAnsi (cp1252): tipografia estandar del PDF y del DOCX."""
    s = str(s)
    out = []
    for ch in s:
        try:
            ch.encode("cp1252")
            out.append(ch)
        except UnicodeEncodeError:
            base = unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode()
            out.append(base or "?")
    return "".join(out)


def _f(d):
    return d.strftime("%d/%m/%Y") if isinstance(d, (datetime.date, datetime.datetime)) else "fecha no identificada"


def _bs(x):
    return "Bs " + cl.bs(Decimal(str(x)))


def _caja(c):
    return CAJAS.get(c, str(c).title())


def _lista(items, maximo=MAX_ITEMS_EN_CELDA):
    items = [str(i) for i in items]
    if len(items) <= maximo:
        return cl._unir(items)
    return ", ".join(items[:maximo]) + " y %d más" % (len(items) - maximo)


def _filas(hs):
    return _lista(sorted({h["fila"] for h in hs if h.get("fila") is not None}))


def _plural(n, uno, varios):
    return uno if n == 1 else varios


def _campo(c):
    return str(c).lower()


def _dep(d):
    d = str(d or "Depósito")
    return d.replace("DEPOSITO", "Depósito", 1) if d.upper().startswith("DEPOSITO") else d


def _original_texto(valor_actual):
    """'texto \\'23/09/26\\'' -> 23/09/26 ; '09/10/2026' -> 09/10/2026."""
    v = str(valor_actual or "")
    m = re.match(r"^texto (['\"])(.*)\1$", v)
    return m.group(2) if m else v


_CAT = {"OTROS INGRESOS": "Otros ingresos", "GASTO.ADM": "Gastos administrativos (GASTO.ADM)",
        "POSGRADO.PLA": "Posgrado (POSGRADO.PLA)"}

# motivo (Control 5) -> (que ocurrio, que debe revisarse)
_C5_REVISION = {
    "CELDA_CON_FORMULA_NO_SE_MODIFICA": ("La fecha de depósito es una fórmula y no se modifica automáticamente.",
                                         "Verificar la fecha manualmente."),
    "TIPO_DE_CELDA_NO_SOPORTADO": ("La fecha de depósito tiene un tipo de dato que no pudo interpretarse.",
                                   "Verificar la fecha manualmente."),
    "TEXTO_FORMATO_NO_ADMITIDO": ("La fecha registrada no pudo interpretarse de forma segura.",
                                  "Verificar la fecha con el comprobante del depósito y corregirla en el cierre."),
    "TEXTO_ANIO_2_DIGITOS_NO_COINCIDE_CON_ANIO_DEL_CIERRE": (
        "La fecha registrada tiene un año de dos dígitos que no coincide con el año del cierre.",
        "Verificar la fecha con el comprobante del depósito y corregirla en el cierre."),
    "TEXTO_FECHA_INEXISTENTE_EN_CALENDARIO": ("La fecha registrada no existe en el calendario.",
                                              "Verificar la fecha con el comprobante del depósito y corregirla en el cierre."),
    "NUMERO_FUERA_DE_RANGO_DE_FECHAS": ("El valor registrado como fecha de depósito no es una fecha válida.",
                                        "Verificar la fecha con el comprobante del depósito."),
    "NUMERO_SIN_FORMATO_DE_FECHA": ("El valor registrado como fecha de depósito no tiene formato de fecha.",
                                    "Verificar la fecha con el comprobante del depósito."),
    "FECHA_CON_HORA_NO_SE_MODIFICA": ("La fecha de depósito incluye hora y no se modifica automáticamente.",
                                      "Verificar la fecha manualmente."),
    "IMPORTE_NO_NUMERICO": ("El importe del depósito no es un número válido, por lo que no se pudo buscar su voucher.",
                            "Corregir el importe en el cierre."),
    "ASIGNACION_VACIA_SIN_VOUCHER": ("El depósito no tiene código de asignación, por lo que no se puede contrastar con MACROS.",
                                     "Completar la asignación en el cierre."),
    "SIN_VOUCHER_EN_MACROS": ("El voucher informado no fue encontrado en MACROS.",
                              "Confirmar que el depósito figure en MACROS (código de asignación e importe) o verificar la fecha con el comprobante."),
    "VOUCHER_SIN_FECHA_VALIDA": ("El voucher encontrado en MACROS no tiene una fecha válida.",
                                 "Verificar la fecha con el comprobante del depósito."),
    "TEXTO_DISTINTO_A_VOUCHER": ("La fecha registrada no coincide con la fecha del voucher en MACROS.",
                                 "Verificar la fecha con el comprobante del depósito y corregirla en el cierre si corresponde."),
    "FECHA_DISTINTA_A_VOUCHER_Y_SU_INVERSION_TAMPOCO_COINCIDE": (
        "La fecha registrada no coincide con la fecha del voucher en MACROS.",
        "Verificar la fecha con el comprobante del depósito y corregirla en el cierre si corresponde."),
}


# ---------------------------------------------------------------------------
# traduccion de hallazgos -> entradas humanas
# ---------------------------------------------------------------------------
# Cada entrada: {"caja","fecha","bucket": revisar|correccion|info, "orden", "corto", "ocurrio", "revisar", ...}

def _ent(h, bucket, orden, corto, ocurrio, revisar=None, **extra):
    e = {"caja": h["caja"], "fecha": h["fecha"], "bucket": bucket, "orden": orden, "corto": corto, "ocurrio": ocurrio,
         "revisar": revisar}
    e.update(extra)
    return e


def _lineas(items):
    return "\n".join("• " + str(i) for i in items) if items else "—"


def _hasta(txt):
    return datetime.datetime.strptime(txt, "%d/%m/%Y").date() if txt else None


def _interpretacion(c, dec_asig, fecha_dec):
    """Texto humano de lo que probablemente ocurrio (solo presentacion del resultado de candidatos_voucher)."""
    v, cands = c["veredicto"], c["candidatos"]
    if v == cv.NINGUNO:
        txt = "No se encontró en MACROS ningún registro razonable con el mismo importe (%s)." % _bs(c["declarado"]["importe"])
        hasta, esp = _hasta(c.get("macros_hasta")), _hasta(c.get("fecha_esperada"))
        if hasta and esp and esp > hasta:
            return txt + " MACROS llega solo hasta el %s, por lo que el voucher puede no estar cargado todavía." % c["macros_hasta"]
        return txt + " Puede que el importe del cierre esté mal escrito o que el depósito no figure en MACROS."
    if v == cv.VARIOS:
        return ("Hay %d registros posibles en MACROS con el mismo importe; no se puede determinar automáticamente cuál corresponde "
                "al depósito." % (len(cands) + c.get("omitidos", 0)))
    p = cands[0]
    if p["fuerza"] == cv.FUERTE:
        txt = ("Es probable que el depósito corresponda al voucher de MACROS del %s por %s (asignación %s): coinciden el importe, la "
               "cuenta y la caja, y la fecha está dentro de lo esperado." % (p["fecha"], _bs(p["importe"]), p["asignacion"]))
    else:
        if p["fuerza"] == cv.PARCIAL:
            txt = ("Existe un único registro en MACROS con el mismo importe y una fecha cercana a la esperada (%s, asignación %s)"
                   % (p["fecha"], p["asignacion"]))
        else:
            txt = ("El único registro razonable con el mismo importe está en otra fecha (%s, asignación %s)"
                   % (p["fecha"] or "sin fecha válida", p["asignacion"]))
            if c.get("ventana") and p["fecha"]:
                txt += ", fuera de la ventana esperada (%s a %s)" % (c["ventana"][0], c["ventana"][1])
        motivos = ["no coincide la %s" % x for x in p.get("difiere", []) if x != "fecha"]
        motivos += ["no se pudo verificar la %s por falta de datos" % x for x in p.get("no_verificable", [])]
        txt += (" y " + cl._unir(motivos)) if motivos else ""
        txt += ". Por eso no puede darse como probable: podría ser el depósito o una coincidencia casual y requiere revisión manual."
        if c.get("importe_en_macros") == 1:
            txt += " Ese importe no aparece en ningún otro registro de MACROS."
    if p["fecha_invertida"]:
        txt += " El día y el mes de la fecha declarada (%s) parecen estar invertidos: MACROS indica %s." % (fecha_dec, p["fecha"])
    elif p["fecha_distinta"]:
        txt += " La fecha declarada (%s) difiere de la de MACROS (%s)." % (fecha_dec, p["fecha"])
    if p["asignacion_distinta"]:
        txt += " La asignación escrita (%s) no coincide con la del voucher (%s)." % (dec_asig, p["asignacion"])
        otro = c.get("asignacion_en_macros") or []
        if otro:
            o = otro[0]
            txt += " En MACROS, la asignación escrita en el cierre corresponde a otro abono (%s, %s)." % (o["fecha"], _bs(o["importe"]))
    elif not dec_asig:
        txt += " El cierre no declara asignación."
    return txt


def _accion_cajero(c):
    if c["veredicto"] == cv.PROBABLE:
        return ("Confirmar con el comprobante bancario que ese voucher corresponde al depósito y, si es así, corregir manualmente en "
                "el cierre lo que difiere (fecha y/o asignación). El auditor no modifica el cierre.")
    if c["veredicto"] == cv.POSIBLE:
        return ("Verificar con el comprobante bancario si ese registro corresponde al depósito (no todos los criterios coinciden o "
                "pudieron verificarse) y, si es así, corregir manualmente en el cierre lo que difiere. El auditor no modifica el cierre.")
    if c["veredicto"] == cv.VARIOS:
        return ("Revisar con el comprobante bancario cuál de los registros corresponde al depósito y corregir el cierre si es "
                "necesario. El auditor no modifica el cierre.")
    return ("Verificar el importe, la fecha y la asignación con el comprobante bancario; si el depósito es reciente, confirmar que "
            "MACROS ya lo tenga cargado. El auditor no modifica el cierre.")


def _bloque_voucher(h, d, sfc, normalizada=False):
    """Detalle humano de un deposito sin voucher exacto: datos declarados, candidatos, coincidencias/diferencias,
    interpretacion probable y accion requerida."""
    c = d["candidatos"]
    dec = c["declarado"]
    etiqueta = c["veredicto"]              # exactamente uno de los cuatro veredictos, sin calificativos
    filas = [[p["fecha"] or "sin fecha válida", _bs(p["importe"]), p["asignacion"] or "—", p["cuenta"] or "—", p["caja"] or "—",
              _lineas(p["coincidencias"]), _lineas(p["diferencias"])] for p in c["candidatos"]]
    interp = _interpretacion(c, dec["asignacion"], dec["fecha"] or "sin fecha")
    if c.get("omitidos"):
        interp += " Se muestran los %d candidatos más relevantes; hay %d más." % (len(c["candidatos"]), c["omitidos"])
    if normalizada:
        interp += (" La normalización del formato de la fecha del cierre no depende de esta búsqueda, que es solo evidencia "
                   "adicional.")
    return {
        "titulo": "%s · %s · %s · %s (celda %s)" % (_caja(h["caja"]), _f(h["fecha"]), sfc, _dep(d["deposito"]), d["celda"]),
        "veredicto": etiqueta,
        "declarados": ("Datos declarados en el cierre: importe %s; fecha de depósito %s; asignación %s; cuenta/banco %s." %
                       (_bs(dec["importe"]), dec["fecha"] or "sin fecha", dec["asignacion"] or "sin asignación",
                        dec["cuenta"] or "no identificada")),
        "columnas": list(COLS_VOUCHER),
        "filas": filas,
        "interpretacion": "Interpretación probable: " + interp,
        "accion": "Acción requerida al cajero: " + _accion_cajero(c),
    }


def _por_c5(h, d, sfc):
    """Entradas del Control 5 para una fila de deposito."""
    clase, aplicado = d["clase"], d["aplicado"]
    dep = "%s (%s, celda %s)" % (_dep(d["deposito"]), sfc, d["celda"])
    orig = _original_texto(d["valor_actual"])
    nuevo = d["fecha_nueva"]

    def correccion(motivo):
        return _ent(h, "correccion", 50, "Fecha de depósito normalizada automáticamente", None, None,
                    campo="Fecha de depósito\n(%s, celda %s)" % (sfc, d["celda"]), original=orig, normalizado=nuevo, motivo=motivo)

    def pendiente(que):
        razon = ("no se pudo aplicar la corrección automática" if d.get("error_aplicar")
                 else "no se generó una copia corregida")
        return _ent(h, "revisar", 45, "Fecha de depósito pendiente de normalizar",
                    "%s: %s, pero %s." % (dep, que, razon),
                    "Normalizar manualmente la fecha de depósito en el cierre.")

    if clase == "DENTRO_TOLERANCIA_1_DIA":
        return [_ent(h, "info", 60, "Fecha de depósito con un día de diferencia respecto de MACROS",
                     "La fecha del %s fue registrada como %s y el voucher en MACROS indica %s. La diferencia de un día está dentro "
                     "de la tolerancia permitida, por lo que la fecha no se modificó." % (dep[0].lower() + dep[1:], d["fecha_en_cierre"],
                                                                                          d["fecha_macros"]))]
    if clase == "CONVERTIR_TEXTO_DENTRO_TOLERANCIA_1_DIA":
        if not aplicado:
            return [pendiente("la fecha %s está escrita como texto y debía normalizarse" % orig)]
        return [correccion("Se normalizó el formato sin cambiar la fecha declarada. El voucher en MACROS indica %s (un día de "
                           "diferencia, dentro de la tolerancia permitida)." % d["fecha_macros"])]
    if clase == "CONVERTIR_TEXTO_SIN_VOUCHER_MACROS":
        if not aplicado:
            ents = [pendiente("la fecha %s está escrita como texto y debía normalizarse" % orig)]
        else:
            ents = [correccion("Se normalizó el formato sin cambiar la fecha declarada. No se encontró el voucher en MACROS, por lo "
                               "que la fecha no pudo contrastarse.")]
        if d.get("candidatos"):
            ents[0]["voucher"] = _bloque_voucher(h, d, sfc, normalizada=True)
        return ents
    if clase == "CONVERTIR_TEXTO_A_FECHA":
        if not aplicado:
            return [pendiente("la fecha %s está escrita como texto y debía convertirse" % orig)]
        return [correccion("Se convirtió el texto en una fecha real. Coincide con el voucher de MACROS (%s)." % d["fecha_voucher"])]
    if clase == "NORMALIZAR_INVERSION_DDMM":
        if not aplicado:
            return [pendiente("el día y el mes de la fecha %s parecen invertidos" % orig)]
        return [correccion("Se invirtieron el día y el mes: la fecha registrada era %s y el voucher de MACROS confirma %s." %
                           (orig, d["fecha_voucher"]))]
    if clase == "VACIA":
        return [_ent(h, "revisar", 40, "Falta una fecha de depósito", "Falta la fecha de depósito: %s." % dep,
                     "Completar la fecha manualmente (el auditor no la completa).")]
    motivo = d["motivo"] or ""
    if motivo.startswith("VOUCHER_NO_UNICO"):
        que, rev = ("Existe más de un voucher idéntico en MACROS; no se puede determinar cuál corresponde.",
                    "Verificar en MACROS cuál voucher corresponde y confirmar la fecha.")
    else:
        que, rev = _C5_REVISION.get(motivo, ("La fecha de depósito no pudo verificarse de forma segura.",
                                             "Verificar la fecha con el comprobante del depósito."))
    extra = " Fecha del voucher en MACROS: %s." % d["fecha_voucher"] if d.get("fecha_voucher") else ""
    valor = " Fecha registrada: %s." % orig if orig else ""
    if d.get("candidatos") and motivo in ("SIN_VOUCHER_EN_MACROS", "ASIGNACION_VACIA_SIN_VOUCHER"):
        bloque = _bloque_voucher(h, d, sfc)
        return [_ent(h, "revisar", 41, que.rstrip("."),
                     "%s: %s%s%s Búsqueda de candidatos en MACROS: %s (ver detalle en “%s”)." % (dep, que, valor, extra,
                                                                                               bloque["veredicto"], T_SEC2B),
                     rev, voucher=bloque)]
    return [_ent(h, "revisar", 41, que.rstrip("."), "%s: %s%s%s" % (dep, que, valor, extra), rev)]


def traducir(hallazgos):
    """Lista de entradas humanas. Agrupa filas repetidas del mismo problema dentro de un mismo cierre."""
    out, grupos = [], {}
    sin_macros_c5 = {(h["caja"], h["fecha"]) for h in hallazgos if (h.get("datos") or {}).get("tipo") == "c5_sin_macros"}
    for h in hallazgos:
        d = h.get("datos")
        if d is None or h["resultado"] == cl.OK and d.get("tipo") != "c5":
            continue                                    # cuadres y controles correctos: no aparecen
        t = d["tipo"]
        sfc = h["sfc"] or ""
        clave = (h["caja"], h["fecha"], t, sfc)
        if t == "c1_no_cuadra":
            dif = Decimal(d["diferencia"])
            out.append(_ent(h, "revisar", 10, "Diferencia de %s en el cuadre de %s" % (_bs(abs(dif)), sfc),
                            "Existe una diferencia de %s entre el total calculado y el registrado en %s." % (_bs(abs(dif)), sfc),
                            "Revisar los importes de %s antes de procesar el cierre." % sfc))
        elif t == "c1_faltan_campos":
            faltan = _lista([_campo(c) for c in d["campos"]])
            out.append(_ent(h, "revisar", 11, "No se pudo verificar el cuadre de %s" % sfc,
                            "No se pudo verificar el cuadre de %s porque falta el dato: %s." % (sfc, faltan),
                            "Completar el dato en el cierre y volver a auditar."))
        elif t == "c1_hoja_no_encontrada":
            out.append(_ent(h, "revisar", 12, "Falta la hoja %s" % sfc,
                            "El cierre no tiene la hoja %s, por lo que no se pudo verificar el cuadre." % sfc,
                            "Confirmar que el archivo del cierre esté completo."))
        elif t == "c2_encabezado_no_reconocido":
            out.append(_ent(h, "revisar", 20, "Comunicaciones Internas de %s no revisadas" % sfc,
                            "No se pudo revisar la hoja de Comunicaciones Internas de %s porque no se reconoció su estructura." % sfc,
                            "Revisar la hoja manualmente."))
        elif t == "c2_hoja_no_encontrada":
            out.append(_ent(h, "revisar", 21, "Falta la hoja de Comunicaciones Internas de %s" % sfc,
                            "El cierre no tiene la hoja de Comunicaciones Internas de %s." % sfc,
                            "Confirmar que el archivo del cierre esté completo."))
        elif t == "c2_faltan":
            grupos.setdefault(clave + (tuple(d["faltan"]),), []).append(h)
        elif t == "c2_especiales":
            cat = _CAT.get(d["categoria"], d["categoria"])
            n = len(d["filas"])
            base = "%s: %d %s (%s) por %s en %s" % (cat, n, _plural(n, "fila", "filas"), _lista(d["filas"]), _bs(d["importe"]), sfc)
            if d["completo"]:
                out.append(_ent(h, "info", 62, "%s registrado en %s" % (cat, sfc),
                                base + ", con cuenta contable y asignación completas."))
            else:
                out.append(_ent(h, "revisar", 22, "%s por completar en %s" % (cat, sfc),
                                base + ", sin cuenta contable y/o asignación.",
                                "Completar la cuenta contable y la asignación manualmente (el auditor no las completa)."))
        elif t in ("c3_banco_no_reconocido", "c3_sufijo_no_reconocido", "c3_cuenta_distinta", "c4_formato_asignacion",
                   "c6_alquiler_sin_importe"):
            grupos.setdefault(clave, []).append(h)
        elif t == "c6_fecha_del_cierre":
            n = d["cantidad"]
            out.append(_ent(h, "info", 63, "Alquileres sin fecha propia",
                            "%d %s no tenía una fecha propia válida; se usó la fecha del cierre." %
                            (n, _plural(n, "alquiler", "alquileres"))))
        elif t == "c5_hoja_no_reconocida":
            out.append(_ent(h, "revisar", 42, "Fechas de depósito de %s no revisadas" % sfc,
                            "No se pudieron revisar las fechas de depósito de %s porque no se reconoció la hoja." % sfc,
                            "Revisar la hoja manualmente."))
        elif t == "c5_sin_macros":
            causa = "no se pudo leer el archivo MACROS" if d.get("ilegible") else "no se dispuso del archivo MACROS"
            out.append(_ent(h, "revisar", 43, "Fechas de depósito sin contrastar con MACROS",
                            "Las fechas de depósito no se contrastaron con los vouchers porque %s. Solo se normalizó el formato "
                            "de las fechas escritas como texto válido." % causa,
                            "Contar con el MACROS oficial del mes y volver a auditar."))
        elif t == "c5":
            out += _por_c5(h, d, sfc)
        elif t == "cierre_ilegible":
            out.append(_ent(h, "revisar", 1, "No se pudo leer el archivo del cierre", "No se pudo leer el archivo del cierre.",
                            "Verificar que el archivo esté completo y volver a auditar."))
        elif t == "motor_no_lee":
            pendiente = d.get("archivo_validado") == "original_con_cambios_pendientes"
            out.append(_ent(h, "revisar", 70, "Validación con CAJAS GABO no concluida",
                            "La validación cruzada con CAJAS GABO no pudo leer el cierre" +
                            (" porque aún tiene fechas por normalizar." if pendiente else "."),
                            "Revisar el cierre antes de procesarlo con CAJAS GABO."))
        elif t == "motor_observaciones":
            partes = []
            msg = str(d.get("mensaje") or "")
            hasta = _f(datetime.date.fromisoformat(d["fecha_maxima_macros"])) if d.get("fecha_maxima_macros") else None
            if d.get("estado_maestro") not in (None, "MAESTRO_APTO", "SIN_MACROS"):
                if d.get("codigo_bloqueo") == "MACROS_NO_CUBRE_FECHA_DEPOSITO":
                    req = d.get("fecha_requerida_deposito")
                    partes.append("MACROS no cubre todavía la fecha requerida para validar el cierre%s." % (
                        " (los depósitos llegan hasta el %s y MACROS hasta el %s)" % (_f(datetime.date.fromisoformat(req)), hasta)
                        if req and hasta else ""))
                elif msg.startswith("MAESTRO_ILEGIBLE"):
                    partes.append("El motor de CAJAS GABO no pudo leer el archivo mensual de datos (MACROS / ATC), por lo que no "
                                  "pudo confirmar su cobertura para este cierre.")
                elif msg.startswith("MAESTRO_SIN_FECHAS_REGISTRADAS"):
                    partes.append("MACROS no tiene movimientos registrados, por lo que no pudo confirmarse su cobertura.")
                elif hasta and d.get("fecha_cierre") and d["fecha_cierre"] > d["fecha_maxima_macros"]:
                    partes.append("MACROS no cubre todavía la fecha del cierre (llega hasta el %s)." % hasta)
                elif msg.startswith("No se pudo leer el cierre para determinar"):
                    partes.append("El motor de CAJAS GABO no pudo determinar si el cierre tuvo movimiento ATC.")
                elif "información ATC" in msg:
                    partes.append("El archivo mensual no tiene todavía la información ATC de la fecha del cierre.")
                else:
                    partes.append("El motor de CAJAS GABO no pudo confirmar la cobertura de los datos mensuales para este cierre.")
            if d.get("estado_maestro") == "SIN_MACROS" and (h["caja"], h["fecha"]) not in sin_macros_c5:
                partes.append("No se pudo comprobar la cobertura de MACROS porque no hay un MACROS del mes.")
            for o in d.get("observaciones") or []:
                if str(o).startswith("FECHA_DEPOSITO_ANOMALA"):
                    partes.append("Un depósito tiene una fecha de un año distinto al del cierre (%s)." % str(o).split(": ")[-1])
                else:
                    partes.append("La validación con CAJAS GABO dejó una observación sobre las fechas de depósito.")
            if partes:
                out.append(_ent(h, "revisar", 71, partes[0].rstrip("."), " ".join(partes),
                                "Verificar antes de procesar el cierre con CAJAS GABO."))
        # motor_ok y demas: sin entrada
    for clave, hs in grupos.items():
        caja, fecha, t, sfc = clave[:4]
        h0, d0 = hs[0], hs[0]["datos"]
        n = len(hs)
        filas = _filas(hs)
        pl = _plural(n, "la fila", "las filas")
        if t == "c2_faltan":
            faltan = cl._unir([_campo(c) for c in d0["faltan"]])
            out.append(_ent(h0, "revisar", 23, "Datos incompletos en Comunicaciones Internas de %s" % sfc,
                            "En %s, %s %s de Comunicaciones Internas no tienen %s." % (sfc, pl, filas, faltan),
                            "Completar los datos faltantes en el cierre."))
        elif t == "c3_banco_no_reconocido":
            bancos = _lista(sorted({x["datos"]["banco"] for x in hs}))
            out.append(_ent(h0, "revisar", 24, "Banco no reconocido en %s" % sfc,
                            "En %s, %s %s: el banco %s no figura en la tabla oficial de bancos." % (sfc, pl, filas, bancos),
                            "Revisar el banco y la cuenta contable registrada."))
        elif t == "c3_sufijo_no_reconocido":
            bancos = _lista(sorted({x["datos"]["banco"] for x in hs}))
            out.append(_ent(h0, "revisar", 25, "Banco con descripción no reconocida en %s" % sfc,
                            "En %s, %s %s: %s no corresponde a una cuenta de la tabla oficial." % (sfc, pl, filas, bancos),
                            "Revisar el banco y la cuenta contable registrada."))
        elif t == "c3_cuenta_distinta":
            det = "; ".join("fila %s (banco %s): cuenta registrada %s, esperada %s" %
                            (x["fila"], x["datos"]["banco"], x["datos"]["registrada"], x["datos"]["esperada"]) for x in hs[:MAX_ITEMS_EN_CELDA])
            mas = " y %d más" % (n - MAX_ITEMS_EN_CELDA) if n > MAX_ITEMS_EN_CELDA else ""
            out.append(_ent(h0, "revisar", 26, "Cuenta contable distinta a la esperada en %s" % sfc,
                            "En %s la cuenta contable no coincide con la esperada para el banco: %s%s." % (sfc, det, mas),
                            "Revisar la cuenta contable registrada."))
        elif t == "c4_formato_asignacion":
            det = "; ".join("fila %s (banco %s): asignación %s, se esperaba %s" %
                            (x["fila"], x["datos"]["banco"], x["datos"]["asignacion"], x["datos"]["esperado"]) for x in hs[:MAX_ITEMS_EN_CELDA])
            mas = " y %d más" % (n - MAX_ITEMS_EN_CELDA) if n > MAX_ITEMS_EN_CELDA else ""
            out.append(_ent(h0, "revisar", 27, "Asignación con formato inesperado en %s" % sfc,
                            "En %s la asignación no tiene el formato esperado para el banco: %s%s." % (sfc, det, mas),
                            "Revisar el código de asignación."))
        elif t == "c6_alquiler_sin_importe":
            out.append(_ent(h0, "revisar", 28, "Alquiler sin importe legible",
                            "En %s, %s %s: hay %s sin importe legible." % (sfc, pl, filas, _plural(n, "un alquiler", "alquileres")),
                            "Revisar el importe manualmente."))
    return out


# ---------------------------------------------------------------------------
# modelo
# ---------------------------------------------------------------------------

def _conclusion(n, cajas, req, corr, motor_ok):
    partes = ["Se %s %d %s de caja (%s)." % (_plural(n, "revisó", "revisaron"), n, _plural(n, "cierre", "cierres"), cajas)]
    if req == 0:
        partes.append("Ningún cierre presenta observaciones que requieran intervención.")
    elif n - req > req:
        partes.append("La mayoría no presenta observaciones que requieran intervención.")
    else:
        partes.append("%d de %d cierres presentan observaciones que requieren intervención." % (req, n))
    if req:
        partes.append("Los casos detallados en la sección “Casos que requieren revisión” deben ser verificados antes de "
                      "considerar concluido el control.")
    if corr:
        partes.append("Las normalizaciones automáticas realizadas (%d) corrigieron únicamente la fecha de depósito y no alteraron "
                      "los importes registrados." % corr)
    else:
        partes.append("No se realizaron correcciones automáticas.")
    if not motor_ok:
        partes.append("La validación cruzada con el motor de CAJAS GABO no estuvo disponible en esta ejecución.")
    return " ".join(partes)


def construir_modelo(ejecucion, cierres, hallazgos, alquileres, motor=None, macros=None):
    """Contenido del informe (estructuras simples de texto). `cierres`: [{fecha, caja, estado, avisos, archivo}]."""
    ents = traducir(hallazgos)
    por_cierre = {}
    for e in ents:
        por_cierre.setdefault((e["caja"], e["fecha"]), []).append(e)

    filas_cierres, estados = [], {}
    for c in sorted(cierres, key=lambda c: (c["fecha"] or datetime.date.min, c["caja"])):
        es = sorted(por_cierre.get((c["caja"], c["fecha"]), []), key=lambda e: e["orden"])
        rev = [e for e in es if e["bucket"] == "revisar"]
        otros = [e for e in es if e["bucket"] != "revisar"]
        if rev:
            estado = REQUIERE
            motivo = rev[0]["corto"] + (" (y %d más)" % (len(rev) - 1) if len(rev) > 1 else "")
        elif otros:
            estado = INFORMATIVA
            motivo = otros[0]["corto"] + (" (y %d más)" % (len(otros) - 1) if len(otros) > 1 else "")
        else:
            estado, motivo = SIN, "Sin observaciones."
        estados[(c["caja"], c["fecha"])] = estado
        filas_cierres.append([_caja(c["caja"]), _f(c["fecha"]), estado, motivo])

    def clave(e):
        return (e["fecha"] or datetime.date.min, e["caja"], e["orden"])

    revisar = [[_caja(e["caja"]), _f(e["fecha"]), e["ocurrio"], e["revisar"]] for e in sorted(
        (e for e in ents if e["bucket"] == "revisar"), key=clave)]
    vouchers = [e["voucher"] for e in sorted((e for e in ents if e.get("voucher")), key=clave)]
    correcciones = [[_caja(e["caja"]), _f(e["fecha"]), e["campo"], e["original"], e["normalizado"], e["motivo"]] for e in sorted(
        (e for e in ents if e["bucket"] == "correccion"), key=clave)]
    informativas = [[_caja(e["caja"]), _f(e["fecha"]), e["ocurrio"] + "\n" + ACCION_NINGUNA] for e in sorted(
        (e for e in ents if e["bucket"] == "info"), key=clave)]

    n = len(cierres)
    n_sin = sum(1 for v in estados.values() if v == SIN)
    n_req = sum(1 for v in estados.values() if v == REQUIERE)
    n_inf = sum(1 for v in estados.values() if v == INFORMATIVA)
    fechas = [c["fecha"] for c in cierres if c["fecha"]]
    periodo = ("%s al %s" % (_f(min(fechas)), _f(max(fechas))) if fechas and min(fechas) != max(fechas)
               else (_f(fechas[0]) if fechas else "no identificado"))
    cajas = cl._unir(sorted({"Caja " + _caja(c["caja"]) for c in cierres})) or "-"
    motor_ok = bool(motor and motor.get("disponible"))

    por_dia = {}
    for a in alquileres:
        k = (a["fecha"], a["caja"])
        por_dia[k] = por_dia.get(k, Decimal("0")) + a["importe"]
    total_alq = sum(por_dia.values(), Decimal("0"))

    if not motor:
        txt_motor = "La validación cruzada con el motor de CAJAS GABO no se ejecutó."
    elif not motor_ok:
        txt_motor = "La validación cruzada con el motor de CAJAS GABO no estuvo disponible en esta ejecución."
    else:
        con_obs = sum(1 for e in ents if e["bucket"] == "revisar" and e["orden"] >= 70)
        txt_motor = ("La validación cruzada con el motor de CAJAS GABO se realizó sobre la copia que se entrega de cada cierre "
                     "(%d %s con observaciones)." % (con_obs, _plural(con_obs, "cierre", "cierres")))

    return {
        "titulo": TITULO,
        "meta": [["Fecha y hora de ejecución", limpio(ejecucion.strftime("%d/%m/%Y %H:%M"))],
                 ["Periodo revisado", periodo],
                 ["Cajas revisadas", cajas],
                 ["Cantidad total de cierres", str(n)]],
        "resumen": [["Cierres revisados", str(n)], ["Sin observaciones", str(n_sin)], ["Requieren revisión", str(n_req)],
                    ["Informativos", str(n_inf)], ["Correcciones automáticas", str(len(correcciones))]],
        "leyenda": ["%s: el cierre no presenta hallazgos." % SIN,
                    "%s: una persona debe verificar algo antes de dar por concluido el control." % REQUIERE,
                    "%s: hechos que conviene conocer, incluidas las correcciones automáticas; no requieren acción." % INFORMATIVA],
        "cierres": filas_cierres,
        "revisar": revisar,
        "vouchers": vouchers,
        "vouchers_intro": INTRO_VOUCHERS,
        "vouchers_leyenda": list(LEYENDA_VOUCHERS),
        "correcciones": correcciones,
        "informativas": informativas,
        "conclusion": _conclusion(n, cajas, n_req, len(correcciones), motor_ok),
        "anexo": {
            "macros": [str(m) for m in (macros or [])],
            "macros_texto": "Archivos MACROS utilizados: " + (", ".join(str(m) for m in macros) if macros else "ninguno."),
            "motor": txt_motor,
            "alquileres": [[_f(k[0]), _caja(k[1]), cl.bs(v)] for k, v in sorted(por_dia.items(), key=lambda kv: (kv[0][0], kv[0][1]))],
            "total_alquileres": "Bs " + cl.bs(total_alq),
            "alquileres_texto": ("Alquileres del periodo: %d %s con alquileres, total Bs %s." %
                                 (len({k[0] for k in por_dia}), _plural(len({k[0] for k in por_dia}), "día", "días"), cl.bs(total_alq))
                                 if por_dia else "No se registraron alquileres en los cierres revisados."),
            "detalle": "El detalle técnico completo (resultados, hallazgos, cierres normalizados y recursos usados) está en el "
                       "archivo AUDITORIA_CIERRES_RESULTADOS.zip.",
        },
        "vacios": {"revisar": "No hay casos que requieran revisión.", "correcciones": "No se realizaron correcciones automáticas.",
                   "informativas": "No hay observaciones informativas."},
        "estados": {"SIN": SIN, "REQUIERE": REQUIERE, "INFORMATIVA": INFORMATIVA},
    }


def _sanear(o):
    if isinstance(o, str):
        return limpio(o)
    if isinstance(o, list):
        return [_sanear(x) for x in o]
    if isinstance(o, dict):
        return {k: _sanear(v) for k, v in o.items()}
    return o


def modelo_limpio(*a, **k):
    return _sanear(construir_modelo(*a, **k))


def textos_del_modelo(m):
    """Todas las cadenas visibles del informe, en orden (para comparar PDF y DOCX)."""
    out = [m["titulo"]]
    for k, v in m["meta"]:
        out += [k, v]
    for k, v in m["resumen"]:
        out += [k, v]
    out += m["leyenda"]
    for cols, filas in ((COLS_CIERRES, m["cierres"]), (COLS_REVISAR, m["revisar"]), (COLS_CORRECCIONES, m["correcciones"]),
                        (COLS_INFO, m["informativas"])):
        out += list(cols)
        for f in filas:
            out += [str(c) for c in f]
    if m["vouchers"]:
        out += [T_SEC2B, m["vouchers_intro"]] + list(m["vouchers_leyenda"])
        for b in m["vouchers"]:
            out += [b["titulo"], b["veredicto"], b["declarados"]]
            if b["filas"]:
                out += list(b["columnas"])
                for f in b["filas"]:
                    out += [str(c) for c in f]
            out += [b["interpretacion"], b["accion"]]
    out += [T_SEC1, T_SEC2, T_SEC3, T_SEC4, T_SEC5, T_ANEXO, m["conclusion"]]
    a = m["anexo"]
    out += [a["macros_texto"], a["motor"], a["alquileres_texto"], a["detalle"]]
    if a["alquileres"]:
        out += list(COLS_ALQUILERES)
        for f in a["alquileres"]:
            out += f
        out += ["Total del periodo", a["total_alquileres"]]
    return out


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------

def _p(txt):
    return escape(txt).replace("\n", "<br/>")


def escribir_pdf(m, ruta):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas
    from reportlab.platypus import CondPageBreak, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ancho = A4[0] - 4 * cm - 12          # el marco de SimpleDocTemplate descuenta 6 pt por lado
    NEGRO, GRIS, GRIS_CLARO = colors.HexColor("#1a1a1a"), colors.HexColor("#8c8c8c"), colors.HexColor("#ececec")
    TINTE = {SIN: colors.HexColor("#eef6ee"), REQUIERE: colors.HexColor("#fdf1dc"), INFORMATIVA: colors.HexColor("#e9f0f8")}

    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9, leading=12, textColor=NEGRO, alignment=TA_LEFT)
    est = {
        "titulo": ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=6),
        "h": ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=11.5, leading=14, spaceBefore=14, spaceAfter=5),
        "meta_k": ParagraphStyle("mk", parent=base, fontName="Helvetica-Bold", fontSize=9),
        "meta_v": base,
        "celda": ParagraphStyle("c", parent=base, fontSize=8.3, leading=10.6),
        "celda_b": ParagraphStyle("cb", parent=base, fontName="Helvetica-Bold", fontSize=8.3, leading=10.6),
        "cab": ParagraphStyle("cab", parent=base, fontName="Helvetica-Bold", fontSize=8.3, leading=10.6),
        "num": ParagraphStyle("n", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=20, alignment=1),
        "et": ParagraphStyle("e", parent=base, fontSize=7.6, leading=9.2, alignment=1),
        "peq": ParagraphStyle("p", parent=base, fontSize=8, leading=10.4, textColor=colors.HexColor("#333333")),
        "vacio": ParagraphStyle("v", parent=base, fontName="Helvetica-Oblique"),
    }

    def tabla(cols, filas, anchos, estado_col=None, negrita_col=None):
        data = [[Paragraph(_p(c), est["cab"]) for c in cols]]
        for f in filas:
            fila = []
            for i, c in enumerate(f):
                st = est["celda_b"] if i in (estado_col, negrita_col) else est["celda"]
                fila.append(Paragraph(_p(str(c)), st))
            data.append(fila)
        t = Table(data, colWidths=[a * ancho for a in anchos], repeatRows=1, hAlign="LEFT")
        ts = [("BACKGROUND", (0, 0), (-1, 0), GRIS_CLARO), ("LINEBELOW", (0, 0), (-1, 0), 0.9, NEGRO),
              ("GRID", (0, 0), (-1, -1), 0.4, GRIS), ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
              ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
        if estado_col is not None:
            for i, f in enumerate(filas, 1):
                ts.append(("BACKGROUND", (estado_col, i), (estado_col, i), TINTE.get(f[estado_col], colors.white)))
        t.setStyle(TableStyle(ts))
        return t

    class Numerado(canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for st in self._paginas:
                self.__dict__.update(st)
                self.setFont("Helvetica", 7.5)
                self.setFillColor(colors.HexColor("#555555"))
                self.setStrokeColor(GRIS)
                self.line(2 * cm, 1.55 * cm, A4[0] - 2 * cm, 1.55 * cm)
                self.drawString(2 * cm, 1.15 * cm, PIE)
                self.drawRightString(A4[0] - 2 * cm, 1.15 * cm, "Página %d de %d" % (self._pageNumber, total))
                super().showPage()
            super().save()

    doc = SimpleDocTemplate(ruta, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.9 * cm, bottomMargin=2.1 * cm,
                            title=limpio(m["titulo"]), author="CAJAS GABO", subject="Auditoría de cierres de caja")
    s = [Paragraph(_p(m["titulo"]), est["titulo"])]

    meta = Table([[Paragraph(_p(k), est["meta_k"]), Paragraph(_p(v), est["meta_v"])] for k, v in m["meta"]],
                 colWidths=[0.30 * ancho, 0.70 * ancho], hAlign="LEFT")
    meta.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("LINEBELOW", (0, -1), (-1, -1), 0.6, GRIS)]))
    s += [meta, Paragraph(_p(T_SEC1), est["h"])]      # (el encabezado 1 esta siempre en la primera pagina)

    cards = Table([[Paragraph(_p(v), est["num"]) for _k, v in m["resumen"]],
                   [Paragraph(_p(k), est["et"]) for k, _v in m["resumen"]]], colWidths=[ancho / 5.0] * 5, hAlign="LEFT")
    cards.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, NEGRO), ("INNERGRID", (0, 0), (-1, -1), 0.4, GRIS),
                               ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f6f6f6")),
                               ("BOX", (2, 0), (2, -1), 1.8, NEGRO),                       # "Requieren revisión" destaca sin color
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, 0), 7),
                               ("BOTTOMPADDING", (0, 0), (-1, 0), 1), ("TOPPADDING", (0, 1), (-1, 1), 1),
                               ("BOTTOMPADDING", (0, 1), (-1, 1), 7)]))
    s += [cards, Spacer(1, 6)]
    for linea in m["leyenda"]:
        s.append(Paragraph(_p(linea), est["peq"]))
    s.append(Spacer(1, 8))
    s.append(tabla(COLS_CIERRES, m["cierres"], [0.14, 0.15, 0.24, 0.47], estado_col=2))

    def seccion(titulo, cols, filas, anchos, vacio, **kw):
        s.append(CondPageBreak(110))
        s.append(Paragraph(_p(titulo), est["h"]))
        s.append(tabla(cols, filas, anchos, **kw) if filas else Paragraph(_p(vacio), est["vacio"]))

    seccion(T_SEC2, COLS_REVISAR, m["revisar"], [0.12, 0.12, 0.42, 0.34], m["vacios"]["revisar"])
    if m["vouchers"]:
        s.append(CondPageBreak(120))
        s.append(Paragraph(_p(T_SEC2B), est["h"]))
        s.append(Paragraph(_p(m["vouchers_intro"]), est["peq"]))
        for linea in m["vouchers_leyenda"]:
            s.append(Paragraph(_p(linea), est["peq"]))
        for b in m["vouchers"]:
            s.append(Spacer(1, 6))
            cabeza = [Paragraph(_p(b["titulo"]), est["celda_b"]), Paragraph(_p(b["veredicto"]), est["celda_b"]),
                      Paragraph(_p(b["declarados"]), est["celda"]), Spacer(1, 3)]
            if b["filas"]:
                s.append(KeepTogether(cabeza + [tabla(b["columnas"], b["filas"], ANCHOS_VOUCHER)]))
            else:
                s.append(KeepTogether(cabeza))
            s.append(Spacer(1, 3))
            s.append(Paragraph(_p(b["interpretacion"]), est["celda"]))
            s.append(Paragraph(_p(b["accion"]), est["celda_b"]))
    seccion(T_SEC3, COLS_CORRECCIONES, m["correcciones"], [0.10, 0.13, 0.15, 0.13, 0.14, 0.35], m["vacios"]["correcciones"])
    seccion(T_SEC4, COLS_INFO, m["informativas"], [0.13, 0.13, 0.74], m["vacios"]["informativas"])
    s.append(CondPageBreak(90))
    s.append(Paragraph(_p(T_SEC5), est["h"]))
    s.append(Paragraph(_p(m["conclusion"]), est["meta_v"]))

    a = m["anexo"]
    s.append(CondPageBreak(90))
    s.append(Paragraph(_p(T_ANEXO), est["h"]))
    s.append(Paragraph(_p(a["macros_texto"]), est["peq"]))
    s.append(Spacer(1, 2))
    s.append(Paragraph(_p(a["motor"]), est["peq"]))
    s.append(Spacer(1, 2))
    s.append(Paragraph(_p(a["alquileres_texto"]), est["peq"]))
    s.append(Spacer(1, 2))
    s.append(Paragraph(_p(a["detalle"]), est["peq"]))
    if a["alquileres"]:
        s.append(Spacer(1, 5))
        filas = a["alquileres"] + [["Total del periodo", "", a["total_alquileres"]]]
        s.append(KeepTogether([tabla(COLS_ALQUILERES, filas, [0.30, 0.35, 0.35])]))
    doc.build(s, canvasmaker=Numerado)
    return ruta


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

def escribir_docx(m, ruta):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    TINTE = {SIN: "EEF6EE", REQUIERE: "FDF1DC", INFORMATIVA: "E9F0F8"}
    ancho_cm = 17.0

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(2.1)
    doc.core_properties.title = m["titulo"]
    doc.core_properties.author = "CAJAS GABO"

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(9.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(3)

    def sombrear(celda, hexcolor):
        tcPr = celda._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def texto(celda, txt, negrita=False, tam=8.5, alinear=None):
        celda.text = ""
        partes = txt.split("\n")
        for i, parte in enumerate(partes):
            p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            if alinear is not None:
                p.alignment = alinear
            r = p.add_run(parte)
            r.bold = negrita
            r.font.size = Pt(tam)

    def fijar(tabla, anchos, repetir_cabecera):
        tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
        tabla.autofit = False
        mar = OxmlElement("w:tblCellMar")                      # aire vertical/horizontal dentro de las celdas
        for lado, tw in (("top", 45), ("left", 90), ("bottom", 45), ("right", 90)):
            e = OxmlElement("w:%s" % lado)
            e.set(qn("w:w"), str(tw))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tabla._tbl.tblPr.append(mar)
        for col, a in zip(tabla.columns, anchos):
            col.width = Cm(ancho_cm * a)
        for fila in tabla.rows:
            trPr = fila._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            for celda, a in zip(fila.cells, anchos):
                celda.width = Cm(ancho_cm * a)
        if repetir_cabecera:
            th = OxmlElement("w:tblHeader")
            tabla.rows[0]._tr.get_or_add_trPr().append(th)

    def tabla(cols, filas, anchos, estado_col=None, negrita_col=None):
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        for i, c in enumerate(cols):
            texto(t.rows[0].cells[i], c, negrita=True)
            sombrear(t.rows[0].cells[i], "ECECEC")
        for f in filas:
            celdas = t.add_row().cells
            for i, c in enumerate(f):
                texto(celdas[i], str(c), negrita=(i in (estado_col, negrita_col)))
                if estado_col is not None and i == estado_col:
                    sombrear(celdas[i], TINTE.get(f[i], "FFFFFF"))
        fijar(t, anchos, True)
        return t

    def titulo(txt, nivel=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12 if nivel else 0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(txt)
        r.bold = True
        r.font.size = Pt(15 if nivel == 0 else 11.5)
        return p

    def parrafo(txt, tam=9.5, cursiva=False):
        p = doc.add_paragraph()
        r = p.add_run(txt)
        r.italic = cursiva
        r.font.size = Pt(tam)
        return p

    titulo(m["titulo"], 0)
    meta = doc.add_table(rows=0, cols=2)
    for k, v in m["meta"]:
        c = meta.add_row().cells
        texto(c[0], k, negrita=True, tam=9.5)
        texto(c[1], v, tam=9.5)
    fijar(meta, [0.30, 0.70], False)
    titulo(T_SEC1)

    cards = doc.add_table(rows=2, cols=5)
    cards.style = "Table Grid"
    for i, (k, v) in enumerate(m["resumen"]):
        texto(cards.rows[0].cells[i], v, negrita=True, tam=17, alinear=1)
        texto(cards.rows[1].cells[i], k, tam=8, alinear=1)
        for f in (0, 1):
            sombrear(cards.rows[f].cells[i], "F6F6F6")
    fijar(cards, [0.2] * 5, False)
    doc.add_paragraph()
    for linea in m["leyenda"]:
        parrafo(linea, tam=8.5)
    tabla(COLS_CIERRES, m["cierres"], [0.14, 0.15, 0.24, 0.47], estado_col=2)

    def seccion(t, cols, filas, anchos, vacio):
        titulo(t)
        if filas:
            tabla(cols, filas, anchos)
        else:
            parrafo(vacio, cursiva=True)

    seccion(T_SEC2, COLS_REVISAR, m["revisar"], [0.12, 0.12, 0.42, 0.34], m["vacios"]["revisar"])
    if m["vouchers"]:
        titulo(T_SEC2B)
        parrafo(m["vouchers_intro"], tam=8.5)
        for linea in m["vouchers_leyenda"]:
            parrafo(linea, tam=8.5)
        for b in m["vouchers"]:
            p = parrafo(b["titulo"], tam=9)
            p.runs[0].bold = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p = parrafo(b["veredicto"], tam=9)
            p.runs[0].bold = True
            p.paragraph_format.keep_with_next = True
            p = parrafo(b["declarados"], tam=8.5)
            p.paragraph_format.keep_with_next = True
            if b["filas"]:
                tabla(b["columnas"], b["filas"], ANCHOS_VOUCHER)
            parrafo(b["interpretacion"], tam=8.5)
            p = parrafo(b["accion"], tam=8.5)
            p.runs[0].bold = True
    seccion(T_SEC3, COLS_CORRECCIONES, m["correcciones"], [0.10, 0.13, 0.15, 0.13, 0.14, 0.35], m["vacios"]["correcciones"])
    seccion(T_SEC4, COLS_INFO, m["informativas"], [0.13, 0.13, 0.74], m["vacios"]["informativas"])
    titulo(T_SEC5)
    parrafo(m["conclusion"])
    a = m["anexo"]
    titulo(T_ANEXO)
    parrafo(a["macros_texto"], tam=8.5)
    parrafo(a["motor"], tam=8.5)
    parrafo(a["alquileres_texto"], tam=8.5)
    parrafo(a["detalle"], tam=8.5)
    if a["alquileres"]:
        tabla(COLS_ALQUILERES, a["alquileres"] + [["Total del periodo", "", a["total_alquileres"]]], [0.30, 0.35, 0.35])

    # pie de pagina: texto + "Pagina X de Y" (campos de Word)
    pie = sec.footer.paragraphs[0]
    pie.text = ""
    r = pie.add_run(PIE + "\t\tPágina ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    def campo(instr):
        run = pie.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        for tipo, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
            if tipo:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), tipo)
                run._r.append(fc)
            elif txt == instr:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = " %s " % instr
                run._r.append(it)
            else:
                t = OxmlElement("w:t")
                t.text = txt
                run._r.append(t)

    campo("PAGE")
    r2 = pie.add_run(" de ")
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    campo("NUMPAGES")
    doc.save(ruta)
    return ruta
