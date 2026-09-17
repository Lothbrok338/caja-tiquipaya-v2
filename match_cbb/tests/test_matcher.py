"""Tests del motor de matching Cochabamba vs BCP."""

import datetime as dt
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import matcher as M

CFG = M.Config()


# --------------------------------------------------------------------------- #
# Fabricas de datos
# --------------------------------------------------------------------------- #

def reg(fila, fecha_hora, monto, es_qr=True, estudiante="ALUMNO PRUEBA", factura="1"):
    return M.RegistroCBB(
        fila_excel=fila,
        crudo={"Nro": fila, "Numero Factura": factura, "Nombre Estudiante": estudiante,
               "Monto": monto, "Fecha": fecha_hora, "Tipo Pago": "QR" if es_qr else "Tarjeta D/C"},
        fecha_hora=fecha_hora,
        centavos=M.a_centavos(monto),
        es_qr=es_qr,
    )


def mov(fila, fecha_hora, importe, oper, glosa="QR DE PRUEBA", cd=None):
    return M.MovimientoBCP(
        fila_excel=fila,
        crudo={},
        fecha_hora=fecha_hora,
        centavos=M.a_centavos(importe),
        nro_oper=oper,
        nro_oper_norm=M.norm_id(oper),
        glosa=glosa,
        cd_confirmacion=cd,
        es_qr="QR" in glosa.upper(),
    )


def t(h, m, s=0, dia=1):
    return dt.datetime(2026, 9, dia, h, m, s)


# --------------------------------------------------------------------------- #
# 1-2. Coincidencia exacta y de 1 segundo
# --------------------------------------------------------------------------- #

def test_coincidencia_exacta():
    res = M.emparejar([reg(1, t(10, 49, 11), 2630)], [mov(2, t(10, 49, 11), 2630, 301902)], CFG)
    assert len(res) == 1
    assert res[0].estado == M.ESTADO_SEGURO
    assert res[0].diferencia_seg == 0
    assert res[0].movimiento.nro_oper == 301902


def test_diferencia_de_un_segundo():
    res = M.emparejar([reg(1, t(10, 49, 11), 2630)], [mov(2, t(10, 49, 10), 2630, 301902)], CFG)
    assert res[0].estado == M.ESTADO_SEGURO
    assert res[0].diferencia_seg == 1.0


# --------------------------------------------------------------------------- #
# 3-4. Mismo monto repetido el mismo dia, asignado 1 a 1 por hora
# --------------------------------------------------------------------------- #

def test_mismo_monto_repetido_elige_el_mas_cercano():
    registros = [reg(1, t(10, 0, 0), 500)]
    movimientos = [mov(2, t(10, 0, 40), 500, "A"), mov(3, t(10, 0, 1), 500, "B")]
    res = M.emparejar(registros, movimientos, CFG)
    assert res[0].movimiento.nro_oper == "B"
    assert res[0].diferencia_seg == 1.0
    assert res[0].candidatos == 2


def test_asignacion_uno_a_uno_por_hora():
    registros = [reg(1, t(10, 0, 0), 500), reg(2, t(10, 5, 0), 500)]
    movimientos = [mov(10, t(10, 5, 1), 500, "TARDE"), mov(11, t(10, 0, 1), 500, "TEMPRANO")]
    res = M.emparejar(registros, movimientos, CFG)
    asignado = {r.registro.fila_excel: r.movimiento.nro_oper for r in res}
    assert asignado == {1: "TEMPRANO", 2: "TARDE"}
    assert all(r.diferencia_seg == 1.0 for r in res)


def test_asignacion_global_minima_no_solo_el_primero():
    # Si se tomara "el primero que aparece", el registro 1 robaria el movimiento
    # del registro 2 y este quedaria mal emparejado.
    registros = [reg(1, t(10, 0, 0), 800), reg(2, t(10, 0, 4), 800)]
    movimientos = [mov(10, t(10, 0, 3), 800, "X"), mov(11, t(10, 0, 0), 800, "Y")]
    res = M.emparejar(registros, movimientos, CFG)
    asignado = {r.registro.fila_excel: r.movimiento.nro_oper for r in res}
    assert asignado == {1: "Y", 2: "X"}


# --------------------------------------------------------------------------- #
# 5. Un movimiento BCP no puede usarse dos veces
# --------------------------------------------------------------------------- #

def test_bcp_nunca_se_reutiliza():
    registros = [reg(1, t(10, 0, 0), 500), reg(2, t(10, 0, 2), 500)]
    movimientos = [mov(10, t(10, 0, 1), 500, 777)]
    res = M.emparejar(registros, movimientos, CFG)
    usados = [r.movimiento.nro_oper for r in res if r.movimiento]
    assert usados == [777]
    sin_match = [r for r in res if r.movimiento is None]
    assert len(sin_match) == 1
    assert sin_match[0].estado == M.ESTADO_SIN_MATCH
    assert "ya fue asignado a otro pago" in sin_match[0].motivo


# --------------------------------------------------------------------------- #
# 6-7. Sin candidato y candidato fuera de la ventana
# --------------------------------------------------------------------------- #

def test_sin_candidato():
    res = M.emparejar([reg(1, t(10, 0, 0), 500)], [mov(10, t(10, 0, 0), 999, 1)], CFG)
    assert res[0].estado == M.ESTADO_SIN_MATCH
    assert res[0].movimiento is None
    assert res[0].candidatos == 0


def test_candidato_fuera_de_ventana_no_se_asigna():
    res = M.emparejar([reg(1, t(10, 0, 0), 500)], [mov(10, t(10, 3, 0), 500, 1)], CFG)
    assert res[0].estado == M.ESTADO_SIN_MATCH
    assert res[0].movimiento is None
    assert "180 s" in res[0].motivo  # informa el mas cercano sin asignarlo


def test_escalera_de_umbrales():
    casos = {3: M.ESTADO_SEGURO, 20: M.ESTADO_PROBABLE, 100: M.ESTADO_REVISAR, 200: M.ESTADO_SIN_MATCH}
    for segundos, esperado in casos.items():
        res = M.emparejar(
            [reg(1, t(10, 0, 0), 500)],
            [mov(10, t(10, 0, 0) + dt.timedelta(seconds=segundos), 500, 1)],
            CFG,
        )
        assert res[0].estado == esperado, f"{segundos} s deberia ser {esperado}"


# --------------------------------------------------------------------------- #
# 8. Cd. Confirmacion de otra ciudad no invalida el match
# --------------------------------------------------------------------------- #

def test_cd_confirmacion_sucre_no_impide_el_match():
    res = M.emparejar(
        [reg(1, t(11, 38, 56), 660)],
        [mov(10, t(11, 38, 55), 660, 364703, cd="SUCRE")],
        CFG,
    )
    assert res[0].estado == M.ESTADO_SEGURO
    assert res[0].movimiento.nro_oper == 364703
    assert M._conflicto_ciudad(res[0], CFG) is True
    # El valor original se muestra tal cual, nunca se reemplaza.
    fila = M.a_filas(res, CFG)[0]
    assert fila["BCP_Cd_Confirmacion"] == "SUCRE"
    assert fila["Conflicto_Ciudad"] == "SI"
    assert fila["Estado_match"] == M.ESTADO_SEGURO


# --------------------------------------------------------------------------- #
# 9. Nro Oper. duplicado en BCP genera alerta
# --------------------------------------------------------------------------- #

def test_nro_oper_duplicado_genera_alerta():
    movimientos = [mov(10, t(10, 0, 0), 500, 301902), mov(11, t(12, 0, 0), 900, "301902")]
    alertas = M.validar([], movimientos, [], CFG)
    detalles = [d for _, d, _ in alertas]
    assert any("Nro Oper. duplicados" in d for d in detalles)
    assert any("301902" in d for d in detalles)


def test_nro_oper_normalizado():
    assert M.norm_id(301902) == "301902"
    assert M.norm_id("301902") == "301902"
    assert M.norm_id(" 301902 ") == "301902"
    assert M.norm_id(301902.0) == "301902"
    assert M.norm_id("301902.0") == "301902"


# --------------------------------------------------------------------------- #
# 10. Una celda "Fecha" suelta arriba del reporte no confunde al detector
# --------------------------------------------------------------------------- #

def _crear_cbb(ruta, con_membrete=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RPagosEnLinea"
    if con_membrete:
        ws.cell(row=3, column=24, value="Fecha:")
        ws.cell(row=3, column=32, value="16/09/2026 15:48")
        ws.cell(row=6, column=11, value="INGRESO DIARIO")
        ws.cell(row=10, column=7, value="SEDE COCHABAMBA")
    fila = 13
    encabezados = ["Nro", "Fecha", "Número Factura", "Nit/C.I.", "Razon Social",
                   "Nombre Estudiante", "Tipo Pago", "Monto", "Canal de Pago", "Estado"]
    for c, nombre in enumerate(encabezados, start=1):
        ws.cell(row=fila, column=c, value=nombre)
    datos = [
        (1, dt.datetime(2026, 9, 1, 10, 49, 11), "16180", "7863318", "GARRON",
         "GARRON CESPEDES AYRIN", "QR", 2630, "Pago En Línea", "Válido"),
        (2, dt.datetime(2026, 9, 1, 11, 0, 0), "16181", "123", "X",
         "OTRO ALUMNO", "Tarjeta D/C", 151, "Pago En Línea", "Válido"),
    ]
    for i, registro in enumerate(datos, start=fila + 1):
        for c, valor in enumerate(registro, start=1):
            ws.cell(row=i, column=c, value=valor)
    ws.cell(row=fila + 4, column=2, value="Monto QR:")
    ws.cell(row=fila + 4, column=6, value="2,630.00")
    wb.save(ruta)


def _crear_bcp(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    encabezados = ["Fecha", "Hora", "Glosa", "Tipo", "Suc. Age.", "Usuario",
                   "Importe", "Saldo", "Nro Oper. cod. Bancarizacion", "Cd, Confirmacion"]
    for c, nombre in enumerate(encabezados, start=1):
        ws.cell(row=1, column=c, value=nombre)
    datos = [
        (dt.datetime(2026, 9, 1), dt.time(10, 49, 10), "QR DE Garron Otalora", 2401,
         201204, "AK0", 2630, 5803394.75, 301902, None),
        (dt.datetime(2026, 9, 1), dt.time(11, 3, 25), "CHEQUE 00015925", 3001,
         301328, "HD1", -2532.6, 5814449.15, 70, None),
        (dt.datetime(2026, 9, 1), dt.time(11, 38, 55), "QR DE DONAIRE JONATH", 2401,
         201204, "AK0", 660, 5815711.15, 364703, "SUCRE"),
    ]
    for i, registro in enumerate(datos, start=2):
        for c, valor in enumerate(registro, start=1):
            ws.cell(row=i, column=c, value=valor)
    wb.save(ruta)


def test_no_confunde_la_celda_fecha_del_membrete(tmp_path):
    ruta = tmp_path / "cbb.xlsx"
    _crear_cbb(ruta)
    hojas = M.leer_libro(str(ruta))
    fila = M.detectar_fila_encabezado(hojas[0][1], M.REQUERIDOS_CBB)
    assert fila == 12  # 0-based: fila 13 de Excel, no la fila 3 del membrete


def test_encabezado_requiere_todos_los_campos():
    filas = [["Fecha:", "16/09/2026 15:48"], ["Fecha", "Monto"]]
    assert M.detectar_fila_encabezado(filas, M.REQUERIDOS_CBB) is None


# --------------------------------------------------------------------------- #
# Identificacion de archivos y flujo completo
# --------------------------------------------------------------------------- #

def test_identifica_los_archivos_en_cualquier_orden(tmp_path):
    ruta_cbb, ruta_bcp = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _crear_cbb(ruta_cbb)
    _crear_bcp(ruta_bcp)
    for orden in ([str(ruta_cbb), str(ruta_bcp)], [str(ruta_bcp), str(ruta_cbb)]):
        cbb, bcp, previo = M.identificar_tablas(orden)
        assert cbb.ruta == str(ruta_cbb) and cbb.hoja == "RPagosEnLinea"
        assert bcp.ruta == str(ruta_bcp) and bcp.hoja == "Hoja1"
        assert previo is None


def test_error_claro_si_falta_un_archivo(tmp_path):
    ruta_bcp = tmp_path / "b.xlsx"
    _crear_bcp(ruta_bcp)
    with pytest.raises(ValueError, match="No pude identificar ambas tablas"):
        M.identificar_tablas([str(ruta_bcp)])


def test_flujo_completo_no_modifica_los_originales(tmp_path):
    ruta_cbb, ruta_bcp = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _crear_cbb(ruta_cbb)
    _crear_bcp(ruta_bcp)
    antes = {r: (os.path.getsize(r), open(r, "rb").read()) for r in (ruta_cbb, ruta_bcp)}

    salida = tmp_path / "MATCH_CBB.xlsx"
    datos = M.ejecutar([str(ruta_cbb), str(ruta_bcp)], str(salida), CFG)

    assert {r: (os.path.getsize(r), open(r, "rb").read()) for r in (ruta_cbb, ruta_bcp)} == antes
    assert salida.exists()
    wb = openpyxl.load_workbook(salida)
    import reporte

    visibles = [n for n in wb.sheetnames if wb[n].sheet_state == "visible"]
    assert visibles == ["RESUMEN_VISUAL", "PARA_PEGAR_CBB", "REVISAR_MANUAL", "SIN_MATCH",
                        "TARJETA", "INGRESOS_NORMALIZADOS"]
    ocultas = [n for n in wb.sheetnames if wb[n].sheet_state == "hidden"]
    assert ocultas == [reporte.HOJA_CANDIDATOS, reporte.HOJA_MANUALES]
    assert len(datos["registros"]) == 2  # el pie de totales no se cuenta como registro
    estados = {r.registro.crudo["Numero Factura"]: r.estado for r in datos["resultados"]}
    assert estados == {"16180": M.ESTADO_SEGURO, "16181": M.ESTADO_NO_QR}


def test_no_fuerza_tarjeta_contra_qr():
    registros = [reg(1, t(10, 0, 0), 500, es_qr=False)]
    movimientos = [mov(10, t(10, 0, 0), 500, 1)]
    res = M.emparejar(registros, movimientos, CFG)
    assert res[0].estado == M.ESTADO_NO_QR
    assert res[0].movimiento is None


def test_bcp_negativo_nunca_se_asigna():
    res = M.emparejar([reg(1, t(10, 0, 0), 500)], [mov(10, t(10, 0, 0), -500, 1, glosa="QR DEVOLUCION")], CFG)
    assert res[0].estado == M.ESTADO_SIN_MATCH


def test_movimiento_sin_qr_en_glosa_se_excluye():
    movimientos = [mov(10, t(10, 0, 0), 500, 1, glosa="CHEQUE 00015925")]
    assert M.emparejar([reg(1, t(10, 0, 0), 500)], movimientos, CFG)[0].movimiento is None
    sin_filtro = M.Config(solo_qr_bcp=False)
    assert M.emparejar([reg(1, t(10, 0, 0), 500)], movimientos, sin_filtro)[0].movimiento is not None


# --------------------------------------------------------------------------- #
# Ambiguedad
# --------------------------------------------------------------------------- #

def test_candidato_libre_cercano_degrada_a_revisar():
    registros = [reg(1, t(10, 10, 0), 500)]
    movimientos = [mov(10, t(10, 10, 2), 500, "A"), mov(11, t(10, 10, 3), 500, "B")]
    res = M.emparejar(registros, movimientos, CFG)
    assert res[0].movimiento.nro_oper == "A"
    assert res[0].diferencia_seg == 2.0
    assert res[0].segundo_delta == 3.0
    assert res[0].margen == 1.0
    assert res[0].ambiguo is True
    assert res[0].estado == M.ESTADO_REVISAR


def test_candidato_alternativo_ocupado_no_genera_ambiguedad():
    registros = [reg(1, t(10, 10, 0), 500), reg(2, t(10, 10, 5), 500)]
    movimientos = [mov(10, t(10, 10, 1), 500, "A"), mov(11, t(10, 10, 4), 500, "B")]
    res = M.emparejar(registros, movimientos, CFG)
    assert [r.estado for r in res] == [M.ESTADO_SEGURO, M.ESTADO_SEGURO]
    assert [r.ambiguo for r in res] == [False, False]


# --------------------------------------------------------------------------- #
# Conversiones
# --------------------------------------------------------------------------- #

def test_redondeo_de_segundo_elimina_el_desfase_del_serial():
    # 46266.450822997685 -> 10:49:11.107 en el reporte real de Cochabamba.
    crudo = M.a_datetime(46266.450822997685)
    assert crudo.microsecond != 0
    assert M.redondear_a_segundo(crudo) == dt.datetime(2026, 9, 1, 10, 49, 11)


def test_conversion_de_montos():
    assert M.a_centavos(2630) == 263000
    assert M.a_centavos(2630.5) == 263050
    assert M.a_centavos("2,630.50") == 263050
    assert M.a_centavos("633,334.97") == 63333497
    assert M.a_centavos("") is None
    assert M.a_centavos("abc") is None


def test_conversion_de_horas():
    assert M.a_time(dt.time(10, 49, 10)) == dt.time(10, 49, 10)
    assert M.a_time(dt.datetime(2026, 9, 1, 10, 49, 10)) == dt.time(10, 49, 10)
    assert M.a_time("10:49:10") == dt.time(10, 49, 10)
    assert M.a_time(0.45082) == dt.time(10, 49, 11)
    assert M.a_time(None) is None


# --------------------------------------------------------------------------- #
# Presentacion (reporte.py): no cambia ningun resultado, solo como se muestra
# --------------------------------------------------------------------------- #

def test_hojas_operativas_traen_los_casos_correctos(tmp_path):
    import reporte

    ruta_cbb, ruta_bcp = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _crear_cbb(ruta_cbb)
    _crear_bcp(ruta_bcp)
    salida = tmp_path / "MATCH_CBB.xlsx"
    M.ejecutar([str(ruta_cbb), str(ruta_bcp)], str(salida), CFG)

    wb = openpyxl.load_workbook(salida)
    assert [c.value for c in wb["PARA_PEGAR_CBB"][1]][:3] == ["Estado", "Origen del match", "Factura"]
    filas = list(wb["PARA_PEGAR_CBB"].iter_rows(min_row=2, values_only=True))
    assert [f[0] for f in filas] == [M.ESTADO_SEGURO]  # sin casos REVISAR no hay ranuras
    assert filas[0][1] == reporte.ORIGEN_AUTOMATICO
    ingresos = wb["INGRESOS_NORMALIZADOS"]
    assert [c.value for c in ingresos[1]] == [
        "Nro", "Fecha", "Número Factura", "Nit/C.I.", "Razon Social",
        "Nombre Estudiante", "Tipo Pago", "Monto", "Canal de Pago", "Estado", "Nro Oper BCP",
    ]
    assert ingresos.max_row == 3  # los dos registros del reporte, sin pie de totales
    tarjeta = list(wb["TARJETA"].iter_rows(min_row=2, values_only=True))
    assert len(tarjeta) == 1 and tarjeta[0][-1] == "Pago con tarjeta - fuera del cruce QR"
    for hoja in wb.sheetnames:
        if wb[hoja].sheet_state == "visible":
            assert "NO_QR" not in str(list(wb[hoja].iter_rows(values_only=True)))


def test_candidatos_para_revision_incluyen_todo_el_dia_y_monto():
    import reporte

    # El motor solo mira glosas QR dentro de la ventana; la revision humana ve mas.
    res = M.emparejar(
        [reg(1, t(10, 0, 0), 500)],
        [mov(10, t(10, 0, 50), 500, "A"), mov(11, t(18, 0, 0), 500, "B"),
         mov(12, t(10, 0, 10), 900, "OTRO_MONTO")],
        CFG,
    )[0]
    assert res.estado == M.ESTADO_REVISAR
    candidatos = reporte.candidatos_mismo_dia_monto(res, [
        mov(10, t(10, 0, 50), 500, "A"), mov(11, t(18, 0, 0), 500, "B"),
        mov(12, t(10, 0, 10), 900, "OTRO_MONTO"),
    ])
    assert [m.nro_oper for m, _ in candidatos] == ["A", "B"]  # ordenados por cercania
    assert [round(d) for _, d in candidatos] == [50, 28800]


def test_etiqueta_del_desplegable_distingue_candidatos():
    import reporte

    etiqueta = reporte._etiqueta_candidato(
        mov(10, t(19, 1, 18), 2596, 940311, glosa="QR DE CARRERA RECALD"), 42.0
    )
    assert etiqueta == "940311 | 19:01:18 | Bs 2.596,00 | QR DE CARRERA RECALD | Δ 42 s"


def test_zona_de_decision_tiene_desplegable_y_formulas(tmp_path):
    import reporte

    registros = [reg(1, t(10, 0, 0), 500)]
    movimientos = [mov(10, t(10, 0, 50), 500, "A"), mov(11, t(10, 1, 30), 500, "B")]
    resultados = M.emparejar(registros, movimientos, CFG)
    assert resultados[0].estado == M.ESTADO_REVISAR

    salida = tmp_path / "MATCH_CBB.xlsx"
    reporte.escribir(resultados, registros, movimientos, [], CFG, str(salida))

    wb = openpyxl.load_workbook(salida)
    ws = wb["REVISAR_MANUAL"]
    assert [c.value for c in ws[1]][9:] == [
        "Candidato elegido", "Nro Oper elegido", "Fecha/hora BCP elegida", "Importe BCP elegido",
        "Glosa BCP elegida", "Cd. Confirmación actual", "Dif. seg elegido", "Decisión", "Control",
    ]
    # J preseleccionado con el candidato que eligio el algoritmo.
    assert ws["J2"].value.startswith("A | ")
    assert ws["E2"].value == ws["J2"].value
    assert ws["Q2"].value == "PENDIENTE"
    # K:P se calculan a partir de J.
    for celda in ("K2", "L2", "M2", "N2", "O2", "P2"):
        assert ws[celda].value.startswith("=IFERROR(IF(ISBLANK(INDEX(")
        assert "$J2" in ws[celda].value

    listas = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    rangos = {str(dv.sqref): dv.formula1 for dv in listas}
    assert rangos["J2"].startswith(f"'{reporte.HOJA_CANDIDATOS}'!$A$")
    assert rangos["Q2"] == '"CONFIRMAR MATCH,DESCARTAR,PENDIENTE"'

    aux = wb[reporte.HOJA_CANDIDATOS]
    assert aux.sheet_state == "hidden"
    assert [f[0] for f in aux.iter_rows(min_row=2, values_only=True)] == [
        ws["J2"].value, "B | 10:01:30 | Bs 500,00 | QR DE PRUEBA | Δ 90 s",
    ]


def test_ranuras_manuales_dependen_de_la_decision(tmp_path):
    import reporte

    registros = [reg(1, t(10, 0, 0), 500)]
    movimientos = [mov(10, t(10, 0, 50), 500, "A"), mov(11, t(10, 1, 30), 500, "B")]
    resultados = M.emparejar(registros, movimientos, CFG)
    salida = tmp_path / "MATCH_CBB.xlsx"
    reporte.escribir(resultados, registros, movimientos, [], CFG, str(salida))

    wb = openpyxl.load_workbook(salida)
    pegar, manuales = wb["PARA_PEGAR_CBB"], wb[reporte.HOJA_MANUALES]

    # Un caso REVISAR: una ranura, toda por formula y sin valores fijos.
    assert pegar.max_row == 2
    assert pegar["A2"].value.startswith("=IFERROR(INDEX('_MANUALES'!$F$2:$F$2,MATCH(1,")
    # La ranura solo se llena si la decision es CONFIRMAR MATCH y no hay duplicado.
    assert manuales["A2"].value == "='REVISAR_MANUAL'!Q2"
    assert manuales["D2"].value == '=IF(AND($A2="CONFIRMAR MATCH",$B2<>"",$C2=0),1,0)'
    assert "COUNTIF(PARA_PEGAR_CBB!$K$2:$K$1,$B2)" in manuales["C2"].value
    assert manuales["F2"].value == reporte.ESTADO_MANUAL
    assert manuales["G2"].value == reporte.ORIGEN_MANUAL
    # El control de duplicados avisa en REVISAR_MANUAL.
    assert reporte.ALERTA_DUPLICADO in wb["REVISAR_MANUAL"]["R2"].value


def test_separa_el_sin_match_por_corte_del_extracto():
    import reporte

    inicio, fin = t(8, 0, 0), t(22, 0, 0)
    despues = M.Resultado(reg(1, t(22, 30, 0), 500), None, None, M.ESTADO_SIN_MATCH, 0, None, None, False, "")
    dentro = M.Resultado(reg(2, t(12, 0, 0), 500), None, None, M.ESTADO_SIN_MATCH, 0, None, None, False, "")
    assert reporte._tipo_sin_match(despues, inicio, fin) == reporte.SIN_MATCH_CORTE
    assert reporte._tipo_sin_match(dentro, inicio, fin) == reporte.SIN_MATCH_REAL


def test_observacion_humana_menciona_cd_confirmacion_existente():
    import reporte

    con_cd = M.emparejar([reg(1, t(11, 38, 56), 660)], [mov(10, t(11, 38, 55), 660, 1, cd="SUCRE")], CFG)[0]
    sin_cd = M.emparejar([reg(1, t(11, 38, 56), 660)], [mov(10, t(11, 38, 55), 660, 1)], CFG)[0]
    assert reporte._observacion_pegar(sin_cd) == "Listo para usar"
    assert "SUCRE" in reporte._observacion_pegar(con_cd)


def test_delta_con_signo_distingue_el_orden_de_los_hechos():
    import reporte

    despues = M.emparejar([reg(1, t(10, 0, 5), 500)], [mov(10, t(10, 0, 0), 500, 1)], CFG)[0]
    antes = M.emparejar([reg(1, t(10, 0, 0), 500)], [mov(10, t(10, 0, 5), 500, 1)], CFG)[0]
    assert reporte._delta_con_signo(despues) == 5.0
    assert reporte._delta_con_signo(antes) == -5.0
    assert despues.diferencia_seg == antes.diferencia_seg == 5.0


def test_registro_sin_fecha_queda_incompleto():
    registro = M.RegistroCBB(5, {"Numero Factura": "1"}, None, 50000, True, ["FECHA_INVALIDA"])
    res = M.emparejar([registro], [mov(10, t(10, 0, 0), 500, 1)], CFG)
    assert res[0].estado == M.ESTADO_INCOMPLETO
    assert res[0].movimiento is None


# --------------------------------------------------------------------------- #
# Confirmacion manual de SIN_MATCH y columna Nro Oper de INGRESOS_NORMALIZADOS
# --------------------------------------------------------------------------- #

def _libro_con_sin_match(tmp_path, nombre="QUICKVALLE.xlsx"):
    """Un SIN_MATCH con candidato fuera de ventana y otro sin ningun candidato."""
    import reporte

    registros = [
        reg(1, t(10, 0, 0), 500, factura="A1", estudiante="CON CANDIDATO"),
        reg(2, t(11, 0, 0), 700, factura="A2", estudiante="SIN CANDIDATO"),
        reg(3, t(12, 0, 0), 900, factura="A3", estudiante="AUTOMATICO"),
    ]
    movimientos = [
        mov(10, t(16, 0, 0), 500, 452646),   # mismo dia y monto, pero a 6 horas
        mov(11, t(12, 0, 1), 900, 301902),   # match seguro
    ]
    resultados = M.emparejar(registros, movimientos, CFG)
    salida = tmp_path / nombre
    reporte.escribir(resultados, registros, movimientos, [], CFG, str(salida))
    return salida, resultados


def _evaluar(ruta):
    """Resuelve las formulas del libro con un motor de Excel."""
    formulas = pytest.importorskip("formulas")
    solucion = formulas.ExcelModel().loads(str(ruta)).finish().calculate()
    valores = {}
    for clave, valor in solucion.items():
        partes = clave.split("!")
        if len(partes) < 2:
            continue
        hoja = partes[-2].split("]")[-1].strip("'").upper()
        try:
            valores[(hoja, partes[-1].strip("'"))] = valor.value[0, 0]
        except Exception:
            pass
    return valores


def _decidir(origen, destino, celdas):
    import shutil

    shutil.copy(origen, destino)
    wb = openpyxl.load_workbook(destino)
    for hoja, celda, valor in celdas:
        wb[hoja][celda] = valor
    wb.save(destino)
    return destino


def test_sin_match_sin_candidato_no_ofrece_desplegable(tmp_path):
    import reporte

    salida, _ = _libro_con_sin_match(tmp_path)
    ws = openpyxl.load_workbook(salida)["SIN_MATCH"]
    candidato = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH)
    listas = {str(dv.sqref) for dv in ws.data_validations.dataValidation
              if dv.formula1.startswith(f"'{reporte.HOJA_CANDIDATOS}'")}
    assert listas == {f"{candidato}2"}  # solo el que tiene candidato real
    # Sin candidato no quedan celdas vacias: Excel leeria una celda vacia como 0.
    for desplazamiento in range(1, 7):
        assert ws.cell(row=3, column=reporte.COLUMNA_CANDIDATO_SIN_MATCH + desplazamiento).value == '=""'


def get_letra(indice):
    from openpyxl.utils import get_column_letter

    return get_column_letter(indice)


def test_sin_match_confirmado_entra_y_pendiente_no(tmp_path):
    import reporte

    salida, _ = _libro_con_sin_match(tmp_path)
    candidato = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH)
    decision = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH + 7)
    etiqueta = openpyxl.load_workbook(salida)[reporte.HOJA_CANDIDATOS]["A2"].value

    pendiente = _evaluar(salida)
    assert str(pendiente.get(("PARA_PEGAR_CBB", "A3"), "")).strip() == ""

    confirmado = _evaluar(_decidir(salida, tmp_path / "ok.xlsx", [
        ("SIN_MATCH", f"{candidato}2", etiqueta),
        ("SIN_MATCH", f"{decision}2", reporte.DECISION_CONFIRMAR),
    ]))
    assert confirmado[("PARA_PEGAR_CBB", "A3")] == reporte.ESTADO_MANUAL
    assert confirmado[("PARA_PEGAR_CBB", "B3")] == reporte.ORIGEN_MANUAL
    assert confirmado[("PARA_PEGAR_CBB", "C3")] == "A1"
    assert confirmado[("PARA_PEGAR_CBB", "K3")] == 452646


def test_sin_match_sin_candidato_no_se_puede_confirmar(tmp_path):
    import reporte

    salida, _ = _libro_con_sin_match(tmp_path)
    decision = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH + 7)
    control = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH + 8)

    valores = _evaluar(_decidir(salida, tmp_path / "invento.xlsx", [
        ("SIN_MATCH", f"{decision}3", reporte.DECISION_CONFIRMAR),
    ]))
    assert str(valores[("SIN_MATCH", f"{control}3")]).strip() == reporte.ALERTA_SIN_CANDIDATO
    # Ninguna ranura se llena: no se inventa ningun movimiento.
    for fila in (3, 4, 5, 6, 7):
        assert str(valores.get(("PARA_PEGAR_CBB", f"A{fila}"), "")).strip() == ""


def test_nro_oper_duplicado_queda_bloqueado(tmp_path):
    import reporte

    salida, _ = _libro_con_sin_match(tmp_path)
    candidato = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH)
    decision = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH + 7)
    control = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH + 8)

    wb = openpyxl.load_workbook(salida)
    # El candidato pasa a ser el Nro Oper que ya usa el match automatico.
    wb[reporte.HOJA_CANDIDATOS]["B2"] = wb["PARA_PEGAR_CBB"]["K2"].value
    wb.save(salida)

    etiqueta = openpyxl.load_workbook(salida)[reporte.HOJA_CANDIDATOS]["A2"].value
    valores = _evaluar(_decidir(salida, tmp_path / "dup.xlsx", [
        ("SIN_MATCH", f"{candidato}2", etiqueta),
        ("SIN_MATCH", f"{decision}2", reporte.DECISION_CONFIRMAR),
    ]))
    assert str(valores[("SIN_MATCH", f"{control}2")]).strip() == reporte.ALERTA_DUPLICADO
    assert str(valores.get(("PARA_PEGAR_CBB", "A3"), "")).strip() == ""


def test_columna_nro_oper_de_ingresos_normalizados(tmp_path):
    import reporte

    salida, _ = _libro_con_sin_match(tmp_path)
    candidato = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH)
    decision = get_letra(reporte.COLUMNA_CANDIDATO_SIN_MATCH + 7)
    etiqueta = openpyxl.load_workbook(salida)[reporte.HOJA_CANDIDATOS]["A2"].value

    # A1 sin match pendiente, A2 sin candidato, A3 match automatico.
    pendiente = _evaluar(salida)
    assert pendiente[("INGRESOS_NORMALIZADOS", "K4")] == 301902  # automatico
    assert str(pendiente[("INGRESOS_NORMALIZADOS", "K2")]).strip() == ""  # pendiente
    assert str(pendiente[("INGRESOS_NORMALIZADOS", "K3")]).strip() == ""  # sin candidato

    confirmado = _evaluar(_decidir(salida, tmp_path / "k.xlsx", [
        ("SIN_MATCH", f"{candidato}2", etiqueta),
        ("SIN_MATCH", f"{decision}2", reporte.DECISION_CONFIRMAR),
    ]))
    assert confirmado[("INGRESOS_NORMALIZADOS", "K2")] == 452646  # confirmado a mano
    assert confirmado[("INGRESOS_NORMALIZADOS", "K4")] == 301902


def test_columna_nro_oper_vacia_para_tarjeta(tmp_path):
    import reporte

    registros = [reg(1, t(10, 0, 0), 500, es_qr=False, factura="T1")]
    resultados = M.emparejar(registros, [], CFG)
    salida = tmp_path / "QUICKVALLE.xlsx"
    reporte.escribir(resultados, registros, [], [], CFG, str(salida))
    assert str(_evaluar(salida)[("INGRESOS_NORMALIZADOS", "K2")]).strip() == ""


# --------------------------------------------------------------------------- #
# Cierre historico: lo ya cerrado no se vuelve a decidir
# --------------------------------------------------------------------------- #

def cerrado(factura, oper, fecha_hora, monto):
    import cierre as C

    return C.RegistroCerrado(
        factura=M.norm_id(factura), fecha=fecha_hora.date(), centavos=M.a_centavos(monto),
        nro_oper=oper, nro_oper_norm=M.norm_id(oper), fecha_hora_bcp=fecha_hora,
        importe_centavos=M.a_centavos(monto), glosa="QR DE PRUEBA", cd_confirmacion=None,
        estado_previo=M.ESTADO_SEGURO,
    )


def test_match_historico_no_cambia_aunque_aparezca_uno_mejor():
    import cierre as C

    registros = [reg(1, t(10, 0, 0), 500, factura="F1")]
    movimientos = [
        mov(10, t(10, 0, 30), 500, "HISTORICO"),   # el que se cerro, a 30 s
        mov(11, t(10, 0, 0), 500, "MEJOR"),        # aparece uno exacto en el BCP nuevo
    ]
    congelado = C.aplicar(registros, movimientos, [cerrado("F1", "HISTORICO", t(10, 0, 30), 500)])

    assert [r.estado for r in congelado.resultados] == [C.ESTADO_HISTORICO]
    assert congelado.resultados[0].movimiento.nro_oper == "HISTORICO"
    assert congelado.facturas == {"F1"}
    # Y el registro sale del universo, asi que el motor nunca ve el candidato mejor.
    por_cruzar = [r for r in registros
                  if M.norm_id(r.crudo.get("Numero Factura")) not in congelado.facturas]
    assert por_cruzar == []


def test_nro_oper_historico_queda_reservado():
    import cierre as C

    registros = [reg(1, t(10, 0, 0), 500, factura="F1"), reg(2, t(10, 0, 1), 500, factura="F2")]
    movimientos = [mov(10, t(10, 0, 0), 500, "RESERVADO")]
    congelado = C.aplicar(registros, movimientos, [cerrado("F1", "RESERVADO", t(10, 0, 0), 500)])

    assert congelado.opers_reservados == {"RESERVADO"}
    disponibles = [m for m in movimientos if m.nro_oper_norm not in congelado.opers_reservados]
    assert disponibles == []
    # F2 queda sin match en vez de robar el movimiento ya cerrado.
    resto = M.emparejar([registros[1]], disponibles, CFG)
    assert resto[0].estado == M.ESTADO_SIN_MATCH


def test_pendiente_anterior_se_vuelve_a_intentar():
    import cierre as C

    registros = [reg(1, t(10, 0, 0), 500, factura="F1"), reg(2, t(11, 0, 0), 700, factura="PEND")]
    movimientos = [mov(10, t(10, 0, 0), 500, "CERRADO"), mov(11, t(11, 0, 1), 700, "NUEVO")]
    # El cierre anterior solo trae F1: lo que quedo pendiente no viene con Nro Oper.
    congelado = C.aplicar(registros, movimientos, [cerrado("F1", "CERRADO", t(10, 0, 0), 500)])

    por_cruzar = [r for r in registros
                  if M.norm_id(r.crudo.get("Numero Factura")) not in congelado.facturas]
    disponibles = [m for m in movimientos if m.nro_oper_norm not in congelado.opers_reservados]
    nuevos = M.emparejar(por_cruzar, disponibles, CFG)
    assert [r.registro.crudo["Numero Factura"] for r in nuevos] == ["PEND"]
    assert nuevos[0].estado == M.ESTADO_SEGURO
    assert nuevos[0].movimiento.nro_oper == "NUEVO"


def test_inconsistencia_historica_alerta_pero_no_reasigna():
    import cierre as C

    registros = [reg(1, t(10, 0, 0), 500, factura="F1")]
    # El Nro Oper cerrado ya no esta en el extracto nuevo, y hay otro candidato.
    movimientos = [mov(10, t(10, 0, 0), 500, "OTRO")]
    congelado = C.aplicar(registros, movimientos, [cerrado("F1", "DESAPARECIDO", t(10, 0, 0), 500)])

    resultado = congelado.resultados[0]
    assert resultado.estado == C.ESTADO_HISTORICO
    assert resultado.movimiento.nro_oper == "DESAPARECIDO"  # se conserva, no se reasigna
    assert C.ALERTA_INCONSISTENTE in resultado.motivo
    assert any("no existe en el extracto BCP actual" in detalle for _, detalle, _ in congelado.alertas)


def test_cambio_de_importe_historico_genera_alerta():
    import cierre as C

    registros = [reg(1, t(10, 0, 0), 500, factura="F1")]
    movimientos = [mov(10, t(10, 0, 0), 900, "HISTORICO")]  # el importe cambio
    congelado = C.aplicar(registros, movimientos, [cerrado("F1", "HISTORICO", t(10, 0, 0), 500)])

    assert C.ALERTA_INCONSISTENTE in congelado.resultados[0].motivo
    assert any("importe" in detalle for _, detalle, _ in congelado.alertas)


def test_pago_cerrado_que_ya_no_esta_en_el_reporte():
    import cierre as C

    congelado = C.aplicar([], [], [cerrado("F1", "HISTORICO", t(10, 0, 0), 500)])
    assert congelado.resultados == []
    assert congelado.opers_reservados == {"HISTORICO"}  # sigue reservado
    assert any("ya no aparece" in detalle for _, detalle, _ in congelado.alertas)


def test_identifica_el_cierre_anterior_por_estructura(tmp_path):
    ruta_cbb, ruta_bcp = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _crear_cbb(ruta_cbb)
    _crear_bcp(ruta_bcp)
    previo = tmp_path / "QUICKVALLE.xlsx"
    M.ejecutar([str(ruta_cbb), str(ruta_bcp)], str(previo), CFG)

    # El QUICKVALLE lleva dentro su copia normalizada del reporte: no debe
    # confundirse con el archivo de ingresos.
    cbb, bcp, cierre_tabla = M.identificar_tablas([str(previo), str(ruta_bcp), str(ruta_cbb)])
    assert cbb.ruta == str(ruta_cbb)
    assert bcp.ruta == str(ruta_bcp)
    assert cierre_tabla is not None and cierre_tabla.hoja == "PARA_PEGAR_CBB"


def test_tres_archivos_congelan_el_cierre_anterior(tmp_path):
    import cierre as C

    ruta_cbb, ruta_bcp = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _crear_cbb(ruta_cbb)
    _crear_bcp(ruta_bcp)
    previo = tmp_path / "QUICKVALLE.xlsx"
    primera = M.ejecutar([str(ruta_cbb), str(ruta_bcp)], str(previo), CFG)
    assert [r.estado for r in primera["resultados"]] == [M.ESTADO_SEGURO, M.ESTADO_NO_QR]

    segunda = M.ejecutar([str(ruta_cbb), str(ruta_bcp), str(previo)],
                         str(tmp_path / "nuevo.xlsx"), CFG)
    estados = {r.registro.crudo["Numero Factura"]: r.estado for r in segunda["resultados"]}
    assert estados["16180"] == C.ESTADO_HISTORICO  # ya no se vuelve a decidir
    assert estados["16181"] == M.ESTADO_NO_QR
    assert segunda["congelado"].opers_reservados == {"301902"}
