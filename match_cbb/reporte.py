"""Presentacion de QUICKVALLE.xlsx: la misma informacion, legible para una persona.

Este modulo solo da formato. No decide ningun match ni altera ningun resultado.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from typing import Any, Callable, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import cierre as C
import matcher as M

# Paleta suave y consistente: (encabezado, relleno de fila)
VERDE = ("2E7D32", "E8F5E9")
AZUL = ("1565C0", "E3F2FD")
NARANJA = ("E65100", "FFF8E1")
ROJO = ("B71C1C", "FFEBEE")
ROJO_INTENSO = "FFCDD2"
GRIS = ("616161", "F5F5F5")
AZUL_OSCURO = ("1F4E79", "FFFFFF")
AZUL_OSCURO_TARJETA = ("1F4E79", "DCE6F1")
INDIGO = ("283593", "E8EAF6")
TEAL = ("00695C", "E0F2F1")
INDIGO_EDITABLE = "C5CAE9"
ROJO_TEXTO = "C62828"

COLOR_POR_ESTADO = {
    M.ESTADO_SEGURO: VERDE[1],
    M.ESTADO_PROBABLE: AZUL[1],
    M.ESTADO_REVISAR: NARANJA[1],
    M.ESTADO_SIN_MATCH: ROJO[1],
    M.ESTADO_INCOMPLETO: ROJO[1],
    M.ESTADO_NO_QR: GRIS[1],
    C.ESTADO_HISTORICO: TEAL[1],
}

OBSERVACION_TARJETA = "Pago con tarjeta - fuera del cruce QR"

FECHA_HORA = "DD/MM/YYYY HH:MM:SS"
HORA = "HH:MM:SS"
MONEDA = "#,##0.00"
BORDE = Border(*(Side(style="thin", color="D0D0D0"),) * 4)

SIN_MATCH_CORTE = "SIN_MATCH_CORTE_BCP"
SIN_MATCH_REAL = "SIN_MATCH_REAL"

HOJA_CANDIDATOS = "_CANDIDATOS"
HOJA_MANUALES = "_MANUALES"
DECISION_CONFIRMAR = "CONFIRMAR MATCH"
DECISIONES = (DECISION_CONFIRMAR, "DESCARTAR", "PENDIENTE")
DECISION_POR_DEFECTO = "PENDIENTE"

ORIGEN_AUTOMATICO = "AUTOMÁTICO"
ORIGEN_HISTORICO = "CIERRE ANTERIOR"
ORIGEN_MANUAL = "CONFIRMADO MANUALMENTE"
ESTADO_MANUAL = "MATCH_MANUAL_CONFIRMADO"
OBSERVACION_MANUAL = "Confirmado manualmente"
OBSERVACION_HISTORICO = "Congelado del cierre anterior"
ALERTA_DUPLICADO = "NRO OPER YA UTILIZADO"
ALERTA_SIN_CANDIDATO = "SIN CANDIDATO: no se puede confirmar"
VERDE_MANUAL = "F1F8E9"

# En _MANUALES, la primera columna que PARA_PEGAR_CBB copia en sus ranuras.
PRIMERA_COLUMNA_SALIDA_MANUAL = 6

# Columna "Candidato elegido" en cada hoja con zona de decision.
COLUMNA_CANDIDATO_REVISAR = 10   # J
COLUMNA_CANDIDATO_SIN_MATCH = 14  # N

ETIQUETAS_ESTADISTICA = {
    "minimo": "Mínimo (seg)", "mediana": "Mediana (seg)", "p90": "Percentil 90 (seg)",
    "p95": "Percentil 95 (seg)", "maximo": "Máximo (seg)",
}


class Columna:
    """Una columna de una hoja operativa: titulo humano, valor y formato."""

    def __init__(self, titulo: str, valor: Callable[[Any], Any], formato: str | None = None, ancho: int | None = None):
        self.titulo = titulo
        self.valor = valor
        self.formato = formato
        self.ancho = ancho


# --------------------------------------------------------------------------- #
# Datos derivados (solo para mostrar)
# --------------------------------------------------------------------------- #

def _delta_con_signo(res: M.Resultado) -> float | None:
    """Positivo = la factura se emitio despues del abono bancario."""
    if res.movimiento is None or res.registro.fecha_hora is None or res.movimiento.fecha_hora is None:
        return None
    return round((res.registro.fecha_hora - res.movimiento.fecha_hora).total_seconds(), 3)


def _cd(res: M.Resultado) -> Any:
    return res.movimiento.cd_confirmacion if res.movimiento else None


def _bs(valor: float | None) -> str:
    """Importe en formato boliviano: 2.596,00"""
    if valor is None:
        return ""
    return f"{valor:,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")


def _origen(res: M.Resultado) -> str:
    return ORIGEN_HISTORICO if res.estado == C.ESTADO_HISTORICO else ORIGEN_AUTOMATICO


def _observacion_pegar(res: M.Resultado) -> str:
    if res.estado == C.ESTADO_HISTORICO:
        return res.motivo or OBSERVACION_HISTORICO
    cd = str(_cd(res)).strip() if _cd(res) not in (None, "") else ""
    base = "Listo para usar" if res.estado == M.ESTADO_SEGURO else "Probable, conviene revisar"
    if cd:
        return f"{base} — ya tenía Cd. Confirmación ({cd})"
    return base


def _tipo_sin_match(res: M.Resultado, inicio: dt.datetime | None, fin: dt.datetime | None) -> str:
    """Distingue el pago que el extracto no alcanza a cubrir del que deberia estar y no esta."""
    fecha = res.registro.fecha_hora
    if fecha is None or inicio is None or fin is None:
        return SIN_MATCH_REAL
    return SIN_MATCH_CORTE if fecha > fin or fecha < inicio else SIN_MATCH_REAL


# --------------------------------------------------------------------------- #
# Escritura generica de hojas
# --------------------------------------------------------------------------- #

def _escribir_tabla(
    ws,
    columnas: Sequence[Columna],
    filas: Sequence[Any],
    *,
    encabezado: tuple[str, str],
    relleno: Callable[[Any], str | None] | None = None,
    fila_inicio: int = 1,
) -> None:
    fuente_encabezado = Font(bold=True, color="FFFFFF", size=11)
    relleno_encabezado = PatternFill("solid", fgColor=encabezado[0])

    for c, columna in enumerate(columnas, start=1):
        celda = ws.cell(row=fila_inicio, column=c, value=columna.titulo)
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE

    for i, fila in enumerate(filas, start=fila_inicio + 1):
        color = relleno(fila) if relleno else None
        for c, columna in enumerate(columnas, start=1):
            celda = ws.cell(row=i, column=c, value=columna.valor(fila))
            celda.border = BORDE
            if columna.formato:
                celda.number_format = columna.formato
            if color:
                celda.fill = PatternFill("solid", fgColor=color)

    ws.freeze_panes = ws.cell(row=fila_inicio + 1, column=1)
    if filas:
        ws.auto_filter.ref = (
            f"A{fila_inicio}:{get_column_letter(len(columnas))}{fila_inicio + len(filas)}"
        )
    ws.row_dimensions[fila_inicio].height = 30
    _ajustar_anchos(ws, columnas, filas, fila_inicio)


def _ajustar_anchos(ws, columnas: Sequence[Columna], filas: Sequence[Any], fila_inicio: int) -> None:
    for c, columna in enumerate(columnas, start=1):
        if columna.ancho:
            ws.column_dimensions[get_column_letter(c)].width = columna.ancho
            continue
        largos = [len(columna.titulo)]
        for fila in filas[:400]:
            valor = columna.valor(fila)
            if isinstance(valor, dt.datetime):
                largos.append(19)
            elif valor is not None:
                largos.append(len(f"{valor:,.2f}") if columna.formato == MONEDA and isinstance(valor, (int, float))
                              else len(str(valor)))
        ws.column_dimensions[get_column_letter(c)].width = min(max(max(largos) + 2, 10), 42)


# --------------------------------------------------------------------------- #
# Hojas operativas
# --------------------------------------------------------------------------- #

def _columnas_comunes() -> list[Columna]:
    return [
        Columna("Factura", lambda r: r.registro.crudo.get("Numero Factura"), ancho=12),
        Columna("Nombre Estudiante", lambda r: r.registro.crudo.get("Nombre Estudiante")),
        Columna("Fecha y hora CBB", lambda r: r.registro.fecha_hora, FECHA_HORA, ancho=21),
        Columna("Monto Bs", lambda r: M.centavos_a_float(r.registro.centavos), MONEDA, ancho=13),
    ]


def _hoja_para_pegar(ws, resultados: Sequence[M.Resultado], n_manuales: int) -> None:
    """Historicos y automaticos fijos, mas una ranura por cada caso a decidir.

    Las ranuras se llenan solas cuando se marca CONFIRMAR MATCH en REVISAR_MANUAL
    o en SIN_MATCH, y quedan vacias en cualquier otro caso.
    """
    columnas = [
        Columna("Estado", lambda r: r.estado, ancho=24),
        Columna("Origen del match", _origen, ancho=22),
        *_columnas_comunes(),
        Columna("Tipo Pago", lambda r: r.registro.crudo.get("Tipo Pago"), ancho=12),
        Columna("Canal de Pago", lambda r: r.registro.crudo.get("Canal de Pago"), ancho=15),
        Columna("Fecha y hora BCP", lambda r: r.movimiento.fecha_hora if r.movimiento else None, FECHA_HORA, ancho=21),
        Columna("Dif. seg", lambda r: r.diferencia_seg, "0", ancho=9),
        Columna("Nro Oper BCP", lambda r: r.movimiento.nro_oper if r.movimiento else None, ancho=14),
        Columna("Glosa BCP", lambda r: r.movimiento.glosa.strip() if r.movimiento else None, ancho=24),
        Columna("Cd. Confirmación BCP", lambda r: _cd(r), ancho=18),
        Columna("Observación", _observacion_pegar, ancho=44),
    ]
    _escribir_tabla(ws, columnas, resultados, encabezado=AZUL_OSCURO,
                    relleno=lambda r: COLOR_POR_ESTADO.get(r.estado))
    _ranuras_manuales(ws, columnas, primera_fila=len(resultados) + 2, cantidad=n_manuales)


def _ranuras_manuales(ws, columnas: Sequence[Columna], primera_fila: int, cantidad: int) -> None:
    if not cantidad:
        return
    aux = f"'{HOJA_MANUALES}'"
    orden = f"{aux}!$E$2:$E${cantidad + 1}"
    for k in range(1, cantidad + 1):
        fila = primera_fila + k - 1
        for indice, columna in enumerate(columnas):
            letra = get_column_letter(PRIMERA_COLUMNA_SALIDA_MANUAL + indice)
            rango = f"{aux}!${letra}$2:${letra}${cantidad + 1}"
            celda = ws.cell(row=fila, column=indice + 1)
            celda.value = f'=IFERROR(INDEX({rango},MATCH({k},{orden},0)),"")'
            celda.border = BORDE
            celda.fill = PatternFill("solid", fgColor=VERDE_MANUAL)
            if columna.formato:
                celda.number_format = columna.formato
    ultima = primera_fila + cantidad - 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{ultima}"


@dataclass
class CasoManual:
    """Un caso que la persona puede confirmar, venga de REVISAR_MANUAL o de SIN_MATCH."""

    resultado: M.Resultado
    hoja: str
    fila: int
    columna_candidato: int


def _hoja_manuales(ws, casos: Sequence[CasoManual], filas_automaticas: int) -> None:
    """Puente entre la decision humana y PARA_PEGAR_CBB.

    Una fila por caso. Lee la decision y el candidato elegido, descarta los
    duplicados de Nro Oper y numera los validos para que las ranuras de
    PARA_PEGAR_CBB los tomen en orden, sin dejar huecos.
    """
    ws.append(["Decision", "NroOperElegido", "Duplicados", "Valido", "Orden",
               "Estado", "Origen", "Factura", "Nombre", "FechaHoraCBB", "Monto",
               "TipoPago", "Canal", "FechaHoraBCP", "DifSeg", "NroOper", "Glosa",
               "CdConfirmacion", "Observacion"])
    opers_automaticos = f"PARA_PEGAR_CBB!$K$2:$K${filas_automaticas + 1}"

    for i, caso in enumerate(casos, start=2):
        hoja, fila = caso.hoja, caso.fila
        col = caso.columna_candidato  # columna "Candidato elegido" en la hoja de origen
        col_oper = get_column_letter(col + 1)
        fecha_bcp = get_column_letter(col + 2)
        glosa = get_column_letter(col + 4)
        cd = get_column_letter(col + 5)
        dif = get_column_letter(col + 6)
        decision = get_column_letter(col + 7)
        previos = (f'+COUNTIFS($A$2:$A{i - 1},"{DECISION_CONFIRMAR}",$B$2:$B{i - 1},$B{i})'
                   if i > 2 else "")
        ws.append([
            f"='{hoja}'!{decision}{fila}",
            f"='{hoja}'!{col_oper}{fila}",
            f'=IF($B{i}="",0,COUNTIF({opers_automaticos},$B{i}){previos})',
            # Sin candidato elegido no hay match posible, aunque se marque confirmar.
            f'=IF(AND($A{i}="{DECISION_CONFIRMAR}",$B{i}<>"",$C{i}=0),1,0)',
            f'=IF($D{i}=1,SUM($D$2:$D{i}),"")',
            ESTADO_MANUAL,
            ORIGEN_MANUAL,
            f"='{hoja}'!A{fila}",
            f"='{hoja}'!B{fila}",
            f"='{hoja}'!C{fila}",
            f"='{hoja}'!D{fila}",
            caso.resultado.registro.crudo.get("Tipo Pago"),
            caso.resultado.registro.crudo.get("Canal de Pago"),
            f"='{hoja}'!{fecha_bcp}{fila}",
            f"='{hoja}'!{dif}{fila}",
            f"='{hoja}'!{col_oper}{fila}",
            f"='{hoja}'!{glosa}{fila}",
            f"='{hoja}'!{cd}{fila}",
            OBSERVACION_MANUAL,
        ])
        ws.cell(row=i, column=10).number_format = FECHA_HORA
        ws.cell(row=i, column=11).number_format = MONEDA
        ws.cell(row=i, column=14).number_format = FECHA_HORA
    ws.sheet_state = "hidden"


def candidatos_mismo_dia_monto(
    res: M.Resultado, movimientos: Sequence[M.MovimientoBCP]
) -> list[tuple[M.MovimientoBCP, float]]:
    """Todos los movimientos BCP del mismo dia y mismo importe, por cercania horaria.

    Es una vista mas amplia que la del motor (que ademas exige glosa QR y ventana de
    tiempo): aqui la persona decide, asi que ve todo lo que podria corresponder.
    """
    reg = res.registro
    if reg.fecha_hora is None or reg.centavos is None:
        return []
    candidatos = [
        (m, abs((reg.fecha_hora - m.fecha_hora).total_seconds()))
        for m in movimientos
        if m.fecha == reg.fecha and m.centavos == reg.centavos and m.fecha_hora is not None
    ]
    return sorted(candidatos, key=lambda par: (par[1], par[0].fila_excel))


def _etiqueta_candidato(mov: M.MovimientoBCP, delta: float) -> str:
    """Etiqueta legible del desplegable: sin esto la lista solo mostraria el Nro Oper."""
    return (f"{mov.nro_oper} | {mov.fecha_hora:%H:%M:%S} | Bs {_bs(M.centavos_a_float(mov.centavos))}"
            f" | {mov.glosa.strip()[:22]} | Δ {delta:.0f} s")


def _hoja_candidatos(ws, candidatos: dict[int, list[tuple[M.MovimientoBCP, float]]]) -> dict[int, tuple[int, int]]:
    """Tabla auxiliar oculta que alimenta el desplegable y las formulas de la zona de decision.

    Devuelve, por fila de Excel del pago de Cochabamba, el rango de sus candidatos.
    """
    ws.append(["Etiqueta", "Nro Oper", "FechaHora", "Importe", "Glosa", "Cd Confirmacion", "Dif seg"])
    rangos: dict[int, tuple[int, int]] = {}
    fila = 2
    for clave, lista in candidatos.items():
        inicio = fila
        for mov, delta in lista:
            ws.append([
                _etiqueta_candidato(mov, delta),
                mov.nro_oper,
                mov.fecha_hora,
                M.centavos_a_float(mov.centavos),
                mov.glosa.strip(),
                # Cadena vacia y no None: INDEX sobre una celda vacia devolveria 0.
                "" if mov.cd_confirmacion is None else mov.cd_confirmacion,
                round(delta),
            ])
            ws.cell(row=fila, column=3).number_format = FECHA_HORA
            ws.cell(row=fila, column=4).number_format = MONEDA
            fila += 1
        if fila > inicio:
            rangos[clave] = (inicio, fila - 1)
    ws.sheet_state = "hidden"
    return rangos


def _columnas_decision(preseleccion: Callable[[Any], Any]) -> list[Columna]:
    """El mismo bloque de decision en REVISAR_MANUAL y en SIN_MATCH."""
    return [
        Columna("Candidato elegido", preseleccion, ancho=52),
        Columna("Nro Oper elegido", lambda r: None, ancho=16),
        Columna("Fecha/hora BCP elegida", lambda r: None, FECHA_HORA, ancho=21),
        Columna("Importe BCP elegido", lambda r: None, MONEDA, ancho=15),
        Columna("Glosa BCP elegida", lambda r: None, ancho=26),
        Columna("Cd. Confirmación actual", lambda r: None, ancho=18),
        Columna("Dif. seg elegido", lambda r: None, "0", ancho=12),
        Columna("Decisión", lambda r: DECISION_POR_DEFECTO, ancho=18),
        Columna("Control", lambda r: None, ancho=30),
    ]


def _etiqueta_asignada(res: M.Resultado, lista: Sequence[tuple[M.MovimientoBCP, float]]) -> str | None:
    if res.movimiento is None:
        return None
    for mov, delta in lista:
        if mov.fila_excel == res.movimiento.fila_excel:
            return _etiqueta_candidato(mov, delta)
    return None


def _observacion_candidatos(cantidad: int) -> str:
    if cantidad == 0:
        return "Sin candidatos del mismo día y monto"
    if cantidad == 1:
        return "Candidato único"
    return f"Existen {cantidad} candidatos - revisar selección"


def _hoja_revisar(
    ws,
    resultados: Sequence[M.Resultado],
    candidatos: dict[int, list[tuple[M.MovimientoBCP, float]]],
    rangos: dict[int, tuple[int, int]],
    fila_manuales: int,
) -> None:
    def lista(res):
        return candidatos.get(res.registro.fila_excel, [])

    analisis = [
        *_columnas_comunes(),
        Columna("Candidato automático BCP", lambda r: _etiqueta_asignada(r, lista(r)), ancho=52),
        Columna("Dif. seg automática", lambda r: r.diferencia_seg, "0", ancho=11),
        Columna("Delta con signo", _delta_con_signo, "+0;-0;0", ancho=15),
        Columna("Cantidad candidatos", lambda r: len(lista(r)), "0", ancho=12),
        Columna("Observación", lambda r: _observacion_candidatos(len(lista(r))), ancho=34),
    ]
    # El candidato que propuso el algoritmo viene preseleccionado.
    decision = _columnas_decision(lambda r: _etiqueta_asignada(r, lista(r)))

    _escribir_tabla(ws, analisis + decision, resultados, encabezado=NARANJA, relleno=lambda r: NARANJA[1])
    _pintar_zona_decision(ws, len(analisis), len(decision), len(resultados))
    _formulas_decision(ws, resultados, rangos, len(analisis) + 1, fila_manuales)

    col_dif = next(i for i, c in enumerate(analisis, start=1) if c.titulo == "Dif. seg automática")
    for i, res in enumerate(resultados, start=2):
        if res.diferencia_seg is not None and res.diferencia_seg > 30:
            celda = ws.cell(row=i, column=col_dif)
            celda.font = Font(bold=True, color=ROJO_TEXTO)
            celda.fill = PatternFill("solid", fgColor=ROJO_INTENSO)


def _hoja_sin_match(
    ws,
    resultados: Sequence[M.Resultado],
    tipos: dict[int, str],
    candidatos: dict[int, list[tuple[M.MovimientoBCP, float]]],
    rangos: dict[int, tuple[int, int]],
    fila_manuales: int,
) -> None:
    def lista(res):
        return candidatos.get(res.registro.fila_excel, [])

    analisis = [
        *_columnas_comunes(),
        Columna("Tipo Pago", lambda r: r.registro.crudo.get("Tipo Pago"), ancho=12),
        Columna("Canal de Pago", lambda r: r.registro.crudo.get("Canal de Pago"), ancho=15),
        Columna("Tipo sin match", lambda r: tipos[r.registro.fila_excel], ancho=22),
        Columna("Candidato más cercano — fecha y hora",
                lambda r: r.candidato_cercano.fecha_hora if r.candidato_cercano else None, FECHA_HORA, ancho=21),
        Columna("Candidato más cercano — Nro Oper",
                lambda r: r.candidato_cercano.nro_oper if r.candidato_cercano else None, ancho=16),
        Columna("Candidato más cercano — Glosa",
                lambda r: r.candidato_cercano.glosa.strip() if r.candidato_cercano else None, ancho=24),
        Columna("Candidato más cercano — Dif. seg", lambda r: r.candidato_cercano_delta, "0", ancho=16),
        Columna("Cantidad candidatos", lambda r: len(lista(r)), "0", ancho=12),
        Columna("Motivo", lambda r: r.motivo, ancho=42),
    ]
    # El motor lo dejo SIN_MATCH: aqui no se preselecciona nada, la persona elige.
    decision = _columnas_decision(lambda r: None)

    _escribir_tabla(ws, analisis + decision, resultados, encabezado=ROJO,
                    relleno=lambda r: ROJO_INTENSO if tipos[r.registro.fila_excel] == SIN_MATCH_REAL else ROJO[1])
    _pintar_zona_decision(ws, len(analisis), len(decision), len(resultados))
    _formulas_decision(ws, resultados, rangos, len(analisis) + 1, fila_manuales)


def _pintar_zona_decision(ws, n_analisis: int, n_decision: int, n_filas: int) -> None:
    """La zona de decision se distingue del analisis; el candidato y la decision
    son los dos campos que se manipulan."""
    editables = (n_analisis + 1, n_analisis + n_decision - 1)
    for columna in range(n_analisis + 1, n_analisis + n_decision + 1):
        ws.cell(row=1, column=columna).fill = PatternFill("solid", fgColor=INDIGO[0])
        for fila in range(2, n_filas + 2):
            celda = ws.cell(row=fila, column=columna)
            celda.fill = PatternFill("solid", fgColor=
                                     INDIGO_EDITABLE if columna in editables else INDIGO[1])
            if columna in editables:
                celda.font = Font(bold=True)


def _formulas_decision(ws, resultados: Sequence[M.Resultado], rangos: dict[int, tuple[int, int]],
                       primera_columna: int, fila_manuales: int) -> None:
    """Las columnas del candidato elegido se resuelven con INDEX/MATCH sobre la
    tabla auxiliar, segun lo que se elija en el desplegable."""
    col_elegido = get_column_letter(primera_columna)
    col_decision = get_column_letter(primera_columna + 7)
    aux = f"'{HOJA_CANDIDATOS}'"
    validacion_decision = DataValidation(
        type="list", formula1='"' + ",".join(DECISIONES) + '"', allow_blank=False
    )
    ws.add_data_validation(validacion_decision)

    for i, res in enumerate(resultados, start=2):
        rango = rangos.get(res.registro.fila_excel)
        if rango:
            inicio, fin = rango
            validacion = DataValidation(
                type="list", formula1=f"{aux}!$A${inicio}:$A${fin}", allow_blank=True
            )
            ws.add_data_validation(validacion)
            validacion.add(ws.cell(row=i, column=primera_columna))

            for desplazamiento, columna_aux in enumerate("BCDEFG", start=1):
                indice = (f"INDEX({aux}!${columna_aux}${inicio}:${columna_aux}${fin},"
                          f"MATCH(${col_elegido}{i},{aux}!$A${inicio}:$A${fin},0))")
                # ISBLANK: sin esto una celda vacia (Cd. Confirmacion) se mostraria como 0.
                ws.cell(row=i, column=primera_columna + desplazamiento).value = (
                    f'=IFERROR(IF(ISBLANK({indice}),"",{indice}),"")'
                )
        else:
            # Sin candidatos no hay nada que elegir. Se deja cadena vacia y no celda
            # vacia: Excel lee una celda vacia como 0, y ese 0 valdria como Nro Oper.
            for desplazamiento in range(1, 7):
                ws.cell(row=i, column=primera_columna + desplazamiento).value = '=""'
        validacion_decision.add(ws.cell(row=i, column=primera_columna + 7))
        ws.cell(row=i, column=primera_columna + 8).value = (
            f"=IF('{HOJA_MANUALES}'!$C{fila_manuales + i - 2}>0,\"{ALERTA_DUPLICADO}\","
            f'IF({col_decision}{i}="{DECISION_CONFIRMAR}",'
            f'IF({col_elegido}{i}="","{ALERTA_SIN_CANDIDATO}","Incorporado a PARA_PEGAR_CBB"),'
            f'IF({col_decision}{i}="DESCARTAR","Descartado","")))'
        )

    if not resultados:
        return
    columna_control = get_column_letter(primera_columna + 8)
    ws.conditional_formatting.add(
        f"{columna_control}2:{columna_control}{len(resultados) + 1}",
        FormulaRule(
            formula=[f'OR(${columna_control}2="{ALERTA_DUPLICADO}",${columna_control}2="{ALERTA_SIN_CANDIDATO}")'],
            fill=PatternFill("solid", fgColor=ROJO_INTENSO),
            font=Font(bold=True, color=ROJO_TEXTO),
        ),
    )


def _hoja_tarjeta(ws, resultados: Sequence[M.Resultado]) -> None:
    columnas = [
        *_columnas_comunes(),
        Columna("Tipo Pago", lambda r: r.registro.crudo.get("Tipo Pago"), ancho=14),
        Columna("Canal de Pago", lambda r: r.registro.crudo.get("Canal de Pago"), ancho=15),
        Columna("Estado", lambda r: r.registro.crudo.get("Estado"), ancho=12),
        Columna("Observación", lambda r: OBSERVACION_TARJETA, ancho=38),
    ]
    _escribir_tabla(ws, columnas, resultados, encabezado=GRIS, relleno=lambda r: GRIS[1])


def _hoja_ingresos(ws, registros: Sequence[M.RegistroCBB], ultima_fila_pegar: int) -> None:
    """El reporte original de Cochabamba ya normalizado, mas el Nro Oper asignado.

    La columna K se resuelve buscando la factura en PARA_PEGAR_CBB, asi que refleja
    tanto los matches automaticos e historicos como las confirmaciones manuales, y
    queda vacia para lo pendiente, lo descartado y las tarjetas.
    """
    columnas = [
        Columna("Nro", lambda r: r.crudo.get("Nro"), "0", ancho=8),
        Columna("Fecha", lambda r: r.fecha_hora, "DD/MM/YYYY", ancho=13),
        Columna("Número Factura", lambda r: r.crudo.get("Numero Factura"), ancho=15),
        Columna("Nit/C.I.", lambda r: r.crudo.get("Nit/C.I."), ancho=14),
        Columna("Razon Social", lambda r: r.crudo.get("Razon Social"), ancho=24),
        Columna("Nombre Estudiante", lambda r: r.crudo.get("Nombre Estudiante"), ancho=36),
        Columna("Tipo Pago", lambda r: r.crudo.get("Tipo Pago"), ancho=13),
        Columna("Monto", lambda r: M.centavos_a_float(r.centavos), MONEDA, ancho=13),
        Columna("Canal de Pago", lambda r: r.crudo.get("Canal de Pago"), ancho=16),
        Columna("Estado", lambda r: r.crudo.get("Estado"), ancho=11),
        Columna("Nro Oper BCP", lambda r: None, ancho=14),
    ]
    _escribir_tabla(ws, columnas, registros, encabezado=AZUL_OSCURO)

    facturas = f"PARA_PEGAR_CBB!$C$2:$C${ultima_fila_pegar}"
    opers = f"PARA_PEGAR_CBB!$K$2:$K${ultima_fila_pegar}"
    for i in range(2, len(registros) + 2):
        ws.cell(row=i, column=len(columnas)).value = (
            f'=IFERROR(INDEX({opers},MATCH($C{i},{facturas},0)),"")'
        )


# --------------------------------------------------------------------------- #
# Dashboard
# --------------------------------------------------------------------------- #

def _hoja_resumen(ws, datos: dict[str, Any]) -> None:
    ws.sheet_view.showGridLines = False
    for c, ancho in zip("ABCDEFGH", (22, 16, 22, 16, 22, 16, 22, 16)):
        ws.column_dimensions[c].width = ancho

    _titulo(ws, 1, "CRUCE COCHABAMBA vs BCP", 18, AZUL_OSCURO[0], alto=32)
    _titulo(ws, 2, f"Período analizado: {datos['periodo']}", 12, "5B7C99", alto=20)
    _titulo(ws, 3, f"Generado: {dt.datetime.now():%d/%m/%Y %H:%M:%S}", 10, "8095A8", alto=18)

    fila = 5
    _seccion(ws, fila, "UNIVERSO ANALIZADO")
    fila = _tarjetas(ws, fila + 1, [
        ("Registros Cochabamba", datos["total_cbb"], GRIS),
        ("QR Cochabamba", datos["total_qr"], GRIS),
        ("Tarjeta", datos["total_tarjeta"], GRIS),
        ("Movimientos BCP", datos["total_bcp"], GRIS),
    ])

    fila += 1
    _seccion(ws, fila, "RESULTADO DEL CRUCE")
    fila = _tarjetas(ws, fila + 1, [
        ("MATCH SEGURO", datos["n_seguro"], VERDE),
        ("MATCH PROBABLE", datos["n_probable"], AZUL),
        ("REVISAR", datos["n_revisar"], NARANJA),
        ("SIN MATCH", datos["n_sin_match"], ROJO),
    ])

    fila += 1
    _tarjetas(ws, fila, [("% ENCONTRADO SOBRE QR COCHABAMBA", datos["pct"], VERDE, "0.00%")], ancho_tarjeta=8)
    fila += 3

    # Se recalcula solo con las decisiones que se tomen en REVISAR_MANUAL y SIN_MATCH.
    validos = f"'{HOJA_MANUALES}'!$D$2:$D${max(datos['n_casos_manuales'], 1) + 1}"
    _seccion(ws, fila, "LISTO PARA PEGAR")
    fila = _tarjetas(ws, fila + 1, [
        ("CONFIRMADOS HISTÓRICOS", datos["n_historicos"], TEAL),
        ("AUTOMÁTICOS", datos["n_automaticos"], AZUL_OSCURO_TARJETA),
        ("CONFIRMADOS MANUALMENTE", f"=SUM({validos})", VERDE),
        ("TOTAL LISTO PARA PEGAR", f"={datos['n_fijos']}+SUM({validos})", VERDE),
    ], ancho_tarjeta=2)
    fila += 1

    _seccion(ws, fila, "IMPORTES (Bs)")
    fila += 1
    importes = [
        ("Total QR Cochabamba", datos["bs_qr"], None),
        ("MATCH SEGURO", datos["bs_seguro"], VERDE[1]),
        ("MATCH PROBABLE", datos["bs_probable"], AZUL[1]),
        ("REVISAR", datos["bs_revisar"], NARANJA[1]),
        ("SIN MATCH", datos["bs_sin_match"], ROJO[1]),
    ]
    for concepto, valor, color in importes:
        _fila_dato(ws, fila, concepto, valor, MONEDA, color, negrita=color is None)
        fila += 1

    fila += 1
    _seccion(ws, fila, "LEYENDA DE COLORES")
    fila += 1
    for texto, color in (("VERDE — MATCH SEGURO", VERDE[1]), ("AZUL — MATCH PROBABLE", AZUL[1]),
                         ("AMARILLO — REVISAR", NARANJA[1]), ("ROJO — SIN MATCH", ROJO[1]),
                         ("GRIS — TARJETA (fuera del cruce QR)", GRIS[1])):
        celda = ws.cell(row=fila, column=1, value=texto)
        celda.fill = PatternFill("solid", fgColor=color)
        celda.border = BORDE
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=3)
        for c in (2, 3):
            ws.cell(row=fila, column=c).fill = PatternFill("solid", fgColor=color)
            ws.cell(row=fila, column=c).border = BORDE
        fila += 1

    fila += 1
    _seccion(ws, fila, "DISTRIBUCIÓN DE DIFERENCIAS DE TIEMPO")
    fila += 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
    ws.merge_cells(start_row=fila, start_column=7, end_row=fila, end_column=8)
    for columna, titulo in ((1, "Diferencia entre la factura y el abono"), (7, "Casos")):
        ws.cell(row=fila, column=columna, value=titulo).font = Font(bold=True, color="FFFFFF")
    for c in range(1, 9):
        ws.cell(row=fila, column=c).fill = PatternFill("solid", fgColor=AZUL_OSCURO[0])
        ws.cell(row=fila, column=c).border = BORDE
    fila += 1
    for rango, casos in datos["distribucion"].items():
        _fila_dato(ws, fila, rango, casos, "0", "FFFFFF")
        fila += 1
    for etiqueta, valor in datos["estadisticas"].items():
        _fila_dato(ws, fila, ETIQUETAS_ESTADISTICA.get(etiqueta, etiqueta), valor, "0.##", GRIS[1])
        fila += 1

    fila += 1
    _seccion(ws, fila, "VALIDACIONES")
    fila += 1
    if datos["alertas"]:
        for nivel, detalle, cantidad in datos["alertas"]:
            color = ROJO[1] if nivel == "CRITICO" else (NARANJA[1] if nivel == "ALERTA" else GRIS[1])
            _fila_dato(ws, fila, f"[{nivel}] {detalle}", cantidad, "0", color)
            fila += 1
    else:
        _fila_dato(ws, fila, "Sin observaciones", 0, "0", VERDE[1])


def _titulo(ws, fila: int, texto: str, tamano: int, color: str, alto: int) -> None:
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=8)
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(bold=True, size=tamano, color=color)
    celda.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[fila].height = alto


def _seccion(ws, fila: int, texto: str) -> None:
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(bold=True, size=11, color=AZUL_OSCURO[0])
    ws.row_dimensions[fila].height = 20


def _tarjetas(ws, fila: int, tarjetas: Sequence[tuple], ancho_tarjeta: int = 2) -> int:
    """Bloques con color: etiqueta arriba, numero grande abajo."""
    ws.row_dimensions[fila].height = 18
    ws.row_dimensions[fila + 1].height = 30
    for i, tarjeta in enumerate(tarjetas):
        etiqueta, valor, paleta = tarjeta[0], tarjeta[1], tarjeta[2]
        formato = tarjeta[3] if len(tarjeta) > 3 else "#,##0"
        col = 1 + i * ancho_tarjeta
        ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + ancho_tarjeta - 1)
        ws.merge_cells(start_row=fila + 1, start_column=col, end_row=fila + 1, end_column=col + ancho_tarjeta - 1)

        celda = ws.cell(row=fila, column=col, value=etiqueta)
        celda.font = Font(bold=True, size=9, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=paleta[0])
        celda.alignment = Alignment(horizontal="center", vertical="center")

        valor_celda = ws.cell(row=fila + 1, column=col, value=valor)
        valor_celda.font = Font(bold=True, size=16, color=paleta[0])
        valor_celda.fill = PatternFill("solid", fgColor=paleta[1])
        valor_celda.alignment = Alignment(horizontal="center", vertical="center")
        valor_celda.number_format = formato

        for f in (fila, fila + 1):
            for c in range(col, col + ancho_tarjeta):
                ws.cell(row=f, column=c).border = BORDE
                if f == fila:
                    ws.cell(row=f, column=c).fill = PatternFill("solid", fgColor=paleta[0])
                else:
                    ws.cell(row=f, column=c).fill = PatternFill("solid", fgColor=paleta[1])
    return fila + 2


def _fila_dato(ws, fila: int, concepto: str, valor: Any, formato: str, color: str | None, negrita: bool = False) -> None:
    """Concepto ancho (A:F) y valor a la derecha (G:H), para que los textos largos quepan."""
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
    ws.merge_cells(start_row=fila, start_column=7, end_row=fila, end_column=8)
    celda = ws.cell(row=fila, column=1, value=concepto)
    celda.font = Font(bold=negrita, size=10)
    valor_celda = ws.cell(row=fila, column=7, value=valor)
    valor_celda.number_format = formato
    valor_celda.font = Font(bold=True, size=10)
    valor_celda.alignment = Alignment(horizontal="right")
    for c in range(1, 9):
        ws.cell(row=fila, column=c).border = BORDE
        if color:
            ws.cell(row=fila, column=c).fill = PatternFill("solid", fgColor=color)


# --------------------------------------------------------------------------- #
# Entrada principal
# --------------------------------------------------------------------------- #

def escribir(
    resultados: Sequence[M.Resultado],
    registros: Sequence[M.RegistroCBB],
    movimientos: Sequence[M.MovimientoBCP],
    alertas: Sequence[tuple[str, str, int]],
    cfg: M.Config,
    salida: str,
) -> str:
    por_estado = {e: [r for r in resultados if r.estado == e] for e in
                  (C.ESTADO_HISTORICO, M.ESTADO_SEGURO, M.ESTADO_PROBABLE,
                   M.ESTADO_REVISAR, M.ESTADO_NO_QR)}
    sin_match = [r for r in resultados if r.estado in (M.ESTADO_SIN_MATCH, M.ESTADO_INCOMPLETO)]
    qr = [r for r in resultados if r.estado != M.ESTADO_NO_QR]

    fechas_bcp = [m.fecha_hora for m in movimientos if m.fecha_hora is not None]
    inicio, fin = (min(fechas_bcp), max(fechas_bcp)) if fechas_bcp else (None, None)
    tipos = {r.registro.fila_excel: _tipo_sin_match(r, inicio, fin) for r in sin_match}

    conteo, stats = M.distribucion(resultados)

    def bs(items: Sequence[M.Resultado]) -> float:
        return round(sum(M.centavos_a_float(r.registro.centavos) or 0.0 for r in items), 2)

    fechas_cbb = [r.fecha for r in registros if r.fecha is not None]
    historicos = por_estado[C.ESTADO_HISTORICO]
    automaticos = por_estado[M.ESTADO_SEGURO] + por_estado[M.ESTADO_PROBABLE]
    revisar = por_estado[M.ESTADO_REVISAR]
    encontrados = len(historicos) + len(automaticos) + len(revisar)

    # Lo que ya no se decide: historicos congelados y matches automaticos.
    fijos = sorted(historicos + automaticos, key=lambda r: r.registro.fila_excel)
    # Lo que la persona puede confirmar: primero REVISAR, luego SIN_MATCH.
    casos = (
        [CasoManual(r, "REVISAR_MANUAL", i, COLUMNA_CANDIDATO_REVISAR)
         for i, r in enumerate(revisar, start=2)]
        + [CasoManual(r, "SIN_MATCH", i, COLUMNA_CANDIDATO_SIN_MATCH)
           for i, r in enumerate(sin_match, start=2)]
    )
    candidatos = {r.registro.fila_excel: candidatos_mismo_dia_monto(r, movimientos)
                  for r in revisar + sin_match}

    wb = Workbook()
    wb.remove(wb.active)

    _hoja_resumen(wb.create_sheet("RESUMEN_VISUAL"), {
        "periodo": (f"{min(fechas_cbb):%d/%m/%Y} a {max(fechas_cbb):%d/%m/%Y}" if fechas_cbb else "-"),
        "total_cbb": len(registros),
        "total_qr": len(qr),
        "total_tarjeta": len(por_estado[M.ESTADO_NO_QR]),
        "total_bcp": len(movimientos),
        "n_seguro": len(por_estado[M.ESTADO_SEGURO]),
        "n_probable": len(por_estado[M.ESTADO_PROBABLE]),
        "n_revisar": len(revisar),
        "n_sin_match": len(sin_match),
        "pct": (encontrados / len(qr)) if qr else 0.0,
        "bs_qr": bs(qr),
        "bs_seguro": bs(por_estado[M.ESTADO_SEGURO]),
        "bs_probable": bs(por_estado[M.ESTADO_PROBABLE]),
        "bs_revisar": bs(revisar),
        "bs_sin_match": bs(sin_match),
        "distribucion": conteo,
        "estadisticas": stats,
        "alertas": alertas,
        "n_historicos": len(historicos),
        "n_automaticos": len(automaticos),
        "n_fijos": len(fijos),
        "n_casos_manuales": len(casos),
    })

    _hoja_para_pegar(wb.create_sheet("PARA_PEGAR_CBB"), fijos, len(casos))

    rangos = _hoja_candidatos(wb.create_sheet(HOJA_CANDIDATOS), candidatos)
    _hoja_revisar(wb.create_sheet("REVISAR_MANUAL"), revisar, candidatos, rangos, fila_manuales=2)
    _hoja_sin_match(wb.create_sheet("SIN_MATCH"), sin_match, tipos, candidatos, rangos,
                    fila_manuales=2 + len(revisar))
    _hoja_manuales(wb.create_sheet(HOJA_MANUALES), casos, len(fijos))

    _hoja_tarjeta(wb.create_sheet("TARJETA"), por_estado[M.ESTADO_NO_QR])
    _hoja_ingresos(wb.create_sheet("INGRESOS_NORMALIZADOS"), registros,
                   ultima_fila_pegar=len(fijos) + len(casos) + 1)

    for oculta in (HOJA_CANDIDATOS, HOJA_MANUALES):
        wb.move_sheet(oculta, offset=len(wb.sheetnames))

    wb.save(salida)
    return salida
