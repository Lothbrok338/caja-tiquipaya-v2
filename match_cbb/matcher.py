"""Matcher entre el reporte de ingresos de Cochabamba y el extracto BCP nacional.

Llave de cruce: misma fecha + mismo monto + hora mas cercana, con asignacion 1 a 1
(scipy.optimize.linear_sum_assignment) dentro de cada grupo fecha+monto.

Los dos archivos de entrada se abren SIEMPRE en modo lectura. El unico archivo que
se escribe es el Excel de salida (MATCH_CBB.xlsx por defecto).
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import statistics
import sys
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Iterable, Sequence

import numpy as np
from scipy.optimize import linear_sum_assignment

NOMBRE_SALIDA = "MATCH_CBB.xlsx"

# Encabezados que deben aparecer SIMULTANEAMENTE para identificar cada tabla.
REQUERIDOS_CBB = ("Fecha", "Numero Factura", "Monto", "Nombre Estudiante")
REQUERIDOS_BCP = ("Fecha", "Hora", "Importe", "Nro Oper")

MAX_FILAS_ESCANEO = 60

# Variantes aceptadas por columna (la primera que coincida gana).
COLUMNAS_CBB = {
    "Nro": ("Nro", "Nro.", "N"),
    "Fecha": ("Fecha",),
    "Numero Factura": ("Numero Factura", "Nro Factura", "Factura"),
    "Nit/C.I.": ("Nit C I", "Nit CI", "Nit"),
    "Razon Social": ("Razon Social",),
    "Nombre Estudiante": ("Nombre Estudiante", "Estudiante"),
    "Tipo Pago": ("Tipo Pago",),
    "Monto": ("Monto",),
    "Canal de Pago": ("Canal de Pago", "Canal"),
    "Estado": ("Estado",),
}

COLUMNAS_BCP = {
    "Fecha": ("Fecha",),
    "Hora": ("Hora",),
    "Glosa": ("Glosa", "Descripcion"),
    "Importe": ("Importe", "Monto"),
    "Nro Oper": ("Nro Oper", "Numero Operacion", "Nro Operacion"),
    "Cd Confirmacion": ("Cd Confirmacion", "Cod Confirmacion", "Codigo Confirmacion"),
}

# Columnas obligatorias: si faltan, no se puede cruzar.
OBLIGATORIAS_CBB = ("Fecha", "Numero Factura", "Monto")
OBLIGATORIAS_BCP = ("Fecha", "Hora", "Importe", "Nro Oper")

ESTADO_SEGURO = "MATCH_SEGURO"
ESTADO_PROBABLE = "MATCH_PROBABLE"
ESTADO_REVISAR = "REVISAR"
ESTADO_SIN_MATCH = "SIN_MATCH"
ESTADO_INCOMPLETO = "DATOS_INCOMPLETOS"
ESTADO_NO_QR = "NO_QR"

# Costo usado para prohibir un par fuera de la ventana maxima sin romper el Hungarian.
COSTO_PROHIBIDO = 1e9


@dataclass(frozen=True)
class Config:
    """Umbrales y reglas. Todo centralizado aqui para poder ajustarlo sin tocar el motor."""

    seguro_seg: float = 5.0
    probable_seg: float = 30.0
    revisar_seg: float = 120.0
    # Si un movimiento BCP libre queda a esta distancia del asignado, el caso es ambiguo.
    margen_ambiguo_seg: float = 2.0
    degradar_ambiguos: bool = True
    solo_qr_bcp: bool = True
    patron_qr_glosa: str = "QR"
    tipos_pago_qr: tuple[str, ...] = ("QR",)
    ciudad_propia: tuple[str, ...] = ("COCHABAMBA", "CBBA", "CBB")


# --------------------------------------------------------------------------- #
# Normalizacion
# --------------------------------------------------------------------------- #

def norm_texto(valor: Any) -> str:
    """Mayusculas, sin tildes, solo alfanumerico separado por espacios."""
    if valor is None:
        return ""
    s = str(valor)
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"[^A-Za-z0-9]+", " ", s.upper())
    return re.sub(r"\s+", " ", s).strip()


def norm_id(valor: Any) -> str:
    """Normaliza un identificador: 301902, '301902', ' 301902 ' y 301902.0 son lo mismo."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor))
        return repr(valor)
    s = str(valor).strip()
    if not s:
        return ""
    s = re.sub(r"\s+", "", s)
    if re.fullmatch(r"-?\d+\.0+", s):
        s = s.split(".")[0]
    return s


def _contiene_palabras(celda_norm: str, requerido_norm: str) -> bool:
    """True si el requerido aparece como secuencia de palabras completas dentro de la celda."""
    if not celda_norm or not requerido_norm:
        return False
    patron = r"(?:^| )" + re.escape(requerido_norm) + r"(?:$| )"
    return re.search(patron, celda_norm) is not None


def a_datetime(valor: Any) -> dt.datetime | None:
    """Convierte a datetime conservando la hora. None si no es una fecha valida."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, dt.datetime):
        return valor
    if isinstance(valor, dt.date):
        return dt.datetime(valor.year, valor.month, valor.day)
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        # Serial de Excel (sistema 1900).
        try:
            base = dt.datetime(1899, 12, 30)
            return base + dt.timedelta(days=float(valor))
        except (OverflowError, ValueError):
            return None
    texto = str(valor).strip()
    if not texto:
        return None
    formatos = (
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    )
    for fmt in formatos:
        try:
            return dt.datetime.strptime(texto, fmt)
        except ValueError:
            continue
    return None


def redondear_a_segundo(valor: dt.datetime | None) -> dt.datetime | None:
    """Ambos reportes muestran segundos enteros.

    El serial de Excel del reporte de Cochabamba arrastra un desfase constante de
    +0.107 s al convertirse a datetime; sin este redondeo toda diferencia sale con
    decimales que no existen en el dato original.
    """
    if valor is None or valor.microsecond == 0:
        return valor
    return (valor + dt.timedelta(microseconds=500000)).replace(microsecond=0)


def a_time(valor: Any) -> dt.time | None:
    """Convierte a hora. Acepta time, datetime, timedelta, fraccion de dia y texto."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, dt.datetime):
        return valor.time()
    if isinstance(valor, dt.time):
        return valor
    if isinstance(valor, dt.timedelta):
        total = int(valor.total_seconds()) % 86400
        return dt.time(total // 3600, (total % 3600) // 60, total % 60)
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        frac = float(valor) % 1.0
        total = int(round(frac * 86400)) % 86400
        return dt.time(total // 3600, (total % 3600) // 60, total % 60)
    texto = str(valor).strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return dt.datetime.strptime(texto, fmt).time()
        except ValueError:
            continue
    return None


def a_centavos(valor: Any) -> int | None:
    """Monto a centavos enteros: evita comparar floats. None si no es numerico."""
    if valor is None or valor == "" or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return int(round(float(valor) * 100))
    texto = str(valor).strip()
    if not texto:
        return None
    texto = texto.replace(" ", "")
    # Formato boliviano/ingles: 1,234.56 -> el ultimo separador manda.
    if "," in texto and "." in texto:
        texto = texto.replace(",", "") if texto.rfind(".") > texto.rfind(",") else texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".") if len(texto.split(",")[-1]) <= 2 else texto.replace(",", "")
    try:
        return int(round(float(texto) * 100))
    except ValueError:
        return None


def centavos_a_float(centavos: int | None) -> float | None:
    return None if centavos is None else centavos / 100.0


# --------------------------------------------------------------------------- #
# Lectura de libros (solo lectura, nunca escribe sobre el original)
# --------------------------------------------------------------------------- #

def leer_libro(ruta: str) -> list[tuple[str, list[list[Any]]]]:
    """Devuelve [(nombre_hoja, filas)] con valores nativos de Python."""
    ext = os.path.splitext(ruta)[1].lower()
    if ext == ".xls":
        return _leer_xls(ruta)
    if ext in (".xlsx", ".xlsm"):
        return _leer_xlsx(ruta)
    raise ValueError(f"Extension no soportada: {ruta} (se esperaba .xls, .xlsx o .xlsm)")


def _leer_xls(ruta: str) -> list[tuple[str, list[list[Any]]]]:
    import xlrd

    libro = xlrd.open_workbook(ruta)
    hojas = []
    for nombre in libro.sheet_names():
        hoja = libro.sheet_by_name(nombre)
        filas = []
        for r in range(hoja.nrows):
            fila = []
            for c in range(hoja.ncols):
                celda = hoja.cell(r, c)
                if celda.ctype == xlrd.XL_CELL_DATE:
                    fila.append(_xls_fecha(celda.value, libro.datemode))
                elif celda.ctype == xlrd.XL_CELL_EMPTY:
                    fila.append(None)
                elif celda.ctype == xlrd.XL_CELL_BOOLEAN:
                    fila.append(bool(celda.value))
                else:
                    fila.append(celda.value)
            filas.append(fila)
        hojas.append((nombre, filas))
    return hojas


def _xls_fecha(valor: float, datemode: int) -> Any:
    import xlrd

    try:
        if valor < 1:
            return xlrd.xldate_as_datetime(valor, datemode).time()
        return xlrd.xldate_as_datetime(valor, datemode)
    except (ValueError, OverflowError, xlrd.XLDateError):
        return valor


def _leer_xlsx(ruta: str) -> list[tuple[str, list[list[Any]]]]:
    import openpyxl

    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        return [
            (nombre, [list(fila) for fila in libro[nombre].iter_rows(values_only=True)])
            for nombre in libro.sheetnames
        ]
    finally:
        libro.close()


# --------------------------------------------------------------------------- #
# Deteccion de encabezados e identificacion de archivos
# --------------------------------------------------------------------------- #

def detectar_fila_encabezado(filas: Sequence[Sequence[Any]], requeridos: Iterable[str]) -> int | None:
    """Indice (0-based) de la primera fila que contiene TODOS los encabezados requeridos.

    Exigir varios encabezados a la vez evita confundir la tabla con celdas sueltas
    del membrete (por ejemplo un 'Fecha:' de cabecera del reporte).
    """
    req = [norm_texto(r) for r in requeridos]
    for i, fila in enumerate(filas[:MAX_FILAS_ESCANEO]):
        celdas = [norm_texto(v) for v in fila]
        if all(any(c == r or _contiene_palabras(c, r) for c in celdas) for r in req):
            return i
    return None


def mapear_columnas(encabezado: Sequence[Any], definicion: dict[str, tuple[str, ...]]) -> dict[str, int]:
    """Canonico -> indice de columna, segun las variantes declaradas."""
    celdas = [norm_texto(v) for v in encabezado]
    mapa: dict[str, int] = {}
    usados: set[int] = set()
    for canonico, variantes in definicion.items():
        idx = _buscar_columna(celdas, variantes, usados)
        if idx is not None:
            mapa[canonico] = idx
            usados.add(idx)
    return mapa


def _buscar_columna(celdas: list[str], variantes: tuple[str, ...], usados: set[int]) -> int | None:
    for variante in variantes:
        nv = norm_texto(variante)
        for i, celda in enumerate(celdas):
            if i not in usados and celda == nv:
                return i
    for variante in variantes:
        nv = norm_texto(variante)
        for i, celda in enumerate(celdas):
            if i not in usados and _contiene_palabras(celda, nv):
                return i
    return None


@dataclass
class TablaDetectada:
    ruta: str
    hoja: str
    fila_encabezado: int
    filas: list[list[Any]]
    columnas: dict[str, int]


def identificar_tablas(rutas: Sequence[str]) -> tuple[TablaDetectada, TablaDetectada]:
    """Identifica cual archivo es Cochabamba y cual es BCP por estructura, no por nombre.

    Revisa todas las hojas de todos los libros entregados, en cualquier orden.
    """
    candidatos_cbb: list[TablaDetectada] = []
    candidatos_bcp: list[TablaDetectada] = []

    for ruta in rutas:
        for hoja, filas in leer_libro(ruta):
            fila_cbb = detectar_fila_encabezado(filas, REQUERIDOS_CBB)
            if fila_cbb is not None:
                cols = mapear_columnas(filas[fila_cbb], COLUMNAS_CBB)
                if all(c in cols for c in OBLIGATORIAS_CBB):
                    candidatos_cbb.append(TablaDetectada(ruta, hoja, fila_cbb, filas, cols))
                    continue
            fila_bcp = detectar_fila_encabezado(filas, REQUERIDOS_BCP)
            if fila_bcp is not None:
                cols = mapear_columnas(filas[fila_bcp], COLUMNAS_BCP)
                if all(c in cols for c in OBLIGATORIAS_BCP):
                    candidatos_bcp.append(TablaDetectada(ruta, hoja, fila_bcp, filas, cols))

    if not candidatos_cbb or not candidatos_bcp:
        raise ValueError(
            "No pude identificar ambas tablas por estructura.\n"
            f"  Cochabamba (requiere {list(REQUERIDOS_CBB)}): "
            f"{[(c.ruta, c.hoja) for c in candidatos_cbb] or 'NO ENCONTRADA'}\n"
            f"  BCP (requiere {list(REQUERIDOS_BCP)}): "
            f"{[(c.ruta, c.hoja) for c in candidatos_bcp] or 'NO ENCONTRADA'}\n"
            "Revisa que hayas entregado el reporte de ingresos Cochabamba y el extracto BCP."
        )
    if len(candidatos_cbb) > 1 or len(candidatos_bcp) > 1:
        raise ValueError(
            "Identificacion ambigua: mas de una hoja califica.\n"
            f"  Cochabamba: {[(c.ruta, c.hoja) for c in candidatos_cbb]}\n"
            f"  BCP: {[(c.ruta, c.hoja) for c in candidatos_bcp]}"
        )
    if candidatos_cbb[0].ruta == candidatos_bcp[0].ruta and candidatos_cbb[0].hoja == candidatos_bcp[0].hoja:
        raise ValueError("La misma hoja califica como Cochabamba y como BCP; no puedo continuar.")
    return candidatos_cbb[0], candidatos_bcp[0]


# --------------------------------------------------------------------------- #
# Registros
# --------------------------------------------------------------------------- #

@dataclass
class RegistroCBB:
    fila_excel: int
    crudo: dict[str, Any]
    fecha_hora: dt.datetime | None
    centavos: int | None
    es_qr: bool
    problemas: list[str] = field(default_factory=list)

    @property
    def fecha(self) -> dt.date | None:
        return None if self.fecha_hora is None else self.fecha_hora.date()

    @property
    def utilizable(self) -> bool:
        return self.fecha_hora is not None and self.centavos is not None


@dataclass
class MovimientoBCP:
    fila_excel: int
    crudo: dict[str, Any]
    fecha_hora: dt.datetime | None
    centavos: int | None
    nro_oper: Any
    nro_oper_norm: str
    glosa: str
    cd_confirmacion: Any
    es_qr: bool

    @property
    def fecha(self) -> dt.date | None:
        return None if self.fecha_hora is None else self.fecha_hora.date()

    @property
    def elegible(self) -> bool:
        return self.fecha_hora is not None and self.centavos is not None and self.centavos > 0


def cargar_cochabamba(tabla: TablaDetectada, cfg: Config) -> list[RegistroCBB]:
    cols = tabla.columnas
    registros: list[RegistroCBB] = []
    for i in range(tabla.fila_encabezado + 1, len(tabla.filas)):
        fila = tabla.filas[i]
        crudo = {c: _celda(fila, idx) for c, idx in cols.items()}
        if not _es_fila_datos_cbb(crudo):
            continue

        fecha_hora = redondear_a_segundo(a_datetime(crudo.get("Fecha")))
        centavos = a_centavos(crudo.get("Monto"))
        tipo_pago = norm_texto(crudo.get("Tipo Pago"))
        canal = norm_texto(crudo.get("Canal de Pago"))
        qr_norm = [norm_texto(t) for t in cfg.tipos_pago_qr]
        es_qr = any(t in (tipo_pago, canal) for t in qr_norm) or any(
            _contiene_palabras(tipo_pago, t) or _contiene_palabras(canal, t) for t in qr_norm
        )

        problemas = []
        if fecha_hora is None:
            problemas.append("FECHA_INVALIDA")
        elif fecha_hora.time() == dt.time(0, 0, 0):
            problemas.append("SIN_HORA")
        if centavos is None:
            problemas.append("MONTO_NO_NUMERICO")
        elif centavos <= 0:
            problemas.append("MONTO_NO_POSITIVO")
        if not norm_id(crudo.get("Numero Factura")):
            problemas.append("SIN_NUMERO_FACTURA")
        estado = norm_texto(crudo.get("Estado"))
        if estado and estado != "VALIDO":
            problemas.append(f"ESTADO_{estado}")

        registros.append(RegistroCBB(i + 1, crudo, fecha_hora, centavos, es_qr, problemas))
    return registros


def _es_fila_datos_cbb(crudo: dict[str, Any]) -> bool:
    """Descarta filas vacias, repeticiones del encabezado y el pie de totales."""
    if all(v is None or str(v).strip() == "" for v in crudo.values()):
        return False
    if norm_texto(crudo.get("Fecha")) == "FECHA" and norm_texto(crudo.get("Monto")) == "MONTO":
        return False
    tiene_nro = isinstance(crudo.get("Nro"), (int, float)) and not isinstance(crudo.get("Nro"), bool)
    tiene_factura = bool(norm_id(crudo.get("Numero Factura")))
    return tiene_nro or tiene_factura


def cargar_bcp(tabla: TablaDetectada, cfg: Config) -> list[MovimientoBCP]:
    cols = tabla.columnas
    movimientos: list[MovimientoBCP] = []
    for i in range(tabla.fila_encabezado + 1, len(tabla.filas)):
        fila = tabla.filas[i]
        crudo = {c: _celda(fila, idx) for c, idx in cols.items()}
        if all(v is None or str(v).strip() == "" for v in crudo.values()):
            continue
        if norm_texto(crudo.get("Fecha")) == "FECHA" and norm_texto(crudo.get("Importe")) == "IMPORTE":
            continue

        fecha = a_datetime(crudo.get("Fecha"))
        hora = a_time(crudo.get("Hora"))
        fecha_hora = None
        if fecha is not None:
            fecha_hora = redondear_a_segundo(
                dt.datetime.combine(fecha.date(), hora) if hora is not None else fecha
            )

        glosa = "" if crudo.get("Glosa") is None else str(crudo["Glosa"])
        movimientos.append(
            MovimientoBCP(
                fila_excel=i + 1,
                crudo=crudo,
                fecha_hora=fecha_hora,
                centavos=a_centavos(crudo.get("Importe")),
                nro_oper=crudo.get("Nro Oper"),
                nro_oper_norm=norm_id(crudo.get("Nro Oper")),
                glosa=glosa,
                cd_confirmacion=crudo.get("Cd Confirmacion"),
                es_qr=_contiene_palabras(norm_texto(glosa), norm_texto(cfg.patron_qr_glosa)),
            )
        )
    return movimientos


def _celda(fila: Sequence[Any], idx: int) -> Any:
    return fila[idx] if idx < len(fila) else None


# --------------------------------------------------------------------------- #
# Matching
# --------------------------------------------------------------------------- #

@dataclass
class Resultado:
    registro: RegistroCBB
    movimiento: MovimientoBCP | None
    diferencia_seg: float | None
    estado: str
    candidatos: int
    segundo_delta: float | None
    margen: float | None
    ambiguo: bool
    motivo: str


def emparejar(
    registros: Sequence[RegistroCBB],
    movimientos: Sequence[MovimientoBCP],
    cfg: Config = Config(),
) -> list[Resultado]:
    """Asigna 1 a 1 cada registro QR de Cochabamba con un movimiento BCP.

    Un movimiento BCP nunca se usa dos veces: la asignacion se resuelve con el
    algoritmo hungaro dentro de cada grupo fecha+monto, minimizando la diferencia
    horaria total.
    """
    elegibles = [
        m for m in movimientos
        if m.elegible and (m.es_qr or not cfg.solo_qr_bcp)
    ]
    grupos_bcp: dict[tuple[dt.date, int], list[MovimientoBCP]] = {}
    for mov in elegibles:
        grupos_bcp.setdefault((mov.fecha, mov.centavos), []).append(mov)

    resultados: list[Resultado] = []
    grupos_cbb: dict[tuple[Any, Any], list[RegistroCBB]] = {}
    for reg in registros:
        if not reg.es_qr:
            resultados.append(_resultado_simple(reg, ESTADO_NO_QR, "Tipo de pago no QR"))
            continue
        if not reg.utilizable:
            resultados.append(_resultado_simple(reg, ESTADO_INCOMPLETO, ", ".join(reg.problemas) or "Fecha o monto invalido"))
            continue
        grupos_cbb.setdefault((reg.fecha, reg.centavos), []).append(reg)

    for clave, grupo in grupos_cbb.items():
        resultados.extend(_emparejar_grupo(grupo, grupos_bcp.get(clave, []), cfg))

    resultados.sort(key=lambda r: r.registro.fila_excel)
    return resultados


def _resultado_simple(reg: RegistroCBB, estado: str, motivo: str) -> Resultado:
    return Resultado(reg, None, None, estado, 0, None, None, False, motivo)


def _emparejar_grupo(
    grupo: list[RegistroCBB],
    candidatos: list[MovimientoBCP],
    cfg: Config,
) -> list[Resultado]:
    if not candidatos:
        return [_resultado_simple(r, ESTADO_SIN_MATCH, "Sin movimiento BCP en la misma fecha y monto") for r in grupo]

    deltas = np.array(
        [[abs((r.fecha_hora - m.fecha_hora).total_seconds()) for m in candidatos] for r in grupo],
        dtype=float,
    )
    # Fuera de la ventana maxima el par se prohibe; asi el hungaro maximiza los pares validos.
    costos = np.where(deltas <= cfg.revisar_seg, deltas, COSTO_PROHIBIDO + deltas)
    filas, columnas = linear_sum_assignment(costos)

    asignado_por_cbb: dict[int, int] = {}
    for i, j in zip(filas, columnas):
        if deltas[i, j] <= cfg.revisar_seg:
            asignado_por_cbb[int(i)] = int(j)
    ocupados = set(asignado_por_cbb.values())

    resultados = []
    for i, reg in enumerate(grupo):
        en_ventana = [j for j in range(len(candidatos)) if deltas[i, j] <= cfg.revisar_seg]
        ordenados = sorted(en_ventana, key=lambda j: (deltas[i, j], candidatos[j].fila_excel))
        segundo = float(deltas[i, ordenados[1]]) if len(ordenados) > 1 else None

        j = asignado_por_cbb.get(i)
        if j is None:
            if en_ventana:
                motivo = "Candidatos ya asignados a otro pago"
            else:
                cercano = int(np.argmin(deltas[i]))
                motivo = (
                    f"Ningun movimiento BCP dentro de {cfg.revisar_seg:g} s; el mas cercano "
                    f"con misma fecha y monto esta a {deltas[i, cercano]:.0f} s "
                    f"(Nro Oper. {candidatos[cercano].nro_oper}, no asignado)"
                )
            resultados.append(
                Resultado(reg, None, None, ESTADO_SIN_MATCH, len(en_ventana), segundo, None, False, motivo)
            )
            continue

        delta = float(deltas[i, j])
        margen = None if segundo is None else round(segundo - delta, 3)
        libre_cercano = any(
            k not in ocupados and deltas[i, k] <= delta + cfg.margen_ambiguo_seg
            for k in en_ventana
            if k != j
        )
        estado = _clasificar(delta, cfg)
        motivo = ""
        if libre_cercano and cfg.degradar_ambiguos and estado in (ESTADO_SEGURO, ESTADO_PROBABLE):
            motivo = f"Degradado a REVISAR: otro movimiento BCP libre a <= {cfg.margen_ambiguo_seg:g} s"
            estado = ESTADO_REVISAR
        resultados.append(
            Resultado(reg, candidatos[j], round(delta, 3), estado, len(en_ventana), segundo, margen, libre_cercano, motivo)
        )
    return resultados


def _clasificar(delta: float, cfg: Config) -> str:
    if delta <= cfg.seguro_seg:
        return ESTADO_SEGURO
    if delta <= cfg.probable_seg:
        return ESTADO_PROBABLE
    if delta <= cfg.revisar_seg:
        return ESTADO_REVISAR
    return ESTADO_SIN_MATCH


# --------------------------------------------------------------------------- #
# Validaciones y estadisticas
# --------------------------------------------------------------------------- #

def validar(
    registros: Sequence[RegistroCBB],
    movimientos: Sequence[MovimientoBCP],
    resultados: Sequence[Resultado],
    cfg: Config,
) -> list[tuple[str, str, int]]:
    """Controles de integridad. Devuelve [(nivel, detalle, cantidad)]."""
    alertas: list[tuple[str, str, int]] = []

    def agregar(nivel: str, detalle: str, cantidad: int) -> None:
        if cantidad:
            alertas.append((nivel, detalle, cantidad))

    problemas: dict[str, int] = {}
    for reg in registros:
        for p in reg.problemas:
            problemas[p] = problemas.get(p, 0) + 1
    for nombre, cantidad in sorted(problemas.items()):
        agregar("ALERTA", f"Cochabamba: filas con {nombre}", cantidad)

    duplicados_cbb: dict[tuple, int] = {}
    for reg in registros:
        clave = (reg.fecha_hora, reg.centavos, norm_texto(reg.crudo.get("Nombre Estudiante")))
        duplicados_cbb[clave] = duplicados_cbb.get(clave, 0) + 1
    agregar("ALERTA", "Cochabamba: registros duplicados (misma fecha/hora, monto y estudiante)",
            sum(v - 1 for v in duplicados_cbb.values() if v > 1))

    facturas: dict[str, int] = {}
    for reg in registros:
        f = norm_id(reg.crudo.get("Numero Factura"))
        if f:
            facturas[f] = facturas.get(f, 0) + 1
    agregar("ALERTA", "Cochabamba: numeros de factura repetidos",
            sum(v - 1 for v in facturas.values() if v > 1))

    agregar("INFO", "BCP: movimientos con importe negativo (cargos, excluidos del cruce)",
            sum(1 for m in movimientos if m.centavos is not None and m.centavos < 0))
    agregar("ALERTA", "BCP: movimientos sin fecha u hora valida",
            sum(1 for m in movimientos if m.fecha_hora is None))
    agregar("ALERTA", "BCP: movimientos sin Nro Oper.",
            sum(1 for m in movimientos if not m.nro_oper_norm))
    agregar("ALERTA", "BCP: importes no numericos",
            sum(1 for m in movimientos if m.centavos is None))

    opers: dict[str, int] = {}
    for mov in movimientos:
        if mov.nro_oper_norm:
            opers[mov.nro_oper_norm] = opers.get(mov.nro_oper_norm, 0) + 1
    repetidos = {k: v for k, v in opers.items() if v > 1}
    agregar("ALERTA", f"BCP: Nro Oper. duplicados en el extracto ({_muestra(repetidos)})", len(repetidos))

    usados: dict[int, int] = {}
    for res in resultados:
        if res.movimiento is not None:
            usados[res.movimiento.fila_excel] = usados.get(res.movimiento.fila_excel, 0) + 1
    agregar("CRITICO", "BCP: un mismo movimiento fue asignado a mas de un pago",
            sum(1 for v in usados.values() if v > 1))

    agregar("ALERTA", "Casos con mas de un candidato dentro de la ventana",
            sum(1 for r in resultados if r.candidatos > 1))
    agregar("ALERTA", "Casos ambiguos (otro movimiento BCP libre casi igual de cerca)",
            sum(1 for r in resultados if r.ambiguo))

    fechas_bcp = [m.fecha for m in movimientos if m.fecha is not None]
    if fechas_bcp:
        ini, fin = min(fechas_bcp), max(fechas_bcp)
        fuera = sum(1 for r in registros if r.fecha is not None and not (ini <= r.fecha <= fin))
        agregar("ALERTA", f"Cochabamba: pagos fuera del rango del extracto BCP ({ini} a {fin})", fuera)

    conflictos = sum(1 for r in resultados if _conflicto_ciudad(r, cfg))
    agregar("INFO", "Matches con Cd. Confirmacion de otra ciudad (no afecta el match)", conflictos)
    return alertas


def _muestra(repetidos: dict[str, int], limite: int = 8) -> str:
    if not repetidos:
        return "ninguno"
    items = sorted(repetidos.items(), key=lambda kv: (-kv[1], kv[0]))[:limite]
    texto = ", ".join(f"{k} x{v}" for k, v in items)
    return texto + (" ..." if len(repetidos) > limite else "")


def _conflicto_ciudad(res: Resultado, cfg: Config) -> bool:
    if res.movimiento is None:
        return False
    valor = norm_texto(res.movimiento.cd_confirmacion)
    return bool(valor) and valor not in cfg.ciudad_propia


BUCKETS = (
    ("0 s", 0.0, 0.0),
    ("1 s", 0.001, 1.0),
    ("2 s", 1.001, 2.0),
    ("3-5 s", 2.001, 5.0),
    ("6-10 s", 5.001, 10.0),
    ("11-30 s", 10.001, 30.0),
    ("31-60 s", 30.001, 60.0),
    ("61-120 s", 60.001, 120.0),
    ("> 120 s", 120.001, float("inf")),
)


def distribucion(resultados: Sequence[Resultado]) -> tuple[dict[str, int], dict[str, float]]:
    deltas = [r.diferencia_seg for r in resultados if r.diferencia_seg is not None]
    conteo = {nombre: 0 for nombre, _, _ in BUCKETS}
    for d in deltas:
        for nombre, bajo, alto in BUCKETS:
            if bajo <= d <= alto:
                conteo[nombre] += 1
                break
    if not deltas:
        return conteo, {}
    ordenados = sorted(deltas)
    resumen = {
        "minimo": ordenados[0],
        "mediana": statistics.median(ordenados),
        "p90": _percentil(ordenados, 90),
        "p95": _percentil(ordenados, 95),
        "maximo": ordenados[-1],
    }
    return conteo, resumen


def _percentil(ordenados: list[float], p: float) -> float:
    if not ordenados:
        return 0.0
    k = (len(ordenados) - 1) * p / 100.0
    bajo, alto = int(k), min(int(k) + 1, len(ordenados) - 1)
    return ordenados[bajo] + (ordenados[alto] - ordenados[bajo]) * (k - bajo)


# --------------------------------------------------------------------------- #
# Salida
# --------------------------------------------------------------------------- #

COLUMNAS_SALIDA = [
    "Nro", "FechaHora_CBB", "Número Factura", "Nit/C.I.", "Razon Social",
    "Nombre Estudiante", "Tipo Pago", "Monto", "Canal de Pago", "Estado",
    "BCP_Fecha", "BCP_Hora", "BCP_FechaHora", "BCP_Importe", "BCP_Glosa",
    "BCP_Nro_Oper", "BCP_Cd_Confirmacion", "Diferencia_segundos", "Estado_match",
    "Cantidad_candidatos", "Segundo_mejor_delta_seg", "Margen_vs_segundo_candidato",
    "Ambiguo", "Conflicto_Ciudad", "Motivo", "Fila_Excel_CBB", "Fila_Excel_BCP",
]


def a_filas(resultados: Sequence[Resultado], cfg: Config) -> list[dict[str, Any]]:
    filas = []
    for res in resultados:
        reg, mov = res.registro, res.movimiento
        fila = {
            "Nro": reg.crudo.get("Nro"),
            "FechaHora_CBB": reg.fecha_hora,
            "Número Factura": reg.crudo.get("Numero Factura"),
            "Nit/C.I.": reg.crudo.get("Nit/C.I."),
            "Razon Social": reg.crudo.get("Razon Social"),
            "Nombre Estudiante": reg.crudo.get("Nombre Estudiante"),
            "Tipo Pago": reg.crudo.get("Tipo Pago"),
            "Monto": centavos_a_float(reg.centavos),
            "Canal de Pago": reg.crudo.get("Canal de Pago"),
            "Estado": reg.crudo.get("Estado"),
            "BCP_Fecha": mov.fecha if mov else None,
            "BCP_Hora": mov.fecha_hora.time() if mov and mov.fecha_hora else None,
            "BCP_FechaHora": mov.fecha_hora if mov else None,
            "BCP_Importe": centavos_a_float(mov.centavos) if mov else None,
            "BCP_Glosa": mov.glosa if mov else None,
            "BCP_Nro_Oper": mov.nro_oper if mov else None,
            "BCP_Cd_Confirmacion": mov.cd_confirmacion if mov else None,
            "Diferencia_segundos": res.diferencia_seg,
            "Estado_match": res.estado,
            "Cantidad_candidatos": res.candidatos,
            "Segundo_mejor_delta_seg": res.segundo_delta,
            "Margen_vs_segundo_candidato": res.margen,
            "Ambiguo": "SI" if res.ambiguo else "",
            "Conflicto_Ciudad": "SI" if _conflicto_ciudad(res, cfg) else "",
            "Motivo": res.motivo,
            "Fila_Excel_CBB": reg.fila_excel,
            "Fila_Excel_BCP": mov.fila_excel if mov else None,
        }
        filas.append(fila)
    return filas


def construir_resumen(
    registros: Sequence[RegistroCBB],
    movimientos: Sequence[MovimientoBCP],
    resultados: Sequence[Resultado],
    alertas: Sequence[tuple[str, str, int]],
    cfg: Config,
) -> list[tuple[str, Any]]:
    qr = [r for r in resultados if r.estado != ESTADO_NO_QR]
    por_estado = {e: [r for r in qr if r.estado == e] for e in
                  (ESTADO_SEGURO, ESTADO_PROBABLE, ESTADO_REVISAR, ESTADO_SIN_MATCH, ESTADO_INCOMPLETO)}
    total_qr = len(qr)
    encontrados = len(por_estado[ESTADO_SEGURO]) + len(por_estado[ESTADO_PROBABLE]) + len(por_estado[ESTADO_REVISAR])

    def bs(items: Sequence[Resultado]) -> float:
        return round(sum(centavos_a_float(r.registro.centavos) or 0.0 for r in items), 2)

    conteo, stats = distribucion(resultados)
    fechas = [r.fecha for r in registros if r.fecha is not None]
    fechas_b = [m.fecha for m in movimientos if m.fecha is not None]

    filas: list[tuple[str, Any]] = [
        ("PARAMETROS", ""),
        ("Umbral MATCH_SEGURO (s)", cfg.seguro_seg),
        ("Umbral MATCH_PROBABLE (s)", cfg.probable_seg),
        ("Umbral REVISAR (s)", cfg.revisar_seg),
        ("Margen de ambiguedad (s)", cfg.margen_ambiguo_seg),
        ("Solo glosa QR en BCP", "SI" if cfg.solo_qr_bcp else "NO"),
        ("", ""),
        ("UNIVERSO", ""),
        ("Registros Cochabamba totales", len(registros)),
        ("QR Cochabamba", total_qr),
        ("No QR", len(registros) - total_qr),
        ("Rango fechas Cochabamba", f"{min(fechas)} a {max(fechas)}" if fechas else "-"),
        ("Movimientos BCP totales", len(movimientos)),
        ("QR positivos BCP", sum(1 for m in movimientos if m.elegible and m.es_qr)),
        ("Rango fechas BCP", f"{min(fechas_b)} a {max(fechas_b)}" if fechas_b else "-"),
        ("", ""),
        ("RESULTADO DEL CRUCE", ""),
        ("MATCH_SEGURO", len(por_estado[ESTADO_SEGURO])),
        ("MATCH_PROBABLE", len(por_estado[ESTADO_PROBABLE])),
        ("REVISAR", len(por_estado[ESTADO_REVISAR])),
        ("SIN_MATCH", len(por_estado[ESTADO_SIN_MATCH])),
        ("DATOS_INCOMPLETOS", len(por_estado[ESTADO_INCOMPLETO])),
        ("Porcentaje de Cochabamba encontrado", f"{(encontrados / total_qr * 100):.2f}%" if total_qr else "-"),
        ("", ""),
        ("IMPORTES (Bs)", ""),
        ("Total Bs Cochabamba QR", bs(qr)),
        ("Total Bs match seguro", bs(por_estado[ESTADO_SEGURO])),
        ("Total Bs probable", bs(por_estado[ESTADO_PROBABLE])),
        ("Total Bs revisar", bs(por_estado[ESTADO_REVISAR])),
        ("Total Bs sin match", bs(por_estado[ESTADO_SIN_MATCH] + por_estado[ESTADO_INCOMPLETO])),
        ("", ""),
        ("DISTRIBUCION DE DIFERENCIA DE SEGUNDOS", ""),
    ]
    filas.extend((nombre, cantidad) for nombre, cantidad in conteo.items())
    if stats:
        filas.append(("", ""))
        filas.extend((f"Diferencia {k}", round(v, 3)) for k, v in stats.items())
    filas.append(("", ""))
    filas.append(("VALIDACIONES", ""))
    if alertas:
        filas.extend((f"[{nivel}] {detalle}", cantidad) for nivel, detalle, cantidad in alertas)
    else:
        filas.append(("Sin observaciones", 0))
    return filas


def exportar(
    resultados: Sequence[Resultado],
    registros: Sequence[RegistroCBB],
    movimientos: Sequence[MovimientoBCP],
    alertas: Sequence[tuple[str, str, int]],
    cfg: Config,
    salida: str,
) -> str:
    import pandas as pd

    filas = a_filas(resultados, cfg)
    todos = pd.DataFrame(filas, columns=COLUMNAS_SALIDA)
    qr = todos[todos["Estado_match"] != ESTADO_NO_QR]
    resumen = pd.DataFrame(
        construir_resumen(registros, movimientos, resultados, alertas, cfg),
        columns=["Concepto", "Valor"],
    )

    hojas = {
        "RESUMEN": resumen,
        "COCHABAMBA_MATCH": qr,
        "MATCH_SEGURO": qr[qr["Estado_match"] == ESTADO_SEGURO],
        "MATCH_PROBABLE": qr[qr["Estado_match"] == ESTADO_PROBABLE],
        "REVISAR": qr[qr["Estado_match"] == ESTADO_REVISAR],
        "SIN_MATCH": qr[qr["Estado_match"].isin([ESTADO_SIN_MATCH, ESTADO_INCOMPLETO])],
        "NO_QR": todos[todos["Estado_match"] == ESTADO_NO_QR],
        "TODOS": todos,
    }
    formatos = {"FechaHora_CBB": "DD/MM/YYYY HH:MM:SS", "BCP_Fecha": "DD/MM/YYYY",
                "BCP_FechaHora": "DD/MM/YYYY HH:MM:SS", "Monto": "#,##0.00", "BCP_Importe": "#,##0.00"}
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        for nombre, df in hojas.items():
            df.to_excel(writer, sheet_name=nombre, index=False)
            _formatear(writer.sheets[nombre], formatos)
    return salida


def _formatear(hoja, formatos: dict[str, str]) -> None:
    columnas = {celda.value: celda.column for celda in hoja[1]}
    for nombre, formato in formatos.items():
        if nombre not in columnas:
            continue
        for (celda,) in hoja.iter_rows(min_row=2, min_col=columnas[nombre], max_col=columnas[nombre]):
            celda.number_format = formato


# --------------------------------------------------------------------------- #
# Ejecucion
# --------------------------------------------------------------------------- #

def ejecutar(rutas: Sequence[str], salida: str = NOMBRE_SALIDA, cfg: Config = Config()) -> dict[str, Any]:
    tabla_cbb, tabla_bcp = identificar_tablas(rutas)
    registros = cargar_cochabamba(tabla_cbb, cfg)
    movimientos = cargar_bcp(tabla_bcp, cfg)
    resultados = emparejar(registros, movimientos, cfg)
    alertas = validar(registros, movimientos, resultados, cfg)
    ruta = exportar(resultados, registros, movimientos, alertas, cfg, salida)
    return {
        "tabla_cbb": tabla_cbb,
        "tabla_bcp": tabla_bcp,
        "registros": registros,
        "movimientos": movimientos,
        "resultados": resultados,
        "alertas": alertas,
        "salida": ruta,
    }


def imprimir_reporte(datos: dict[str, Any], cfg: Config) -> None:
    tabla_cbb: TablaDetectada = datos["tabla_cbb"]
    tabla_bcp: TablaDetectada = datos["tabla_bcp"]
    print(f"Cochabamba: {os.path.basename(tabla_cbb.ruta)} | hoja '{tabla_cbb.hoja}' | "
          f"encabezado fila {tabla_cbb.fila_encabezado + 1} | {len(datos['registros'])} registros")
    print(f"BCP:        {os.path.basename(tabla_bcp.ruta)} | hoja '{tabla_bcp.hoja}' | "
          f"encabezado fila {tabla_bcp.fila_encabezado + 1} | {len(datos['movimientos'])} movimientos")
    print()
    for concepto, valor in construir_resumen(
        datos["registros"], datos["movimientos"], datos["resultados"], datos["alertas"], cfg
    ):
        if concepto == "" and valor == "":
            print()
        elif valor == "":
            print(f"--- {concepto} ---")
        else:
            print(f"{concepto}: {valor}")
    print(f"\nArchivo generado: {datos['salida']}")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Cruza el reporte de ingresos de Cochabamba contra el extracto BCP nacional.",
        epilog="Los archivos de entrada nunca se modifican.",
    )
    parser.add_argument("archivos", nargs="*", help="Los dos archivos, en cualquier orden.")
    parser.add_argument("--cbb", help="Reporte de ingresos Cochabamba (.xls/.xlsx)")
    parser.add_argument("--bcp", help="Extracto BCP (.xlsx)")
    parser.add_argument("--salida", default=NOMBRE_SALIDA, help=f"Excel de resultado (por defecto {NOMBRE_SALIDA})")
    parser.add_argument("--seguro", type=float, default=Config.seguro_seg, help="Umbral MATCH_SEGURO en segundos")
    parser.add_argument("--probable", type=float, default=Config.probable_seg, help="Umbral MATCH_PROBABLE en segundos")
    parser.add_argument("--revisar", type=float, default=Config.revisar_seg, help="Umbral REVISAR en segundos")
    parser.add_argument("--margen-ambiguo", type=float, default=Config.margen_ambiguo_seg,
                        help="Distancia a la que un candidato libre vuelve ambiguo el match")
    parser.add_argument("--sin-degradar-ambiguos", action="store_true",
                        help="No bajar a REVISAR los matches ambiguos")
    parser.add_argument("--sin-filtro-qr", action="store_true",
                        help="No exigir 'QR' en la glosa del BCP")
    args = parser.parse_args(argv)

    rutas = [r for r in (args.cbb, args.bcp) if r] + list(args.archivos)
    if len(rutas) != 2:
        parser.error("Debes indicar exactamente 2 archivos: --cbb X --bcp Y, o los dos como argumentos.")
    faltantes = [r for r in rutas if not os.path.isfile(r)]
    if faltantes:
        parser.error(f"No existe: {', '.join(faltantes)}")

    cfg = Config(
        seguro_seg=args.seguro,
        probable_seg=args.probable,
        revisar_seg=args.revisar,
        margen_ambiguo_seg=args.margen_ambiguo,
        degradar_ambiguos=not args.sin_degradar_ambiguos,
        solo_qr_bcp=not args.sin_filtro_qr,
    )
    try:
        datos = ejecutar(rutas, args.salida, cfg)
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    imprimir_reporte(datos, cfg)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
