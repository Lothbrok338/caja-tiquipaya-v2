"""
test_control_asignaciones.py — CONTROL 1: auditoría de asignaciones (ZUONR)
duplicadas/históricas del SAP GLOBAL mensual, con validación humana en
Excel (REVISION_ASIGNACIONES_<PERIODO>.xlsx, una fila por ocurrencia) y
corrección autorizada del propio GLOBAL.

Usa un SAP GLOBAL sintético construido localmente en este archivo (nunca
datos contables reales). Verifica exclusivamente el comportamiento de
control_asignaciones.py: no repite ninguna regla de excel_io.py/
motor_tiquipaya.py/sap_writer.py/pipeline_tiquipaya.py/run_batch.py/
consolidador_mensual.py.

Uso: python -m unittest tests.test_control_asignaciones -v
"""

import csv
import datetime
import json
import os
import shutil
import sys
import tempfile
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import control_asignaciones as ca


# ---------------------------------------------------------------------------
# Fixtures: SAP GLOBAL sintético mínimo (hoja "1", partidas desde fila 16).
# ---------------------------------------------------------------------------

def _crear_global(ruta, partidas, hoja="1"):
    """partidas: lista de dicts con asignacion/cuenta_mayor/cargo/haber/
    glosa/fecha_valor (ver _partida)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = hoja

    fila = 16
    for p in partidas:
        ws[f"C{fila}"] = p.get("cuenta_mayor", "110101001")
        ws[f"D{fila}"] = p.get("glosa")
        cargo = p.get("cargo")
        haber = p.get("haber")
        if cargo is not None:
            ws[f"E{fila}"] = Decimal(str(cargo))
        if haber is not None:
            ws[f"F{fila}"] = Decimal(str(haber))
        fecha_valor = p.get("fecha_valor")
        if fecha_valor is not None:
            ws[f"O{fila}"] = fecha_valor
        ws[f"R{fila}"] = p.get("asignacion")
        fila += 1

    wb.save(ruta)


def _partida(asignacion, cuenta_mayor="110101001", cargo="100.00", haber=None,
             glosa="RECAUDACION CAJA TIQUIPAYA", fecha_valor=None):
    return {
        "asignacion": asignacion,
        "cuenta_mayor": cuenta_mayor,
        "cargo": cargo,
        "haber": haber,
        "glosa": glosa,
        "fecha_valor": fecha_valor or datetime.date(2026, 8, 5),
    }


def _leer_csv(ruta):
    if not os.path.isfile(ruta):
        return []
    with open(ruta, "r", encoding="utf-8", newline="") as f:
        return [dict(fila) for fila in csv.DictReader(f)]


def _celda_global(ruta_global, columna, fila, hoja="1"):
    wb = openpyxl.load_workbook(ruta_global)
    try:
        return wb[hoja][f"{columna}{fila}"].value
    finally:
        wb.close()


def _marcar_decisiones(ruta_xlsx, decisiones):
    """decisiones: {fila_global: {"VALIDACION_AUDITOR":..., "ASIGNACION_CORRECTA":...,
    "OBSERVACION_AUDITOR":..., "FECHA_VALIDACION":...}} — simula al auditor
    editando el Excel directamente."""
    wb = openpyxl.load_workbook(ruta_xlsx)
    ws = wb[ca._HOJA_REVISION]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    for row in ws.iter_rows(min_row=2):
        fila_global = row[idx["FILA_GLOBAL"]].value
        decision = decisiones.get(fila_global)
        if not decision:
            continue
        for campo, valor in decision.items():
            row[idx[campo]].value = valor
    wb.save(ruta_xlsx)


class _ControlAsignacionesTestBase(unittest.TestCase):
    NOMBRE_GLOBAL_AGOSTO = "SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx"
    NOMBRE_GLOBAL_SEPTIEMBRE = "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx"

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="control_asignaciones_test_")
        self.ruta_historico = os.path.join(self.tmpdir, "HISTORICO_ASIGNACIONES.csv")

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _ruta_global(self, nombre=NOMBRE_GLOBAL_AGOSTO):
        return os.path.join(self.tmpdir, nombre)

    def _ruta_xlsx(self, periodo="AGOSTO_2026"):
        return ca.ruta_revision_xlsx(self.tmpdir, periodo)

    def _ruta_csv_legado(self, periodo="AGOSTO_2026"):
        return ca.ruta_revision_asignaciones(self.tmpdir, periodo)

    def _ejecutar(self, partidas, ruta_global=None, **kwargs):
        ruta_global = ruta_global or self._ruta_global()
        _crear_global(ruta_global, partidas)
        return ca.ejecutar_control(
            ruta_global=ruta_global, ruta_historico=self.ruta_historico, **kwargs
        )


# ---------------------------------------------------------------------------
# Exclusiones — SIN CAMBIOS respecto de la versión anterior (test 25)
# ---------------------------------------------------------------------------

class TestExclusiones(_ControlAsignacionesTestBase):
    def test_sfc101_repetido_excluido(self):
        resumen = self._ejecutar([_partida("SFC101"), _partida("SFC101")])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertEqual(resumen["asignaciones_excluidas"], 2)
        self.assertEqual(resumen["cantidad_hallazgos"], 0)

    def test_sfc102_repetido_excluido(self):
        resumen = self._ejecutar([_partida("SFC102"), _partida("SFC102")])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertEqual(resumen["asignaciones_excluidas"], 2)

    def test_tiquipaya_ago_repetido_excluido(self):
        resumen = self._ejecutar([_partida("TIQUIPAYA AGO"), _partida("TIQUIPAYA AGO")])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertEqual(resumen["asignaciones_excluidas"], 2)

    def test_cuenta_110201008_con_tiquipaya_ago_excluido(self):
        resumen = self._ejecutar([_partida("TIQUIPAYA AGO", cuenta_mayor="110201008")])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertEqual(resumen["asignaciones_excluidas"], 1)
        self.assertEqual(resumen["asignaciones_evaluadas"], 0)

    def test_cuenta_110201008_con_revisar_excluido(self):
        resumen = self._ejecutar([_partida("REVISAR", cuenta_mayor="110201008")])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertEqual(resumen["asignaciones_excluidas"], 1)
        self.assertEqual(resumen["asignaciones_evaluadas"], 0)

    def test_cuenta_110201008_con_cualquier_otra_asignacion_excluido(self):
        resumen = self._ejecutar([
            _partida("OTRA-CUALQUIERA-1", cuenta_mayor="110201008"),
            _partida("OTRA-CUALQUIERA-1", cuenta_mayor="110201008"),
        ])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertEqual(resumen["asignaciones_excluidas"], 2)
        self.assertEqual(resumen["cantidad_hallazgos"], 0)

    def test_cuenta_110201008_normaliza_float_excel(self):
        self.assertTrue(ca.es_excluida("FORTALEZA", 110201008.0))
        self.assertTrue(ca.es_excluida("FORTALEZA", "110201008.0"))
        self.assertTrue(ca.es_excluida("FORTALEZA", "110201008"))
        self.assertFalse(ca.es_excluida("FORTALEZA", "110101001"))

    def test_fortaleza_siempre_se_evalua(self):
        resumen = self._ejecutar([_partida("FORTALEZA")])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertEqual(resumen["asignaciones_evaluadas"], 1)
        self.assertEqual(resumen["asignaciones_excluidas"], 0)


# ---------------------------------------------------------------------------
# Hoja EXACTA "1" y periodo canónico — SIN CAMBIOS
# ---------------------------------------------------------------------------

class TestHojaExacta(_ControlAsignacionesTestBase):
    def test_global_sin_hoja_1_error_tecnico(self):
        ruta_global = self._ruta_global()
        _crear_global(ruta_global, [_partida("ASIG-1")], hoja="OTRA_HOJA")
        resumen = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_HOJA_1_NO_ENCONTRADA", resumen["problemas"])

    def test_nunca_hace_fallback_a_otra_hoja(self):
        ruta_global = self._ruta_global()
        _crear_global(ruta_global, [_partida("ASIG-1"), _partida("ASIG-1")], hoja="Hoja1")
        resumen = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_HOJA_1_NO_ENCONTRADA", resumen["problemas"])
        self.assertNotIn("filas_leidas_global", resumen)


class TestPeriodoCanonico(_ControlAsignacionesTestBase):
    def test_nombre_no_canonico_detiene_el_control(self):
        ruta_global = self._ruta_global("GLOBAL_AGOSTO_2026.xlsx")
        _crear_global(ruta_global, [_partida("ASIG-1"), _partida("ASIG-1")])
        resumen = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NOMBRE_NO_CANONICO", resumen["problemas"])
        self.assertFalse(os.path.isfile(self.ruta_historico))
        revisiones = [f for f in os.listdir(self.tmpdir) if f.startswith("REVISION_ASIGNACIONES_")]
        self.assertEqual(revisiones, [])

    def test_nombre_no_canonico_sin_anio_detiene_el_control(self):
        ruta_global = self._ruta_global("SAP_GLOBAL_TIQ_AGOSTO.xlsx")
        _crear_global(ruta_global, [_partida("ASIG-1")])
        resumen = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NOMBRE_NO_CANONICO", resumen["problemas"])

    def test_nombre_canonico_deriva_periodo_correctamente(self):
        resumen = self._ejecutar([_partida("ASIG-1")])
        self.assertEqual(resumen["periodo"], "AGOSTO_2026")


# ---------------------------------------------------------------------------
# 1-4: se genera el XLSX, una fila por ocurrencia, FILA_GLOBAL correcta,
# y la lista desplegable CORRECTA/INCORRECTA existe.
# ---------------------------------------------------------------------------

class TestGeneracionXlsx(_ControlAsignacionesTestBase):
    def test_1_genera_xlsx_de_revision(self):
        resumen = self._ejecutar([_partida("ASIG-DUP"), _partida("ASIG-DUP")])
        ruta = resumen["ruta_revision"]
        self.assertTrue(ruta.endswith(".xlsx"))
        self.assertTrue(os.path.isfile(ruta))
        wb = openpyxl.load_workbook(ruta)
        self.assertEqual(wb.sheetnames, ["REVISION"])

    def test_2_una_fila_por_cada_ocurrencia_duplicada(self):
        # 3 ocurrencias de la misma asignación -> 3 filas (no 1 agrupada).
        resumen = self._ejecutar([
            _partida("3P66536982"), _partida("3P66536982"), _partida("3P66536982"),
        ])
        filas = ca.cargar_revision_xlsx(resumen["ruta_revision"])
        self.assertEqual(len(filas), 3)
        self.assertEqual(sorted(f["FILA_GLOBAL"] for f in filas), [16, 17, 18])

    def test_3_fila_global_coincide_con_global_real(self):
        resumen = self._ejecutar([
            _partida("ASIG-A"),
            _partida("ASIG-B"),
            _partida("3P66536982"), _partida("3P66536982"),
        ])
        filas = ca.cargar_revision_xlsx(resumen["ruta_revision"])
        # Las ocurrencias duplicadas están en las filas 18 y 19 del GLOBAL
        # (16=ASIG-A, 17=ASIG-B, 18/19=3P66536982). Python lo calcula solo.
        self.assertEqual(sorted(f["FILA_GLOBAL"] for f in filas), [18, 19])
        for f in filas:
            self.assertEqual(f["ASIGNACION_ORIGINAL"], "3P66536982")

    def test_4_lista_desplegable_correcta_incorrecta_existe(self):
        resumen = self._ejecutar([_partida("ASIG-DUP"), _partida("ASIG-DUP")])
        wb = openpyxl.load_workbook(resumen["ruta_revision"])
        ws = wb[ca._HOJA_REVISION]
        validaciones = list(ws.data_validations.dataValidation)
        self.assertEqual(len(validaciones), 1)
        dv = validaciones[0]
        self.assertEqual(dv.type, "list")
        self.assertIn("CORRECTA", dv.formula1)
        self.assertIn("INCORRECTA", dv.formula1)


# ---------------------------------------------------------------------------
# 5-11: cierre, CORRECTA/INCORRECTA, verificación de que el resto del
# GLOBAL no cambia, y que una fila pendiente bloquea TODO.
# ---------------------------------------------------------------------------

class TestCierreYCorreccion(_ControlAsignacionesTestBase):
    def test_5_decisiones_humanas_sobreviven_reejecucion(self):
        resumen1 = self._ejecutar([_partida("ASIG-A"), _partida("ASIG-A")])
        _marcar_decisiones(resumen1["ruta_revision"], {
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "ASIG-A2",
                 "OBSERVACION_AUDITOR": "Ajuste de digitación"},
        })
        ruta_global = self._ruta_global()
        # Reejecutar SIN resolver la fila 16 -> sigue pendiente, pero la
        # decisión de la fila 17 debe seguir presente en el xlsx regenerado.
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen2["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        filas = ca.cargar_revision_xlsx(resumen2["ruta_revision"])
        fila17 = next(f for f in filas if f["FILA_GLOBAL"] == 17)
        self.assertEqual(fila17["VALIDACION_AUDITOR"], "INCORRECTA")
        self.assertEqual(fila17["ASIGNACION_CORRECTA"], "ASIG-A2")
        self.assertEqual(fila17["OBSERVACION_AUDITOR"], "Ajuste de digitación")

    def test_6_correcta_no_modifica_global(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("FORTALEZA"), _partida("FORTALEZA")], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "CORRECTA"},
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen2["estado_validacion"], "CERRADO_CON_VALIDACION_AUDITOR")
        self.assertFalse(resumen2["global_modificado"])
        self.assertEqual(resumen2["correcciones_aplicadas"], 0)
        self.assertEqual(_celda_global(ruta_global, "R", 16), "FORTALEZA")
        self.assertEqual(_celda_global(ruta_global, "R", 17), "FORTALEZA")

    def test_7_incorrecta_sin_asignacion_correcta_queda_pendiente(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("ASIG-A"), _partida("ASIG-A")], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            17: {"VALIDACION_AUDITOR": "INCORRECTA"},  # sin ASIGNACION_CORRECTA
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen2["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        self.assertFalse(resumen2["historico_actualizado"])
        self.assertEqual(_celda_global(ruta_global, "R", 17), "ASIG-A")  # sin tocar

    def test_8_incorrecta_con_asignacion_correcta_modifica_solo_r_fila(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("ASIG-A"), _partida("ASIG-A")], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "ASIG-A2"},
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen2["estado_validacion"], "CERRADO_CON_VALIDACION_AUDITOR")
        self.assertTrue(resumen2["global_modificado"])
        self.assertEqual(resumen2["correcciones_aplicadas"], 1)
        self.assertEqual(_celda_global(ruta_global, "R", 16), "ASIG-A")
        self.assertEqual(_celda_global(ruta_global, "R", 17), "ASIG-A2")

    def test_9_resto_del_global_permanece_identico(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([
            _partida("ASIG-A", cuenta_mayor="110101005", glosa="GLOSA UNO",
                     cargo="123.45", fecha_valor=datetime.date(2026, 8, 1)),
            _partida("ASIG-A", cuenta_mayor="110101006", glosa="GLOSA DOS",
                     cargo="678.90", fecha_valor=datetime.date(2026, 8, 2)),
        ], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "ASIG-A2"},
        })
        ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)

        wb = openpyxl.load_workbook(ruta_global)
        ws = wb["1"]
        self.assertEqual(ws["C16"].value, "110101005")
        self.assertEqual(ws["D16"].value, "GLOSA UNO")
        self.assertEqual(float(ws["E16"].value), 123.45)
        self.assertEqual(ws["C17"].value, "110101006")
        self.assertEqual(ws["D17"].value, "GLOSA DOS")
        self.assertEqual(float(ws["E17"].value), 678.90)
        # Solo R cambió en la fila 17; el resto de columnas, intacto.
        self.assertEqual(ws["R16"].value, "ASIG-A")
        self.assertEqual(ws["R17"].value, "ASIG-A2")
        # No aparecieron filas nuevas ni se corrió la hoja.
        self.assertIsNone(ws["C18"].value)

    def test_10_varias_correcciones_se_aplican_correctamente(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([
            _partida("A1"), _partida("A1"),
            _partida("B1"), _partida("B1"),
            _partida("C1"), _partida("C1"),
        ], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "A2"},
            17: {"VALIDACION_AUDITOR": "CORRECTA"},
            18: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "B2"},
            19: {"VALIDACION_AUDITOR": "CORRECTA"},
            20: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "C2"},
            21: {"VALIDACION_AUDITOR": "CORRECTA"},
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen2["estado_validacion"], "CERRADO_CON_VALIDACION_AUDITOR")
        self.assertEqual(resumen2["correcciones_aplicadas"], 3)
        self.assertEqual(_celda_global(ruta_global, "R", 16), "A2")
        self.assertEqual(_celda_global(ruta_global, "R", 17), "A1")
        self.assertEqual(_celda_global(ruta_global, "R", 18), "B2")
        self.assertEqual(_celda_global(ruta_global, "R", 19), "B1")
        self.assertEqual(_celda_global(ruta_global, "R", 20), "C2")
        self.assertEqual(_celda_global(ruta_global, "R", 21), "C1")

    def test_11_una_fila_pendiente_impide_todas_las_modificaciones(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([
            _partida("A1"), _partida("A1"),
            _partida("B1"), _partida("B1"),
        ], ruta_global=ruta_global)
        # A1 resuelta, B1 (fila 19) queda sin validar.
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "A2"},
            18: {"VALIDACION_AUDITOR": "CORRECTA"},
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen2["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        self.assertFalse(resumen2["historico_actualizado"])
        self.assertEqual(resumen2["correcciones_aplicadas"], 0)
        # NINGUNA corrección se aplicó, ni siquiera la de A1 que sí estaba resuelta.
        self.assertEqual(_celda_global(ruta_global, "R", 17), "A1")
        self.assertFalse(os.path.isfile(self.ruta_historico))


# ---------------------------------------------------------------------------
# 12-13: guardarraíles de aplicación (SHA distinto / ASIGNACION_ORIGINAL
# distinta impiden corrección).
# ---------------------------------------------------------------------------

class TestGuardarrailesAplicacion(_ControlAsignacionesTestBase):
    def test_12_sha_distinto_impide_correcciones(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("A1"), _partida("A1")], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "A2"},
        })
        # El GLOBAL cambia de contenido (sin pasar por CONTROL 1) antes del cierre.
        _crear_global(ruta_global, [_partida("A1"), _partida("A1"), _partida("A1")])

        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        # Las decisiones del xlsx (atadas al SHA anterior) no se reutilizan.
        self.assertEqual(resumen2["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        self.assertEqual(_celda_global(ruta_global, "R", 17), "A1")
        self.assertFalse(os.path.isfile(self.ruta_historico))

    def test_13_asignacion_original_distinta_impide_correccion(self):
        ruta_global = self._ruta_global()
        _crear_global(ruta_global, [_partida("A1"), _partida("B1")])
        with self.assertRaises(ca.CorreccionInvalidaError):
            ca.aplicar_correcciones_global(ruta_global, [(16, "VALOR_QUE_NO_ESTA_EN_LA_CELDA", "A2")])
        # Ninguna corrección se aplicó (ni siquiera de otras filas del lote).
        self.assertEqual(_celda_global(ruta_global, "R", 16), "A1")
        self.assertEqual(_celda_global(ruta_global, "R", 17), "B1")

    def test_13b_lote_con_una_fila_invalida_no_aplica_ninguna(self):
        ruta_global = self._ruta_global()
        _crear_global(ruta_global, [_partida("A1"), _partida("B1")])
        correcciones = [(16, "A1", "A2"), (17, "VALOR_INCORRECTO", "B2")]
        with self.assertRaises(ca.CorreccionInvalidaError):
            ca.aplicar_correcciones_global(ruta_global, correcciones)
        self.assertEqual(_celda_global(ruta_global, "R", 16), "A1")  # tampoco se aplicó esta
        self.assertEqual(_celda_global(ruta_global, "R", 17), "B1")


# ---------------------------------------------------------------------------
# 14-15: una corrección que genera una duplicidad nueva bloquea el cierre
# y preserva todas las decisiones anteriores.
# ---------------------------------------------------------------------------

class TestNuevaDuplicidadPorCorreccion(_ControlAsignacionesTestBase):
    def test_14_correccion_que_genera_nueva_duplicidad_bloquea_cierre(self):
        ruta_global = self._ruta_global()
        # fila16=AAA(dup con 17), fila17=AAA, fila18=BBB (única).
        resumen1 = self._ejecutar([
            _partida("AAA111"), _partida("AAA111"), _partida("BBB222"),
        ], ruta_global=ruta_global)
        self.assertEqual(len(ca.cargar_revision_xlsx(resumen1["ruta_revision"])), 2)

        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "BBB222"},  # colisiona con 18
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)

        self.assertEqual(resumen2["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        self.assertFalse(resumen2["global_modificado"])
        self.assertFalse(os.path.isfile(self.ruta_historico))
        self.assertEqual(_celda_global(ruta_global, "R", 17), "AAA111")  # no se aplicó

    def test_15_decisiones_previas_se_conservan_al_aparecer_nueva_alerta(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([
            _partida("AAA111"), _partida("AAA111"), _partida("BBB222"),
        ], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA", "OBSERVACION_AUDITOR": "ya revisado"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "BBB222"},
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        filas = ca.cargar_revision_xlsx(resumen2["ruta_revision"])
        self.assertEqual(len(filas), 3)  # se agregó la fila 18, nunca se perdieron las otras dos
        fila16 = next(f for f in filas if f["FILA_GLOBAL"] == 16)
        fila17 = next(f for f in filas if f["FILA_GLOBAL"] == 17)
        fila18 = next(f for f in filas if f["FILA_GLOBAL"] == 18)
        self.assertEqual(fila16["VALIDACION_AUDITOR"], "CORRECTA")
        self.assertEqual(fila16["OBSERVACION_AUDITOR"], "ya revisado")
        self.assertEqual(fila17["VALIDACION_AUDITOR"], "INCORRECTA")
        self.assertEqual(fila17["ASIGNACION_CORRECTA"], "BBB222")
        self.assertEqual(fila18["VALIDACION_AUDITOR"], "")  # nueva, aún sin validar


# ---------------------------------------------------------------------------
# 16-20: histórico con asignación final/original, SHA original/final, y
# comparación del mes siguiente contra la asignación final.
# ---------------------------------------------------------------------------

class TestHistoricoYShaFinal(_ControlAsignacionesTestBase):
    def test_16_17_20_historico_asignacion_final_original_y_sha(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("A1"), _partida("A1")], ruta_global=ruta_global)
        sha_original_esperado = resumen1["sha256_archivo"]
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "A2"},
        })
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)

        self.assertEqual(resumen2["sha256_global_original"], sha_original_esperado)
        self.assertNotEqual(resumen2["sha256_global_final"], sha_original_esperado)
        self.assertEqual(resumen2["sha256_global_final"], ca._hash_archivo(ruta_global))

        historico = _leer_csv(self.ruta_historico)
        fila16 = next(f for f in historico if f["fila_global"] == "16")
        fila17 = next(f for f in historico if f["fila_global"] == "17")
        # test 16: asignacion = FINAL.
        self.assertEqual(fila16["asignacion"], "A1")
        self.assertEqual(fila17["asignacion"], "A2")
        # test 17: se conserva el original.
        self.assertEqual(fila16["asignacion_original"], "A1")
        self.assertEqual(fila17["asignacion_original"], "A1")
        self.assertEqual(fila17["asignacion_final"], "A2")
        # test 20: SHA original/final registrados en cada fila.
        self.assertEqual(fila16["sha256_global_original"], sha_original_esperado)
        self.assertEqual(fila17["sha256_global_original"], sha_original_esperado)
        self.assertEqual(fila16["sha256_global_final"], resumen2["sha256_global_final"])
        self.assertEqual(fila17["sha256_global_final"], resumen2["sha256_global_final"])

    def test_18_19_mes_siguiente_compara_contra_asignacion_final(self):
        # AGOSTO: 3P66536982 duplicada, el auditor corrige una ocurrencia
        # a 3P66536983 (CORRECTA la otra) -> el histórico queda con AMBAS
        # asignaciones finales.
        ruta_agosto = self._ruta_global(self.NOMBRE_GLOBAL_AGOSTO)
        resumen1 = self._ejecutar([
            _partida("3P66536982"), _partida("3P66536982"),
        ], ruta_global=ruta_agosto)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA", "OBSERVACION_AUDITOR": "doble depósito legítimo"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "3P66536983"},
        })
        resumen_agosto = ca.ejecutar_control(ruta_global=ruta_agosto, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen_agosto["estado_validacion"], "CERRADO_CON_VALIDACION_AUDITOR")

        # SEPTIEMBRE: aparece 3P66536982 otra vez -> DEBE alertar de nuevo
        # (test 19: una validación CORRECTA anterior no evita la alerta),
        # y 3P66536983 (la asignación FINAL de la corrección) también.
        ruta_septiembre = self._ruta_global(self.NOMBRE_GLOBAL_SEPTIEMBRE)
        _crear_global(ruta_septiembre, [
            _partida("3P66536982", fecha_valor=datetime.date(2026, 9, 3)),
            _partida("3P66536983", fecha_valor=datetime.date(2026, 9, 4)),
        ])
        resumen_sept = ca.ejecutar_control(ruta_global=ruta_septiembre, ruta_historico=self.ruta_historico)

        self.assertEqual(resumen_sept["estado"], "REVISAR_DUPLICADOS_ENCONTRADOS")
        self.assertEqual(resumen_sept["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        filas_sept = ca.cargar_revision_xlsx(resumen_sept["ruta_revision"])
        asignaciones_alertadas = {f["ASIGNACION_ORIGINAL"] for f in filas_sept}
        self.assertEqual(asignaciones_alertadas, {"3P66536982", "3P66536983"})
        for f in filas_sept:
            self.assertIn("AGOSTO_2026", f["ANTECEDENTE_HISTORICO"])


# ---------------------------------------------------------------------------
# 21-22: idempotencia sobre el GLOBAL final ya cerrado / corregido, y
# guardarraíl de GLOBAL modificado sin autorización.
# ---------------------------------------------------------------------------

class TestIdempotenciaGlobalFinal(_ControlAsignacionesTestBase):
    def test_21_rerun_sobre_global_final_ya_cerrado_ya_procesado(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("A1"), _partida("A1")], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "A2"},
        })
        ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)

        resumen3 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen3["estado"], "YA_PROCESADO_SIN_CAMBIOS")

    def test_22_cambio_posterior_no_autorizado_requiere_revision(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("A1"), _partida("A1")], ruta_global=ruta_global)
        _marcar_decisiones(resumen1["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "A2"},
        })
        ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)

        # Alguien edita el GLOBAL (ya cerrado) por fuera de CONTROL 1.
        wb = openpyxl.load_workbook(ruta_global)
        wb["1"]["D16"] = "GLOSA EDITADA A MANO"
        wb.save(ruta_global)

        resumen3 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen3["estado"], "GLOBAL_MODIFICADO_REQUIERE_REVISION")


# ---------------------------------------------------------------------------
# 23: compatibilidad con el CSV de revisión del esquema anterior.
# ---------------------------------------------------------------------------

class TestFallbackCsvLegado(_ControlAsignacionesTestBase):
    def test_23_fallback_csv_anterior_sigue_funcionando(self):
        ruta_global = self._ruta_global()
        _crear_global(ruta_global, [_partida("A1"), _partida("A1")])
        sha = ca._hash_archivo(ruta_global)

        ruta_csv_legado = self._ruta_csv_legado()
        with open(ruta_csv_legado, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "PERIODO", "ASIGNACION", "TIPO_ALERTA", "DETALLE",
                "VALIDACION_AUDITOR", "OBSERVACION_AUDITOR", "FECHA_VALIDACION", "SHA256_GLOBAL",
            ])
            writer.writeheader()
            writer.writerow({
                "PERIODO": "AGOSTO_2026", "ASIGNACION": "A1", "TIPO_ALERTA": "DUPLICADA_MISMO_MES",
                "DETALLE": "detalle previo", "VALIDACION_AUDITOR": "CORRECTA",
                "OBSERVACION_AUDITOR": "importado del csv legado", "FECHA_VALIDACION": "2026-08-20",
                "SHA256_GLOBAL": sha,
            })

        resumen = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen["estado_validacion"], "CERRADO_CON_VALIDACION_AUDITOR")
        self.assertTrue(os.path.isfile(resumen["ruta_revision"]))  # ahora existe el xlsx
        filas = ca.cargar_revision_xlsx(resumen["ruta_revision"])
        self.assertTrue(all(f["VALIDACION_AUDITOR"] == "CORRECTA" for f in filas))
        self.assertTrue(all(f["OBSERVACION_AUDITOR"] == "importado del csv legado" for f in filas))
        # El CSV legado nunca se borra ni se reescribe.
        self.assertTrue(os.path.isfile(ruta_csv_legado))


# ---------------------------------------------------------------------------
# 24: --dry-run no modifica absolutamente nada.
# ---------------------------------------------------------------------------

class TestDryRun(_ControlAsignacionesTestBase):
    def test_24_dry_run_no_modifica_xlsx_global_ni_historico(self):
        ruta_global = self._ruta_global()
        sha_antes = None
        _crear_global(ruta_global, [_partida("A1"), _partida("A1")])
        hash_antes = ca._hash_archivo(ruta_global)

        resumen1 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico, dry_run=True)
        self.assertEqual(resumen1["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        self.assertFalse(os.path.isfile(resumen1["ruta_revision"]))
        self.assertFalse(os.path.isfile(self.ruta_historico))

        # Aunque hipotéticamente ya estuviera todo resuelto, dry-run tampoco
        # debe tocar nada: se simula creando el xlsx fuera de dry-run,
        # marcando decisiones, y volviendo a correr con dry_run=True.
        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        _marcar_decisiones(resumen2["ruta_revision"], {
            16: {"VALIDACION_AUDITOR": "CORRECTA"},
            17: {"VALIDACION_AUDITOR": "INCORRECTA", "ASIGNACION_CORRECTA": "A2"},
        })
        hash_xlsx_antes = ca._hash_archivo(resumen2["ruta_revision"])
        hash_global_antes = ca._hash_archivo(ruta_global)

        resumen3 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico, dry_run=True)
        self.assertTrue(resumen3["dry_run"])
        self.assertFalse(resumen3["historico_actualizado"])
        self.assertFalse(resumen3["global_modificado"])
        self.assertFalse(os.path.isfile(self.ruta_historico))
        self.assertEqual(ca._hash_archivo(ruta_global), hash_global_antes)
        self.assertEqual(ca._hash_archivo(resumen2["ruta_revision"]), hash_xlsx_antes)

    def test_dry_run_sin_duplicados_no_escribe_historico(self):
        resumen = self._ejecutar([_partida("ASIG-OK")], dry_run=True)
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        self.assertFalse(resumen["historico_actualizado"])
        self.assertFalse(os.path.isfile(self.ruta_historico))


# ---------------------------------------------------------------------------
# Idempotencia SHA-256 exacta / archivo_global con SHA distinto / revisión
# de un SHA anterior que no se aplica en silencio — SIN CAMBIOS de fondo.
# ---------------------------------------------------------------------------

class TestIdempotencia(_ControlAsignacionesTestBase):
    def test_mismo_sha256_ya_procesado_sin_cambios(self):
        ruta_global = self._ruta_global()
        _crear_global(ruta_global, [_partida("ASIG-1")])
        primero = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(primero["estado"], "OK_SIN_DUPLICADOS")

        segundo = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(segundo["estado"], "YA_PROCESADO_SIN_CAMBIOS")

    def test_mismo_archivo_global_sha_diferente_requiere_revision(self):
        nombre = self.NOMBRE_GLOBAL_AGOSTO
        ca.guardar_historico(self.ruta_historico, [{
            "asignacion": "ASIG-1", "fecha_valor": "2026-08-01",
            "cuenta_mayor": "110101001", "glosa": "GLOSA", "monto": "10.00",
            "archivo_global": nombre, "fila_sap": "16",
            "sha256_archivo": "sha_diferente_al_actual",
            "fecha_incorporacion": "2026-08-31T00:00:00",
            "alerta_duplicado": "SIN_ALERTA", "validacion_auditor": "",
            "observacion_auditor": "", "fecha_validacion": "",
            "asignacion_original": "ASIG-1", "asignacion_final": "ASIG-1",
            "fila_global": "16", "sha256_global_original": "sha_diferente_al_actual",
            "sha256_global_final": "sha_diferente_al_actual",
        }])
        resumen = self._ejecutar([_partida("ASIG-1")], ruta_global=self._ruta_global(nombre))
        self.assertEqual(resumen["estado"], "GLOBAL_MODIFICADO_REQUIERE_REVISION")
        self.assertEqual(resumen["sha256_historico_existente"], "sha_diferente_al_actual")
        self.assertFalse(resumen["historico_actualizado"])

    def test_revision_de_sha_anterior_no_se_aplica_en_silencio(self):
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([_partida("ASIG-A"), _partida("ASIG-A")], ruta_global=ruta_global)
        # El xlsx queda "atado" a un SHA viejo (simulando que el GLOBAL
        # cambió después de generarlo, sin que el auditor haya alcanzado
        # a validar sobre el contenido correcto).
        wb = openpyxl.load_workbook(resumen1["ruta_revision"])
        ws = wb[ca._HOJA_REVISION]
        header = [c.value for c in ws[1]]
        idx_sha = header.index("SHA256_GLOBAL")
        idx_val = header.index("VALIDACION_AUDITOR")
        for row in ws.iter_rows(min_row=2):
            row[idx_sha].value = "sha_anterior_distinto"
            row[idx_val].value = "CORRECTA"
        wb.save(resumen1["ruta_revision"])

        resumen2 = ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen2["estado_validacion"], "PENDIENTE_VALIDACION_AUDITOR")
        filas = ca.cargar_revision_xlsx(resumen1["ruta_revision"])
        self.assertTrue(all(f["VALIDACION_AUDITOR"] == "" for f in filas))


# ---------------------------------------------------------------------------
# Reproceso e inmutabilidad — SIN CAMBIOS de fondo.
# ---------------------------------------------------------------------------

class TestReprocesoEInmutabilidad(_ControlAsignacionesTestBase):
    def test_global_origen_no_cambia_de_hash_sin_duplicados(self):
        ruta_global = self._ruta_global()
        _crear_global(ruta_global, [_partida("ASIG-1")])
        hash_antes = ca._hash_archivo(ruta_global)
        ca.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico)
        hash_despues = ca._hash_archivo(ruta_global)
        self.assertEqual(hash_antes, hash_despues)

    def test_historico_conserva_filas_anteriores_al_agregar_global_valido(self):
        fila_previa = {
            "asignacion": "ASIG-VIEJA", "fecha_valor": "2026-07-01",
            "cuenta_mayor": "110101001", "glosa": "GLOSA VIEJA", "monto": "20.00",
            "archivo_global": "SAP_GLOBAL_TIQ_JULIO_2026.xlsx", "fila_sap": "16",
            "sha256_archivo": "sha_julio", "fecha_incorporacion": "2026-07-31T00:00:00",
            "alerta_duplicado": "SIN_ALERTA", "validacion_auditor": "",
            "observacion_auditor": "", "fecha_validacion": "",
            "asignacion_original": "ASIG-VIEJA", "asignacion_final": "ASIG-VIEJA",
            "fila_global": "16", "sha256_global_original": "sha_julio",
            "sha256_global_final": "sha_julio",
        }
        ca.guardar_historico(self.ruta_historico, [fila_previa])

        self._ejecutar([_partida("ASIG-NUEVA")])

        historico_final = _leer_csv(self.ruta_historico)
        self.assertEqual(len(historico_final), 2)
        self.assertIn(fila_previa["asignacion"], [f["asignacion"] for f in historico_final])
        self.assertEqual(historico_final[0], fila_previa)

    def test_historico_esquema_antiguo_se_lee_sin_perder_informacion(self):
        columnas_antiguas = [
            "asignacion", "fecha_valor", "cuenta_mayor", "glosa", "monto",
            "archivo_global", "fila_sap", "sha256_archivo", "fecha_incorporacion",
        ]
        fila_vieja = {
            "asignacion": "ASIG-VIEJA", "fecha_valor": "2026-07-01",
            "cuenta_mayor": "110101001", "glosa": "GLOSA VIEJA", "monto": "20.00",
            "archivo_global": "SAP_GLOBAL_TIQ_JULIO_2026.xlsx", "fila_sap": "16",
            "sha256_archivo": "sha_julio", "fecha_incorporacion": "2026-07-31T00:00:00",
        }
        with open(self.ruta_historico, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=columnas_antiguas)
            writer.writeheader()
            writer.writerow(fila_vieja)

        resumen = self._ejecutar([_partida("ASIG-NUEVA")])
        self.assertEqual(resumen["estado"], "OK_SIN_DUPLICADOS")
        historico_final = _leer_csv(self.ruta_historico)
        fila_vieja_final = next(f for f in historico_final if f["asignacion"] == "ASIG-VIEJA")
        for col, valor in fila_vieja.items():
            self.assertEqual(fila_vieja_final[col], valor)
        self.assertEqual(fila_vieja_final["asignacion_final"], "")  # columna nueva: vacía, nunca inventada


# ---------------------------------------------------------------------------
# Preservación del JSON de detalle "por par" (compatibilidad) + campos
# nuevos del resumen.
# ---------------------------------------------------------------------------

class TestDetalleJson(_ControlAsignacionesTestBase):
    def test_json_conserva_hallazgos_por_par_y_campos_nuevos(self):
        glosa_larga = "RECAUDACION CAJA SFC101 - GLOSA COMPLETA DE PRUEBA CON DETALLE"
        ruta_global = self._ruta_global()
        resumen1 = self._ejecutar([
            _partida("ASIG-Y", glosa=glosa_larga), _partida("ASIG-Y", glosa=glosa_larga),
        ], ruta_global=ruta_global, ruta_detalle_json=os.path.join(self.tmpdir, "detalle.json"))

        with open(resumen1["detalle_json"], encoding="utf-8") as f:
            detalle = json.load(f)
        self.assertEqual(detalle["hallazgos"][0]["glosa"], glosa_larga)
        self.assertIn("sha256_global_original", detalle)
        self.assertIn("sha256_global_final", detalle)
        self.assertIn("filas_revisadas", detalle)
        self.assertEqual(detalle["filas_revisadas"], 2)


# ---------------------------------------------------------------------------
# Sin dependencias externas / sin acoplamiento con el V2 diario
# ---------------------------------------------------------------------------

class TestSinDependenciasExternas(unittest.TestCase):
    """Verifica el CÓDIGO EJECUTABLE (imports reales vía AST), no el texto
    de docstrings/comentarios."""

    def setUp(self):
        import ast
        ruta_modulo = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "control_asignaciones.py",
        )
        with open(ruta_modulo, "r", encoding="utf-8") as f:
            arbol = ast.parse(f.read(), filename=ruta_modulo)

        self.modulos_importados = set()
        for nodo in ast.walk(arbol):
            if isinstance(nodo, ast.Import):
                for alias in nodo.names:
                    self.modulos_importados.add(alias.name.split(".")[0])
            elif isinstance(nodo, ast.ImportFrom) and nodo.module:
                self.modulos_importados.add(nodo.module.split(".")[0])

    def test_ninguna_dependencia_google_drive(self):
        prohibidos = {"google", "googleapiclient", "gspread", "pydrive", "pydrive2"}
        self.assertEqual(self.modulos_importados & prohibidos, set())

    def test_ninguna_dependencia_base64(self):
        self.assertNotIn("base64", self.modulos_importados)

    def test_ningun_import_del_motor_diario(self):
        modulos_v2_diario = {
            "motor_tiquipaya", "pipeline_tiquipaya", "sap_writer",
            "excel_io", "run_batch", "consolidador_mensual",
        }
        self.assertEqual(self.modulos_importados & modulos_v2_diario, set())


if __name__ == "__main__":
    unittest.main()
