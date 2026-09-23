"""test_caja_america.py — Soporte de CAJA AMERICA sobre el núcleo de Caja
Tiquipaya: configuración por caja, separación del ATC por la columna CAJA y
regla POSGRADO RESERVA.

Ningún dato contable real: todos los importes, cuentas y códigos son
inventados para estas pruebas.

Los fixtures viven aquí (y no en tests/xlsx_fixtures.py) a propósito: ese
módulo lo comparten las 432 pruebas históricas de Tiquipaya y no se toca.

Cobertura (A-G del encargo):
  A. ATC histórico sin columna CAJA + TIQUIPAYA -> resultado histórico.
  B. ATC histórico sin columna CAJA + AMERICA   -> falla cerrado.
  C. ATC nuevo, misma fecha con las dos cajas   -> cada una ve solo lo suyo.
  D. AMERICA con reserva 0.00                   -> sin partida CxP.
  E. AMERICA 2000/100/1900                      -> la regla completa.
  F. reserva > ATC bruto                        -> bloqueado, sin negativos.
  G. TIQUIPAYA idéntica (literales e invariantes que no deben moverse).
"""

import os
import sys
import unittest
from decimal import Decimal

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config_cajas as cfg  # noqa: E402
import excel_io as io  # noqa: E402
import motor_tiquipaya as motor  # noqa: E402
import run_batch  # noqa: E402

from tests.xlsx_fixtures import crear_cierre, crear_maestro_unico  # noqa: E402


FECHA = "2026-09-15"          # septiembre -> POSTG-SEPT (no "SEP")
NOMBRE_CIERRE = "CIERRE 15-09-2026.xlsm"

CUENTA_ATC_NETO = "110103012"
CUENTA_ATC_COMISION = "110201008"


# ---------------------------------------------------------------------------
# Fixtures locales
# ---------------------------------------------------------------------------

def _hoja_resumen(ws, datos):
    """Resumen SFC con el mismo layout que excel_io exige, más la fila
    POSGRADO RESERVA cuando el fixture la pide."""
    ws.append(["TOTAL MOVIMIENTO DEL DIA", datos["total_movimiento"]])
    ws.append(["COBROS ATC", datos["cobros_atc"]])
    ws.append(["TOTAL COMUNICACIONES INTERNAS", datos.get("total_ci", "0.00")])
    ws.append(["DOLARES", datos.get("dolares", "0.00")])
    if "posgrado_reserva" in datos:
        ws.append(["POSGRADO RESERVA", datos["posgrado_reserva"]])
    ws.append([None, None, None, None, None])
    ws.append(["COMPOSICION DE DEPOSITOS", "IMPORTE Bs", "FECHA DE DEPOSITO",
               "ASIGNACION", "BANCO"])
    for dep in datos.get("depositos", []):
        ws.append([dep.get("deposito", "DEPOSITO"), dep["importe"],
                   dep.get("fecha"), dep.get("asignacion"), dep.get("banco", "BNB")])


def crear_cierre_america(ruta, sfc107, sfc108):
    """CIERRE de América: hojas SFC107/SFC108 (+ sus Comunicaciones
    Internas). TOTAL MOVIMIENTO DEL DIA ya viene NETO de reserva, igual
    que en el archivo real corregido."""
    wb = openpyxl.Workbook()
    ws107 = wb.active
    ws107.title = "SFC107"
    _hoja_resumen(ws107, sfc107)

    ws108 = wb.create_sheet("SFC108")
    _hoja_resumen(ws108, sfc108)

    for sfc, datos in (("SFC107", sfc107), ("SFC108", sfc108)):
        ws_ci = wb.create_sheet(f"COMUNICACIONES INTERNAS {sfc}")
        ws_ci.append(["N°", "N° DE FACTURA", "TOTAL C.I.",
                      "CUENTA CONTABLE BANCO", "ASIGNACION", "BANCO"])
        for i, f in enumerate(datos.get("ci", []), start=1):
            ws_ci.append([i, f.get("factura", f"FAC-{i}"), f["total"],
                          f.get("cuenta"), f.get("asignacion"), f.get("banco", "BNB")])

    wb.save(ruta)


def crear_maestro_con_caja(ruta, atc_filas, macros_filas=()):
    """Maestro mensual con el ATC en formato NUEVO: la hoja sigue
    llamándose "ATC TIQUIPAYA" y trae la columna CAJA (A..G del encargo).

    atc_filas: [(fecha, tipo, cuenta, detalle, monto, asignacion, caja), ...]
    """
    wb = openpyxl.Workbook()
    ws_macros = wb.active
    ws_macros.title = "Tablas Dinamicas Profesional"
    ws_macros.append(["Fecha", "Código de Asignación", "Créditos"])
    for fila in macros_filas:
        ws_macros.append(list(fila))

    ws_atc = wb.create_sheet("ATC TIQUIPAYA")
    ws_atc.append(["FECHA", "TIPO", "CUENTA CONTABLE", "DETALLE", "MONTO",
                   "ASIGNACION", "CAJA"])
    for fila in atc_filas:
        ws_atc.append(list(fila))

    wb.save(ruta)


def _fila_atc(tipo, monto, caja, fecha=FECHA):
    cuenta = CUENTA_ATC_NETO if "NETO" in tipo else CUENTA_ATC_COMISION
    return (fecha, tipo, cuenta, f"{tipo} {caja}", monto, f"ATC-{caja}", caja)


def _atc_dos_cajas(neto_tiq="900.00", com_tiq="100.00",
                   neto_ame="1950.00", com_ame="50.00"):
    """Las 4 filas que el mismo día puede traer en el maestro nuevo."""
    return [
        _fila_atc("BANCO (NETO)", neto_tiq, "TIQUIPAYA"),
        _fila_atc("COMISIÓN ATC", com_tiq, "TIQUIPAYA"),
        _fila_atc("BANCO (NETO)", neto_ame, "AMERICA"),
        _fila_atc("COMISIÓN ATC", com_ame, "AMERICA"),
    ]


def _procesar_america(tmpdir, sfc107, sfc108, atc_filas=None):
    ruta_cierre = os.path.join(tmpdir, NOMBRE_CIERRE)
    ruta_maestro = os.path.join(tmpdir, "MAESTRO SEPTIEMBRE.xlsm")
    crear_cierre_america(ruta_cierre, sfc107, sfc108)
    crear_maestro_con_caja(ruta_maestro, atc_filas if atc_filas is not None
                           else _atc_dos_cajas())
    resultado = motor.ejecutar_v2(ruta_cierre, ruta_maestro, ruta_maestro,
                                  caja="america")
    return resultado, motor.construir_asiento(resultado)


def _partidas_por_origen(asiento, origen):
    return [p for p in asiento["partidas"] if p["origen"] == origen]


class _TmpMixin:
    def setUp(self):
        import tempfile
        self._tmp = tempfile.TemporaryDirectory()
        self.tmpdir = self._tmp.name
        self.addCleanup(self._tmp.cleanup)


# ---------------------------------------------------------------------------
# A / B / C — separación del ATC por la columna CAJA
# ---------------------------------------------------------------------------

class TestAtcPorCaja(_TmpMixin, unittest.TestCase):

    def test_A_atc_historico_sin_columna_caja_tiquipaya_igual_que_siempre(self):
        """A. Maestro histórico (sin columna CAJA): TIQUIPAYA lee exactamente
        lo de siempre, sin filtrar nada."""
        ruta = os.path.join(self.tmpdir, "MAESTRO SEPTIEMBRE.xlsm")
        crear_maestro_unico(ruta, [], [
            (FECHA, "BANCO (NETO)", CUENTA_ATC_NETO, "NETO", "900.00", "ATC-1"),
            (FECHA, "COMISIÓN ATC", CUENTA_ATC_COMISION, "COM", "100.00", "ATC-2"),
        ])

        por_defecto = io.leer_atc_mensual(ruta)
        explicito = io.leer_atc_mensual(ruta, caja="tiquipaya")

        self.assertEqual(por_defecto, explicito)
        self.assertEqual(por_defecto["modo"], "PRECONCILIADO")
        registro = por_defecto["por_fecha"][FECHA]
        self.assertEqual(registro["neto"]["monto"], "900.00")
        self.assertEqual(registro["comision"]["monto"], "100.00")

    def test_B_atc_historico_sin_columna_caja_america_falla_cerrado(self):
        """B. De un maestro sin diferenciador de caja NO se puede inferir
        América: error explícito, nunca un silencio ni un default."""
        ruta = os.path.join(self.tmpdir, "MAESTRO SEPTIEMBRE.xlsm")
        crear_maestro_unico(ruta, [], [
            (FECHA, "BANCO (NETO)", CUENTA_ATC_NETO, "NETO", "900.00", "ATC-1"),
            (FECHA, "COMISIÓN ATC", CUENTA_ATC_COMISION, "COM", "100.00", "ATC-2"),
        ])

        with self.assertRaises(ValueError) as ctx:
            io.leer_atc_mensual(ruta, caja="america")

        mensaje = str(ctx.exception)
        self.assertIn("CAJA", mensaje)
        self.assertIn("america", mensaje)

    def test_B2_atc_legado_sin_columna_caja_america_falla_cerrado(self):
        """El formato LEGADO tampoco tiene columna CAJA: mismo fallo."""
        ruta = os.path.join(self.tmpdir, "ATC_LEGADO.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["FECHA", "TIPO", "MONTO"])
        ws.append([FECHA, "BANCO (NETO)", "900.00"])
        ws.append([FECHA, "COMISIÓN ATC", "100.00"])
        wb.save(ruta)

        self.assertEqual(io.leer_atc_mensual(ruta)["modo"], "LEGADO")
        with self.assertRaises(ValueError):
            io.leer_atc_mensual(ruta, caja="america")

    def test_C_misma_fecha_dos_cajas_cada_una_recibe_solo_lo_suyo(self):
        """C. Las 4 filas del mismo día no se ven entre sí: ni se suman, ni
        se pisan, ni disparan el control de duplicados."""
        ruta = os.path.join(self.tmpdir, "MAESTRO SEPTIEMBRE.xlsm")
        crear_maestro_con_caja(ruta, _atc_dos_cajas())

        tiq = io.leer_atc_mensual(ruta, caja="tiquipaya")["por_fecha"][FECHA]
        ame = io.leer_atc_mensual(ruta, caja="america")["por_fecha"][FECHA]

        self.assertEqual(tiq["neto"]["monto"], "900.00")
        self.assertEqual(tiq["comision"]["monto"], "100.00")
        self.assertEqual(tiq["neto"]["detalle"], "BANCO (NETO) TIQUIPAYA")

        self.assertEqual(ame["neto"]["monto"], "1950.00")
        self.assertEqual(ame["comision"]["monto"], "50.00")
        self.assertEqual(ame["neto"]["detalle"], "BANCO (NETO) AMERICA")

    def test_C2_fila_con_caja_vacia_falla_cerrado(self):
        """Una fila con datos pero sin CAJA es ambigua: nunca se asigna por
        defecto a ninguna caja."""
        ruta = os.path.join(self.tmpdir, "MAESTRO SEPTIEMBRE.xlsm")
        crear_maestro_con_caja(ruta, [
            (FECHA, "BANCO (NETO)", CUENTA_ATC_NETO, "NETO", "900.00", "ATC-1", None),
            (FECHA, "COMISIÓN ATC", CUENTA_ATC_COMISION, "COM", "100.00", "ATC-2", "TIQUIPAYA"),
        ])

        with self.assertRaises(ValueError) as ctx:
            io.leer_atc_mensual(ruta, caja="tiquipaya")
        self.assertIn("CAJA", str(ctx.exception))

    def test_C3_maestro_nuevo_no_rompe_el_control_de_duplicados(self):
        """Dos NETO de la MISMA caja y fecha siguen siendo un error."""
        ruta = os.path.join(self.tmpdir, "MAESTRO SEPTIEMBRE.xlsm")
        crear_maestro_con_caja(ruta, [
            _fila_atc("BANCO (NETO)", "900.00", "AMERICA"),
            _fila_atc("BANCO (NETO)", "800.00", "AMERICA"),
        ])

        with self.assertRaises(ValueError) as ctx:
            io.leer_atc_mensual(ruta, caja="america")
        self.assertIn("NETO duplicada", str(ctx.exception))


# ---------------------------------------------------------------------------
# Lectura del cierre de América
# ---------------------------------------------------------------------------

class TestLecturaCierreAmerica(_TmpMixin, unittest.TestCase):

    def test_lee_sfc107_sfc108_y_reserva(self):
        ruta = os.path.join(self.tmpdir, NOMBRE_CIERRE)
        crear_cierre_america(
            ruta,
            {"total_movimiento": "1900.00", "cobros_atc": "2000.00",
             "posgrado_reserva": "100.00"},
            {"total_movimiento": "500.00", "cobros_atc": "0.00",
             "posgrado_reserva": "0.00"},
        )

        cierre = io.leer_cierre(ruta, caja="america")

        self.assertEqual(cierre["caja"], "america")
        self.assertEqual(cierre["sfc107"]["total_movimiento"], "1900.00")
        self.assertEqual(cierre["sfc107"]["cobros_atc"], "2000.00")
        self.assertEqual(cierre["sfc107"]["posgrado_reserva"], "100.00")
        self.assertEqual(cierre["sfc108"]["posgrado_reserva"], "0.00")
        self.assertNotIn("sfc101", cierre)

    def test_america_sin_campo_reserva_falla_cerrado(self):
        """Sin el campo no se asume 0.00: TOTAL MOVIMIENTO ya viene neto y
        un 0.00 inventado descuadraría el cierre."""
        ruta = os.path.join(self.tmpdir, NOMBRE_CIERRE)
        crear_cierre_america(
            ruta,
            {"total_movimiento": "1900.00", "cobros_atc": "2000.00"},
            {"total_movimiento": "0.00", "cobros_atc": "0.00"},
        )

        with self.assertRaises(ValueError) as ctx:
            io.leer_cierre(ruta, caja="america")
        self.assertIn("POSGRADO RESERVA", str(ctx.exception))

    def test_caja_desconocida_falla_cerrado(self):
        for valor in ("sucre", "AMERICA LATINA", 7):
            with self.assertRaises(ValueError):
                cfg.resolver_caja(valor)


# ---------------------------------------------------------------------------
# D / E / F — regla POSGRADO RESERVA
# ---------------------------------------------------------------------------

class TestReservaPosgrado(_TmpMixin, unittest.TestCase):

    def test_D_america_reserva_cero(self):
        """D. Sin reserva, América se comporta como cualquier caja: dos
        HABER normales en 110101003, ninguna partida 210103003."""
        resultado, asiento = _procesar_america(
            self.tmpdir,
            {"total_movimiento": "1500.00", "cobros_atc": "2000.00",
             "posgrado_reserva": "0.00"},
            {"total_movimiento": "500.00", "cobros_atc": "0.00",
             "posgrado_reserva": "0.00"},
        )

        self.assertEqual(resultado["estado"], "OK")
        self.assertEqual(resultado["diferencia"], "0.00")
        self.assertEqual(asiento["estado"], "OK")
        self.assertEqual(asiento["problemas"], [])

        haber_normales = (_partidas_por_origen(asiento, "UNIVERSO_SFC107")
                          + _partidas_por_origen(asiento, "UNIVERSO_SFC108"))
        self.assertEqual(len(haber_normales), 2)
        for p in haber_normales:
            self.assertEqual(p["cuenta_mayor"], "110101003")

        self.assertEqual(_partidas_por_origen(asiento, "RESERVA_POSGRADO"), [])
        self.assertNotIn("210103003", [p["cuenta_mayor"] for p in asiento["partidas"]])
        self.assertEqual(asiento["total_cargo"], asiento["total_haber"])

    def test_E_america_2000_100_1900(self):
        """E. El caso del encargo, completo."""
        resultado, asiento = _procesar_america(
            self.tmpdir,
            {"total_movimiento": "1900.00", "cobros_atc": "2000.00",
             "posgrado_reserva": "100.00"},
            {"total_movimiento": "0.00", "cobros_atc": "0.00",
             "posgrado_reserva": "0.00"},
        )

        # El ATC BRUTO con el que se valida NETO + COMISION sigue siendo el
        # cobro real de 2000: la reserva no lo toca.
        self.assertEqual(resultado["componentes"]["atc_bruto"], "2000.00")
        self.assertEqual(resultado["componentes"]["reserva_posgrado"], "100.00")
        # ATC COMPUTABLE (solo para el cuadre).
        self.assertEqual(resultado["componentes"]["atc_computable"], "1900.00")

        # El universo NO vuelve a descontar la reserva: TOTAL MOVIMIENTO ya
        # viene neto (1900), y el HABER normal sale de ahí.
        self.assertEqual(resultado["universo_original"], "1900.00")
        self.assertEqual(resultado["universo_ajustado"], "1900.00")
        self.assertEqual(resultado["recaudacion_explicada"], "1900.00")
        self.assertEqual(resultado["diferencia"], "0.00")
        self.assertEqual(resultado["estado"], "OK")

        self.assertEqual(asiento["estado"], "OK")
        self.assertEqual(asiento["problemas"], [])

        haber_107 = _partidas_por_origen(asiento, "UNIVERSO_SFC107")[0]
        haber_108 = _partidas_por_origen(asiento, "UNIVERSO_SFC108")[0]
        self.assertEqual(haber_107["cuenta_mayor"], "110101003")
        self.assertEqual(haber_107["haber"], "1900.00")   # NO 1800.00
        self.assertEqual(haber_107["texto_posicion"], "RECAUDACION CAJA SFC107")
        self.assertEqual(haber_107["asignacion"], "SFC107")
        self.assertEqual(haber_108["haber"], "0.00")
        self.assertEqual(haber_108["texto_posicion"], "RECAUDACION CAJA SFC108")

        reservas = _partidas_por_origen(asiento, "RESERVA_POSGRADO")
        self.assertEqual(len(reservas), 1)
        reserva = reservas[0]
        self.assertEqual(reserva["cuenta_mayor"], "210103003")
        self.assertEqual(reserva["haber"], "100.00")
        self.assertEqual(reserva["cargo"], "0.00")
        self.assertEqual(reserva["texto_posicion"], "RESERVA POSGRADO")
        self.assertEqual(reserva["asignacion"], "POSTG-SEPT")

        # El DEBE conserva el ATC completo (NETO + COMISION = 2000).
        neto = _partidas_por_origen(asiento, "ATC_NETO")[0]
        comision = _partidas_por_origen(asiento, "ATC_COMISION")[0]
        self.assertEqual(neto["cargo"], "1950.00")
        self.assertEqual(comision["cargo"], "50.00")
        self.assertEqual(
            Decimal(neto["cargo"]) + Decimal(comision["cargo"]), Decimal("2000.00")
        )

        self.assertEqual(asiento["total_cargo"], "2000.00")
        self.assertEqual(asiento["total_haber"], "2000.00")
        self.assertEqual(asiento["diferencia"], "0.00")

    def test_E2_reserva_repartida_en_las_dos_hojas_una_sola_partida(self):
        """Dos hojas con reserva -> UNA partida CxP por el total (misma
        cuenta y misma asignación: dos líneas serían indistinguibles)."""
        resultado, asiento = _procesar_america(
            self.tmpdir,
            {"total_movimiento": "1000.00", "cobros_atc": "1100.00",
             "posgrado_reserva": "100.00"},
            {"total_movimiento": "850.00", "cobros_atc": "900.00",
             "posgrado_reserva": "50.00"},
            atc_filas=[
                _fila_atc("BANCO (NETO)", "1900.00", "AMERICA"),
                _fila_atc("COMISIÓN ATC", "100.00", "AMERICA"),
            ],
        )

        self.assertEqual(resultado["componentes"]["atc_bruto"], "2000.00")
        self.assertEqual(resultado["componentes"]["reserva_posgrado"], "150.00")
        self.assertEqual(resultado["componentes"]["atc_computable"], "1850.00")
        self.assertEqual(resultado["diferencia"], "0.00")

        reservas = _partidas_por_origen(asiento, "RESERVA_POSGRADO")
        self.assertEqual(len(reservas), 1)
        self.assertEqual(reservas[0]["haber"], "150.00")
        self.assertEqual(asiento["total_cargo"], asiento["total_haber"])

    def test_F_reserva_mayor_que_atc_bruto_bloquea(self):
        """F. Fail-closed: nunca se genera un asiento (ni una partida
        negativa) para forzar el cuadre."""
        resultado, asiento = _procesar_america(
            self.tmpdir,
            {"total_movimiento": "1000.00", "cobros_atc": "100.00",
             "posgrado_reserva": "500.00"},
            {"total_movimiento": "0.00", "cobros_atc": "0.00",
             "posgrado_reserva": "0.00"},
            atc_filas=[
                _fila_atc("BANCO (NETO)", "90.00", "AMERICA"),
                _fila_atc("COMISIÓN ATC", "10.00", "AMERICA"),
            ],
        )

        self.assertEqual(resultado["estado"], "BLOQUEADO_EXCEPCION")
        self.assertGreaterEqual(resultado["excepciones_bloqueantes"], 1)
        tipos = [e["tipo"] for e in resultado["excepciones"]]
        self.assertIn("RESERVA_POSGRADO_MAYOR_QUE_ATC_BRUTO", tipos)

        self.assertEqual(asiento["estado"], "NO_ASIENTO")
        self.assertEqual(asiento["partidas"], [])

    def test_F2_reserva_negativa_bloquea(self):
        resultado, asiento = _procesar_america(
            self.tmpdir,
            {"total_movimiento": "1900.00", "cobros_atc": "2000.00",
             "posgrado_reserva": "-100.00"},
            {"total_movimiento": "0.00", "cobros_atc": "0.00",
             "posgrado_reserva": "0.00"},
        )

        self.assertEqual(resultado["estado"], "BLOQUEADO_EXCEPCION")
        tipos = [e["tipo"] for e in resultado["excepciones"]]
        self.assertIn("RESERVA_POSGRADO_NEGATIVA", tipos)
        self.assertEqual(asiento["estado"], "NO_ASIENTO")
        self.assertEqual(asiento["partidas"], [])

    def test_F3_ninguna_partida_negativa_en_ningun_escenario(self):
        """Invariante transversal: ningún asiento emitido puede traer un
        importe negativo."""
        escenarios = [
            ("0.00", "2000.00", "1500.00"),
            ("100.00", "2000.00", "1900.00"),
            ("2000.00", "2000.00", "0.00"),   # reserva == bruto (límite válido)
        ]
        for reserva, bruto, movimiento in escenarios:
            with self.subTest(reserva=reserva):
                neto = str(Decimal(bruto) - Decimal("50.00"))
                _resultado, asiento = _procesar_america(
                    self.tmpdir,
                    {"total_movimiento": movimiento, "cobros_atc": bruto,
                     "posgrado_reserva": reserva},
                    {"total_movimiento": "0.00", "cobros_atc": "0.00",
                     "posgrado_reserva": "0.00"},
                    atc_filas=[
                        _fila_atc("BANCO (NETO)", neto, "AMERICA"),
                        _fila_atc("COMISIÓN ATC", "50.00", "AMERICA"),
                    ],
                )
                for p in asiento["partidas"]:
                    self.assertGreaterEqual(Decimal(p["cargo"]), 0)
                    self.assertGreaterEqual(Decimal(p["haber"]), 0)
                if asiento["partidas"]:
                    self.assertEqual(asiento["total_cargo"], asiento["total_haber"])

    def test_asignacion_postg_cubre_los_12_meses_con_sept(self):
        esperado = {
            "01": "POSTG-ENE", "02": "POSTG-FEB", "03": "POSTG-MAR",
            "04": "POSTG-ABR", "05": "POSTG-MAY", "06": "POSTG-JUN",
            "07": "POSTG-JUL", "08": "POSTG-AGO", "09": "POSTG-SEPT",
            "10": "POSTG-OCT", "11": "POSTG-NOV", "12": "POSTG-DIC",
        }
        for mes, asignacion in esperado.items():
            self.assertEqual(
                cfg.asignacion_reserva_posgrado(f"2026-{mes}-15"), asignacion
            )
        self.assertIsNone(cfg.asignacion_reserva_posgrado(None))


# ---------------------------------------------------------------------------
# G — TIQUIPAYA no se movió
# ---------------------------------------------------------------------------

class TestTiquipayaIntacta(_TmpMixin, unittest.TestCase):

    def test_G_literales_de_tiquipaya(self):
        caja = cfg.resolver_caja()
        self.assertIs(caja, cfg.TIQUIPAYA)
        self.assertEqual(caja.codigo, "tiquipaya")
        self.assertEqual(caja.sfcs, ("SFC101", "SFC102"))
        self.assertEqual(caja.cuenta_haber, "110101001")
        self.assertEqual(caja.nombre_sap, "CAJA TIQUIPAYA")
        self.assertEqual(caja.claves_sfc, ("sfc101", "sfc102"))
        self.assertEqual(caja.texto_haber("SFC101"), "RECAUDACION CAJA SFC101")
        self.assertEqual(caja.texto_haber("SFC102"), "RECAUDACION CAJA SFC102")
        self.assertEqual(caja.origenes_universo,
                         ("UNIVERSO_SFC101", "UNIVERSO_SFC102"))
        self.assertFalse(caja.reserva_posgrado)
        self.assertIsNone(caja.cuenta_reserva_posgrado)

    def test_G_america_literales(self):
        caja = cfg.resolver_caja("america")
        self.assertEqual(caja.sfcs, ("SFC107", "SFC108"))
        self.assertEqual(caja.cuenta_haber, "110101003")
        self.assertEqual(caja.nombre_sap, "CAJA AMERICA")
        self.assertEqual(caja.atc_caja, "AMERICA")
        self.assertTrue(caja.reserva_posgrado)
        self.assertEqual(caja.cuenta_reserva_posgrado, "210103003")

    def test_G_cabecera_sap_por_caja(self):
        self.assertEqual(
            run_batch.construir_metadata_cabecera(FECHA)["referencia"],
            "CAJA TIQUIPAYA",
        )
        self.assertEqual(
            run_batch.construir_metadata_cabecera(FECHA, "tiquipaya")["referencia"],
            "CAJA TIQUIPAYA",
        )
        self.assertEqual(
            run_batch.construir_metadata_cabecera(FECHA, "america")["referencia"],
            "CAJA AMERICA",
        )
        # El texto de cabecera (G10) no depende de la caja.
        self.assertEqual(
            run_batch.construir_metadata_cabecera(FECHA, "america")["texto_cabecera"],
            run_batch.construir_metadata_cabecera(FECHA)["texto_cabecera"],
        )

    def test_G_cierre_de_tiquipaya_no_gana_claves_nuevas(self):
        """El dict de Tiquipaya conserva EXACTAMENTE su forma histórica
        (salvo la identidad "caja", que no se publica en ningún artefacto,
        y "facturas_anuladas"/"facturas_anuladas_detalle" del BLOQUE 1,
        universales y en 0.00/[] cuando el cierre no las trae): sin
        posgrado_reserva en las hojas SFC."""
        ruta = os.path.join(self.tmpdir, NOMBRE_CIERRE)
        crear_cierre(
            ruta,
            {"total_movimiento": "1000.00", "cobros_atc": "0.00", "dolares": "0.00"},
            {"total_movimiento": "500.00", "cobros_atc": "0.00", "dolares": "0.00"},
        )

        cierre = io.leer_cierre(ruta)

        self.assertEqual(set(cierre), {"fecha_cierre", "caja", "sfc101",
                                       "sfc102", "comunicaciones_internas"})
        self.assertEqual(cierre["caja"], "tiquipaya")
        for clave in ("sfc101", "sfc102"):
            self.assertEqual(
                set(cierre[clave]),
                {"total_movimiento", "cobros_atc", "total_ci", "dolares",
                 "depositos", "facturas_anuladas", "facturas_anuladas_detalle"},
            )
            self.assertEqual(cierre[clave]["facturas_anuladas"], "0.00")
            self.assertEqual(cierre[clave]["facturas_anuladas_detalle"], [])

    def test_G_componentes_de_tiquipaya_sin_claves_de_reserva(self):
        ruta_cierre = os.path.join(self.tmpdir, NOMBRE_CIERRE)
        ruta_maestro = os.path.join(self.tmpdir, "MAESTRO SEPTIEMBRE.xlsm")
        crear_cierre(
            ruta_cierre,
            {"total_movimiento": "1000.00", "cobros_atc": "0.00", "dolares": "0.00"},
            {"total_movimiento": "0.00", "cobros_atc": "0.00", "dolares": "0.00"},
        )
        crear_maestro_unico(ruta_maestro, [], [])

        resultado = motor.ejecutar_v2(ruta_cierre, ruta_maestro, ruta_maestro)

        self.assertEqual(
            set(resultado["componentes"]),
            {"vouchers", "ci_operativas", "atc_bruto", "dolares", "facturas_anuladas"},
        )
        self.assertEqual(resultado["componentes"]["facturas_anuladas"], "0.00")
        self.assertNotIn("reserva_posgrado", resultado["detalle"])
        # Las claves por SFC del detalle siguen siendo las históricas.
        for clave in ("sfc101_total", "sfc102_total", "sfc101_haber",
                      "sfc102_haber", "alquileres_sfc101", "alquileres_sfc102"):
            self.assertIn(clave, resultado["detalle"])

    def test_G_partida_de_reserva_rechazada_en_caja_sin_la_regla(self):
        """Defensa en profundidad: si una partida CxP de reserva se colara
        en un asiento de Tiquipaya, la validación la rechaza."""
        partida = {
            "sociedad": "BO01", "cuenta_mayor": "210103003",
            "texto_posicion": "RESERVA POSGRADO", "cargo": "0.00",
            "haber": "100.00", "centro_beneficio": "10010101",
            "fecha_valor": FECHA, "asignacion": "POSTG-SEPT",
            "origen": "RESERVA_POSGRADO", "sfc_origen": None,
            "codigo_informado_original": None,
        }
        problemas = motor._validar_partidas(
            [partida], Decimal("0.00"), Decimal("100.00"), Decimal("-100.00")
        )
        self.assertIn("RESERVA_POSGRADO_NO_APLICA", problemas)


if __name__ == "__main__":
    unittest.main()
