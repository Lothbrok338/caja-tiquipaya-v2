"""
test_control_cxc_cxp.py — CONTROL 3: seguimiento acumulado de CxC/CxP
(cuentas transitorias 110201002/210103002/110201003/210103003/110201004/
210103004) del SAP GLOBAL mensual, con llave CUENTA+ASIGNACION, histórico
técnico persistente e idempotencia por periodo.

Usa un SAP GLOBAL sintético construido localmente en este archivo (nunca
datos contables reales). Verifica exclusivamente el comportamiento de
control_cxc_cxp.py: no repite ninguna regla de excel_io.py/
motor_tiquipaya.py/sap_writer.py/pipeline_tiquipaya.py/run_batch.py/
consolidador_mensual.py/control_asignaciones.py.

Uso: python -m unittest tests.test_control_cxc_cxp -v
"""

import hashlib
import json
import os
import shutil
import sys
import tempfile
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import control_cxc_cxp as c3


# ---------------------------------------------------------------------------
# Fixtures: SAP GLOBAL sintético mínimo (hoja "1", partidas desde fila 16).
# ---------------------------------------------------------------------------

def _crear_global(ruta, partidas, hoja="1"):
    """partidas: lista de dicts con cuenta/asignacion/debe/haber/glosa/
    fecha_valor (ver _partida). Cargo/Haber solo se escriben si no son
    None, para poder simular celdas realmente vacías."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = hoja

    fila = 16
    for p in partidas:
        ws[f"C{fila}"] = p.get("cuenta")
        ws[f"D{fila}"] = p.get("glosa", "GLOSA DE PRUEBA")
        if p.get("debe") is not None:
            ws[f"E{fila}"] = Decimal(str(p["debe"]))
        if p.get("haber") is not None:
            ws[f"F{fila}"] = Decimal(str(p["haber"]))
        ws[f"O{fila}"] = p.get("fecha_valor")
        ws[f"R{fila}"] = p.get("asignacion")
        fila += 1

    wb.save(ruta)


def _partida(cuenta, asignacion, debe=None, haber=None, glosa="GLOSA DE PRUEBA"):
    return {"cuenta": cuenta, "asignacion": asignacion, "debe": debe, "haber": haber, "glosa": glosa}


def _hash_archivo(ruta):
    with open(ruta, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


class _Control3TestBase(unittest.TestCase):
    CXC_UNIDADES = "110201002"
    CXP_UNIDADES = "210103002"
    CXC_EMPRESAS = "110201003"
    CXP_EMPRESAS = "210103003"
    CXC_PARTICULARES = "110201004"
    CXP_PARTICULARES = "210103004"

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="control_cxc_cxp_test_")
        self.ruta_historico = os.path.join(self.tmpdir, "HISTORICO_CXC_CXP.csv")

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _ruta_global(self, nombre):
        return os.path.join(self.tmpdir, nombre)

    def _ejecutar(self, partidas, nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx", **kwargs):
        ruta_global = self._ruta_global(nombre_global)
        _crear_global(ruta_global, partidas)
        kwargs.setdefault("ruta_salida_xlsx", os.path.join(self.tmpdir, "CONTROL.xlsx"))
        kwargs.setdefault("ruta_salida_json", os.path.join(self.tmpdir, "CONTROL.json"))
        resumen = c3.ejecutar_control(
            ruta_global=ruta_global, ruta_historico=self.ruta_historico, **kwargs
        )
        return resumen, ruta_global

    def _historico(self):
        return c3.cargar_historico(self.ruta_historico)

    def _crear_global_fijo(self, partidas, nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx"):
        """Crea el archivo UNA sola vez y devuelve su ruta, para pruebas
        que necesitan reprocesar EXACTAMENTE el mismo GLOBAL (mismo
        SHA-256) varias veces. openpyxl incrusta una marca de tiempo en
        core.xml: recrear "el mismo" archivo vía _crear_global() en cada
        llamada produce SHA-256 distintos, así que estas pruebas deben
        reutilizar la MISMA ruta física en vez de regenerarla."""
        ruta_global = self._ruta_global(nombre_global)
        _crear_global(ruta_global, partidas)
        return ruta_global

    def _ejecutar_sobre(self, ruta_global, **kwargs):
        kwargs.setdefault("ruta_salida_xlsx", os.path.join(self.tmpdir, "CONTROL.xlsx"))
        kwargs.setdefault("ruta_salida_json", os.path.join(self.tmpdir, "CONTROL.json"))
        return c3.ejecutar_control(ruta_global=ruta_global, ruta_historico=self.ruta_historico, **kwargs)


# ---------------------------------------------------------------------------
# 1-4: reglas CxC
# ---------------------------------------------------------------------------

class TestReglasCxC(_Control3TestBase):
    def test_1_cxc_abierta(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")])
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["saldo"], "10000.00")
        self.assertEqual(fila["estado"], "ABIERTO")

    def test_2_cxc_cerrada_mismo_mes(self):
        self._ejecutar([
            _partida(self.CXC_EMPRESAS, "MISMOMES", debe="500.00"),
            _partida(self.CXC_EMPRESAS, "MISMOMES", haber="500.00"),
        ])
        fila = self._historico()[(self.CXC_EMPRESAS, "MISMOMES")]
        self.assertEqual(fila["saldo"], "0.00")
        self.assertEqual(fila["estado"], "CERRADO")

    def test_3_cxc_cerrada_mes_posterior(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="6000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        r3, _ = self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="4000.00")],
                                nombre_global="SAP_GLOBAL_TIQ_OCTUBRE_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["debe_acumulado"], "10000.00")
        self.assertEqual(fila["haber_acumulado"], "10000.00")
        self.assertEqual(fila["saldo"], "0.00")
        self.assertEqual(fila["estado"], "CERRADO")
        self.assertEqual(r3["estado"], "OK")

    def test_4_cxc_saldo_negativo_revisar(self):
        self._ejecutar([_partida(self.CXC_UNIDADES, "SOBRECOMP", debe="5000.00", haber="6000.00")])
        fila = self._historico()[(self.CXC_UNIDADES, "SOBRECOMP")]
        self.assertEqual(fila["saldo"], "-1000.00")
        self.assertEqual(fila["estado"], "REVISAR")


# ---------------------------------------------------------------------------
# 5-7: reglas CxP
# ---------------------------------------------------------------------------

class TestReglasCxP(_Control3TestBase):
    def test_5_cxp_abierta(self):
        self._ejecutar([_partida(self.CXP_UNIDADES, "PROV1", haber="5000.00")])
        fila = self._historico()[(self.CXP_UNIDADES, "PROV1")]
        self.assertEqual(fila["saldo"], "5000.00")
        self.assertEqual(fila["estado"], "ABIERTO")

    def test_6_cxp_cerrada(self):
        self._ejecutar([_partida(self.CXP_UNIDADES, "PROV1", haber="5000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        self._ejecutar([_partida(self.CXP_UNIDADES, "PROV1", debe="5000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        fila = self._historico()[(self.CXP_UNIDADES, "PROV1")]
        self.assertEqual(fila["saldo"], "0.00")
        self.assertEqual(fila["estado"], "CERRADO")

    def test_7_cxp_saldo_negativo_revisar(self):
        # CXP: saldo = haber_acum - debe_acum. debe > haber -> negativo.
        self._ejecutar([_partida(self.CXP_EMPRESAS, "SOBRECOMP", debe="800.00", haber="300.00")])
        fila = self._historico()[(self.CXP_EMPRESAS, "SOBRECOMP")]
        self.assertEqual(fila["saldo"], "-500.00")
        self.assertEqual(fila["estado"], "REVISAR")


# ---------------------------------------------------------------------------
# 8-9, 21: llave CUENTA+ASIGNACION
# ---------------------------------------------------------------------------

class TestLlave(_Control3TestBase):
    def test_8_misma_llave_acumula_sin_fila_nueva(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "ABC123", debe="1000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "ABC123", debe="500.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "ABC123", debe="200.00")],
                        nombre_global="SAP_GLOBAL_TIQ_OCTUBRE_2026.xlsx")
        historico = self._historico()
        claves_abc123 = [k for k in historico if k[1] == "ABC123"]
        self.assertEqual(len(claves_abc123), 1)
        fila = historico[(self.CXC_EMPRESAS, "ABC123")]
        self.assertEqual(fila["debe_acumulado"], "1700.00")

    def test_9_misma_asignacion_cuentas_distintas_son_entidades_distintas(self):
        self._ejecutar([
            _partida(self.CXC_EMPRESAS, "ABC123", debe="100.00"),
            _partida(self.CXC_PARTICULARES, "ABC123", debe="200.00"),
        ])
        historico = self._historico()
        self.assertIn((self.CXC_EMPRESAS, "ABC123"), historico)
        self.assertIn((self.CXC_PARTICULARES, "ABC123"), historico)
        self.assertEqual(historico[(self.CXC_EMPRESAS, "ABC123")]["debe_acumulado"], "100.00")
        self.assertEqual(historico[(self.CXC_PARTICULARES, "ABC123")]["debe_acumulado"], "200.00")

    def test_21_cuentas_fuera_del_universo_se_ignoran(self):
        resumen, _ = self._ejecutar([
            _partida(self.CXC_EMPRESAS, "DENTRO", debe="100.00"),
            _partida("999999999", "FUERA", debe="99999.00"),
            _partida("110101001", "OTRA_CUENTA_CUALQUIERA", debe="1.00"),
        ])
        historico = self._historico()
        self.assertEqual(len(historico), 1)
        self.assertIn((self.CXC_EMPRESAS, "DENTRO"), historico)
        self.assertEqual(resumen["llaves_evaluadas"], 1)


# ---------------------------------------------------------------------------
# 10: movimientos parciales mantienen ABIERTO
# ---------------------------------------------------------------------------

class TestMovimientosParciales(_Control3TestBase):
    def test_10_movimiento_parcial_mantiene_abierto(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        r2, _ = self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="6000.00")],
                                nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["saldo"], "4000.00")
        self.assertEqual(fila["estado"], "ABIERTO")
        self.assertEqual(r2["abiertas"], 1)


# ---------------------------------------------------------------------------
# 11-15: OBSERVACION_SISTEMA determinística
# ---------------------------------------------------------------------------

class TestObservacionSistema(_Control3TestBase):
    def test_11_observacion_indica_apertura(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["observacion_sistema"],
                          "Cuenta abierta AGOSTO 2026. Saldo pendiente Bs 10.000,00.")

    def test_12_observacion_indica_continuidad(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="6000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(
            fila["observacion_sistema"],
            "Cuenta abierta AGOSTO 2026. Continúa abierta a SEPTIEMBRE 2026. Saldo pendiente Bs 4.000,00.",
        )

    def test_13_observacion_indica_cierre(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["observacion_sistema"],
                          "Cuenta abierta AGOSTO 2026. Cerrada SEPTIEMBRE 2026.")

    def test_13b_observacion_abre_y_cierra_mismo_mes(self):
        self._ejecutar([
            _partida(self.CXC_EMPRESAS, "MISMOMES", debe="100.00"),
            _partida(self.CXC_EMPRESAS, "MISMOMES", haber="100.00"),
        ], nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "MISMOMES")]
        self.assertEqual(fila["observacion_sistema"], "Cuenta abierta AGOSTO 2026. Cerrada AGOSTO 2026.")

    def test_14_observacion_indica_reapertura(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="2000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_OCTUBRE_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(
            fila["observacion_sistema"],
            "Cuenta abierta AGOSTO 2026. Cerrada SEPTIEMBRE 2026. Reabierta OCTUBRE 2026. "
            "Saldo pendiente Bs 2.000,00.",
        )

    def test_15_observacion_indica_sobrecompensacion(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="10500.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["saldo"], "-500.00")
        self.assertEqual(fila["estado"], "REVISAR")
        self.assertEqual(
            fila["observacion_sistema"],
            "Cuenta abierta AGOSTO 2026. Sobrecompensación detectada SEPTIEMBRE 2026. "
            "Saldo Bs -500,00. REVISAR.",
        )

    def test_15b_observacion_negativa_sin_apertura_previa(self):
        self._ejecutar([_partida(self.CXP_UNIDADES, "NUEVA1", debe="500.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        fila = self._historico()[(self.CXP_UNIDADES, "NUEVA1")]
        self.assertEqual(fila["saldo"], "-500.00")
        self.assertEqual(
            fila["observacion_sistema"],
            "Saldo negativo detectado AGOSTO 2026 sin apertura previa registrada. Saldo Bs -500,00. REVISAR.",
        )


# ---------------------------------------------------------------------------
# 16-19: OBSERVACION_AUDITOR y puente JSON
# ---------------------------------------------------------------------------

class TestObservacionAuditorYPuenteJSON(_Control3TestBase):
    def _sha(self, ruta_global):
        return _hash_archivo(ruta_global)

    def _escribir_observaciones_json(self, ruta, periodo, sha256_global, observaciones):
        with open(ruta, "w", encoding="utf-8") as f:
            json.dump({"periodo": periodo, "sha256_global": sha256_global, "observaciones": observaciones}, f)
        return ruta

    def test_16_observacion_auditor_se_conserva_entre_meses(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        historico = self._historico()
        historico[(self.CXC_EMPRESAS, "FORTALEZA")]["observacion_auditor"] = "Esperando transferencia."
        c3.guardar_historico(self.ruta_historico, historico)

        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", haber="6000.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["observacion_auditor"], "Esperando transferencia.")
        self.assertEqual(fila["saldo"], "4000.00")

    def test_17_puente_json_modifica_unicamente_observacion_auditor(self):
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")])
        self._ejecutar_sobre(ruta_global)
        sha = self._sha(ruta_global)
        obs_json = self._escribir_observaciones_json(
            os.path.join(self.tmpdir, "OBS.json"), "AGOSTO_2026", sha,
            [{"cuenta": self.CXC_EMPRESAS, "asignacion": "FORTALEZA", "observacion_auditor": "Nota del auditor."}],
        )
        resumen = self._ejecutar_sobre(ruta_global, ruta_observaciones_json=obs_json)
        self.assertEqual(resumen["estado"], "YA_PROCESADO_SIN_CAMBIOS")
        self.assertEqual(resumen["observaciones_aplicadas"], 1)
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["observacion_auditor"], "Nota del auditor.")
        self.assertEqual(fila["saldo"], "10000.00")
        self.assertEqual(fila["estado"], "ABIERTO")
        self.assertEqual(fila["debe_acumulado"], "10000.00")

    def test_18_json_no_puede_crear_llave_inexistente(self):
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")])
        self._ejecutar_sobre(ruta_global)
        sha = self._sha(ruta_global)
        obs_json = self._escribir_observaciones_json(
            os.path.join(self.tmpdir, "OBS.json"), "AGOSTO_2026", sha,
            [{"cuenta": self.CXC_EMPRESAS, "asignacion": "NO_EXISTE", "observacion_auditor": "x"}],
        )
        resumen = self._ejecutar_sobre(ruta_global, ruta_observaciones_json=obs_json)
        self.assertEqual(resumen["observaciones_aplicadas"], 0)
        self.assertTrue(any("LLAVE_INEXISTENTE" in p for p in resumen["problemas_observaciones_json"]))
        self.assertNotIn((self.CXC_EMPRESAS, "NO_EXISTE"), self._historico())

    def test_19_json_no_puede_modificar_saldo_estado_importes(self):
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")])
        self._ejecutar_sobre(ruta_global)
        sha = self._sha(ruta_global)
        # El JSON no expone ningún campo para saldo/estado/importes/asignación:
        # cualquier clave extra que no sea observacion_auditor se ignora.
        obs_json = self._escribir_observaciones_json(
            os.path.join(self.tmpdir, "OBS.json"), "AGOSTO_2026", sha,
            [{"cuenta": self.CXC_EMPRESAS, "asignacion": "FORTALEZA", "observacion_auditor": "Nota.",
              "saldo": "999999.99", "estado": "CERRADO", "debe_acumulado": "0.00"}],
        )
        resumen = self._ejecutar_sobre(ruta_global, ruta_observaciones_json=obs_json)
        self.assertEqual(resumen["observaciones_aplicadas"], 1)
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["saldo"], "10000.00")
        self.assertEqual(fila["estado"], "ABIERTO")
        self.assertEqual(fila["debe_acumulado"], "10000.00")
        self.assertEqual(fila["observacion_auditor"], "Nota.")

    def test_19b_json_periodo_o_sha_distinto_rechaza_todo_el_lote(self):
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="10000.00")])
        self._ejecutar_sobre(ruta_global)
        obs_json = self._escribir_observaciones_json(
            os.path.join(self.tmpdir, "OBS.json"), "AGOSTO_2026", "sha_incorrecto",
            [{"cuenta": self.CXC_EMPRESAS, "asignacion": "FORTALEZA", "observacion_auditor": "Nota."}],
        )
        resumen = self._ejecutar_sobre(ruta_global, ruta_observaciones_json=obs_json)
        self.assertEqual(resumen["observaciones_aplicadas"], 0)
        self.assertIn("SHA256_GLOBAL_NO_COINCIDE", resumen["problemas_observaciones_json"])
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["observacion_auditor"], "")


# ---------------------------------------------------------------------------
# 20: asignación faltante
# ---------------------------------------------------------------------------

class TestAsignacionFaltante(_Control3TestBase):
    def test_20_asignacion_faltante_se_reporta_y_no_se_acumula(self):
        resumen, _ = self._ejecutar([
            _partida(self.CXP_UNIDADES, None, haber="300.00", glosa="SIN ASIGNACION"),
            _partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00"),
        ])
        self.assertEqual(resumen["asignaciones_faltantes"], 1)
        detalle = resumen["detalle_asignaciones_faltantes"][0]
        self.assertEqual(detalle["tipo"], "ASIGNACION_FALTANTE_CXC_CXP")
        self.assertEqual(detalle["cuenta"], self.CXP_UNIDADES)
        self.assertEqual(detalle["glosa"], "SIN ASIGNACION")
        self.assertEqual(detalle["haber"], "300.00")
        self.assertEqual(detalle["debe"], "0.00")
        self.assertIn("fila_global", detalle)
        historico = self._historico()
        self.assertEqual(len(historico), 1)
        self.assertNotIn((self.CXP_UNIDADES, None), historico)
        self.assertNotIn((self.CXP_UNIDADES, ""), historico)


# ---------------------------------------------------------------------------
# 22-23: validaciones técnicas del GLOBAL (sin fallback)
# ---------------------------------------------------------------------------

class TestValidacionesTecnicas(_Control3TestBase):
    def test_22_nombre_global_no_canonico_bloquea(self):
        resumen, _ = self._ejecutar(
            [_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")],
            nombre_global="GLOBAL_AGOSTO.xlsx",
        )
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NOMBRE_NO_CANONICO", resumen["problemas"])
        self.assertEqual(self._historico(), {})

    def test_23_hoja_distinta_de_1_bloquea(self):
        ruta_global = self._ruta_global("SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        _crear_global(ruta_global, [_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")], hoja="OTRA_HOJA")
        resumen = c3.ejecutar_control(
            ruta_global=ruta_global, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "CONTROL.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "CONTROL.json"),
        )
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_HOJA_1_NO_ENCONTRADA", resumen["problemas"])
        self.assertEqual(self._historico(), {})

    def test_global_no_encontrado(self):
        resumen = c3.ejecutar_control(
            ruta_global=os.path.join(self.tmpdir, "NO_EXISTE.xlsx"),
            ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "CONTROL.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "CONTROL.json"),
        )
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NO_ENCONTRADO", resumen["problemas"])


# ---------------------------------------------------------------------------
# 24-26: idempotencia
# ---------------------------------------------------------------------------

class TestIdempotencia(_Control3TestBase):
    def test_24_mismo_periodo_mismo_sha_ya_procesado(self):
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")])
        self._ejecutar_sobre(ruta_global)
        resumen2 = self._ejecutar_sobre(ruta_global)
        self.assertEqual(resumen2["estado"], "YA_PROCESADO_SIN_CAMBIOS")

    def test_25_mismo_periodo_sha_distinto_global_modificado(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        # Un archivo DISTINTO (contenido distinto -> SHA distinto) pero
        # registrado bajo el MISMO nombre/periodo canónico.
        ruta_v2 = self._ruta_global("SAP_GLOBAL_TIQ_AGOSTO_2026_V2.xlsx")
        _crear_global(ruta_v2, [_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="999999.00")])
        resumen = c3.ejecutar_control(
            ruta_global=ruta_v2, ruta_historico=self.ruta_historico,
            nombre_archivo_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx",
            ruta_salida_xlsx=os.path.join(self.tmpdir, "CONTROL2.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "CONTROL2.json"),
        )
        self.assertEqual(resumen["estado"], "GLOBAL_MODIFICADO_REQUIERE_REVISION")
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["debe_acumulado"], "100.00")  # histórico intacto

    def test_26_rerun_no_duplica_acumulados(self):
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")])
        for _ in range(3):
            self._ejecutar_sobre(ruta_global)
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["debe_acumulado"], "100.00")


# ---------------------------------------------------------------------------
# 27: GLOBAL permanece byte-idéntico
# ---------------------------------------------------------------------------

class TestGlobalReadOnly(_Control3TestBase):
    def test_27_global_permanece_identico(self):
        ruta_global = self._ruta_global("SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        _crear_global(ruta_global, [_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")])
        sha_antes = _hash_archivo(ruta_global)
        mtime_antes = os.path.getmtime(ruta_global)

        c3.ejecutar_control(
            ruta_global=ruta_global, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "CONTROL.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "CONTROL.json"),
        )

        self.assertEqual(_hash_archivo(ruta_global), sha_antes)
        self.assertEqual(os.path.getmtime(ruta_global), mtime_antes)


# ---------------------------------------------------------------------------
# 28: --dry-run no escribe nada
# ---------------------------------------------------------------------------

class TestDryRun(_Control3TestBase):
    def test_28_dry_run_no_escribe_nada(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        ruta_json = os.path.join(self.tmpdir, "CONTROL.json")
        resumen, _ = self._ejecutar(
            [_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")],
            ruta_salida_xlsx=ruta_xlsx, ruta_salida_json=ruta_json, dry_run=True,
        )
        self.assertEqual(resumen["dry_run"], True)
        self.assertFalse(resumen["historico_actualizado"])
        self.assertFalse(os.path.isfile(self.ruta_historico))
        self.assertFalse(os.path.isfile(ruta_xlsx))
        self.assertFalse(os.path.isfile(ruta_json))
        self.assertFalse(os.path.isfile(c3._ruta_libro_periodos(self.ruta_historico)))

    def test_28b_dry_run_no_modifica_historico_existente(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        with open(self.ruta_historico, encoding="utf-8") as f:
            contenido_antes = f.read()

        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="500.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx", dry_run=True)

        with open(self.ruta_historico, encoding="utf-8") as f:
            contenido_despues = f.read()
        self.assertEqual(contenido_antes, contenido_despues)


# ---------------------------------------------------------------------------
# 29-31: Excel humano
# ---------------------------------------------------------------------------

class TestExcelHumano(_Control3TestBase):
    def test_29_excel_una_sola_hoja_control(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")], ruta_salida_xlsx=ruta_xlsx)
        wb = openpyxl.load_workbook(ruta_xlsx)
        self.assertEqual(wb.sheetnames, ["CONTROL"])

    def test_30_excel_una_fila_por_cuenta_asignacion(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        self._ejecutar([
            _partida(self.CXC_EMPRESAS, "ABC123", debe="100.00"),
            _partida(self.CXC_EMPRESAS, "ABC123", debe="50.00"),
            _partida(self.CXC_PARTICULARES, "XYZ789", debe="200.00"),
        ], ruta_salida_xlsx=ruta_xlsx)
        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb["CONTROL"]
        self.assertEqual(ws.max_row, 3)  # encabezado + 2 llaves distintas

    def test_31_orden_revisar_abierto_cerrado(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        self._ejecutar([
            _partida(self.CXC_EMPRESAS, "CERRADA", debe="100.00", haber="100.00"),
            _partida(self.CXC_EMPRESAS, "ABIERTA", debe="100.00"),
            _partida(self.CXC_EMPRESAS, "REVISAR1", debe="50.00", haber="100.00"),
        ], ruta_salida_xlsx=ruta_xlsx)
        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb["CONTROL"]
        header = [c.value for c in ws[1]]
        idx_estado = header.index("ESTADO")
        estados_en_orden = [row[idx_estado].value for row in ws.iter_rows(min_row=2)]
        self.assertEqual(estados_en_orden, ["REVISAR", "ABIERTO", "CERRADO"])

    def test_columnas_visibles_del_excel(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        self._ejecutar([_partida(self.CXC_UNIDADES, "X1", debe="10.00")], ruta_salida_xlsx=ruta_xlsx)
        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb["CONTROL"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header, [
            "PERIODO_CONTROL", "CUENTA", "TIPO", "ASIGNACION", "DEBE_MES", "HABER_MES",
            "DEBE_ACUMULADO", "HABER_ACUMULADO", "SALDO", "ESTADO",
            "OBSERVACION_SISTEMA", "OBSERVACION_AUDITOR",
        ])

    def test_tipo_visible_para_las_6_cuentas(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        self._ejecutar([
            _partida(self.CXC_UNIDADES, "A1", debe="1.00"),
            _partida(self.CXP_UNIDADES, "A2", haber="1.00"),
            _partida(self.CXC_EMPRESAS, "A3", debe="1.00"),
            _partida(self.CXP_EMPRESAS, "A4", haber="1.00"),
            _partida(self.CXC_PARTICULARES, "A5", debe="1.00"),
            _partida(self.CXP_PARTICULARES, "A6", haber="1.00"),
        ], ruta_salida_xlsx=ruta_xlsx)
        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb["CONTROL"]
        header = [c.value for c in ws[1]]
        idx_tipo = header.index("TIPO")
        tipos = {row[idx_tipo].value for row in ws.iter_rows(min_row=2)}
        self.assertEqual(tipos, {
            "CxC UNIDADES", "CxP UNIDADES", "CxC EMPRESAS", "CxP EMPRESAS",
            "CxC PARTICULARES", "CxP PARTICULARES",
        })


# ---------------------------------------------------------------------------
# Snapshot mensual: el Excel incluye TODA la llave del histórico (ABIERTO
# sigue visible aunque no haya movimiento nuevo ese mes).
# ---------------------------------------------------------------------------

class TestSnapshotMensual(_Control3TestBase):
    def test_llave_sin_movimiento_sigue_visible_en_excel(self):
        self._ejecutar([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")],
                        nombre_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx")
        ruta_xlsx_sep = os.path.join(self.tmpdir, "CONTROL_SEP.xlsx")
        self._ejecutar([_partida(self.CXC_UNIDADES, "OTRA", debe="1.00")],
                        nombre_global="SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx", ruta_salida_xlsx=ruta_xlsx_sep)
        wb = openpyxl.load_workbook(ruta_xlsx_sep)
        ws = wb["CONTROL"]
        self.assertEqual(ws.max_row, 3)  # encabezado + FORTALEZA (sin mov.) + OTRA
        header = [c.value for c in ws[1]]
        idx_asig = header.index("ASIGNACION")
        idx_debe_mes = header.index("DEBE_MES")
        filas = {row[idx_asig].value: row[idx_debe_mes].value for row in ws.iter_rows(min_row=2)}
        self.assertEqual(float(filas["FORTALEZA"]), 0.0)


# ---------------------------------------------------------------------------
# Asignación faltante como fila EXCEPCIONAL en el Excel (nunca en el
# histórico): ESTADO=REVISAR, ASIGNACION="ASIGNACION FALTANTE", una fila
# por ocurrencia, arriba junto a REVISAR.
# ---------------------------------------------------------------------------

class TestAsignacionFaltanteExcel(_Control3TestBase):
    def test_asignacion_faltante_fila_excepcional_en_excel(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        resumen, _ = self._ejecutar([
            _partida(self.CXP_UNIDADES, None, haber="300.00", glosa="SIN ASIGNACION"),
        ], ruta_salida_xlsx=ruta_xlsx)
        fila_global_esperada = resumen["detalle_asignaciones_faltantes"][0]["fila_global"]

        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb["CONTROL"]
        header = [c.value for c in ws[1]]
        filas = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(filas), 1)
        fila = dict(zip(header, filas[0]))
        self.assertEqual(fila["ESTADO"], "REVISAR")
        self.assertEqual(fila["ASIGNACION"], "ASIGNACION FALTANTE")
        self.assertEqual(fila["CUENTA"], self.CXP_UNIDADES)
        self.assertEqual(fila["TIPO"], "CxP UNIDADES")
        self.assertEqual(float(fila["HABER_MES"]), 300.0)
        self.assertEqual(float(fila["DEBE_MES"]), 0.0)
        self.assertEqual(
            fila["OBSERVACION_SISTEMA"],
            f"Asignación faltante en fila GLOBAL {fila_global_esperada}. No incorporada al histórico.",
        )
        # Nunca crea una llave del histórico.
        self.assertEqual(self._historico(), {})

    def test_asignacion_faltante_multiples_ocurrencias_filas_separadas(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        self._ejecutar([
            _partida(self.CXC_EMPRESAS, None, debe="100.00"),
            _partida(self.CXC_EMPRESAS, None, debe="200.00"),
        ], ruta_salida_xlsx=ruta_xlsx)
        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb["CONTROL"]
        self.assertEqual(ws.max_row, 3)  # encabezado + 2 filas excepcionales, nunca fusionadas
        header = [c.value for c in ws[1]]
        idx_debe_mes = header.index("DEBE_MES")
        importes = sorted(float(row[idx_debe_mes].value) for row in ws.iter_rows(min_row=2))
        self.assertEqual(importes, [100.0, 200.0])

    def test_asignacion_faltante_aparece_arriba_junto_a_revisar(self):
        ruta_xlsx = os.path.join(self.tmpdir, "CONTROL.xlsx")
        self._ejecutar([
            _partida(self.CXC_EMPRESAS, "CERRADA", debe="100.00", haber="100.00"),
            _partida(self.CXC_EMPRESAS, "ABIERTA", debe="100.00"),
            _partida(self.CXC_EMPRESAS, "REVISAR1", debe="50.00", haber="100.00"),
            _partida(self.CXC_EMPRESAS, None, debe="10.00"),
        ], ruta_salida_xlsx=ruta_xlsx)
        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb["CONTROL"]
        header = [c.value for c in ws[1]]
        idx_asig = header.index("ASIGNACION")
        idx_estado = header.index("ESTADO")
        filas = [(row[idx_asig].value, row[idx_estado].value) for row in ws.iter_rows(min_row=2)]
        self.assertEqual(filas, [
            ("ASIGNACION FALTANTE", "REVISAR"),
            ("REVISAR1", "REVISAR"),
            ("ABIERTA", "ABIERTO"),
            ("CERRADA", "CERRADO"),
        ])


# ---------------------------------------------------------------------------
# Consistencia HISTORICO_CXC_CXP.csv <-> HISTORICO_CXC_CXP_PERIODOS.json
# ante una ejecución interrumpida a mitad de camino (ninguno de los dos
# escenarios debe producir doble acumulación ni un periodo marcado como
# aplicado sin histórico real).
# ---------------------------------------------------------------------------

class TestConsistenciaHistoricoPeriodos(_Control3TestBase):
    """El libro de periodos (HISTORICO_CXC_CXP_PERIODOS.json) es la
    fuente AUTORITATIVA E INMUTABLE de qué periodo+SHA ya fue aplicado —
    nunca se invalida por un movimiento posterior de la MISMA llave en
    otro periodo (periodo_ultimo_movimiento/sha256_global_ultimo son
    mutables y se sobrescriben). Mecanismo de recuperación:
    PENDIENTE -> escritura atómica del histórico -> APLICADO."""

    def test_periodo_antiguo_no_se_reacumula_tras_movimiento_posterior_de_la_misma_llave(self):
        # 1) procesar AGOSTO
        ruta_ago = self._crear_global_fijo(
            [_partida(self.CXC_EMPRESAS, "ABC", debe="1000.00")], "SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx",
        )
        self._ejecutar_sobre(ruta_ago)

        # 2) la MISMA CUENTA+ASIGNACION se mueve en SEPTIEMBRE
        #    (sobrescribe periodo_ultimo_movimiento/sha256_global_ultimo
        #    de esa fila, que ya NO dicen "AGOSTO_2026").
        ruta_sep = self._crear_global_fijo(
            [_partida(self.CXC_EMPRESAS, "ABC", debe="500.00")], "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx",
        )
        self._ejecutar_sobre(ruta_sep)
        fila_antes = self._historico()[(self.CXC_EMPRESAS, "ABC")]
        self.assertEqual(fila_antes["debe_acumulado"], "1500.00")

        # 3) reejecutar AGOSTO (mismo archivo/SHA de antes)
        resumen = self._ejecutar_sobre(ruta_ago)

        # 4) YA_PROCESADO_SIN_CAMBIOS
        self.assertEqual(resumen["estado"], "YA_PROCESADO_SIN_CAMBIOS")

        # 5) saldo exactamente igual antes/después del rerun
        fila_despues = self._historico()[(self.CXC_EMPRESAS, "ABC")]
        self.assertEqual(fila_antes, fila_despues)

        # 6) AGOSTO no se acumuló dos veces (seguiría en 1500, nunca 2500)
        self.assertEqual(fila_despues["debe_acumulado"], "1500.00")
        self.assertEqual(fila_despues["saldo"], "1500.00")

        # El libro de periodos conserva AGOSTO como APLICADO con su SHA
        # original, sin que SEPTIEMBRE lo haya tocado.
        libro = c3._cargar_libro_periodos(self.ruta_historico)
        self.assertEqual(libro["AGOSTO_2026"]["estado"], "APLICADO")
        self.assertEqual(libro["AGOSTO_2026"]["sha256_global"], _hash_archivo(ruta_ago))

    def test_recuperacion_pendiente_sin_historico_escrito_reprocesa_una_vez(self):
        # Simula un corte ANTES de que el histórico llegara a escribirse:
        # el libro de periodos quedó en PENDIENTE y el histórico todavía
        # no refleja nada de este periodo+SHA.
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")])
        sha = _hash_archivo(ruta_global)
        c3._guardar_libro_periodos(self.ruta_historico, {
            "AGOSTO_2026": {"sha256_global": sha, "estado": "PENDIENTE", "fecha_inicio": "2026-01-01T00:00:00"},
        })
        self.assertFalse(os.path.isfile(self.ruta_historico))

        resumen = self._ejecutar_sobre(ruta_global)
        self.assertEqual(resumen["estado"], "OK")  # reprocesa de verdad, nunca YA_PROCESADO ciego
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["debe_acumulado"], "100.00")
        libro = c3._cargar_libro_periodos(self.ruta_historico)
        self.assertEqual(libro["AGOSTO_2026"]["estado"], "APLICADO")

        # Rerun posterior: ya idempotente, sin duplicar.
        resumen2 = self._ejecutar_sobre(ruta_global)
        self.assertEqual(resumen2["estado"], "YA_PROCESADO_SIN_CAMBIOS")
        fila2 = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila2["debe_acumulado"], "100.00")

    def test_recuperacion_pendiente_con_historico_ya_escrito_sella_aplicado_sin_reacumular(self):
        # Corrida normal completa (queda APLICADO).
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")])
        self._ejecutar_sobre(ruta_global)
        sha = _hash_archivo(ruta_global)

        # Simula un corte DESPUÉS de escribir el histórico pero ANTES de
        # sellar APLICADO: se regresa manualmente el registro a PENDIENTE
        # (el histórico real YA tiene la acumulación aplicada).
        c3._guardar_libro_periodos(self.ruta_historico, {
            "AGOSTO_2026": {"sha256_global": sha, "estado": "PENDIENTE", "fecha_inicio": "2026-01-01T00:00:00"},
        })

        resumen = self._ejecutar_sobre(ruta_global)
        self.assertEqual(resumen["estado"], "YA_PROCESADO_SIN_CAMBIOS")  # nunca reacumula
        fila = self._historico()[(self.CXC_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["debe_acumulado"], "100.00")  # sigue en 100, no 200

        libro = c3._cargar_libro_periodos(self.ruta_historico)
        self.assertEqual(libro["AGOSTO_2026"]["estado"], "APLICADO")
        self.assertTrue(libro["AGOSTO_2026"].get("recuperado"))

    def test_idempotencia_no_depende_de_fila_historica_del_periodo(self):
        # Periodo SIN ninguna llave válida (solo una partida con
        # ASIGNACION FALTANTE, que nunca crea fila en el histórico): la
        # idempotencia debe basarse ÚNICAMENTE en el libro de periodos,
        # nunca en que exista una fila histórica para ese periodo.
        ruta_global = self._crear_global_fijo(
            [_partida(self.CXC_EMPRESAS, None, debe="10.00")], "SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx",
        )
        resumen1 = self._ejecutar_sobre(ruta_global)
        self.assertEqual(resumen1["estado"], "OK")
        self.assertEqual(resumen1["llaves_evaluadas"], 0)
        self.assertEqual(self._historico(), {})  # ninguna fila creada

        libro = c3._cargar_libro_periodos(self.ruta_historico)
        self.assertEqual(libro["AGOSTO_2026"]["estado"], "APLICADO")

        # A pesar de que _historico_ya_refleja_periodo daría False (no
        # hay ninguna fila), el rerun debe ser idempotente igual.
        resumen2 = self._ejecutar_sobre(ruta_global)
        self.assertEqual(resumen2["estado"], "YA_PROCESADO_SIN_CAMBIOS")

    def test_sha_distinto_prevalece_sobre_chequeo_de_historico(self):
        # Aunque el histórico ya refleje AGOSTO_2026 con OTRO sha (de una
        # corrida anterior legítima), un GLOBAL con SHA distinto para el
        # mismo periodo sigue bloqueando (nunca se decide por el
        # histórico cuando el libro registra explícitamente otro SHA).
        ruta_global = self._crear_global_fijo([_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="100.00")])
        self._ejecutar_sobre(ruta_global)

        ruta_v2 = self._ruta_global("SAP_GLOBAL_TIQ_AGOSTO_2026_V2.xlsx")
        _crear_global(ruta_v2, [_partida(self.CXC_EMPRESAS, "FORTALEZA", debe="999999.00")])
        resumen = c3.ejecutar_control(
            ruta_global=ruta_v2, ruta_historico=self.ruta_historico,
            nombre_archivo_global="SAP_GLOBAL_TIQ_AGOSTO_2026.xlsx",
            ruta_salida_xlsx=os.path.join(self.tmpdir, "CONTROL2.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "CONTROL2.json"),
        )
        self.assertEqual(resumen["estado"], "GLOBAL_MODIFICADO_REQUIERE_REVISION")


if __name__ == "__main__":
    unittest.main()
