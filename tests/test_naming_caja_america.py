"""test_naming_caja_america.py — Identidad/naming de archivos SAP y GLOBAL
por caja (Caja Tiquipaya vs Caja América).

Alcance ÚNICO de esta suite: nombres de archivo y cabecera (L10). No repite
ninguna regla contable, de lectura de Excel, ni la regla POSGRADO RESERVA
ya cubierta por tests/test_caja_america.py y
tests/test_control_america_postg.py — el test O de aquí solo CONFIRMA (sin
reimplementar) que CONTROL 3 sigue acumulando POSTG-<MES> cuando se
ejecuta explícitamente como caja="america".

NO se toca v3/publicacion.py ni v3/consolidador_mensual_v3.py en esta
tarea: el naming que run_batch produce localmente (SAP_DD-MM-YYYY.xlsx)
permanece sin prefijo de caja a propósito (legacy V2, ver
run_batch._nombre_sap_esperado).

Ningún dato contable real. Uso: python -m unittest tests.test_naming_caja_america -v
"""

import datetime
import os
import shutil
import sys
import tempfile
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import config_cajas as cfg  # noqa: E402
import consolidador_mensual as cm  # noqa: E402
import control_asignaciones as ca  # noqa: E402
import control_cxc_cxp as c3  # noqa: E402
import pipeline_tiquipaya as pipeline  # noqa: E402
import run_batch  # noqa: E402

from tests.xlsx_fixtures import crear_plantilla_sap  # noqa: E402
from tests.test_control_asignaciones import (  # noqa: E402
    _crear_global as _crear_global_c1,
    _partida as _partida_c1,
)
from tests.test_control_cxc_cxp import _partida as _partida_c3  # noqa: E402


# ---------------------------------------------------------------------------
# Fixture local: SAP diario sintético con cabecera parametrizable por caja
# (tests/test_consolidador_mensual._crear_sap_diario la deja fija en "CAJA
# TIQUIPAYA"; aquí necesitamos también "CAJA AMERICA" — se define aparte
# para no tocar ese archivo de tests existente).
# ---------------------------------------------------------------------------

def _crear_sap_diario_caja(ruta, referencia, importe="100.00", tipo_asiento="DB"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1"
    ws["B10"] = "BO01"
    ws["C10"] = tipo_asiento
    ws["H10"] = "BOB"
    ws["L10"] = referencia

    fecha_valor = datetime.date(2026, 9, 5)
    fila = 16
    for cargo, haber, cuenta in (
        (importe, "0.00", "110101001"),
        ("0.00", importe, "210201005"),
    ):
        ws[f"B{fila}"] = "BO01"
        ws[f"C{fila}"] = cuenta
        ws[f"D{fila}"] = "PARTIDA DE PRUEBA"
        ws[f"E{fila}"] = Decimal(cargo)
        ws[f"F{fila}"] = Decimal(haber)
        ws[f"L{fila}"] = "10010101"
        ws[f"O{fila}"] = fecha_valor
        ws[f"R{fila}"] = "ASIG-1"
        fila += 1

    wb.save(ruta)


class _NamingTestBase(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="naming_america_test_")
        self.sap_dir = os.path.join(self.tmpdir, "sap_diarios")
        os.makedirs(self.sap_dir)
        self.ruta_plantilla = os.path.join(self.tmpdir, "Plantilla_SAP_maestra.xlsx")
        crear_plantilla_sap(self.ruta_plantilla)

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _args(self, **overrides):
        class _Args:
            pass
        a = _Args()
        a.anio = overrides.get("anio", 2026)
        a.mes = overrides.get("mes", 9)
        a.sap_dir = overrides.get("sap_dir", self.sap_dir)
        a.plantilla = overrides.get("plantilla", self.ruta_plantilla)
        a.salida = overrides.get("salida", os.path.join(self.tmpdir, "SALIDA_GLOBAL.xlsx"))
        a.archivos_lista = overrides.get("archivos_lista", None)
        a.force = overrides.get("force", False)
        a.caja = overrides.get("caja", None)
        return a


# ---------------------------------------------------------------------------
# A, B — pipeline_tiquipaya.nombre_resultado_json
# ---------------------------------------------------------------------------

class TestNombreResultadoJson(unittest.TestCase):

    def test_A_default_es_tiq(self):
        self.assertEqual(
            pipeline.nombre_resultado_json("2026-09-15"),
            "RESULTADO_TIQ_15-09-2026.json",
        )

    def test_A2_tiquipaya_explicito_identico_al_default(self):
        self.assertEqual(
            pipeline.nombre_resultado_json("2026-09-15", "tiquipaya"),
            pipeline.nombre_resultado_json("2026-09-15"),
        )

    def test_B_america(self):
        self.assertEqual(
            pipeline.nombre_resultado_json("2026-09-15", "america"),
            "RESULTADO_AME_15-09-2026.json",
        )

    def test_run_batch_usa_resultado_ame_para_america(self):
        """run_batch.procesar_cierre, cuando procesa América, escribe el
        RESULTADO con el nombre RESULTADO_AME (SAP_DD-MM-YYYY.xlsx local
        se conserva sin prefijo: legacy V2, no se toca en esta tarea)."""
        # No ejecuta el pipeline completo (fuera de alcance aquí: motor ya
        # cubierto en tests/test_caja_america.py); solo confirma el nombre
        # que construir_metadata_cabecera/nombre_resultado_json producirían
        # para la misma llamada que hace run_batch.procesar_cierre.
        self.assertEqual(
            pipeline.nombre_resultado_json("2026-09-15", "america"),
            "RESULTADO_AME_15-09-2026.json",
        )
        self.assertEqual(
            run_batch.construir_metadata_cabecera("2026-09-15", "america")["referencia"],
            "CAJA AMERICA",
        )

    def test_caja_desconocida_falla_cerrado(self):
        with self.assertRaises(ValueError):
            pipeline.nombre_resultado_json("2026-09-15", "brasil")


# ---------------------------------------------------------------------------
# C, D — consolidador_mensual.nombre_sap_global
# ---------------------------------------------------------------------------

class TestNombreSapGlobal(unittest.TestCase):

    def test_C_tiquipaya(self):
        self.assertEqual(
            cm.nombre_sap_global(2026, 9),
            "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx",
        )
        self.assertEqual(
            cm.nombre_sap_global(2026, 9, "tiquipaya"),
            cm.nombre_sap_global(2026, 9),
        )

    def test_D_america(self):
        self.assertEqual(
            cm.nombre_sap_global(2026, 9, "america"),
            "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx",
        )

    def test_nombre_resultado_global_json_por_caja(self):
        self.assertEqual(
            cm.nombre_resultado_json(2026, 9),
            "RESULTADO_GLOBAL_TIQ_SEPTIEMBRE_2026.json",
        )
        self.assertEqual(
            cm.nombre_resultado_json(2026, 9, "america"),
            "RESULTADO_GLOBAL_AME_SEPTIEMBRE_2026.json",
        )


# ---------------------------------------------------------------------------
# E, F, G, H — aislamiento cruzado y cabecera en el consolidador
# ---------------------------------------------------------------------------

class TestConsolidadorAislamientoPorCaja(_NamingTestBase):

    def test_E_america_incluye_ame_no_tiq(self):
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_AME_05-09-2026.xlsx"), "CAJA AMERICA"
        )
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_TIQ_05-09-2026.xlsx"), "CAJA TIQUIPAYA"
        )
        args = self._args(caja="america",
                          salida=os.path.join(self.tmpdir, "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx"))
        resultado = cm.ejecutar_consolidacion(args)
        self.assertEqual(resultado["estado"], "VALIDADO_PENDIENTE_PUBLICACION")
        self.assertEqual(resultado["sap_incluidos"], ["SAP_AME_05-09-2026.xlsx"])

    def test_F_tiquipaya_incluye_tiq_no_ame(self):
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_AME_05-09-2026.xlsx"), "CAJA AMERICA"
        )
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_TIQ_05-09-2026.xlsx"), "CAJA TIQUIPAYA"
        )
        args = self._args(caja=None,
                          salida=os.path.join(self.tmpdir, "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx"))
        resultado = cm.ejecutar_consolidacion(args)
        self.assertEqual(resultado["estado"], "VALIDADO_PENDIENTE_PUBLICACION")
        self.assertEqual(resultado["sap_incluidos"], ["SAP_TIQ_05-09-2026.xlsx"])

    def test_F2_tiquipaya_explicito_mismo_resultado(self):
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_TIQ_05-09-2026.xlsx"), "CAJA TIQUIPAYA"
        )
        args = self._args(caja="tiquipaya",
                          salida=os.path.join(self.tmpdir, "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx"))
        resultado = cm.ejecutar_consolidacion(args)
        self.assertEqual(resultado["sap_incluidos"], ["SAP_TIQ_05-09-2026.xlsx"])

    def test_G_cabecera_global_tiquipaya_l10(self):
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_TIQ_05-09-2026.xlsx"), "CAJA TIQUIPAYA"
        )
        salida = os.path.join(self.tmpdir, "SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        cm.ejecutar_consolidacion(self._args(salida=salida))

        wb = openpyxl.load_workbook(salida)
        try:
            self.assertEqual(wb["1"]["L10"].value, "CAJA TIQUIPAYA")
        finally:
            wb.close()

    def test_H_cabecera_global_america_l10(self):
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_AME_05-09-2026.xlsx"), "CAJA AMERICA"
        )
        salida = os.path.join(self.tmpdir, "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx")
        cm.ejecutar_consolidacion(self._args(caja="america", salida=salida))

        wb = openpyxl.load_workbook(salida)
        try:
            self.assertEqual(wb["1"]["L10"].value, "CAJA AMERICA")
        finally:
            wb.close()

    def test_america_rechaza_sap_diario_con_cabecera_tiquipaya(self):
        """Un archivo con el nombre SAP_AME_... pero cabecera L10 aún
        'CAJA TIQUIPAYA' (dato mal generado) bloquea, no se cuela."""
        _crear_sap_diario_caja(
            os.path.join(self.sap_dir, "SAP_AME_05-09-2026.xlsx"), "CAJA TIQUIPAYA"
        )
        args = self._args(caja="america",
                          salida=os.path.join(self.tmpdir, "SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx"))
        resultado = cm.ejecutar_consolidacion(args)
        self.assertEqual(resultado["estado"], "ERROR_REVISAR")
        self.assertTrue(any("CABECERA_L10" in b for b in resultado["blockers"]))

    def test_guardarrieles_salida_bloquea_nombre_sap_diario_de_la_caja(self):
        with self.assertRaises(RuntimeError):
            cm.validar_guardarrieles_salida(
                os.path.join(self.sap_dir, "SAP_AME_05-09-2026.xlsx"),
                self.ruta_plantilla, [], force=False, caja="america",
            )
        # El mismo nombre NO es problema para tiquipaya (prefijo distinto).
        cm.validar_guardarrieles_salida(
            os.path.join(self.sap_dir, "SAP_AME_05-09-2026.xlsx"),
            self.ruta_plantilla, [], force=False, caja="tiquipaya",
        )


# ---------------------------------------------------------------------------
# I, J, K — CONTROL 1
# ---------------------------------------------------------------------------

class TestControl1Naming(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="control1_naming_test_")
        self.ruta_historico = os.path.join(self.tmpdir, "HISTORICO_ASIGNACIONES.csv")

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _global(self, nombre):
        ruta = os.path.join(self.tmpdir, nombre)
        _crear_global_c1(ruta, [_partida_c1("ASIG-1")])
        return ruta

    def test_I_america_con_global_ame_nombre_valido(self):
        ruta = self._global("SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx")
        resumen = ca.ejecutar_control(ruta_global=ruta, ruta_historico=self.ruta_historico,
                                      caja="america")
        self.assertNotEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertNotIn("GLOBAL_NOMBRE_NO_CANONICO", resumen.get("problemas", []))

    def test_J_america_con_global_tiq_falla_cerrado(self):
        ruta = self._global("SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        resumen = ca.ejecutar_control(ruta_global=ruta, ruta_historico=self.ruta_historico,
                                      caja="america")
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NOMBRE_NO_CANONICO", resumen["problemas"])
        self.assertFalse(os.path.isfile(self.ruta_historico))

    def test_K_default_con_global_tiq_comportamiento_historico(self):
        ruta = self._global("SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        resumen = ca.ejecutar_control(ruta_global=ruta, ruta_historico=self.ruta_historico)
        self.assertNotEqual(resumen["estado"], "ERROR_TECNICO")

    def test_K2_default_con_global_ame_falla_cerrado(self):
        """Simétrico a J: tiquipaya (o el default) rechaza un GLOBAL_AME."""
        ruta = self._global("SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx")
        resumen = ca.ejecutar_control(ruta_global=ruta, ruta_historico=self.ruta_historico)
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NOMBRE_NO_CANONICO", resumen["problemas"])

    def test_mensaje_de_error_referencia_la_caja_correcta(self):
        ruta = self._global("SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx")
        resumen = ca.ejecutar_control(ruta_global=ruta, ruta_historico=self.ruta_historico,
                                      caja="america")
        self.assertIn("SAP_GLOBAL_AME", resumen["mensaje"])


# ---------------------------------------------------------------------------
# L, M, N — CONTROL 3
# ---------------------------------------------------------------------------

class TestControl3Naming(unittest.TestCase):
    CXP_EMPRESAS = "210103003"
    POSTG_SEPT = "POSTG-SEPT"

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="control3_naming_test_")
        self.ruta_historico = os.path.join(self.tmpdir, "HISTORICO_CXC_CXP.csv")

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _global_c3(self, nombre, partidas):
        ruta = os.path.join(self.tmpdir, nombre)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "1"
        fila = 16
        for p in partidas:
            ws[f"C{fila}"] = p.get("cuenta")
            ws[f"D{fila}"] = p.get("glosa", "GLOSA DE PRUEBA")
            if p.get("debe") is not None:
                ws[f"E{fila}"] = Decimal(str(p["debe"]))
            if p.get("haber") is not None:
                ws[f"F{fila}"] = Decimal(str(p["haber"]))
            ws[f"O{fila}"] = p.get("fecha_valor")
            ws[f"R{fila}"] = p.get("asignacion")
            fila += 1
        wb.save(ruta)
        return ruta

    def test_L_america_con_global_ame_procesa_normalmente(self):
        """CONTROL 3 con caja=america sobre un GLOBAL_AME: sigue
        acumulando 210103003 / POSTG-SEPT tal cual antes."""
        ruta = self._global_c3("SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx", [
            _partida_c3(self.CXP_EMPRESAS, self.POSTG_SEPT, haber="350.00"),
        ])
        resumen = c3.ejecutar_control(
            ruta_global=ruta, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "C.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "C.json"),
            caja="america",
        )
        self.assertEqual(resumen["estado"], "OK")
        fila = c3.cargar_historico(self.ruta_historico)[(self.CXP_EMPRESAS, self.POSTG_SEPT)]
        self.assertEqual(fila["saldo"], "350.00")
        self.assertEqual(fila["estado"], "ABIERTO")

    def test_M_america_con_global_tiq_falla_cerrado(self):
        ruta = self._global_c3("SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx", [
            _partida_c3(self.CXP_EMPRESAS, self.POSTG_SEPT, haber="100.00"),
        ])
        resumen = c3.ejecutar_control(
            ruta_global=ruta, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "C.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "C.json"),
            caja="america",
        )
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NOMBRE_NO_CANONICO", resumen["problemas"])
        self.assertFalse(os.path.isfile(self.ruta_historico))

    def test_N_default_con_global_tiq_comportamiento_historico(self):
        # 210103003 es CxP: SALDO = HABER - DEBE, por eso HABER (no DEBE)
        # abre un saldo positivo.
        ruta = self._global_c3("SAP_GLOBAL_TIQ_SEPTIEMBRE_2026.xlsx", [
            _partida_c3(self.CXP_EMPRESAS, "FORTALEZA", haber="10000.00"),
        ])
        resumen = c3.ejecutar_control(
            ruta_global=ruta, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "C.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "C.json"),
        )
        self.assertEqual(resumen["estado"], "OK")
        fila = c3.cargar_historico(self.ruta_historico)[(self.CXP_EMPRESAS, "FORTALEZA")]
        self.assertEqual(fila["saldo"], "10000.00")
        self.assertEqual(fila["estado"], "ABIERTO")

    def test_N2_default_con_global_ame_falla_cerrado(self):
        ruta = self._global_c3("SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx", [
            _partida_c3(self.CXP_EMPRESAS, self.POSTG_SEPT, haber="100.00"),
        ])
        resumen = c3.ejecutar_control(
            ruta_global=ruta, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "C.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "C.json"),
        )
        self.assertEqual(resumen["estado"], "ERROR_TECNICO")
        self.assertIn("GLOBAL_NOMBRE_NO_CANONICO", resumen["problemas"])

    def test_universo_de_cuentas_no_se_toco(self):
        self.assertIn(self.CXP_EMPRESAS, c3._CUENTAS_CONTROL)
        self.assertEqual(c3._CUENTAS_CONTROL[self.CXP_EMPRESAS], ("CXP", "CxP EMPRESAS"))

    # -----------------------------------------------------------------
    # O — Regresión POSTG (confirmación, no reimplementación)
    # -----------------------------------------------------------------

    def test_O_regresion_postg_abierto_y_cerrado_con_caja_america(self):
        """100+50+200 => SALDO 350 ABIERTO; DEBE 350 posterior => CERRADO.
        Mismo comportamiento que sin `caja`, ahora confirmado también
        pasando caja="america" explícitamente sobre su propio GLOBAL_AME."""
        ruta_sep = self._global_c3("SAP_GLOBAL_AME_SEPTIEMBRE_2026.xlsx", [
            _partida_c3(self.CXP_EMPRESAS, self.POSTG_SEPT, haber="100.00"),
            _partida_c3(self.CXP_EMPRESAS, self.POSTG_SEPT, haber="50.00"),
            _partida_c3(self.CXP_EMPRESAS, self.POSTG_SEPT, haber="200.00"),
        ])
        c3.ejecutar_control(
            ruta_global=ruta_sep, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "C1.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "C1.json"),
            caja="america",
        )
        fila = c3.cargar_historico(self.ruta_historico)[(self.CXP_EMPRESAS, self.POSTG_SEPT)]
        self.assertEqual(fila["debe_acumulado"], "0.00")
        self.assertEqual(fila["haber_acumulado"], "350.00")
        self.assertEqual(fila["saldo"], "350.00")
        self.assertEqual(fila["estado"], "ABIERTO")

        ruta_oct = self._global_c3("SAP_GLOBAL_AME_OCTUBRE_2026.xlsx", [
            _partida_c3(self.CXP_EMPRESAS, self.POSTG_SEPT, debe="350.00"),
        ])
        c3.ejecutar_control(
            ruta_global=ruta_oct, ruta_historico=self.ruta_historico,
            ruta_salida_xlsx=os.path.join(self.tmpdir, "C2.xlsx"),
            ruta_salida_json=os.path.join(self.tmpdir, "C2.json"),
            caja="america",
        )
        fila = c3.cargar_historico(self.ruta_historico)[(self.CXP_EMPRESAS, self.POSTG_SEPT)]
        self.assertEqual(fila["debe_acumulado"], "350.00")
        self.assertEqual(fila["haber_acumulado"], "350.00")
        self.assertEqual(fila["saldo"], "0.00")
        self.assertEqual(fila["estado"], "CERRADO")


if __name__ == "__main__":
    unittest.main()
